import openpyxl,pathlib,sys,datetime,re,filecmp
def formatos_text(f):
    """Formatos (último argumento) de cada TEXT( da fórmula, respeitando aspas e parênteses."""
    out=[]
    for m in re.finditer(r'(?<![A-Za-z_.])TEXT\(',f):
        i=m.end(); nivel=1; aspas=False; args=[""]
        while i<len(f) and nivel:
            ch=f[i]
            if ch=='"': aspas=not aspas
            if not aspas:
                if ch=='(': nivel+=1
                elif ch==')':
                    nivel-=1
                    if nivel==0: break
                elif ch==',' and nivel==1: args.append(""); i+=1; continue
            args[-1]+=ch; i+=1
        out.append(args[-1].strip())
    return out
def data_no_formato(fmt): return fmt.startswith('"') and re.search(r'[dDmMyYaAhHsS]',fmt) is not None
P=pathlib.Path("/home/user/claude-code/projetos/gestao-ia-planilhas/produto")
KITS=[("01-essencial","Kit IA no Trabalho · Essencial (R$ 37)","kit-essencial"),("02-completo","Kit IA no Trabalho · Completo (R$ 197)","kit-completo"),("03-advogados","Kit de Gestão para Advogados (R$ 497)","kit-advogados"),("04-medicos","Kit de Gestão para Médicos (R$ 697)","kit-medicos")]
out=[f"# Inventário das 53 planilhas · {datetime.date.today().strftime('%d/%m/%Y')}, depois da auditoria final-2\n","Gerado dos próprios arquivos de `entrega/` com openpyxl. As sete últimas colunas são de controle e devem ser zero: validação sem bloqueio de erro, célula desbloqueada fora do amarelo, `_xlfn`, regra condicional apontando direto para outra aba (sem INDIRECT), texto literal com mais de 255 caracteres dentro de fórmula (limite do Excel; o LibreOffice não acusa), TEXT com código de data (d, m, a, y, h, s: o código muda com o idioma do Excel) e arquivo de `entrega/` diferente do último gerado (1 = a cópia para a entrega não foi feita).\n"]
T={"f":0,"e":0,"v":0,"cf":0,"abas":0,"arq":0}
for pasta,nome,kit in KITS:
    files=sorted((P/kit/"entrega").glob("[0-2][0-9]-*.xlsx"))
    out.append(f"\n## {pasta} — {nome} · {len(files)} planilhas\n\n| Arquivo | Abas | Fórmulas | Entradas | Listas | Regras cond. | Sem bloqueio | Soltas fora do amarelo | _xlfn | Regra sem INDIRECT | Literal > 255 | TEXT com data | Entrega ≠ gerado |\n|---|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|---:|")
    for f in files:
        wb=openpyxl.load_workbook(f); nf=ne=nv=nsb=nsolta=nx=ncf=ncfx=nlit=ntd=0
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    v=c.value
                    if isinstance(v,str) and v.startswith("="):
                        nf+=1
                        if "_xlfn" in v: nx+=1
                        if any(len(l)>255 for l in re.findall(r'"((?:[^"]|"")*)"',v)): nlit+=1
                        ntd+=sum(data_no_formato(x) for x in formatos_text(v))
                    if c.protection is not None and not c.protection.locked and v is not None or (c.protection is not None and not c.protection.locked and c.fill is not None and c.fill.fgColor is not None and str(c.fill.fgColor.rgb).endswith("FFF4CC")):
                        ne+=1
                        rgb=str(c.fill.fgColor.rgb) if c.fill is not None and c.fill.fgColor is not None else ""
                        if not rgb.endswith("FFF4CC") and not (isinstance(v,str) and v.startswith("=")): nsolta+=0  # entradas sem amarelo contadas abaixo
            for dv in ws.data_validations.dataValidation:
                nv+=1
                if not dv.showErrorMessage: nsb+=1
            for cf in ws.conditional_formatting:
                for r in cf.rules:
                    ncf+=1
                    for fm in (r.formula or []):
                        if "!" in fm and "INDIRECT(" not in fm: ncfx+=1
            # soltas: desbloqueadas sem amarelo (só células com valor ou estilo)
            for row in ws.iter_rows():
                for c in row:
                    if type(c).__name__!="MergedCell" and c.protection is not None and not c.protection.locked:
                        rgb=str(c.fill.fgColor.rgb) if c.fill is not None and c.fill.fgColor is not None else ""
                        if not rgb.endswith("FFF4CC") and (c.value is not None or c.has_style): nsolta+=1
        # a entrega tem de ser o último arquivo gerado (no Completo, 01 a 03 são cópias do Essencial)
        gerado=(P/"kit-essencial"/f.name) if kit=="kit-completo" and f.name[:2] in ("01","02","03") else (P/kit/f.name)
        ndif=0 if gerado.exists() and filecmp.cmp(gerado,f,shallow=False) else 1
        T["td"]=T.get("td",0)+ntd; T["dif"]=T.get("dif",0)+ndif
        out.append(f"| `{f.name}` | {' · '.join(wb.sheetnames)} | {nf} | {ne} | {nv} | {ncf} | {nsb} | {nsolta} | {nx} | {ncfx} | {nlit} | {ntd} | {ndif} |")
        T["f"]+=nf; T["e"]+=ne; T["v"]+=nv; T["cf"]+=ncf; T["abas"]+=len(wb.sheetnames); T["arq"]+=1
out.append(f"\n## Total\n\n{T['arq']} arquivos · {T['abas']} abas · {T['f']:,} fórmulas · {T['e']:,} células de entrada · {T['v']} validações · {T['cf']} regras condicionais · {T.get('td',0)} TEXT com código de data · {T.get('dif',0)} arquivo(s) de entrega diferentes do gerado.".replace(",","."))
pathlib.Path(sys.argv[1]).write_text("\n".join(out)+"\n"); print("\n".join(out[-3:]))
