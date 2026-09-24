#!/usr/bin/env python3
"""Grava os valores calculados (cache) nas planilhas geradas pelo openpyxl, sem reescrever o arquivo pelo LibreOffice.
Por quê: o openpyxl salva as fórmulas sem valor; visualizadores que não recalculam (celular, Google Drive, pré-visualização
do e-mail) mostram as células vazias. O LibreOffice "abrir e salvar" resolve isso, mas reescreve estilos, gráficos e colunas
ocultas. Aqui só o <v> de cada célula com fórmula é acrescentado ao XML original: todo o resto do arquivo fica byte a byte igual.
Uso: python3 cache_valores.py <pasta com as cópias recalculadas (recalc.py)> [arquivo.xlsx ...]
     (sem arquivos: aplica às 20 planilhas desta pasta). Excel recalcula tudo ao abrir (fullCalcOnLoad já está ligado)."""
import sys, pathlib, re, zipfile, datetime, shutil, html
import openpyxl
from openpyxl.utils.datetime import to_excel

PASTA=pathlib.Path(sys.argv[1]); AQUI=pathlib.Path(__file__).resolve().parent
arquivos=[pathlib.Path(a) for a in sys.argv[2:]] or sorted(AQUI.glob("[0-2][0-9]-*.xlsx"))
def valor_xml(v):
    """(atributo t, texto do <v>) para o valor calculado; None se não houver o que gravar."""
    if v is None: return None
    if isinstance(v,bool): return ('t="b"',"1" if v else "0")
    if isinstance(v,(int,float)): return ("",repr(float(v)) if isinstance(v,float) else str(v))
    if isinstance(v,(datetime.datetime,datetime.date)): return ("",repr(float(to_excel(v))))
    if isinstance(v,str):
        if v.startswith("#"): return ('t="e"',html.escape(v,quote=False))
        return ('t="str"',html.escape(v,quote=False))
    return None
CELL=re.compile(r'<c r="([A-Z]+\d+)"([^>]*)>(<f>.*?</f>|<f[^>]*/>)(?:<v>.*?</v>|<v/>)?</c>',re.S)
for arq in arquivos:
    rec=PASTA/arq.name
    vals=openpyxl.load_workbook(rec,data_only=True)
    wbf=openpyxl.load_workbook(arq,read_only=True)   # só para mapear sheetN.xml -> nome da aba
    nomes=wbf.sheetnames
    src=zipfile.ZipFile(arq); tmp=arq.with_suffix(".tmp")
    out=zipfile.ZipFile(tmp,"w",zipfile.ZIP_DEFLATED); n=0
    rels=src.read("xl/_rels/workbook.xml.rels").decode(); wbx=src.read("xl/workbook.xml").decode()
    rid={m.group(2):m.group(1) for m in re.finditer(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="([^"]+)"',wbx)}
    alvo={m.group(2):m.group(1) for m in re.finditer(r'Id="([^"]+)"[^>]*Target="/?(?:xl/)?(worksheets/[^"]+)"',rels)}
    alvo.update({m.group(2):m.group(1) for m in re.finditer(r'Target="/?(?:xl/)?(worksheets/[^"]+)"[^>]*Id="([^"]+)"',rels)})
    por_arquivo={"xl/"+alvo[r]:html.unescape(nome) for r,nome in rid.items() if r in alvo}
    for item in src.infolist():
        data=src.read(item.filename)
        if item.filename in por_arquivo:
            ws=vals[por_arquivo[item.filename]]; x=data.decode()
            def sub(m):
                global n
                vx=valor_xml(ws[m.group(1)].value)
                if vx is None: return m.group(0)
                t,txt=vx; attrs=re.sub(r'\st="[^"]*"','',m.group(2)); n+=1
                return f'<c r="{m.group(1)}"{attrs}{(" "+t) if t else ""}>{m.group(3)}<v>{txt}</v></c>'
            data=CELL.sub(sub,x).encode()
        out.writestr(item,data)
    out.close(); src.close(); shutil.move(tmp,arq); print(f"{arq.name}: {n} valores gravados")
