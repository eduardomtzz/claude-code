#!/usr/bin/env python3
"""Planilha 8 do Kit Completo: Funil de Propostas. Gera 08-funil-de-propostas.xlsx"""
from ssg import *
from openpyxl.chart import BarChart, Reference
from datetime import date
N=400; R0=5; RN=R0+N-1
ETAPAS=[("Contato",0.1),("Reunião feita",0.25),("Proposta enviada",0.5),("Negociação",0.75),("Ganha",1.0),("Perdida",0.0)]
wb=Workbook()
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. As etapas e probabilidades alimentam a previsão ponderada.",merge_to="H")
cfg["A4"]="Empresa ou vendedor"; cfg["B4"]="Prisma Comunicação (exemplo fictício)"
cfg["A5"]="Data de referência (hoje)"; cfg["B5"]=date(2026,9,14)
cfg["C5"]="O exemplo está congelado em 14/09/2026, para os arquivos do kit mostrarem a mesma foto. Ao usar com os seus dados, troque por =HOJE()."; nota(cfg["C5"])
cfg["A6"]="Meta de vendas no trimestre (R$)"; cfg["B6"]=180000
cfg["A7"]="Proposta parada há mais de (dias)"; cfg["B7"]=14
for r in range(4,8): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"],DATA); inp(cfg["B6"],BRL0); inp(cfg["B7"])
cfg["A9"]="Etapas do funil e probabilidade de fechar"; rotulo(cfg["A9"])
hdr(cfg,10,["Etapa","Probabilidade"])
for i,(e,pb) in enumerate(ETAPAS):
    cfg.cell(row=11+i,column=1,value=e).font=F(size=10,color=TINTA); cfg.cell(row=11+i,column=1).border=borda
    c=cfg.cell(row=11+i,column=2,value=pb); inp(c,PCT,center=True)
cfg["A18"]="Origens (até 8)"; rotulo(cfg["A18"])
for i in range(8): inp(cfg.cell(row=19+i,column=1))
cfg["D18"]="Motivos de perda (até 8)"; rotulo(cfg["D18"])
for i in range(8): inp(cfg.cell(row=19+i,column=4))
widths(cfg,(34,16,3,30)); cfg.sheet_view.showGridLines=False
# ---------- Propostas ----------
pr=wb.create_sheet("Propostas")
titulo(pr,"Propostas","Uma linha por oportunidade. Atualize a etapa conforme avança; ao fechar, marque Ganha ou Perdida e a data.",merge_to="N")
hdr(pr,4,["Cliente","Proposta","Origem","Responsável","Valor (R$)","Etapa","Data de entrada","Última movimentação","Fechamento previsto","Data de fechamento","Motivo (se perdida)","Probabilidade","Valor ponderado","Dias parada","Situação"])
HOJE="Config!$B$5"; PAR="Config!$B$7"
for r in range(R0,RN+1):
    for c in (1,2,3,4,5,6,7,8,9,10,11): inp(pr.cell(row=r,column=c))
    for c in (3,4,6,7,8,9,10): pr.cell(row=r,column=c).alignment=Alignment(horizontal="center")
    for c in (7,8,9,10): pr.cell(row=r,column=c).number_format=DATA
    pr.cell(row=r,column=5).number_format=BRL0
    pr.cell(row=r,column=12,value=f'=IF(F{r}="","",IFERROR(INDEX(Config!$B$11:$B$16,MATCH(F{r},Config!$A$11:$A$16,0)),0))'); calc(pr.cell(row=r,column=12),PCT)
    pr.cell(row=r,column=13,value=f'=IF(OR(F{r}="",E{r}=""),"",E{r}*L{r})'); calc(pr.cell(row=r,column=13),BRL0)
    pr.cell(row=r,column=14,value=f'=IF(OR(F{r}="",F{r}="Ganha",F{r}="Perdida"),"",{HOJE}-IF(H{r}="",G{r},H{r}))'); calc(pr.cell(row=r,column=14),"0")
    pr.cell(row=r,column=15,value=f'=IF(F{r}="","",IF(F{r}="Ganha","Fechada",IF(F{r}="Perdida","Perdida",IF(AND(I{r}<>"",I{r}<{HOJE}),"Fechamento vencido",IF(N{r}>{PAR},"Parada","Ativa")))))'); calc(pr.cell(row=r,column=15))
dvs=[lista("=Config!$A$19:$A$26",strict=True), lista("=Config!$A$11:$A$16"), lista("=Config!$D$19:$D$26",strict=True),
     DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True,showErrorMessage=True), DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True,showErrorMessage=True)]
for dv,rng in zip(dvs,[f"C{R0}:C{RN}",f"F{R0}:F{RN}",f"K{R0}:K{RN}",f"G{R0}:J{RN}",f"E{R0}:E{RN}"]): dv.add(rng); pr.add_data_validation(dv)
pr.conditional_formatting.add(f"A{R0}:O{RN}", FormulaRule(formula=[f'$O{R0}="Parada"'], fill=fill("FFF4CC")))
pr.conditional_formatting.add(f"A{R0}:O{RN}", FormulaRule(formula=[f'$O{R0}="Fechamento vencido"'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
pr.conditional_formatting.add(f"A{R0}:O{RN}", FormulaRule(formula=[f'$O{R0}="Fechada"'], font=F(color=VERDE_T,size=10,bold=True)))
pr.conditional_formatting.add(f"A{R0}:O{RN}", FormulaRule(formula=[f'$O{R0}="Perdida"'], font=F(color="8A86A0",size=10)))
widths(pr,(22,30,14,12,13,16,12,13,13,13,20,11,14,9,16)); pr.freeze_panes="C5"; pr.sheet_view.showGridLines=False; pr.auto_filter.ref=f"A4:O{RN}"
# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
p["A1"]='=Config!B4&" · funil em "&TEXT(DAY(Config!B5),"00")&"/"&TEXT(MONTH(Config!B5),"00")&"/"&YEAR(Config!B5)'; p["A1"].font=F(bold=True,size=16,color=UVA); p.merge_cells("A1:H1")
p["A2"]="Nada para preencher aqui: tudo vem de Propostas."; nota(p["A2"]); p.merge_cells("A2:H2")
PE=f"Propostas!$E${R0}:$E${RN}"; PF=f"Propostas!$F${R0}:$F${RN}"; PM=f"Propostas!$M${R0}:$M${RN}"; PO=f"Propostas!$O${R0}:$O${RN}"; PJ=f"Propostas!$J${R0}:$J${RN}"; PC=f"Propostas!$C${R0}:$C${RN}"; PK=f"Propostas!$K${R0}:$K${RN}"; PD=f"Propostas!$D${R0}:$D${RN}"; PA=f"Propostas!$A${R0}:$A${RN}"; PB=f"Propostas!$B${R0}:$B${RN}"; PN=f"Propostas!$N${R0}:$N${RN}"; PG=f"Propostas!$G${R0}:$G${RN}"
ABERTA=f'{PF},"<>Ganha",{PF},"<>Perdida",{PF},"<>"'
kpi(p,4,1,"Em aberto (R$)",f'=SUMIFS({PE},{ABERTA})',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,3,"Previsão ponderada",f'=SUMIFS({PM},{ABERTA})',SOL,UVA,fmt=BRL0)
kpi(p,4,5,"Ganho no trimestre",f'=SUMIFS({PE},{PF},"Ganha",{PJ},">="&DATE(YEAR({HOJE}),3*INT((MONTH({HOJE})-1)/3)+1,1),{PJ},"<"&DATE(YEAR({HOJE})+(3*INT((MONTH({HOJE})-1)/3)+4>12),IF(3*INT((MONTH({HOJE})-1)/3)+4>12,1,3*INT((MONTH({HOJE})-1)/3)+4),1))',VERDE,VERDE_T,fmt=BRL0)
kpi(p,4,7,"% da meta",f'=IFERROR(E5/Config!$B$6,0)',VERDE,VERDE_T,fmt="0%")
p["A7"]="Funil por etapa"; p["A7"].font=F(bold=True,size=13,color=UVA)
hdr(p,8,["Etapa","Propostas","Valor (R$)","Ponderado (R$)","Taxa de passagem","Barra"]); p.merge_cells("F8:H8")
for i,(e,pb) in enumerate(ETAPAS[:4]):
    r=9+i
    p.cell(row=r,column=1,value=e); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=COUNTIFS({PF},A{r})'); calc(p.cell(row=r,column=2))
    p.cell(row=r,column=3,value=f'=SUMIFS({PE},{PF},A{r})'); calc(p.cell(row=r,column=3),BRL0)
    p.cell(row=r,column=4,value=f'=SUMIFS({PM},{PF},A{r})'); calc(p.cell(row=r,column=4),BRL0)
    # taxa de passagem: quantas propostas já passaram desta etapa (estão adiante ou fechadas) sobre o total que chegou nela
    adiante="+".join([f'COUNTIFS({PF},"{x[0]}")' for x in ETAPAS[i+1:]])
    p.cell(row=r,column=5,value=f'=IFERROR(({adiante})/(B{r}+{adiante}),0)'); calc(p.cell(row=r,column=5),"0%")
    p.cell(row=r,column=6,value=f'=REPT("█",ROUND(IFERROR(C{r}/MAX($C$9:$C$12),0)*30,0))'); p.cell(row=r,column=6).font=F(size=10,color=LILAS); p.cell(row=r,column=6).border=borda; p.merge_cells(start_row=r,start_column=6,end_row=r,end_column=8)
p.cell(row=13,column=1,value="Ganhas (total)"); p.cell(row=13,column=2,value=f'=COUNTIFS({PF},"Ganha")'); p.cell(row=13,column=3,value=f'=SUMIFS({PE},{PF},"Ganha")')
p.cell(row=14,column=1,value="Perdidas (total)"); p.cell(row=14,column=2,value=f'=COUNTIFS({PF},"Perdida")'); p.cell(row=14,column=3,value=f'=SUMIFS({PE},{PF},"Perdida")')
p.cell(row=15,column=1,value="Taxa de conversão (ganhas ÷ fechadas)"); p.cell(row=15,column=2,value='=IFERROR(B13/(B13+B14),0)'); p.cell(row=15,column=2).number_format="0%"
p.cell(row=16,column=1,value="Ticket médio das ganhas"); p.cell(row=16,column=2,value='=IFERROR(C13/B13,0)'); p.cell(row=16,column=2).number_format=BRL0
for r in (13,14,15,16):
    calc(p.cell(row=r,column=1),center=False); calc(p.cell(row=r,column=2)); calc(p.cell(row=r,column=3),BRL0)
p["A18"]="O que mexer primeiro"; p["A18"].font=F(bold=True,size=13,color=UVA)
p["A19"]="Propostas abertas com fechamento vencido ou paradas há mais tempo, as de maior valor primeiro."; nota(p["A19"])
hdr(p,20,["#","Cliente","Proposta","Etapa","Valor","Dias parada","Situação","Responsável"])
# chave auxiliar em Propostas coluna P
for r in range(R0,RN+1):
    # chave INTEIRA: situação, dias parada, valor e linha sem sobreposição de casas. A
    # anterior somava valor/1E+6 com ROW()/1E+5 e colidia com R$ 10 de diferença em linhas
    # vizinhas, repetindo uma proposta e escondendo outra.
    # chave inteira: situação · dias (teto 500) · valor em CENTAVOS (teto R$ 999.999,99) · linha.
    # Máximo 3,5E+15, abaixo de 2^53. ROUND(valor,0) empatava 380,00 com 379,99 (rodada 4).
    pr.cell(row=r,column=16,value=f'=IF(OR(F{r}="",O{r}="Fechada",O{r}="Perdida"),0,IF(O{r}="Fechamento vencido",3,IF(O{r}="Parada",2,1))*1E+15+MIN(N(N{r}),500)*1E+12+MIN(ROUND(N(E{r})*100,0),99999999)*1E+4+({RN}+1-ROW()))'); pr.cell(row=r,column=16).font=F(color=CINZA,size=9)
pr.column_dimensions["P"].hidden=True
PP=f"Propostas!$P${R0}:$P${RN}"
for k in range(1,13):
    r=20+k; m=f'MATCH(LARGE({PP},{k}),{PP},0)'; g=f'LARGE({PP},{k})>0'
    p.cell(row=r,column=1,value=k)
    for col,rng in zip((2,3,4,5,6,7,8),(PA,PB,PF,PE,PN,PO,PD)):
        p.cell(row=r,column=col,value=f'=IFERROR(IF({g},INDEX({rng},{m}),""),"")')
    for col in range(1,9): calc(p.cell(row=r,column=col),center=(col not in (2,3)))
    p.cell(row=r,column=5).number_format=BRL0; p.cell(row=r,column=6).number_format="0"
p.conditional_formatting.add("A21:H32", FormulaRule(formula=['$G21="Fechamento vencido"'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
p.conditional_formatting.add("A21:H32", FormulaRule(formula=['$G21="Parada"'], fill=fill("FFF4CC")))
p["A34"]="Por origem e motivo de perda"; p["A34"].font=F(bold=True,size=13,color=UVA)
hdr(p,35,["Origem","Propostas","Ganhas","Conversão","Valor ganho"])
for i in range(8):
    r=36+i; src=f"Config!$A${19+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",COUNTIFS({PC},{src}))'); calc(p.cell(row=r,column=2))
    p.cell(row=r,column=3,value=f'=IF({src}="","",COUNTIFS({PC},{src},{PF},"Ganha"))'); calc(p.cell(row=r,column=3))
    p.cell(row=r,column=4,value=f'=IF({src}="","",IFERROR(C{r}/(C{r}+COUNTIFS({PC},{src},{PF},"Perdida")),0))'); calc(p.cell(row=r,column=4),"0%")
    p.cell(row=r,column=5,value=f'=IF({src}="","",SUMIFS({PE},{PC},{src},{PF},"Ganha"))'); calc(p.cell(row=r,column=5),BRL0)
hdr(p,35,["Motivo de perda","Perdidas","Valor perdido"],start=7)
for i in range(8):
    r=36+i; src=f"Config!$D${19+i}"
    p.cell(row=r,column=7,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=7),center=False)
    p.cell(row=r,column=8,value=f'=IF({src}="","",COUNTIFS({PK},{src},{PF},"Perdida"))'); calc(p.cell(row=r,column=8))
    p.cell(row=r,column=9,value=f'=IF({src}="","",SUMIFS({PE},{PK},{src},{PF},"Perdida"))'); calc(p.cell(row=r,column=9),BRL0)
bc=BarChart(); bc.type="bar"; bc.height=6; bc.width=14; bc.title="Valor por etapa"; bc.style=2
bc.add_data(Reference(p,min_col=3,min_row=8,max_row=12),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=9,max_row=12))
bc.series[0].graphicalProperties.solidFill="3B1F5E"; bc.legend=None; bc.x_axis.majorGridlines=None
p.add_chart(bc,"J4")
widths(p,(26,28,18,14,14,12,18,14,14)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
# exemplos
for i,v in enumerate(["Indicação","Instagram","Site","Evento","Cliente antigo"]): cfg.cell(row=19+i,column=1,value=v)
for i,v in enumerate(["Preço","Prazo","Escolheu concorrente","Adiou o projeto","Sem resposta"]): cfg.cell(row=19+i,column=4,value=v)
ex=[("Loja Verde","Site institucional","Indicação","Ana",24000,"Ganha",date(2026,7,2),date(2026,7,28),date(2026,7,30),date(2026,7,28),""),
    ("Bistrô 42","Campanha de fim de ano","Instagram","Carla",38000,"Ganha",date(2026,7,20),date(2026,8,25),date(2026,8,30),date(2026,8,25),""),
    ("Padaria do Sol","Identidade visual","Indicação","Bruno",9500,"Ganha",date(2026,6,10),date(2026,7,10),date(2026,7,15),date(2026,7,10),""),
    ("Horizonte","Relatório anual","Cliente antigo","Ana",16000,"Ganha",date(2026,8,12),date(2026,9,2),date(2026,9,5),date(2026,9,2),""),
    ("Aurora Móveis","Catálogo digital","Site","Ana",22000,"Negociação",date(2026,8,20),date(2026,9,9),date(2026,9,19),None,""),
    ("Clínica Bem-Estar","Gestão de redes (mensal)","Instagram","Carla",4500,"Proposta enviada",date(2026,8,28),date(2026,9,4),date(2026,9,12),None,""),
    ("Construtora Vale","Vídeo institucional","Evento","Bruno",45000,"Proposta enviada",date(2026,8,15),date(2026,8,26),date(2026,9,10),None,""),
    ("Escola Nova Era","Site + matrículas","Indicação","Ana",31000,"Reunião feita",date(2026,9,3),date(2026,9,8),date(2026,9,30),None,""),
    ("Dra. Lúcia (nutri)","Identidade + Instagram","Instagram","Bruno",7800,"Contato",date(2026,9,11),None,date(2026,10,10),None,""),
    ("Café Central","Cardápio e fachada","Cliente antigo","Carla",6200,"Negociação",date(2026,9,1),date(2026,9,12),date(2026,9,20),None,""),
    ("Tech Solutions","Landing page","Site","Diego",12000,"Perdida",date(2026,7,15),date(2026,8,5),date(2026,8,10),date(2026,8,5),"Preço"),
    ("Academia Forte","Campanha de matrícula","Evento","Carla",18000,"Perdida",date(2026,7,22),date(2026,8,20),date(2026,8,25),date(2026,8,20),"Escolheu concorrente"),
    ("Pet Shop Amigo","Redes sociais (mensal)","Instagram","Carla",3200,"Perdida",date(2026,8,1),date(2026,8,18),date(2026,8,30),date(2026,8,30),"Sem resposta"),
    ("Studio Yoga","Site simples","Site","Diego",8500,"Contato",date(2026,8,10),date(2026,8,12),date(2026,9,15),None,"")]
for i,row in enumerate(ex):
    for c,v in enumerate(row,start=1):
        if v is not None and v!="": pr.cell(row=R0+i,column=c,value=v)
como_usar(wb,"Funil de Propostas",[
 ("O que esta planilha faz","Você registra cada proposta com valor e etapa; ela calcula a previsão ponderada, mostra o funil por etapa, o que está parado ou vencido, a conversão por origem e por que você perde."),
 ("Passo 1","Em Config, preencha a meta do trimestre, o limite de \"parada\" (14 dias é um bom padrão), as origens e os motivos de perda. Ajuste as probabilidades por etapa se a sua experiência for diferente."),
 ("Passo 2","Em Propostas, uma linha por oportunidade: cliente, proposta, origem, responsável, valor, etapa e datas. Sempre que mexer, atualize a última movimentação."),
 ("Passo 3","Ao fechar, mude a etapa para Ganha ou Perdida, preencha a data e, se perdida, o motivo."),
 ("Passo 4","Em Painel, veja o funil, a previsão ponderada, a lista do que mexer primeiro e a conversão por origem."),
 ("Rotina","Sexta-feira, 15 minutos: atualizar etapas e datas, cobrar as paradas. Dia 1 do mês: comparar com a meta."),
 ("Com a IA","Copie \"O que mexer primeiro\" e use \"Escrever 02: e-mail curto que pede algo\" para retomar propostas paradas; copie \"Por origem e motivo de perda\" e use \"Analisar 01: o que esses números dizem\"."),
])
proteger(wb); salvar(wb,"08-funil-de-propostas.xlsx","Funil de Propostas · Kit IA no Trabalho Completo")
