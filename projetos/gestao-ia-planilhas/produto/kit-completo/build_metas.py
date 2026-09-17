#!/usr/bin/env python3
"""Planilha 6 do Kit Completo: Metas do Trimestre. Gera 06-metas-do-trimestre.xlsx"""
from ssg import *
from openpyxl.chart import BarChart, Reference
from datetime import date
NO=5; NK=4  # objetivos × resultados-chave
wb=Workbook()
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche.",merge_to="F")
cfg["A4"]="Equipe ou empresa"; cfg["B4"]="Prisma Comunicação (exemplo fictício)"
cfg["A5"]="Trimestre"; cfg["B5"]="3º trimestre de 2026"
cfg["A6"]="Início do trimestre"; cfg["B6"]=date(2026,7,1)
cfg["A7"]="Fim do trimestre"; cfg["B7"]=date(2026,9,30)
cfg["A8"]="Data de referência (hoje)"; cfg["B8"]="=TODAY()"
cfg["A9"]="Semana atual do trimestre"; cfg["B9"]='=IF(B8<B6,0,MIN(13,INT((B8-B6)/7)+1))'
cfg["A10"]="% do trimestre decorrido"; cfg["B10"]='=MAX(0,MIN(1,(B8-B6)/(B7-B6)))'
for r in range(4,11): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"]); inp(cfg["B6"],DATA); inp(cfg["B7"],DATA); inp(cfg["B8"],DATA); calc(cfg["B9"]); calc(cfg["B10"],PCT)
cfg["C8"]="Para simular o fim do trimestre, troque por uma data."; nota(cfg["C8"])
widths(cfg,(30,28,50)); cfg.sheet_view.showGridLines=False
# ---------- Metas ----------
m=wb.create_sheet("Metas")
titulo(m,"Metas do trimestre","Até 5 objetivos com até 4 resultados-chave cada. Preencha o amarelo; progresso e semáforo são calculados. Atualize o valor atual toda semana.",merge_to="M")
hdr(m,4,["Objetivo","Resultado-chave","Dono","Unidade","Ponto de partida","Meta","Valor atual","Progresso","Esperado até hoje","Semáforo","Sentido","Observação"])
HOJE="Config!$B$8"; DEC="Config!$B$10"
rows=[]
for o in range(NO):
    for k in range(NK):
        r=5+o*NK+k; rows.append(r)
        if k==0:
            inp(m.cell(row=r,column=1)); m.merge_cells(start_row=r,start_column=1,end_row=r+NK-1,end_column=1); m.cell(row=r,column=1).alignment=Alignment(wrap_text=True,vertical="top")
        for c in (2,3,4,5,6,7,11,12): inp(m.cell(row=r,column=c))
        for c in (3,4,5,6,7,11): m.cell(row=r,column=c).alignment=Alignment(horizontal="center")
        m.cell(row=r,column=8,value=f'=IF(B{r}="","",IF(F{r}=E{r},"",IF(K{r}="Menor é melhor",MAX(0,MIN(1,(E{r}-G{r})/(E{r}-F{r}))),MAX(0,MIN(1,(G{r}-E{r})/(F{r}-E{r}))))))'); calc(m.cell(row=r,column=8),PCT)
        m.cell(row=r,column=9,value=f'=IF(B{r}="","",{DEC})'); calc(m.cell(row=r,column=9),PCT)
        m.cell(row=r,column=10,value=f'=IF(OR(B{r}="",H{r}=""),"",IF(H{r}>=1,"Atingido",IF(H{r}>=I{r}-0.1,"No ritmo",IF(H{r}>=I{r}-0.25,"Atenção","Em risco"))))'); calc(m.cell(row=r,column=10))
        m.cell(row=r,column=12).alignment=Alignment(wrap_text=True)
dvsent=lista('"Maior é melhor,Menor é melhor"'); dvsent.add(f"K5:K{4+NO*NK}"); m.add_data_validation(dvsent)
dvnum=DataValidation(type="decimal",allow_blank=True,showErrorMessage=True); dvnum.add(f"E5:G{4+NO*NK}"); m.add_data_validation(dvnum)
RNG=f"A5:L{4+NO*NK}"
m.conditional_formatting.add(f"J5:J{4+NO*NK}", FormulaRule(formula=['J5="Atingido"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10,bold=True)))
m.conditional_formatting.add(f"J5:J{4+NO*NK}", FormulaRule(formula=['J5="No ritmo"'], fill=fill("E6F4EA"), font=F(color=VERDE_T,size=10)))
m.conditional_formatting.add(f"J5:J{4+NO*NK}", FormulaRule(formula=['J5="Atenção"'], fill=fill("FFF4CC"), font=F(color="7A5200",size=10)))
m.conditional_formatting.add(f"J5:J{4+NO*NK}", FormulaRule(formula=['J5="Em risco"'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
for o in range(NO): m.row_dimensions[5+o*NK].height=18
widths(m,(28,40,12,10,12,10,11,11,12,11,16,30)); m.freeze_panes="C5"; m.sheet_view.showGridLines=False
# exemplos
ex=[("Crescer a receita recorrente",[("Fechar 3 contratos mensais novos","Ana","contratos",0,3,2,"Maior é melhor"),("Receita recorrente mensal de R$ 30 mil","Ana","R$",18000,30000,27400,"Maior é melhor"),("Propostas enviadas no trimestre","Carla","propostas",0,24,24,"Maior é melhor")]),
    ("Entregar no prazo",[("Etapas atrasadas por semana (média)","Ana","etapas",4,1,3,"Menor é melhor"),("Projetos entregues na data combinada","Bruno","%",60,90,82,"Maior é melhor")]),
    ("Reduzir retrabalho",[("Rodadas de revisão por peça (média)","Bruno","rodadas",3.2,2,2.8,"Menor é melhor"),("Peças aprovadas de primeira","Carla","%",40,65,60,"Maior é melhor")]),
    ("Organizar a casa",[("Relatório mensal enviado até o dia 5","Ana","meses",0,3,2,"Maior é melhor"),("Horas extras da equipe no mês","Ana","horas",20,8,14,"Menor é melhor")])]
for o,(obj,krs) in enumerate(ex):
    m.cell(row=5+o*NK,column=1,value=obj)
    for k,(kr,dono,un,ini,meta,atual,sent) in enumerate(krs):
        r=5+o*NK+k
        for c,v in zip((2,3,4,5,6,7,11),(kr,dono,un,ini,meta,atual,sent)): m.cell(row=r,column=c,value=v)
# ---------- Semanas ----------
w=wb.create_sheet("Semanas")
titulo(w,"Acompanhamento semanal","Toda semana, copie o valor atual de cada resultado-chave para a coluna da semana. A linha fica com o histórico do trimestre.",merge_to="P")
hdr(w,4,["Resultado-chave","Meta"]+[f"S{i}" for i in range(1,14)])
for i,r in enumerate(rows):
    rr=5+i
    w.cell(row=rr,column=1,value=f'=IF(Metas!B{r}="","",Metas!B{r})'); calc(w.cell(row=rr,column=1),center=False)
    w.cell(row=rr,column=2,value=f'=IF(Metas!B{r}="","",Metas!F{r})'); calc(w.cell(row=rr,column=2))
    for s_ in range(13): inp(w.cell(row=rr,column=3+s_),center=True)
# exemplo: histórico S1..S11 (painel em 13/09/2026 = semana 11 de 13). A S11 é igual ao "Valor atual" da aba Metas.
# Receita recorrente: R$ 18 mil em julho (S1-S4), 22,5 mil em agosto (S5-S9), 27,4 mil em setembro (S10-S11), como no modelo de slides 16.
hist={5:[0,0,0,0,1,1,1,1,1,2,2],                                   # contratos mensais novos
      6:[18000,18000,18000,18000,22500,22500,22500,22500,22500,27400,27400],  # receita recorrente mensal
      7:[2,4,7,9,11,13,15,18,20,22,24],                             # propostas enviadas (acumulado)
      9:[4,4,5,4,3,4,3,3,4,3,3],                                    # etapas atrasadas por semana
      10:[60,60,67,67,75,75,80,80,78,82,82],                         # projetos entregues na data (%)
      13:[3.2,3.1,3.1,3.0,3.0,2.9,2.9,2.9,2.8,2.8,2.8],              # rodadas de revisão por peça
      14:[40,42,45,48,50,52,55,56,58,60,60],                         # peças aprovadas de primeira (%)
      17:[0,0,0,0,0,1,1,1,1,2,2],                                    # relatório mensal enviado até o dia 5
      18:[20,20,20,18,17,16,16,15,15,14,14]}                         # horas extras no mês
for r,vals in hist.items():
    assert vals[-1]==m.cell(row=r,column=7).value, (r,vals[-1],m.cell(row=r,column=7).value)
    for s_,v in enumerate(vals): w.cell(row=r,column=3+s_,value=v)
w.conditional_formatting.add(f"C4:O4", FormulaRule(formula=['COLUMN()-2=Config!$B$9'], fill=fill(SOL), font=F(color=UVA,size=10,bold=True)))
widths(w,[40,10]+[7]*13); w.freeze_panes="C5"; w.sheet_view.showGridLines=False
# ---------- Meses ----------
# A oferta promete acompanhamento mensal. O fechamento de cada mês é a última semana
# preenchida daquele mês (um trimestre tem 13 semanas: 1-4, 5-9, 10-13).
mz=wb.create_sheet("Meses")
titulo(mz,"Acompanhamento mensal","Fechamento de cada mês do trimestre: o último valor lançado nas semanas daquele mês. Nada para preencher aqui — vem da aba Semanas. Serve para a reunião do início do mês.",merge_to="H")
hdr(mz,4,["Resultado-chave","Meta","Mês 1","Mês 2","Mês 3","Ganho no trimestre","Falta para a meta"])
FAIXAS=[("C","F"),("G","K"),("L","O")]   # S1-S4, S5-S9, S10-S13
for i,r in enumerate(rows):
    rr=5+i
    mz.cell(row=rr,column=1,value=f'=IF(Metas!B{r}="","",Metas!B{r})'); calc(mz.cell(row=rr,column=1),center=False)
    mz.cell(row=rr,column=2,value=f'=IF(Metas!B{r}="","",Metas!F{r})'); calc(mz.cell(row=rr,column=2))
    for j,(c0,c1) in enumerate(FAIXAS):
        # último valor preenchido da faixa: LOOKUP com critério sempre verdadeiro pega o
        # último número da linha, sem precisar de função nova
        mz.cell(row=rr,column=3+j,value=f'=IF(Metas!B{r}="","",IFERROR(LOOKUP(9.99E+307,Semanas!${c0}{rr}:${c1}{rr}),""))')
        calc(mz.cell(row=rr,column=3+j))
    mz.cell(row=rr,column=6,value=f'=IF(OR(Metas!B{r}="",C{rr}="",E{rr}=""),"",E{rr}-IFERROR(Metas!E{r},0))'); calc(mz.cell(row=rr,column=6))
    mz.cell(row=rr,column=7,value=f'=IF(OR(Metas!B{r}="",E{rr}=""),"",Metas!F{r}-E{rr})'); calc(mz.cell(row=rr,column=7))
mz.cell(row=5+len(rows)+1,column=1,value="Mês 1 = semanas 1 a 4; Mês 2 = semanas 5 a 9; Mês 3 = semanas 10 a 13. Mês sem nenhuma semana preenchida fica em branco. \"Ganho no trimestre\" compara o fechamento do Mês 3 com o ponto de partida da aba Metas.").font=F(size=9,color=LILAS)
mz.merge_cells(start_row=5+len(rows)+1,start_column=1,end_row=5+len(rows)+1,end_column=7)
widths(mz,(40,12,14,14,14,18,18)); mz.freeze_panes="C5"; mz.sheet_view.showGridLines=False
# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
p["A1"]='=Config!B4&" · "&Config!B5'; p["A1"].font=F(bold=True,size=16,color=UVA); p.merge_cells("A1:H1")
p["A2"]='="Semana "&Config!B9&" de 13 · "&TEXT(Config!B10,"0%")&" do trimestre decorrido. Nada para preencher aqui."'; nota(p["A2"]); p.merge_cells("A2:H2")
MJ=f"Metas!$J$5:$J${4+NO*NK}"; MB=f"Metas!$B$5:$B${4+NO*NK}"; MH=f"Metas!$H$5:$H${4+NO*NK}"
kpi(p,4,1,"Resultados-chave",f'=COUNTIFS({MB},"<>")',LAVANDA,UVA)
kpi(p,4,3,"Atingidos",f'=COUNTIFS({MJ},"Atingido")',VERDE,VERDE_T)
kpi(p,4,5,"Em risco",f'=COUNTIFS({MJ},"Em risco")',VERM,VERM_T)
kpi(p,4,7,"Progresso médio",f'=IFERROR(AVERAGEIFS({MH},{MB},"<>"),0)',SOL,UVA,fmt="0%")
p["A7"]="Por objetivo"; p["A7"].font=F(bold=True,size=13,color=UVA)
hdr(p,8,["Objetivo","Resultados-chave","Progresso médio","Atingidos","Em risco","Barra"]); p.merge_cells("F8:H8")
for o in range(NO):
    r=9+o; a=5+o*NK; b=a+NK-1; src=f"Metas!$A${a}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",COUNTIFS(Metas!$B${a}:$B${b},"<>"))'); calc(p.cell(row=r,column=2))
    p.cell(row=r,column=3,value=f'=IF(OR({src}="",B{r}=0),"",AVERAGEIFS(Metas!$H${a}:$H${b},Metas!$B${a}:$B${b},"<>"))'); calc(p.cell(row=r,column=3),PCT)
    p.cell(row=r,column=4,value=f'=IF({src}="","",COUNTIFS(Metas!$J${a}:$J${b},"Atingido"))'); calc(p.cell(row=r,column=4))
    p.cell(row=r,column=5,value=f'=IF({src}="","",COUNTIFS(Metas!$J${a}:$J${b},"Em risco"))'); calc(p.cell(row=r,column=5))
    p.cell(row=r,column=6,value=f'=IF(C{r}="","",REPT("█",ROUND(C{r}*30,0)))'); p.cell(row=r,column=6).font=F(size=10,color=LILAS); p.cell(row=r,column=6).border=borda; p.merge_cells(start_row=r,start_column=6,end_row=r,end_column=8)
p["A15"]="Todos os resultados-chave"; p["A15"].font=F(bold=True,size=13,color=UVA)
hdr(p,16,["Resultado-chave","Dono","Atual","Meta","Progresso","Esperado","Semáforo","Barra"])
for i,r in enumerate(rows):
    rr=17+i
    p.cell(row=rr,column=1,value=f'=IF(Metas!B{r}="","",Metas!B{r})'); calc(p.cell(row=rr,column=1),center=False)
    p.cell(row=rr,column=2,value=f'=IF(Metas!B{r}="","",Metas!C{r})'); calc(p.cell(row=rr,column=2))
    p.cell(row=rr,column=3,value=f'=IF(Metas!B{r}="","",Metas!G{r})'); calc(p.cell(row=rr,column=3),"#,##0.##")
    p.cell(row=rr,column=4,value=f'=IF(Metas!B{r}="","",Metas!F{r})'); calc(p.cell(row=rr,column=4),"#,##0.##")
    p.cell(row=rr,column=5,value=f'=IF(Metas!B{r}="","",Metas!H{r})'); calc(p.cell(row=rr,column=5),PCT)
    p.cell(row=rr,column=6,value=f'=IF(Metas!B{r}="","",Metas!I{r})'); calc(p.cell(row=rr,column=6),PCT)
    p.cell(row=rr,column=7,value=f'=IF(Metas!B{r}="","",Metas!J{r})'); calc(p.cell(row=rr,column=7))
    p.cell(row=rr,column=8,value=f'=IF(E{rr}="","",REPT("█",ROUND(E{rr}*20,0)))'); p.cell(row=rr,column=8).font=F(size=10,color=LILAS); p.cell(row=rr,column=8).border=borda
R1=17; R2=16+NO*NK
for cor,txt,fnt in ((VERDE,"Atingido",VERDE_T),("E6F4EA","No ritmo",VERDE_T),("FFF4CC","Atenção","7A5200"),(VERM,"Em risco",VERM_T)):
    p.conditional_formatting.add(f"G{R1}:G{R2}", FormulaRule(formula=[f'G{R1}="{txt}"'], fill=fill(cor), font=F(color=fnt,size=10,bold=(txt in ("Atingido","Em risco")))))
widths(p,(40,10,10,10,11,10,11,24)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
como_usar(wb,"Metas do Trimestre",[
 ("O que esta planilha faz","Você define até 5 objetivos com até 4 resultados-chave cada (ponto de partida, meta, valor atual). Ela calcula o progresso, compara com o tempo já decorrido e acende o semáforo: atingido, no ritmo, atenção ou em risco."),
 ("Passo 1","Em Config, preencha o trimestre, as datas de início e fim e deixe a data de referência em =HOJE()."),
 ("Passo 2","Em Metas, escreva cada objetivo e seus resultados-chave. Ponto de partida é o valor no dia 1; meta é onde quer chegar; valor atual é o número de hoje. Diga se maior ou menor é melhor."),
 ("Passo 3","Toda semana, atualize o valor atual e copie para a coluna da semana em Semanas. O Painel mostra o resumo por objetivo e a lista completa com semáforo; a aba Meses fecha cada mês do trimestre a partir das semanas, para a reunião mensal."),
 ("Rotina","Sexta-feira, 10 minutos: atualizar os valores. Primeira segunda do mês: reunião de 20 minutos olhando o Painel."),
 ("Com a IA","Copie a tabela \"Todos os resultados-chave\" e use o prompt \"Analisar 04: meta realista\" ou \"Escrever 04: justificativa para chefe ou cliente\". No fim do trimestre, \"Apresentar 01: roteiro de 8 slides\"."),
])
proteger(wb); salvar(wb,"06-metas-do-trimestre.xlsx","Metas do Trimestre · Kit IA no Trabalho Completo")
