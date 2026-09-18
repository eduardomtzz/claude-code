#!/usr/bin/env python3
"""Planilha 9 do Kit Completo: Horas e Custo por Projeto. Gera 09-horas-e-custo-por-projeto.xlsx"""
from ssg import *
from openpyxl.chart import BarChart, Reference
from datetime import date, timedelta
import random
N=800; R0=5; RN=R0+N-1; NPES=10; NPROJ=12
wb=Workbook()
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Custo-hora = custo mensal da pessoa ÷ horas disponíveis no mês.",merge_to="H")
cfg["A4"]="Empresa ou equipe"; cfg["B4"]="Prisma Comunicação (exemplo fictício)"
cfg["A5"]="Ano"; cfg["B5"]=2026
cfg["A6"]="Mês do painel"; cfg["B6"]="Setembro"
cfg["A7"]="Número do mês"; cfg["B7"]="=MATCH(B6,$H$5:$H$16,0)"
for r in range(4,8): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"]); inp(cfg["B6"],center=True); calc(cfg["B7"])
cfg["H4"]="Meses"; rotulo(cfg["H4"])
for i,m_ in enumerate(MESES): cfg.cell(row=5+i,column=8,value=m_).font=F(size=10,color=TINTA)
dv=lista("=Config!$H$5:$H$16"); dv.add("B6"); cfg.add_data_validation(dv)
cfg["A9"]="Pessoas (até 10)"; rotulo(cfg["A9"])
hdr(cfg,10,["Pessoa","Custo mensal (R$)","Horas disponíveis/mês","Custo-hora (R$)"])
for i in range(NPES):
    r=11+i
    inp(cfg.cell(row=r,column=1)); inp(cfg.cell(row=r,column=2),BRL0,center=True); inp(cfg.cell(row=r,column=3),center=True)
    cfg.cell(row=r,column=4,value=f'=IF(A{r}="","",IF(OR(B{r}="",N(B{r})<=0),"falta o custo mensal",IF(OR(C{r}="",N(C{r})<=0),"faltam as horas disponíveis",B{r}/C{r})))'); calc(cfg.cell(row=r,column=4),BRL)
cfg["A22"]="Projetos (até 12)"; rotulo(cfg["A22"])
hdr(cfg,23,["Projeto","Cliente","Horas orçadas","Valor cobrado (R$)","Status"])
for i in range(NPROJ):
    r=24+i
    for c in (1,2,3,4,5): inp(cfg.cell(row=r,column=c))
    cfg.cell(row=r,column=3).alignment=Alignment(horizontal="center"); cfg.cell(row=r,column=4).number_format=BRL0; cfg.cell(row=r,column=5).alignment=Alignment(horizontal="center")
dvst=lista('"Em andamento,Concluído,Interno,Proposta"'); dvst.add(f"E24:E{23+NPROJ}"); cfg.add_data_validation(dvst)
cfg.cell(row=24+NPROJ+1,column=1,value="Custo mensal: salário + encargos, ou o valor que a pessoa cobra por mês. Horas disponíveis: 160 para tempo integral, menos reuniões e folgas (140 é realista). Status Proposta: projeto ainda não fechado; entra na lista para você já orçar, mas sem horas lançadas.").font=F(size=9,color=LILAS)
widths(cfg,(30,18,20,16,14,3,3,12)); cfg.sheet_view.showGridLines=False
# ---------- Horas ----------
h=wb.create_sheet("Horas")
titulo(h,"Lançamento de horas","Uma linha por pessoa, por dia, por projeto. Pessoa e projeto vêm de listas; custo e mês são calculados.",merge_to="I")
hdr(h,4,["Data","Pessoa","Projeto","Atividade","Horas","Custo (R$)","Mês","Ano","Faturável?"])
for r in range(R0,RN+1):
    for c in (1,2,3,4,5,9): inp(h.cell(row=r,column=c))
    for c in (1,2,3,5,9): h.cell(row=r,column=c).alignment=Alignment(horizontal="center")
    h.cell(row=r,column=1).number_format=DATA; h.cell(row=r,column=5).number_format="0.0"
    # Se o custo-hora da pessoa não é número (custo mensal ou horas em branco em Config),
    # o custo da hora NÃO vira zero: vira recado. Zero fazia o projeto parecer mais
    # lucrativo do que é (achado G-5 da auditoria de 18/09).
    _ch=f'INDEX(Config!$D$11:$D${10+NPES},MATCH(B{r},Config!$A$11:$A${10+NPES},0))'
    h.cell(row=r,column=6,value=f'=IF(OR(B{r}="",E{r}=""),"",'
        f'IFERROR(IF(ISNUMBER({_ch}),E{r}*{_ch},"cadastro de custo incompleto em Config"),'
        f'"pessoa sem custo-hora em Config"))'); calc(h.cell(row=r,column=6),BRL)
    h.cell(row=r,column=7,value=f'=IF(A{r}="","",MONTH(A{r}))'); calc(h.cell(row=r,column=7))
    h.cell(row=r,column=8,value=f'=IF(A{r}="","",YEAR(A{r}))'); calc(h.cell(row=r,column=8))
    # J (oculta): 1 = hora lançada sem custo-hora (F virou recado). É por esta coluna que
    # o Painel conta e propaga a incompletude, sem depender de curinga em COUNTIF.
    h.cell(row=r,column=10,value=f'=IF(AND(B{r}<>"",E{r}<>"",NOT(ISNUMBER(F{r}))),1,0)'); h.cell(row=r,column=10).font=F(color=CINZA,size=9)
# strict=True: pessoa ou projeto digitado fora do cadastro é recusado na hora. Antes
# passava, o MATCH falhava e o IFERROR devolvia custo zero — hora trabalhada virava
# trabalho de graça, sem nenhum aviso.
dvs=[lista(f"=Config!$A$11:$A${10+NPES}",strict=True), lista(f"=Config!$A$24:$A${23+NPROJ}",strict=True), lista('"Sim,Não"'),
     DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True,showErrorMessage=True), DataValidation(type="decimal",operator="between",formula1="0",formula2="24",allow_blank=True,showErrorMessage=True)]
for dv,rng in zip(dvs,[f"B{R0}:B{RN}",f"C{R0}:C{RN}",f"I{R0}:I{RN}",f"A{R0}:A{RN}",f"E{R0}:E{RN}"]): dv.add(rng); h.add_data_validation(dv)
widths(h,(12,16,30,32,8,13,6,7,11)); h.column_dimensions["J"].hidden=True; h.freeze_panes="A5"; h.sheet_view.showGridLines=False; h.auto_filter.ref=f"A4:I{RN}"
# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
p["A1"]='=Config!B4&" · "&Config!B6&" de "&Config!B5'; p["A1"].font=F(bold=True,size=16,color=UVA); p.merge_cells("A1:I1")
p["A2"]="Nada para preencher aqui. Escolha o mês em Config; tudo vem de Horas."; nota(p["A2"]); p.merge_cells("A2:I2")
M="Config!$B$7"; Y="Config!$B$5"
HB=f"Horas!$B${R0}:$B${RN}"; HC=f"Horas!$C${R0}:$C${RN}"; HE=f"Horas!$E${R0}:$E${RN}"; HF=f"Horas!$F${R0}:$F${RN}"; HG=f"Horas!$G${R0}:$G${RN}"; HH=f"Horas!$H${R0}:$H${RN}"; HI=f"Horas!$I${R0}:$I${RN}"; HJ=f"Horas!$J${R0}:$J${RN}"
INC="cadastro incompleto"   # texto que ocupa o lugar de custo e margem enquanto falta custo-hora
kpi(p,4,1,"Horas no mês",f'=SUMIFS({HE},{HG},{M},{HH},{Y})',LAVANDA,UVA,fmt="#,##0.0")
# Custo com lançamento sem custo-hora não é número parcial: é pendência escrita (rodada 4:
# apagar um custo mensal derrubava o custo do mês de 14.967 para 10.207 e a margem do
# projeto SUBIA, sem nada dizer que a soma estava incompleta).
kpi(p,4,3,"Custo no mês",f'=IF(SUMIFS({HJ},{HG},{M},{HH},{Y})>0,"{INC}",SUMIFS({HF},{HG},{M},{HH},{Y}))',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,5,"Horas faturáveis",f'=IFERROR(SUMIFS({HE},{HG},{M},{HH},{Y},{HI},"Sim")/A5,0)',VERDE,VERDE_T,fmt="0%")
kpi(p,4,7,"Ocupação da equipe",f'=IFERROR(A5/SUMIFS(Config!$C$11:$C${10+NPES},Config!$A$11:$A${10+NPES},"<>"),0)',SOL,UVA,fmt="0%")
p["A7"]=f'=IF(SUM({HJ})=0,"","Atenção: "&SUM({HJ})&" lançamento(s) de hora sem custo-hora (cadastro incompleto em Config). O custo desses lançamentos não está contado: o custo real é MAIOR e a margem real é MENOR do que qualquer soma parcial, por isso custo e margem dos projetos e pessoas afetados aparecem como \'{INC}\' até você completar o custo mensal e as horas disponíveis da pessoa.")'
p["A7"].font=F(size=10,bold=True,color=VERM_T); p.merge_cells("A7:J7"); p["A7"].alignment=Alignment(wrap_text=True,vertical="top")
p["A8"]="Por projeto (todo o período)"; p["A8"].font=F(bold=True,size=13,color=UVA)
hdr(p,9,["Projeto","Cliente","Horas orçadas","Horas usadas","% do orçado","Custo (R$)","Valor cobrado","Margem (R$)","Margem (%)","Situação"])
for i in range(NPROJ):
    r=10+i; src=f"Config!$A${24+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",Config!$B${24+i})'); calc(p.cell(row=r,column=2),center=False)
    p.cell(row=r,column=3,value=f'=IF({src}="","",Config!$C${24+i})'); calc(p.cell(row=r,column=3),"0")
    p.cell(row=r,column=4,value=f'=IF({src}="","",SUMIFS({HE},{HC},{src}))'); calc(p.cell(row=r,column=4),"0.0")
    p.cell(row=r,column=5,value=f'=IF(OR({src}="",C{r}=0,C{r}=""),"",D{r}/C{r})'); calc(p.cell(row=r,column=5),"0%")
    p.cell(row=r,column=6,value=f'=IF({src}="","",IF(SUMIFS({HJ},{HC},{src})>0,"{INC}",SUMIFS({HF},{HC},{src})))'); calc(p.cell(row=r,column=6),BRL0)
    p.cell(row=r,column=7,value=f'=IF({src}="","",Config!$D${24+i})'); calc(p.cell(row=r,column=7),BRL0)
    p.cell(row=r,column=8,value=f'=IF(OR({src}="",G{r}=""),"",IF(ISNUMBER(F{r}),G{r}-F{r},"{INC}"))'); calc(p.cell(row=r,column=8),BRL0)
    p.cell(row=r,column=9,value=f'=IF(OR({src}="",G{r}="",G{r}=0,NOT(ISNUMBER(H{r}))),"",H{r}/G{r})'); calc(p.cell(row=r,column=9),"0%")
    p.cell(row=r,column=10,value=f'=IF({src}="","",IF(Config!$E${24+i}="Interno","Interno",IF(Config!$E${24+i}="Proposta","Proposta",IF(NOT(ISNUMBER(F{r})),"Custo incompleto",IF(AND(E{r}<>"",E{r}>1),"Estourou as horas",IF(AND(I{r}<>"",I{r}<0.2),"Margem baixa",IF(AND(E{r}<>"",E{r}>0.85,Config!$E${24+i}="Em andamento"),"Perto do limite","Saudável")))))))'); calc(p.cell(row=r,column=10))
# (a regra vermelha comparava J9 com "Margem baixa" na linha de J10: pintava a linha de baixo)
p.conditional_formatting.add(f"J10:J{9+NPROJ}", FormulaRule(formula=['OR(J10="Estourou as horas",J10="Margem baixa",J10="Custo incompleto")'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
p.conditional_formatting.add(f"J10:J{9+NPROJ}", FormulaRule(formula=['J10="Perto do limite"'], fill=fill("FFF4CC")))
p.conditional_formatting.add(f"J10:J{9+NPROJ}", FormulaRule(formula=['J10="Saudável"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
p.conditional_formatting.add(f"H10:H{9+NPROJ}", FormulaRule(formula=['AND(ISNUMBER(H10),H10<0)'], font=F(color="C8402E",size=10,bold=True)))
r0=11+NPROJ
p.cell(row=r0,column=1,value="Por pessoa no mês").font=F(bold=True,size=13,color=UVA)
hdr(p,r0+1,["Pessoa","Horas lançadas","Disponíveis","Ocupação","Faturáveis","Custo no mês","Custo-hora"])
for i in range(NPES):
    r=r0+2+i; src=f"Config!$A${11+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",SUMIFS({HE},{HB},{src},{HG},{M},{HH},{Y}))'); calc(p.cell(row=r,column=2),"0.0")
    p.cell(row=r,column=3,value=f'=IF({src}="","",Config!$C${11+i})'); calc(p.cell(row=r,column=3),"0")
    p.cell(row=r,column=4,value=f'=IF(OR({src}="",C{r}=0,C{r}=""),"",B{r}/C{r})'); calc(p.cell(row=r,column=4),"0%")
    p.cell(row=r,column=5,value=f'=IF(OR({src}="",B{r}=0),"",SUMIFS({HE},{HB},{src},{HG},{M},{HH},{Y},{HI},"Sim")/B{r})'); calc(p.cell(row=r,column=5),"0%")
    p.cell(row=r,column=6,value=f'=IF({src}="","",IF(SUMIFS({HJ},{HB},{src},{HG},{M},{HH},{Y})>0,"{INC}",SUMIFS({HF},{HB},{src},{HG},{M},{HH},{Y})))'); calc(p.cell(row=r,column=6),BRL0)
    p.cell(row=r,column=7,value=f'=IF({src}="","",Config!$D${11+i})'); calc(p.cell(row=r,column=7),BRL)
p.conditional_formatting.add(f"D{r0+2}:D{r0+11}", FormulaRule(formula=[f'AND(ISNUMBER(D{r0+2}),D{r0+2}>1)'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
p.conditional_formatting.add(f"D{r0+2}:D{r0+11}", FormulaRule(formula=[f'AND(ISNUMBER(D{r0+2}),D{r0+2}<0.6)'], fill=fill("FFF4CC")))
bc=BarChart(); bc.type="bar"; bc.height=7; bc.width=14; bc.title="Horas usadas × orçadas por projeto"; bc.style=2
bc.add_data(Reference(p,min_col=3,max_col=4,min_row=9,max_row=9+NPROJ),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=10,max_row=9+NPROJ))
bc.series[0].graphicalProperties.solidFill="B89BE0"; bc.series[1].graphicalProperties.solidFill="3B1F5E"; bc.legend.position="b"; bc.x_axis.majorGridlines=None
p.add_chart(bc,f"I{r0}")
widths(p,(30,20,12,12,11,13,14,13,11,17)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
# exemplos
pes=[("Ana",9800,140),("Bruno",7200,140),("Carla",6500,140),("Diego",5800,120)]
for i,(a,b,c) in enumerate(pes):
    for k,v in enumerate((a,b,c),start=1): cfg.cell(row=11+i,column=k,value=v)
proj=[("Site novo · Loja Verde","Loja Verde",560,24000,"Em andamento"),("Campanha de fim de ano · Bistrô 42","Bistrô 42",300,38000,"Em andamento"),
      ("Relatório anual · Horizonte","Horizonte",130,16000,"Em andamento"),("Identidade visual · Padaria do Sol","Padaria do Sol",70,9500,"Concluído"),
      ("Redes sociais · Clínica Bem-Estar","Clínica Bem-Estar",40,4500,"Proposta"),("Comercial e propostas","Interno",0,0,"Interno"),("Gestão e reuniões","Interno",0,0,"Interno")]
for i,row in enumerate(proj):
    for k,v in enumerate(row,start=1): cfg.cell(row=24+i,column=k,value=v)
random.seed(11); rows=[]
atv={"Site novo · Loja Verde":["Wireframe","Layout","Textos","Implementação","Reunião com cliente"],"Campanha de fim de ano · Bistrô 42":["Conceito","Peças","Ajustes","Reunião com cliente"],
     "Relatório anual · Horizonte":["Coleta de dados","Redação","Diagramação"],"Identidade visual · Padaria do Sol":["Logo","Aplicações","Manual"],"Redes sociais · Clínica Bem-Estar":["Posts do mês","Aprovação"],
     "Comercial e propostas":["Proposta","Reunião comercial"],"Gestão e reuniões":["Semanal","Financeiro","Planejamento"]}
aloc={"Ana":[("Site novo · Loja Verde",0.2),("Relatório anual · Horizonte",0.3),("Comercial e propostas",0.25),("Gestão e reuniões",0.25)],
      "Bruno":[("Site novo · Loja Verde",0.45),("Campanha de fim de ano · Bistrô 42",0.3),("Identidade visual · Padaria do Sol",0.15),("Gestão e reuniões",0.1)],
      "Carla":[("Campanha de fim de ano · Bistrô 42",0.4),("Site novo · Loja Verde",0.2),("Comercial e propostas",0.3),("Gestão e reuniões",0.1)],
      "Diego":[("Site novo · Loja Verde",0.6),("Comercial e propostas",0.1),("Gestão e reuniões",0.3)]}
# Início e fim de cada projeto iguais a 04-projetos-e-prazos (Config). Antes do início ninguém lança horas nele;
# a parte da agenda que sobra vai para "Comercial e propostas" (julho e agosto foram de propostas, ver 08-funil).
periodo={"Site novo · Loja Verde":(date(2026,8,10),None),"Campanha de fim de ano · Bistrô 42":(date(2026,9,1),None),
         "Relatório anual · Horizonte":(date(2026,9,8),None),"Identidade visual · Padaria do Sol":(date(2026,7,15),date(2026,9,5))}
def aloc_dia(al,d):
    ativos=[(pj,fr) for pj,fr in al if pj not in periodo or (periodo[pj][0]<=d and (periodo[pj][1] is None or d<=periodo[pj][1]))]
    s=sum(fr for _,fr in ativos); out=[(pj,min(fr/s,fr*1.5)) for pj,fr in ativos]
    sobra=1-sum(fr for _,fr in out)
    if sobra>0.01:
        out=[(pj,fr+sobra) if pj=="Comercial e propostas" else (pj,fr) for pj,fr in out]
        if not any(pj=="Comercial e propostas" for pj,_ in out): out.append(("Comercial e propostas",sobra))
    return out
d=date(2026,7,1)
while d<=date(2026,9,12):
    if d.weekday()<5:
        for pessoa,al in aloc.items():
            total=random.choice([6,7,7,8,8,8.5,9])
            for pj,fr in aloc_dia(al,d):
                hrs=round(total*fr*random.uniform(0.7,1.3)*2)/2
                if hrs<=0: continue
                fat="Não" if pj in ("Comercial e propostas","Gestão e reuniões") else "Sim"
                rows.append((d,pessoa,pj,random.choice(atv[pj]),hrs,fat))
    d+=timedelta(days=1)
assert len(rows)<=N, len(rows)
for i,row in enumerate(rows):
    for c,v in zip((1,2,3,4,5,9),row): h.cell(row=R0+i,column=c,value=v)
como_usar(wb,"Horas e Custo por Projeto",[
 ("O que esta planilha faz","Cada pessoa lança as horas por projeto; ela transforma em custo (pelo custo-hora), compara com as horas orçadas e o valor cobrado, mostra a margem de cada projeto e a ocupação de cada pessoa no mês."),
 ("Passo 1","Em Config, cadastre as pessoas com custo mensal E horas disponíveis: os dois são obrigatórios, porque o custo-hora sai da divisão. Faltando um, a planilha escreve o que falta em vez de calcular um custo zerado, e o custo e a margem dos projetos e pessoas com horas dessa pessoa aparecem como \"cadastro incompleto\" no Painel (uma soma parcial mentiria para menos). Cadastre também os projetos com horas orçadas, valor cobrado e status. Projetos internos (comercial, gestão) entram com status Interno; proposta ainda não fechada pode entrar como Proposta, sem horas."),
 ("Passo 2","Em Horas, uma linha por pessoa, por dia, por projeto: data, pessoa, projeto, atividade, horas e se é faturável."),
 ("Passo 3","Em Painel, escolha o mês em Config. Por projeto: horas usadas contra orçadas, custo, margem e situação. Por pessoa: ocupação e horas faturáveis."),
 ("Rotina","Cada pessoa lança no fim do dia (2 minutos) ou na sexta (10 minutos). Dia 1 do mês: olhe margem por projeto antes de precificar o próximo."),
 ("Com a IA","Copie \"Por projeto\" e use \"Analisar 02: comparar dois períodos\" ou \"Analisar 06: perguntas antes de decidir\" para revisar preços. Para justificar horas a um cliente, \"Escrever 04\"."),
])
proteger(wb); salvar(wb,"09-horas-e-custo-por-projeto.xlsx","Horas e Custo por Projeto · Kit IA no Trabalho Completo")
