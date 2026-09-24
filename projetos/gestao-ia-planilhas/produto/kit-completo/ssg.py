"""Helpers comuns das planilhas do Seu Sócio Gestor (mesmo padrão visual do Kit Essencial)."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as L
UVA="3B1F5E"; SOL="FFC83D"; LILAS="7A5AA8"; LAVANDA="F3EEFB"; TINTA="1F1235"; AMARELO="FFF4CC"; BRANCO="FFFFFF"
VERDE="DDF3E7"; VERDE_T="155E3C"; VERM="FBE4E4"; VERM_T="7A1F1F"; CINZA="B0A6C4"
MESES=["Janeiro","Fevereiro","Março","Abril","Maio","Junho","Julho","Agosto","Setembro","Outubro","Novembro","Dezembro"]
BRL='"R$" #,##0.00;[Red]-"R$" #,##0.00'; BRL0='"R$" #,##0;[Red]-"R$" #,##0'; DATA="dd/mm/yyyy"; PCT="0%"
F=lambda **k: Font(name="Arial", **k); fill=lambda c: PatternFill("solid", fgColor=c)
thin=Side(style="thin", color="DCD2EC"); borda=Border(left=thin,right=thin,top=thin,bottom=thin)
KIT="Kit IA no Trabalho · Completo · Seu Sócio Gestor · versão 1.0 (setembro de 2026)"

def titulo(ws,texto,sub=None,merge_to="J"):
    ws["A1"]=texto; ws["A1"].font=F(bold=True,size=16,color=UVA)
    if sub:
        ws["A2"]=sub; ws["A2"].font=F(italic=True,size=10,color=LILAS); ws.merge_cells(f"A2:{merge_to}2")
def hdr(ws,row,vals,start=1,height=26):
    for i,v in enumerate(vals):
        c=ws.cell(row=row,column=start+i,value=v); c.font=F(bold=True,color=BRANCO,size=10); c.fill=fill(UVA)
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); c.border=borda
    ws.row_dimensions[row].height=height
def inp(c,fmt=None,center=False):
    c.fill=fill(AMARELO); c.protection=Protection(locked=False); c.border=borda; c.font=F(color=TINTA,size=10)
    if fmt: c.number_format=fmt
    if center: c.alignment=Alignment(horizontal="center")
def calc(c,fmt=None,center=True):
    c.border=borda; c.font=F(color=TINTA,size=10)
    if center: c.alignment=Alignment(horizontal="center")
    if fmt: c.number_format=fmt
def rotulo(c,bold=True): c.font=F(bold=bold,color=UVA,size=10)
def nota(c): c.font=F(size=9,color=LILAS)
def kpi(ws,row,col,label,formula,bg,fg,fmt='#,##0',span=2):
    a=ws.cell(row=row,column=col,value=label); a.font=F(size=9,bold=True,color=fg); a.fill=fill(bg); a.alignment=Alignment(horizontal="center")
    b=ws.cell(row=row+1,column=col,value=formula); b.font=F(size=16,bold=True,color=fg); b.fill=fill(bg); b.alignment=Alignment(horizontal="center"); b.number_format=fmt
    if span>1:
        ws.merge_cells(start_row=row,start_column=col,end_row=row,end_column=col+span-1); ws.merge_cells(start_row=row+1,start_column=col,end_row=row+1,end_column=col+span-1)
        for k in range(1,span): ws.cell(row=row,column=col+k).fill=fill(bg); ws.cell(row=row+1,column=col+k).fill=fill(bg)
    ws.row_dimensions[row+1].height=30
def lista(formula,allow_blank=True,strict=True):
    return DataValidation(type="list",formula1=formula,allow_blank=allow_blank,showErrorMessage=strict)
def widths(ws,ws_widths):
    for i,w in enumerate(ws_widths,start=1): ws.column_dimensions[L(i)].width=w
def como_usar(wb,nome,linhas,pos=0):
    u=wb.create_sheet("Como usar",pos)
    u["A1"]=nome; u["A1"].font=F(bold=True,size=20,color=UVA)
    u["A2"]=KIT; u["A2"].font=F(size=10,color=LILAS)
    base=linhas+[
     ("Legenda","Células amarelas: você preenche. Brancas: calculadas. Não é preciso mexer em nada fora do amarelo."),
     ("Proteção","Fórmulas protegidas sem senha. Para editar: Revisar > Desproteger planilha (Excel) ou Dados > Proteger intervalos (Google Sheets)."),
     ("Requisitos","Excel 2016 ou mais novo, Microsoft 365 ou Google Sheets (só funções do Excel 2007+; nada de MÍNIMOSES, MÁXIMOSES ou UNIRTEXTO). No celular abre nos aplicativos; para preencher, use o computador."),
     ("Google Sheets","Faça upload no Google Drive e abra com o Google Sheets. Fórmulas, listas, cores, semáforos e gráficos funcionam: as regras de cor que dependem do mês escolhido em Config usam INDIRECT, que é a forma que o Sheets aceita para regra condicional apontar para outra aba."),
     ("Data de referência","O exemplo está congelado em 14/09/2026, para todos os arquivos do kit mostrarem a mesma foto e os números fecharem entre eles. Ao começar a usar com os SEUS dados, troque em Config: a data de referência por =HOJE(), o ano, o mês do painel, o trimestre (onde houver) e a lista de feriados. Trocar só a data de referência deixa o resto apontando para 2026."),
     ("Exemplos","Prisma Comunicação é uma empresa fictícia. Nomes e valores são inventados. Apague-os antes de começar."),
     ("Suporte","suporte@seusociogestor.com.br · resposta em até 5 dias úteis · reembolso em até 7 dias pelo mesmo canal.")]
    for i,(a,b) in enumerate(base,start=4):
        u.cell(row=i,column=1,value=a).font=F(bold=True,color=UVA); u.cell(row=i,column=2,value=b).font=F(color=TINTA)
        u.cell(row=i,column=2).alignment=Alignment(wrap_text=True,vertical="top"); u.cell(row=i,column=1).alignment=Alignment(vertical="top"); u.row_dimensions[i].height=46
    u.column_dimensions["A"].width=20; u.column_dimensions["B"].width=95; u.sheet_view.showGridLines=False
    return u
def proteger(wb):
    for ws in wb.worksheets:
        ws.protection.sheet=True; ws.protection.formatColumns=False; ws.protection.formatRows=False
        ws.protection.selectLockedCells=False; ws.protection.sort=False; ws.protection.autoFilter=False
_REF_OUTRA_ABA=re.compile("(?<![A-Za-z0-9_!$])((?:'[^']+'|[A-Za-z\u00c0-\u00ff][A-Za-z\u00c0-\u00ff0-9_. ]*)![$]?[A-Z]{1,3}[$]?[0-9]+(?::[$]?[A-Z]{1,3}[$]?[0-9]+)?)")
def _indireto(formula):
    """Embrulha referência a OUTRA ABA em INDIRECT, dentro de fórmula de formatação
    condicional. O Google Sheets não aceita referência a outra aba em regra condicional
    (só via INDIRECT); sem isso, os semáforos que comparam com o mês escolhido em Config
    param de acompanhar depois de importar. Excel e LibreOffice aceitam as duas formas.
    Achado da auditoria de 18/09: 120 das 470 regras estavam nessa situação."""
    if "INDIRECT(" in formula.upper(): return formula
    return _REF_OUTRA_ABA.sub(lambda m: 'INDIRECT("%s")' % m.group(1), formula)
def _condicional_para_sheets(wb):
    n=0
    for ws in wb.worksheets:
        for rng in ws.conditional_formatting:
            for r in rng.rules:
                if not r.formula: continue
                novas=[_indireto(f or "") for f in r.formula]
                if novas!=list(r.formula): n+=1
                r.formula=novas
    return n
def _sem_grafico(wb):
    """O texto padrão do "Como usar" promete gráfico no Sheets; num arquivo sem gráfico a
    promessa sai (auditoria final-2: 16 planilhas diziam que "gráficos funcionam" sem ter um)."""
    if any(ws._charts for ws in wb.worksheets) or "Como usar" not in wb.sheetnames: return
    for row in wb["Como usar"].iter_rows(min_col=2,max_col=2):
        c=row[0]
        if isinstance(c.value,str): c.value=c.value.replace("Fórmulas, listas, cores, semáforos e gráficos funcionam","Fórmulas, listas, cores e semáforos funcionam")
def salvar(wb,arquivo,titulo_doc):
    _sem_grafico(wb)
    _condicional_para_sheets(wb)
    wb.properties.creator="Seu Sócio Gestor"; wb.properties.title=titulo_doc; wb.save(arquivo); print("salvo",arquivo)
