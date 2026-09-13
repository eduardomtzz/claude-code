#!/usr/bin/env python3
"""Planilha 7 do Kit Completo: Orçamento Previsto × Realizado. Gera 07-orcamento-previsto-x-realizado.xlsx"""
from ssg import *
from openpyxl.chart import BarChart, LineChart, Reference
from datetime import date
import random
NC=20; N=600; R0=5; RN=R0+N-1
wb=Workbook()
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. As categorias aqui alimentam a aba Previsto e a lista de Realizado.",merge_to="H")
cfg["A4"]="Empresa ou área"; cfg["B4"]="Prisma Comunicação (exemplo fictício)"
cfg["A5"]="Ano"; cfg["B5"]=2026
cfg["A6"]="Mês do painel"; cfg["B6"]="Setembro"
cfg["A7"]="Número do mês"; cfg["B7"]="=MATCH(B6,$E$5:$E$16,0)"
cfg["A8"]="Alerta de estouro acima de (%)"; cfg["B8"]=0.1
for r in range(4,9): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"]); inp(cfg["B6"],center=True); calc(cfg["B7"]); inp(cfg["B8"],PCT)
cfg["E4"]="Meses"; rotulo(cfg["E4"])
for i,m_ in enumerate(MESES): cfg.cell(row=5+i,column=5,value=m_).font=F(size=10,color=TINTA)
dv=lista("=Config!$E$5:$E$16"); dv.add("B6"); cfg.add_data_validation(dv)
cfg["A10"]="Categorias (até 20)"; rotulo(cfg["A10"])
hdr(cfg,11,["Categoria","Tipo","Responsável"])
for i in range(NC):
    r=12+i
    for c in (1,2,3): inp(cfg.cell(row=r,column=c))
    cfg.cell(row=r,column=2).alignment=Alignment(horizontal="center")
dvt=lista('"Receita,Despesa"'); dvt.add(f"B12:B{11+NC}"); cfg.add_data_validation(dvt)
widths(cfg,(34,14,18,3,12)); cfg.sheet_view.showGridLines=False
# ---------- Previsto ----------
pv=wb.create_sheet("Previsto")
titulo(pv,"Orçamento previsto","Quanto você planeja receber ou gastar em cada categoria, mês a mês. Preencha o amarelo; total é calculado.",merge_to="N")
hdr(pv,4,["Categoria"]+[m_[:3] for m_ in MESES]+["Total"])
for i in range(NC):
    r=5+i; src=f"Config!$A${12+i}"
    pv.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(pv.cell(row=r,column=1),center=False)
    for m_ in range(12): inp(pv.cell(row=r,column=2+m_),BRL0)
    pv.cell(row=r,column=14,value=f'=IF(A{r}="","",SUM(B{r}:M{r}))'); calc(pv.cell(row=r,column=14),BRL0)
pv.cell(row=5+NC,column=1,value="Receitas previstas"); pv.cell(row=6+NC,column=1,value="Despesas previstas"); pv.cell(row=7+NC,column=1,value="Resultado previsto")
for m_ in range(13):
    col=2+m_; cl=L(col)
    pv.cell(row=5+NC,column=col,value=f'=SUMIFS({cl}5:{cl}{4+NC},Config!$B$12:$B${11+NC},"Receita")')
    pv.cell(row=6+NC,column=col,value=f'=SUMIFS({cl}5:{cl}{4+NC},Config!$B$12:$B${11+NC},"Despesa")')
    pv.cell(row=7+NC,column=col,value=f'={cl}{5+NC}-{cl}{6+NC}')
    for rr in (5+NC,6+NC,7+NC): calc(pv.cell(row=rr,column=col),BRL0); pv.cell(row=rr,column=col).font=F(bold=True,color=UVA,size=10)
for rr in (5+NC,6+NC,7+NC): rotulo(pv.cell(row=rr,column=1)); pv.cell(row=rr,column=1).border=borda
widths(pv,[30]+[11]*12+[13]); pv.freeze_panes="B5"; pv.sheet_view.showGridLines=False
# ---------- Realizado ----------
re_=wb.create_sheet("Realizado")
titulo(re_,"Realizado","Uma linha por lançamento real (ou o total do mês por categoria, se preferir). Mês e ano são calculados.",merge_to="H")
hdr(re_,4,["Data","Categoria","Descrição","Valor","Observação","Mês","Ano","Tipo"])
for r in range(R0,RN+1):
    for c in (1,2,3,4,5): inp(re_.cell(row=r,column=c))
    re_.cell(row=r,column=1).number_format=DATA; re_.cell(row=r,column=4).number_format=BRL; re_.cell(row=r,column=1).alignment=Alignment(horizontal="center")
    re_.cell(row=r,column=6,value=f'=IF(A{r}="","",MONTH(A{r}))'); calc(re_.cell(row=r,column=6))
    re_.cell(row=r,column=7,value=f'=IF(A{r}="","",YEAR(A{r}))'); calc(re_.cell(row=r,column=7))
    re_.cell(row=r,column=8,value=f'=IF(B{r}="","",IFERROR(INDEX(Config!$B$12:$B${11+NC},MATCH(B{r},Config!$A$12:$A${11+NC},0)),"?"))'); calc(re_.cell(row=r,column=8))
dvc=lista(f"=Config!$A$12:$A${11+NC}",strict=False); dvc.add(f"B{R0}:B{RN}"); re_.add_data_validation(dvc)
dvd=DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True); dvd.add(f"A{R0}:A{RN}"); re_.add_data_validation(dvd)
re_.conditional_formatting.add(f"A{R0}:H{RN}", FormulaRule(formula=[f'$H{R0}="Receita"'], font=F(color=VERDE_T,size=10)))
re_.conditional_formatting.add(f"B{R0}:B{RN}", FormulaRule(formula=[f'$H{R0}="?"'], fill=fill(VERM)))
widths(re_,(12,28,36,14,28,6,7,10)); re_.freeze_panes="A5"; re_.sheet_view.showGridLines=False; re_.auto_filter.ref=f"A4:H{RN}"
# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
p["A1"]='=Config!B4&" · "&Config!B6&" de "&Config!B5'; p["A1"].font=F(bold=True,size=16,color=UVA); p.merge_cells("A1:I1")
p["A2"]="Nada para preencher aqui. Escolha o mês em Config; previsto vem da aba Previsto, realizado da aba Realizado."; nota(p["A2"]); p.merge_cells("A2:I2")
M="Config!$B$7"; Y="Config!$B$5"; AL="Config!$B$8"
RB=f"Realizado!$B${R0}:$B${RN}"; RD=f"Realizado!$D${R0}:$D${RN}"; RF=f"Realizado!$F${R0}:$F${RN}"; RG=f"Realizado!$G${R0}:$G${RN}"; RH=f"Realizado!$H${R0}:$H${RN}"
def real_tipo(t,m=M): return f'SUMIFS({RD},{RH},"{t}",{RF},{m},{RG},{Y})'
def prev_tipo(t,m=M): return f'SUMIFS(INDEX(Previsto!$B$5:$M${4+NC},0,{m}),Config!$B$12:$B${11+NC},"{t}")'
kpi(p,4,1,"Receita prevista",f'={prev_tipo("Receita")}',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,3,"Receita realizada",f'={real_tipo("Receita")}',VERDE,VERDE_T,fmt=BRL0)
kpi(p,4,5,"Despesa prevista",f'={prev_tipo("Despesa")}',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,7,"Despesa realizada",f'={real_tipo("Despesa")}',VERM,VERM_T,fmt=BRL0)
kpi(p,4,9,"Resultado do mês",'=C5-G5',SOL,UVA,fmt=BRL0)
p["A7"]="Por categoria no mês"; p["A7"].font=F(bold=True,size=13,color=UVA)
hdr(p,8,["Categoria","Tipo","Previsto","Realizado","Desvio (R$)","Desvio (%)","Situação","Acumulado previsto","Acumulado realizado"])
for i in range(NC):
    r=9+i; src=f"Config!$A${12+i}"; tp=f"Config!$B${12+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",{tp})'); calc(p.cell(row=r,column=2))
    p.cell(row=r,column=3,value=f'=IF({src}="","",INDEX(Previsto!$B${5+i}:$M${5+i},1,{M}))'); calc(p.cell(row=r,column=3),BRL0)
    p.cell(row=r,column=4,value=f'=IF({src}="","",SUMIFS({RD},{RB},{src},{RF},{M},{RG},{Y}))'); calc(p.cell(row=r,column=4),BRL0)
    p.cell(row=r,column=5,value=f'=IF({src}="","",D{r}-C{r})'); calc(p.cell(row=r,column=5),BRL0)
    p.cell(row=r,column=6,value=f'=IF(OR({src}="",C{r}=0),"",E{r}/C{r})'); calc(p.cell(row=r,column=6),"0%")
    p.cell(row=r,column=7,value=f'=IF({src}="","",IF(B{r}="Despesa",IF(AND(C{r}>0,F{r}>{AL}),"Estourou",IF(D{r}>C{r},"Acima","Dentro")),IF(AND(C{r}>0,F{r}<-{AL}),"Abaixo da meta",IF(D{r}>=C{r},"Meta batida","Perto da meta"))))'); calc(p.cell(row=r,column=7))
    p.cell(row=r,column=8,value=f'=IF({src}="","",SUM(INDEX(Previsto!$B${5+i}:$M${5+i},1,1):INDEX(Previsto!$B${5+i}:$M${5+i},1,{M})))'); calc(p.cell(row=r,column=8),BRL0)
    p.cell(row=r,column=9,value=f'=IF({src}="","",SUMIFS({RD},{RB},{src},{RF},"<="&{M},{RG},{Y}))'); calc(p.cell(row=r,column=9),BRL0)
RNG=f"A9:I{8+NC}"
p.conditional_formatting.add(f"G9:G{8+NC}", FormulaRule(formula=['G9="Estourou"'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
p.conditional_formatting.add(f"G9:G{8+NC}", FormulaRule(formula=['G9="Abaixo da meta"'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
p.conditional_formatting.add(f"G9:G{8+NC}", FormulaRule(formula=['OR(G9="Acima",G9="Perto da meta")'], fill=fill("FFF4CC")))
p.conditional_formatting.add(f"G9:G{8+NC}", FormulaRule(formula=['OR(G9="Dentro",G9="Meta batida")'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
r0=10+NC
p.cell(row=r0,column=1,value="O ano, mês a mês").font=F(bold=True,size=13,color=UVA)
hdr(p,r0+1,["Mês","Receita prevista","Receita realizada","Despesa prevista","Despesa realizada","Resultado previsto","Resultado realizado"])
for i in range(12):
    r=r0+2+i
    p.cell(row=r,column=1,value=MESES[i]); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'={prev_tipo("Receita",i+1)}'); p.cell(row=r,column=3,value=f'={real_tipo("Receita",i+1)}')
    p.cell(row=r,column=4,value=f'={prev_tipo("Despesa",i+1)}'); p.cell(row=r,column=5,value=f'={real_tipo("Despesa",i+1)}')
    p.cell(row=r,column=6,value=f'=B{r}-D{r}'); p.cell(row=r,column=7,value=f'=C{r}-E{r}')
    for c in range(2,8): calc(p.cell(row=r,column=c),BRL0)
p.conditional_formatting.add(f"G{r0+2}:G{r0+13}", FormulaRule(formula=[f'G{r0+2}<0'], font=F(color="C8402E",size=10,bold=True)))
p.cell(row=r0+14,column=1,value="Meses sem lançamento aparecem zerados no realizado. Compare só os meses fechados.").font=F(size=9,color=LILAS)
bc=BarChart(); bc.type="col"; bc.grouping="clustered"; bc.height=7.5; bc.width=20; bc.title="Despesas: previsto × realizado"; bc.style=2
bc.add_data(Reference(p,min_col=4,max_col=5,min_row=r0+1,max_row=r0+13),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=r0+2,max_row=r0+13))
bc.series[0].graphicalProperties.solidFill="B89BE0"; bc.series[1].graphicalProperties.solidFill="3B1F5E"; bc.legend.position="b"; bc.y_axis.majorGridlines=None
p.add_chart(bc,f"I{r0}")
widths(p,(30,12,14,14,14,12,15,16,16)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
# exemplos
cats=[("Serviços · projetos","Receita","Ana"),("Serviços · mensalidades","Receita","Ana"),("Produtos digitais","Receita","Carla"),
      ("Salários e pró-labore","Despesa","Ana"),("Freelancers","Despesa","Bruno"),("Escritório e contas","Despesa","Ana"),("Softwares e assinaturas","Despesa","Diego"),
      ("Marketing e mídia","Despesa","Carla"),("Impostos","Despesa","Ana"),("Equipamentos","Despesa","Diego"),("Treinamento","Despesa","Ana"),("Viagens e eventos","Despesa","Carla"),("Reserva","Despesa","Ana")]
for i,(a,b,c) in enumerate(cats):
    for k,v in enumerate((a,b,c),start=1): cfg.cell(row=12+i,column=k,value=v)
prev={"Serviços · projetos":[52000,55000,58000,60000,60000,62000,64000,64000,66000,70000,72000,60000],
      "Serviços · mensalidades":[18000,18000,20000,20000,22000,22000,24000,24000,26000,26000,28000,28000],
      "Produtos digitais":[1000,1000,1500,1500,2000,2000,2500,2500,3000,3000,3500,4000],
      "Salários e pró-labore":[38000]*6+[41000]*6,"Freelancers":[6000,6000,7000,7000,8000,8000,8000,9000,9000,10000,10000,6000],
      "Escritório e contas":[4200]*12,"Softwares e assinaturas":[1900]*12,"Marketing e mídia":[3000,3000,3500,3500,4000,4000,4500,4500,5000,6000,6000,3000],
      "Impostos":[8500,8700,9000,9200,9400,9600,9900,9900,10200,10700,11000,10000],"Equipamentos":[0,0,4000,0,0,0,3000,0,0,0,0,0],
      "Treinamento":[500,500,500,500,500,500,500,500,500,500,500,500],"Viagens e eventos":[800,800,800,2500,800,800,800,800,3000,800,800,800],"Reserva":[3000]*12}
for i,(a,b,c) in enumerate(cats):
    for m_,v in enumerate(prev[a]): pv.cell(row=5+i,column=2+m_,value=v)
random.seed(3); rows=[]
fator={"Serviços · projetos":0.96,"Serviços · mensalidades":1.02,"Produtos digitais":1.15,"Salários e pró-labore":1.0,"Freelancers":1.22,"Escritório e contas":1.03,"Softwares e assinaturas":1.12,"Marketing e mídia":1.3,"Impostos":0.99,"Equipamentos":1.0,"Treinamento":0.4,"Viagens e eventos":0.9,"Reserva":1.0}
desc={"Serviços · projetos":"Projetos faturados no mês","Serviços · mensalidades":"Contratos mensais","Produtos digitais":"Vendas de templates","Salários e pró-labore":"Folha do mês","Freelancers":"Freelas de motion e redação","Escritório e contas":"Aluguel, luz, internet","Softwares e assinaturas":"Adobe, Google, Notion, hospedagem","Marketing e mídia":"Impulsionamento e mídia paga","Impostos":"DAS e retenções","Equipamentos":"Monitor e notebook","Treinamento":"Curso online","Viagens e eventos":"Deslocamentos a clientes","Reserva":"Transferência para a reserva"}
for m_ in range(1,10):
    for (a,b,c) in cats:
        v=prev[a][m_-1]*fator[a]*(1+random.uniform(-0.04,0.04))
        if v<=0: continue
        rows.append((date(2026,m_,min(28,5+cats.index((a,b,c))*2)),a,desc[a],round(v,2)))
for i,(d,a,ds,v) in enumerate(rows):
    for c,val in zip((1,2,3,4),(d,a,ds,v)): re_.cell(row=R0+i,column=c,value=val)
como_usar(wb,"Orçamento Previsto × Realizado",[
 ("O que esta planilha faz","Você planeja o ano por categoria (previsto) e lança o que aconteceu (realizado). Ela mostra, mês a mês, o desvio em reais e em %, o que estourou, o acumulado do ano e o resultado previsto contra o real."),
 ("Passo 1","Em Config, preencha o ano, o mês do painel, o limite de alerta (10% é um bom padrão) e as categorias com tipo (receita ou despesa)."),
 ("Passo 2","Em Previsto, digite o valor planejado de cada categoria em cada mês. Pode ser o mesmo número o ano inteiro."),
 ("Passo 3","Em Realizado, uma linha por lançamento, ou uma linha por categoria com o total do mês. A categoria vem da lista."),
 ("Passo 4","Em Painel, escolha o mês em Config: por categoria com situação (dentro, acima, estourou) e o ano completo com gráfico."),
 ("Rotina","Uma vez por semana, 10 minutos de lançamentos. Dia 5 de cada mês: fechar o mês anterior e olhar o painel antes de decidir gasto novo."),
 ("Com a IA","Copie \"Por categoria no mês\" e use \"Analisar 03: onde cortar gastos\" ou \"Escrever 05: explicar o mês financeiro\". Para a reunião com o sócio, \"Apresentar 01: roteiro de 8 slides\"."),
])
proteger(wb); salvar(wb,"07-orcamento-previsto-x-realizado.xlsx","Orçamento Previsto × Realizado · Kit IA no Trabalho Completo")
