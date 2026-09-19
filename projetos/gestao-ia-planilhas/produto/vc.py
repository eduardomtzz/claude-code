#!/usr/bin/env python3
"""Ferramentas compartilhadas pelos verifica_coerencia.py dos quatro kits.

Por que existe: cada kit tinha a sua própria cópia de check/num/rows, e a versão do
Essencial e do Completo simplesmente não existia — eram os dois únicos kits sem
verificação automática. Aqui fica a parte comum; as afirmações de cada kit ficam no
seu próprio arquivo.

Regra que vale para todos: a verificação lê CÓPIAS recalculadas pelo LibreOffice em
data_only, nunca os originais entregues.
"""
import pathlib, datetime
import openpyxl

def abrir(pasta, prefixo):
    """Abre a planilha que começa com o prefixo ("04", "17"...) na pasta das cópias."""
    p = pathlib.Path(pasta)
    achados = sorted(p.glob(f"{prefixo}-*.xlsx"))
    if not achados:
        raise SystemExit(f"não achei {prefixo}-*.xlsx em {p}")
    return openpyxl.load_workbook(achados[0], data_only=True)

def num(v):
    """Valor numérico de uma célula. Fórmula que devolve "" vira None no recálculo."""
    if v in (None, "", "—"): return 0
    if isinstance(v, (datetime.datetime, datetime.date)): return 0
    return float(v)

def dt(v):
    return v.date() if isinstance(v, datetime.datetime) else v

def txt(v):
    return "" if v is None else str(v)

def linhas(ws, r0, cols, chave=0):
    """Linhas de uma tabela a partir de r0 até a primeira com a coluna-chave vazia.
    As tabelas dos kits são contíguas; notas e rodapés ficam depois do vazio."""
    out = []
    for r in range(r0, ws.max_row + 1):
        if ws.cell(row=r, column=cols[chave]).value in (None, ""): break
        out.append([ws.cell(row=r, column=c).value for c in cols])
    return out

def rotulo_em(ws, rotulo, col=1, r0=1, r1=200):
    """Linha cujo rótulo COMEÇA com o texto dado. Endereço fixo quebra a cada mudança
    de estrutura; foi assim que uma conferência do kit Advogados passou a mentir."""
    for r in range(r0, r1):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.startswith(rotulo): return r
    raise AssertionError(f"{ws.title}: rótulo {rotulo!r} não encontrado")

class Verificador:
    def __init__(self, nome):
        self.nome = nome; self.ok = 0; self.falhas = []
    def check(self, nome, a, b, tol=0.5):
        if isinstance(a, (int, float)) or isinstance(b, (int, float)):
            bom = abs(num(a) - num(b)) <= tol
        else:
            bom = a == b
        if bom: self.ok += 1
        else: self.falhas.append(f"{nome}: {a!r} != {b!r}")
    def fim(self):
        print(f"{self.nome}: {self.ok} verificações OK, {len(self.falhas)} falhas")
        for f in self.falhas: print("  FALHA:", f)
        return 1 if self.falhas else 0
