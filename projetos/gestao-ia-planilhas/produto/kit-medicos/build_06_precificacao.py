#!/usr/bin/env python3
"""Planilha 6 do Kit de Gestão para Médicos: Precificação de consulta e procedimento. Gera 06-precificacao.xlsx
Preço = tempo (com o retorno embutido, nas consultas) × custo-hora + material, mais impostos e margem.
Compara o preço particular praticado e as tabelas de convênio com o mínimo e o alvo. Custo-hora da planilha 05 (R$ 200,00)."""
from ssg import *
from openpyxl.chart import BarChart, Reference
import dados
NPROC=12; NCONV=4; T0=5; TN=T0+NPROC-1
wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Custo-hora da planilha 05; impostos e margens valem para todos os procedimentos.",merge_to="F")
campos=[("Nome da clínica",f"{dados.CLINICA} (exemplo fictício)",None),("Mês de referência","Setembro de 2026",None),("Data de referência",dados.HOJE,DATA),
 ("Custo da hora de atendimento (R$)",dados.CUSTO_HORA,BRL),("Impostos sobre o que entra (%)",dados.ALIQ,PCT),("Margem mínima sobre o preço (%)",dados.MARGEM,PCT),
 ("Margem alvo sobre o preço (%)",dados.MARGEM_ALVO,PCT),("Retornos por consulta (média)",dados.RETORNO_PROB,"0.00"),("Duração do retorno (minutos)",dados.DUR["Retorno"],"0")]
for i,(a,v,fmt) in enumerate(campos):
    r=4+i; cfg.cell(row=r,column=1,value=a); rotulo(cfg.cell(row=r,column=1)); cfg.cell(row=r,column=2,value=v)
    if r==6: calc(cfg.cell(row=r,column=2),fmt)
    else: inp(cfg.cell(row=r,column=2),fmt,center=fmt is not None)
notas={7:"Copie do Painel da planilha 05 (\"Custo da hora de atendimento\"). No exemplo, R$ 200,00 (R$ 28.000 ÷ 140 h).",8:"A mesma alíquota efetiva da planilha 05 (exemplo: 11 %; confira com o contador). Só imposto: a taxa da maquininha não entra aqui (ela está na 16 e no resultado 18).",
 9:"Abaixo disso o procedimento não vale a pena. É o piso da faixa (a mesma margem da 05).",10:"A margem que a clínica quer de verdade. É o teto da faixa.",
 11:"Consulta que costuma gerar retorno sem cobrança: em média, quantos retornos por consulta. O tempo do retorno entra no custo da consulta.",12:"Minutos de um retorno (planilha 01, Config)."}
for r,t in notas.items(): cfg.cell(row=r,column=3,value=t); nota(cfg.cell(row=r,column=3))
cfg["A14"]="Convênios (até 4, na ordem das colunas da aba Precificação)"; rotulo(cfg["A14"])
for i in range(NCONV): inp(cfg.cell(row=15+i,column=2))
for i,(n,_,_) in enumerate(dados.CONVENIOS): cfg.cell(row=15+i,column=2,value=n)
cfg["A15"]="Nomes"; rotulo(cfg["A15"],bold=False)
cfg["C15"]="Preencha de cima para baixo, sem pular linha. As tabelas de cada convênio ficam na aba Precificação."; nota(cfg["C15"])
widths(cfg,(46,16,90)); cfg.sheet_view.showGridLines=False
CH="Config!$B$7"; IMP="Config!$B$8"; MMIN="Config!$B$9"; MALVO="Config!$B$10"; NRET="Config!$B$11"; DRET="Config!$B$12"
DIVMIN=f"(1-{IMP}-{MMIN})"; DIVALVO=f"(1-{IMP}-{MALVO})"; CFGX='(Config!$B$13<>"Sim")'
# Auditoria final (G02): imposto + margem alvo abaixo de 100 % e margem mínima ≤ margem alvo,
# validado na digitação e guardado nas fórmulas (colagem).
for _c,_f,_m in (("B8",'=IF(ISNUMBER(B8),AND(B8>=0,B8<1,B8+N(B10)<1),FALSE)',"Impostos entre 0 % e 99 %, e impostos + margem alvo abaixo de 100 %."),
                 ("B9",'=IF(ISNUMBER(B9),AND(B9>=0,B9<=N(B10),N(B8)+B9<1),FALSE)',"Margem mínima entre 0 % e a margem alvo, e impostos + margem abaixo de 100 %."),
                 ("B10",'=IF(ISNUMBER(B10),AND(B10>=N(B9),B10<1,N(B8)+B10<1),FALSE)',"Margem alvo entre a margem mínima e 99 %, e impostos + margem alvo abaixo de 100 %.")):
    _dv=DataValidation(type="custom",formula1=_f,allow_blank=False,showErrorMessage=True,errorTitle="Percentual",error=_m); _dv.add(_c); cfg.add_data_validation(_dv)
# Auditoria final-2 (G02): o mesmo conjunto de condições da validação, num predicado só, visível
# em Config!B13. Preço mínimo, preço alvo, margens, situação, resumo e simulação consultam essa
# célula: colar -30 %, texto, branco, mínima acima da alvo ou soma de 100 % suspende tudo.
cfg["A13"]="Impostos e margens conferem? (calculado)"; rotulo(cfg["A13"])
cfg["B13"]='=IF(AND(ISNUMBER(B8),ISNUMBER(B9),ISNUMBER(B10)),IF(AND(B8>=0,B8<1,B9>=0,B9<=B10,B10<1,B8+B10<1),"Sim","Não"),"Não")'; calc(cfg["B13"])
cfg["C13"]="\"Não\" suspende preços, margens, situação e resumo: impostos e margens de 0 % a 99 %, mínima ≤ alvo e impostos + margem alvo abaixo de 100 %."; nota(cfg["C13"])
# ---------- Precificação ----------
s=wb.create_sheet("Precificação",0)
titulo(s,'=Config!$B$4&" · Precificação de consulta e procedimento · "&Config!$B$5',"Amarelo: procedimento, minutos, se gera retorno, material, preço particular praticado e tabelas de convênio. O resto é calculado: custo cheio, preço mínimo, preço alvo e margem por pagador.",merge_to="R")
CONV_H=[f'=IF(Config!$B${15+j}="","",Config!$B${15+j})' for j in range(NCONV)]
heads=["Procedimento","Minutos","Gera retorno?","Material (R$)","Tempo com retorno (min)","Custo cheio (R$)","Preço mínimo (R$)","Preço alvo (R$)","Particular praticado (R$)","Margem no particular","Situação"]
for j in range(NCONV): heads+=[CONV_H[j],f'=IF(Config!$B${15+j}="","","Margem")']
hdr(s,T0-1,heads,height=40)
for r in range(T0,TN+1):
    inp(s.cell(row=r,column=1)); inp(s.cell(row=r,column=2),"0",center=True); inp(s.cell(row=r,column=3),center=True); inp(s.cell(row=r,column=4),BRL,center=True)
    s.cell(row=r,column=5,value=f'=IF(A{r}="","",{f_tempo(f"B{r}",f"C{r}",NRET,DRET)})'); calc(s.cell(row=r,column=5),"0.0")
    # Auditoria final (G01/G02): custo-hora em branco não é custo zero (dava consulta a R$ 4 de
    # custo e 88 % de margem); margens que somam 100 % não dão preço negativo.
    s.cell(row=r,column=6,value=f'=IF(A{r}="","",{f_custo(f"E{r}",f"D{r}",CH)})'); calc(s.cell(row=r,column=6),BRL)
    s.cell(row=r,column=7,value=f'=IF(A{r}="","",IF(NOT(ISNUMBER(F{r})),F{r},IF({CFGX},"margens inválidas (Config)",F{r}/{DIVMIN})))'); calc(s.cell(row=r,column=7),BRL)
    s.cell(row=r,column=8,value=f'=IF(A{r}="","",IF(NOT(ISNUMBER(F{r})),F{r},IF({CFGX},"margens inválidas (Config)",F{r}/{DIVALVO})))'); calc(s.cell(row=r,column=8),BRL)
    inp(s.cell(row=r,column=9),BRL0,center=True)
    s.cell(row=r,column=10,value=f'=IF(OR(A{r}="",I{r}="",NOT({num_ok(f"I{r}")}),N(I{r})=0,NOT(ISNUMBER(F{r})),{CFGX}),"",(I{r}*(1-{IMP})-F{r})/I{r})'); calc(s.cell(row=r,column=10),PCT)
    s.cell(row=r,column=11,value=f'=IF(A{r}="","",IF(I{r}="","Falta o preço particular",IF(NOT({num_ok(f"I{r}")}),"Preço inválido",IF(NOT(ISNUMBER(F{r})),IF(IF(ISNUMBER({CH}),{CH}>=0,FALSE),"Faltam dados da linha","Falta o custo-hora"),IF(OR(NOT(ISNUMBER(G{r})),NOT(ISNUMBER(H{r}))),"Margens inválidas",IF(I{r}=0,"Sem cobrança",IF(I{r}<G{r},"Abaixo do mínimo",IF(I{r}<H{r},"Entre mínimo e alvo","No alvo ou acima"))))))))'); calc(s.cell(row=r,column=11))
    for j in range(NCONV):
        cv=12+2*j; cm=cv+1
        inp(s.cell(row=r,column=cv),BRL0,center=True)
        s.cell(row=r,column=cm,value=f'=IF(OR(A{r}="",{L(cv)}{r}="",NOT({num_ok(f"{L(cv)}{r}")}),N({L(cv)}{r})=0,NOT(ISNUMBER(F{r})),{CFGX}),"",({L(cv)}{r}*(1-{IMP})-F{r})/{L(cv)}{r})'); calc(s.cell(row=r,column=cm),PCT)
    for c in (6,7,8): s.cell(row=r,column=c).font=F(color=UVA,size=10,bold=(c==7))
    # Auditoria final-6 (G04): tabela 0 digitada é convênio que não paga nada: o prejuízo é o custo cheio.
    # Exceção: procedimento sem cobrança também no particular (o retorno, cujo tempo já está na consulta).
    for j in range(NCONV):   # auxiliar oculta: prejuízo por atendimento (R$) quando a tabela não cobre o custo cheio
        cv=12+2*j; cm=cv+1
        s.cell(row=r,column=20+j,value=f'=IF(AND(A{r}<>"",{L(cv)}{r}<>"",{num_ok(f"{L(cv)}{r}")},ISNUMBER(F{r}),NOT({CFGX}),OR({L(cv)}{r}>0,N(I{r})>0)),MIN(0,{L(cv)}{r}*(1-{IMP})-F{r}),0)'); s.cell(row=r,column=20+j).font=F(color=CINZA,size=9)
dvs=lista('"Sim,Não"'); dvs.add(f"C{T0}:C{TN}"); s.add_data_validation(dvs)
s.conditional_formatting.add(f"K{T0}:K{TN}", FormulaRule(formula=[f'OR(K{T0}="Abaixo do mínimo",K{T0}="Falta o custo-hora",K{T0}="Margens inválidas",K{T0}="Faltam dados da linha",K{T0}="Preço inválido",K{T0}="Falta o preço particular")'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
s.conditional_formatting.add(f"K{T0}:K{TN}", FormulaRule(formula=[f'K{T0}="Entre mínimo e alvo"'], fill=fill(AMARELO)))
s.conditional_formatting.add(f"K{T0}:K{TN}", FormulaRule(formula=[f'K{T0}="No alvo ou acima"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
for c in ["J"]+[L(13+2*j) for j in range(NCONV)]:
    s.conditional_formatting.add(f"{c}{T0}:{c}{TN}", FormulaRule(formula=[f'AND(ISNUMBER({c}{T0}),{c}{T0}<0)'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
    s.conditional_formatting.add(f"{c}{T0}:{c}{TN}", FormulaRule(formula=[f'AND(ISNUMBER({c}{T0}),{c}{T0}>=0,{c}{T0}<{MMIN})'], fill=fill(AMARELO)))
    s.conditional_formatting.add(f"{c}{T0}:{c}{TN}", FormulaRule(formula=[f'AND(ISNUMBER({c}{T0}),{c}{T0}>={MMIN})'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
NT=TN+2
s.cell(row=NT,column=1,value="Custo cheio = tempo (com retorno embutido) ÷ 60 × custo-hora + material. Preço mínimo = custo ÷ (1 − impostos − margem mínima); preço alvo, com a margem alvo. Margem = (preço × (1 − impostos) − custo) ÷ preço. Vermelho: não cobre o custo cheio. Amarelo: cobre o custo, mas fica abaixo da margem mínima. Verde: margem mínima ou mais.").font=F(size=9,color=LILAS)
s.merge_cells(start_row=NT,start_column=1,end_row=NT,end_column=19); s.cell(row=NT,column=1).alignment=Alignment(wrap_text=True,vertical="top"); s.row_dimensions[NT].height=32
s.cell(row=NT+1,column=1,value="Retorno tem preço zero de propósito: o tempo dele já está no custo da consulta. Tabela 0 em procedimento cobrado conta como prejuízo do custo cheio; no retorno (particular 0 também) não conta. Tabela de convênio abaixo do custo cheio não é automaticamente \"não aceitar\": com agenda vazia, o que importa é a contribuição (planilha 07).").font=F(size=9,color=LILAS)
s.merge_cells(start_row=NT+1,start_column=1,end_row=NT+1,end_column=19); s.cell(row=NT+1,column=1).alignment=Alignment(wrap_text=True,vertical="top"); s.row_dimensions[NT+1].height=32
# KPIs
SA=f"$A${T0}:$A${TN}"; SK=f"$K${T0}:$K${TN}"; SF=f"$F${T0}:$F${TN}"
R0=NT+3
s.cell(row=R0,column=1,value="Resumo").font=F(bold=True,size=13,color=UVA)
# Auditoria final-2 (G01): sem custo-hora ou com Config inválida, o resumo diz o que falta em vez
# de "0 de 7", zero tabelas e zero prejuízo (as auxiliares T:W viram 0 quando não há margem).
# Auditoria final-4 (F4-G01/G02): linha com dado faltando ou preço em texto também suspende o resumo:
# contar só as linhas completas daria "0 de 7" e "11 tabelas" como se fossem o total.
RESID=f'SUMPRODUCT(($A${T0}:$A${TN}="")*('+'+'.join(f'(${c}${T0}:${c}${TN}<>"")' for c in ['B','C','D','I']+[L(12+2*j) for j in range(NCONV)])+'>0))'   # dado em linha sem procedimento (auditoria final-5)
# Auditoria final-6 (M01): preço particular em branco é pendência ("Falta o preço particular"); zero digitado
# continua valendo como "Sem cobrança". Tabela de convênio 0 conta como abaixo do custo (G04).
NINC=f'(SUMPRODUCT(($A${T0}:$A${TN}<>"")*(ISNUMBER($F${T0}:$F${TN})=FALSE))+SUMPRODUCT(($A${T0}:$A${TN}<>"")*($I${T0}:$I${TN}=""))+{RESID}+'+'+'.join(f'SUMPRODUCT(ISTEXT(${c}${T0}:${c}${TN})+ISNUMBER(${c}${T0}:${c}${TN})*(${c}${T0}:${c}${TN}<0))' for c in ['I']+[L(12+2*j) for j in range(NCONV)])+')'
FALTA=f'IF(NOT(IF(ISNUMBER({CH}),{CH}>=0,FALSE)),"falta o custo-hora em Config",IF({CFGX},"margens inválidas (Config)",IF({NINC}>0,{NINC}&" linha(s) ou preço(s) incompletos",'
res=[("Custo da hora de atendimento (planilha 05)",f'=IF(IF(ISNUMBER({CH}),{CH}>=0,FALSE),{CH},"falta o custo-hora em Config")',BRL),
     ("Hora mínima a cobrar (custo-hora ÷ (1 − impostos − margem mínima))",f"={FALTA}{CH}/{DIVMIN})))",BRL),
     ("Procedimentos com particular abaixo do mínimo",f'={FALTA}COUNTIF({SK},"Abaixo do mínimo")&" de "&COUNTA({SA}))))',"@"),
     ("Tabelas de convênio abaixo do custo cheio + imposto",f'={FALTA}COUNTIF($T${T0}:$W${TN},"<0"))))',"0"),
     ("Maior prejuízo por atendimento em convênio (R$)",f'={FALTA}-MIN($T${T0}:$W${TN}))))',BRL)]
for i,(a,f_,fmt) in enumerate(res):
    r=R0+1+i; s.cell(row=r,column=1,value=a); calc(s.cell(row=r,column=1),center=False); s.cell(row=r,column=5,value=f_); calc(s.cell(row=r,column=5),fmt); s.cell(row=r,column=5).font=F(bold=True,color=UVA,size=10)
    s.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4)
s.cell(row=R0+6,column=1,value="\"Abaixo do custo cheio + imposto\" conta as células de Margem negativas: a tabela do convênio, depois dos impostos, não paga o custo cheio do atendimento. É a mesma conta e o mesmo número do KPI \"Tabelas de convênio abaixo do custo cheio + imposto\" da planilha 08. \"Maior prejuízo\" = margem negativa × tabela do convênio: a célula mais vermelha da tabela, em reais por atendimento (colunas auxiliares T a W, ocultas).").font=F(size=9,color=LILAS)
s.merge_cells(start_row=R0+6,start_column=1,end_row=R0+6,end_column=19); s.cell(row=R0+6,column=1).alignment=Alignment(wrap_text=True,vertical="top"); s.row_dimensions[R0+6].height=28
# simular um procedimento novo
Q0=R0+8
s.cell(row=Q0,column=1,value="Simular um procedimento ou um preço novo").font=F(bold=True,size=13,color=UVA)
sim=[("Nome","Consulta particular estendida (45 min)",None),("Minutos",45,"0"),("Gera retorno? (Sim/Não)","Sim",None),("Material (R$)",4,BRL),("Preço pretendido (R$)",480,BRL0)]
for i,(a,v,fmt) in enumerate(sim):
    r=Q0+1+i; s.cell(row=r,column=1,value=a); rotulo(s.cell(row=r,column=1),bold=False); s.cell(row=r,column=1).border=borda; s.cell(row=r,column=2,value=v); inp(s.cell(row=r,column=2),fmt,center=fmt is not None)
dvq=lista('"Sim,Não"'); dvq.add(f"B{Q0+3}"); s.add_data_validation(dvq)
out=[("Tempo com retorno (min)",f'={f_tempo(f"B{Q0+2}",f"B{Q0+3}",NRET,DRET)}',"0.0"),("Custo cheio (R$)",f'={f_custo(f"B{Q0+6}",f"B{Q0+4}",CH)}',BRL),
     ("Preço mínimo (R$)",f'=IF(NOT(ISNUMBER(B{Q0+7})),B{Q0+7},IF({CFGX},"margens inválidas (Config)",B{Q0+7}/{DIVMIN}))',BRL),("Preço alvo (R$)",f'=IF(NOT(ISNUMBER(B{Q0+7})),B{Q0+7},IF({CFGX},"margens inválidas (Config)",B{Q0+7}/{DIVALVO}))',BRL),
     ("Margem no preço pretendido",f'=IF(OR(NOT({num_ok(f"B{Q0+5}")}),N(B{Q0+5})=0,NOT(ISNUMBER(B{Q0+7})),{CFGX}),"",(B{Q0+5}*(1-{IMP})-B{Q0+7})/B{Q0+5})',PCT),
     ("Leitura",f'=IF(B{Q0+5}="","",IF(NOT({num_ok(f"B{Q0+5}")}),"Preço pretendido inválido",IF(B{Q0+5}=0,"",IF(NOT(ISNUMBER(B{Q0+7})),IF(IF(ISNUMBER({CH}),{CH}>=0,FALSE),"Faltam dados: "&B{Q0+7},"Falta o custo-hora em Config"),IF(OR(NOT(ISNUMBER(B{Q0+8})),NOT(ISNUMBER(B{Q0+9}))),"Margens inválidas em Config",IF(B{Q0+5}<B{Q0+8},"Abaixo do mínimo: cobre parte do custo, não a margem",IF(B{Q0+5}<B{Q0+9},"Entre o mínimo e o alvo","No alvo ou acima")))))))',"@")]
for i,(a,f_,fmt) in enumerate(out):
    r=Q0+6+i; s.cell(row=r,column=1,value=a); calc(s.cell(row=r,column=1),center=False); s.cell(row=r,column=2,value=f_); calc(s.cell(row=r,column=2),fmt)
    if fmt=="@": s.merge_cells(start_row=r,start_column=2,end_row=r,end_column=6); s.cell(row=r,column=2).alignment=Alignment(horizontal="left")
s.cell(row=Q0+8,column=2).font=F(bold=True,color=UVA,size=10)
bc=BarChart(); bc.type="col"; bc.grouping="clustered"; bc.height=7.5; bc.width=20; bc.title="Preço mínimo × alvo × particular praticado"; bc.style=2
bc.add_data(Reference(s,min_col=7,max_col=9,min_row=T0-1,max_row=T0+len(dados.PROCEDIMENTOS)-1),titles_from_data=True); bc.set_categories(Reference(s,min_col=1,min_row=T0,max_row=T0+len(dados.PROCEDIMENTOS)-1))
for sr,cor in zip(bc.series,("B89BE0",UVA,SOL)): sr.graphicalProperties.solidFill=cor
bc.legend.position="b"; bc.y_axis.majorGridlines=None; bc.y_axis.number_format='"R$" #,##0'
s.add_chart(bc,f"H{R0}")
widths(s,(24,8,10,10,11,11,11,11,12,10,18,11,9,11,9,11,9,11,9)); s.freeze_panes="B5"; s.sheet_view.showGridLines=False
for col in "TUVW": s.column_dimensions[col].hidden=True
# ---------- exemplo ----------
for i,(pnome,d,m,ret) in enumerate(dados.PROCEDIMENTOS):
    r=T0+i
    s.cell(row=r,column=1,value=pnome); s.cell(row=r,column=2,value=d); s.cell(row=r,column=3,value="Sim" if ret>0 else "Não"); s.cell(row=r,column=4,value=m)
    s.cell(row=r,column=9,value=dados.preco(pnome,"Particular"))
    for j,(cn,_,_) in enumerate(dados.CONVENIOS):
        v=dados.preco(pnome,cn)
        if v is not None: s.cell(row=r,column=12+2*j,value=v)
como_usar(wb,"Precificação de consulta e procedimento",[
 ("O que esta planilha faz","Calcula o custo cheio de cada consulta e procedimento (tempo × custo-hora + material, com o tempo de retorno embutido nas consultas), o preço mínimo (margem mínima) e o preço alvo (margem alvo), e mostra a margem do preço particular praticado e de cada tabela de convênio."),
 ("Passo 1","Em Config, copie o custo-hora da planilha 05, a alíquota de impostos e as margens mínima e alvo. Ajuste a média de retornos por consulta e a duração do retorno. Liste até 4 convênios."),
 ("Passo 2","Em Precificação, uma linha por procedimento: minutos, se gera retorno, material por atendimento, o preço particular que você pratica hoje e a tabela de cada convênio. Copie os valores da planilha 08 · Tabela de preços: ela é a fonte única da tabela no kit (a 07 e a Config da 01 também copiam de lá)."),
 ("Passo 3","Leia a Situação e as margens: vermelho não cobre o custo cheio; amarelo cobre o custo mas não a margem mínima; verde está na faixa. Use o bloco \"Simular\" para testar um procedimento ou um preço novo antes de mudar a tabela."),
 ("Rotina","Revise quando o custo-hora mudar (05) ou uma vez por semestre. Decidiu mudar um preço? Atualize na ordem: 08 (tabela de preços, a fonte) → 06 (esta) → 07 (simulador) → 01 (Config da agenda). Assim os quatro arquivos continuam com o mesmo número."),
 ("Convênio abaixo do custo","Não decida só por esta tela: com horários vazios, uma tabela baixa ainda pode contribuir; com agenda cheia, toma o lugar de um particular. A planilha 07 faz essa conta com prazo de pagamento e glosa."),
 ("Com a IA","Copie a tabela Precificação (sem os nomes dos convênios, se preferir) e use o prompt \"Preço 02 · Revisar a tabela de preços pela margem\" da biblioteca do kit. A IA ajuda a explicar o número ao sócio; o preço é decisão da clínica."),
])
# Auditoria final-6 (G01): custo-hora aceita só número >= 0 ao digitar; a fórmula também confere.
_dvch=DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True,showErrorMessage=True,errorTitle="Custo-hora",error="Número maior ou igual a zero."); _dvch.add("B7"); cfg.add_data_validation(_dvch)
proteger(wb); salvar(wb,"06-precificacao.xlsx","Precificação de consulta e procedimento · Kit de Gestão para Médicos")
