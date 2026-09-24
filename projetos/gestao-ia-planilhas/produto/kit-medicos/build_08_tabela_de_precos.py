#!/usr/bin/env python3
"""Planilha 8 do Kit de Gestão para Médicos: Tabela de preços e referência por procedimento. Gera 08-tabela-de-precos.xlsx
A tabela interna: preço particular e tabelas de convênio por procedimento contra o mínimo e o alvo (custo-hora da 05, mesma conta da 06),
o valor por hora de cada pagador e o volume do mês fechado (agosto, Painel da 01) para ver o preço médio praticado."""
from ssg import *
from openpyxl.chart import BarChart, Reference
import dados
NPROC=12; NPAG=6; T0=8; TN=T0+NPROC-1     # KPIs nas linhas 4-5, cabeçalho da tabela na linha 7, dados a partir da 8
wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Custo-hora da planilha 05; margens e impostos iguais aos da 06.",merge_to="H")
campos=[("Nome da clínica",f"{dados.CLINICA} (exemplo fictício)",None),("Mês do volume (mês fechado)","Agosto de 2026",None),("Data de referência",dados.HOJE,DATA),
 ("Custo da hora de atendimento (R$)",dados.CUSTO_HORA,BRL),("Impostos sobre o que entra (%)",dados.ALIQ,PCT),("Margem mínima sobre o preço (%)",dados.MARGEM,PCT),("Margem alvo sobre o preço (%)",dados.MARGEM_ALVO,PCT),
 ("Retornos por consulta (média)",dados.RETORNO_PROB,"0.00"),("Duração do retorno (minutos)",dados.DUR["Retorno"],"0")]
for i,(a,v,fmt) in enumerate(campos):
    r=4+i; cfg.cell(row=r,column=1,value=a); rotulo(cfg.cell(row=r,column=1)); cfg.cell(row=r,column=2,value=v)
    if r==6: calc(cfg.cell(row=r,column=2),fmt)
    else: inp(cfg.cell(row=r,column=2),fmt,center=fmt is not None)
# Auditoria final-2 (G01/G02): mesmo predicado da 06, visível em Config!B15; custo-hora em branco
# não é hora mínima zero. Custo cheio, preços, situação, tabelas abaixo do custo e KPIs consultam
# B13/B15 e dizem o que falta em vez de calcular com Config fora do domínio.
cfg["A15"]="Impostos e margens conferem? (calculado)"; rotulo(cfg["A15"])
cfg["B15"]='=IF(AND(ISNUMBER(B8),ISNUMBER(B9),ISNUMBER(B10)),IF(AND(B8>=0,B8<1,B9>=0,B9<=B10,B10<1,B8+B10<1),"Sim","Não"),"Não")'; calc(cfg["B15"])
cfg["C15"]="\"Não\" suspende preços, situação e contagens: impostos e margens de 0 % a 99 %, mínima ≤ alvo e impostos + margem alvo abaixo de 100 %."; nota(cfg["C15"])
for _c,_f,_m in (("B8",'=IF(ISNUMBER(B8),AND(B8>=0,B8<1,B8+N(B10)<1),FALSE)',"Impostos entre 0 % e 99 %, e impostos + margem alvo abaixo de 100 %."),
                 ("B9",'=IF(ISNUMBER(B9),AND(B9>=0,B9<=N(B10),N(B8)+B9<1),FALSE)',"Margem mínima entre 0 % e a margem alvo, e impostos + margem abaixo de 100 %."),
                 ("B10",'=IF(ISNUMBER(B10),AND(B10>=N(B9),B10<1,N(B8)+B10<1),FALSE)',"Margem alvo entre a margem mínima e 99 %, e impostos + margem alvo abaixo de 100 %.")):
    _dv=DataValidation(type="custom",formula1=_f,allow_blank=False,showErrorMessage=True,errorTitle="Percentual",error=_m); _dv.add(_c); cfg.add_data_validation(_dv)
cfg["A13"]="Hora mínima a cobrar (R$)"; cfg["B13"]='=IF(NOT(ISNUMBER(B7)),"falta o custo-hora",IF(B15<>"Sim","margens inválidas (Config)",B7/(1-B8-B9)))'
cfg["A14"]="Hora alvo (R$)"; cfg["B14"]='=IF(NOT(ISNUMBER(B7)),"falta o custo-hora",IF(B15<>"Sim","margens inválidas (Config)",B7/(1-B8-B10)))'
for r in (13,14): rotulo(cfg.cell(row=r,column=1)); calc(cfg.cell(row=r,column=2),BRL); cfg.cell(row=r,column=2).font=F(bold=True,color=UVA,size=10)
cfg["C7"]="Copie do Painel da planilha 05 (\"Custo da hora de atendimento\"). No exemplo, R$ 200,00."; cfg["C8"]="A mesma alíquota efetiva das planilhas 05, 06 e 07 (exemplo: 11 %). Só imposto: a taxa da maquininha não entra aqui (fica na 16 e no resultado 18)."
cfg["C9"]="Piso da faixa (a mesma margem da 05 e da 06)."; cfg["C10"]="Teto da faixa (a mesma da 06)."; cfg["C11"]="Iguais à 06: o tempo do retorno entra no custo da consulta."
cfg["C13"]="Custo-hora ÷ (1 − impostos − margem mínima) = R$ 338,98 no exemplo (a 05 arredonda para R$ 340)."; cfg["C14"]="Custo-hora ÷ (1 − impostos − margem alvo)."
for r in (7,8,9,10,11,13,14): nota(cfg.cell(row=r,column=3))
cfg["A16"]="Pagadores (até 6, na ordem das colunas da Tabela; o primeiro é o particular)"; rotulo(cfg["A16"])
cfg["A23"]="Esta planilha é a FONTE ÚNICA da tabela de preços do kit. A aba Tabela é onde o preço é decidido; 06 (precificação), 07 (simulador) e a Config da planilha 01 (agenda) copiam daqui."; nota(cfg["A23"])
cfg["A24"]="Ordem de atualização quando um preço muda: 1) aqui, na aba Tabela; 2) planilha 06 · Precificação (colunas do particular e dos convênios); 3) planilha 07 · Simulador (quadro 2, valor de tabela); 4) planilha 01 · Config (tabela de preços por pagador, que é o que a Agenda busca). Procedimentos novos: cadastre na 01 (Config), na 02, aqui, na 06 e na 07. Convênio novo: 01, 02, 06, 07, aqui, 09 (categorias de entrada) e 13 (Config)."; nota(cfg["A24"])
for i in range(NPAG): inp(cfg.cell(row=17+i,column=2))
for i,pg in enumerate(dados.PAGADORES): cfg.cell(row=17+i,column=2,value=pg)
cfg["A17"]="Nomes"; rotulo(cfg["A17"],bold=False); cfg["C17"]="Preencha de cima para baixo, sem pular linha. Os mesmos da Config da planilha 01."; nota(cfg["C17"])
widths(cfg,(52,16,90)); cfg.sheet_view.showGridLines=False
CH="Config!$B$7"; IMP="Config!$B$8"; MMIN="Config!$B$9"; MALVO="Config!$B$10"; NRET="Config!$B$11"; DRET="Config!$B$12"; HMIN="Config!$B$13"; HALVO="Config!$B$14"; CFGX='(Config!$B$15<>"Sim")'
# ---------- Tabela ----------
t=wb.create_sheet("Tabela",0)
titulo(t,'=Config!$B$4&" · Tabela de preços e referência · volume de "&Config!$B$5',"Amarelo: procedimento, minutos, retorno, material, valor de cada pagador e o volume do mês (Painel da 01, por procedimento). Branco: custo cheio, mínimo, alvo, situação e valor médio praticado.",merge_to="T")
PAG_H=[f'=IF(Config!$B${17+j}="","",Config!$B${17+j})' for j in range(NPAG)]
heads=["Procedimento","Minutos","Gera retorno?","Material (R$)","Custo cheio (R$)","Preço mínimo (R$)","Preço alvo (R$)"]+PAG_H+["Situação do particular","Tabelas abaixo do custo cheio + imposto","Realizados no mês","Produção no mês (R$)","Valor médio praticado (R$)","Valor médio contra o mínimo","Referência de mercado (opcional)"]
hdr(t,T0-1,heads,height=44)
for r in range(T0,TN+1):
    inp(t.cell(row=r,column=1)); inp(t.cell(row=r,column=2),"0",center=True); inp(t.cell(row=r,column=3),center=True); inp(t.cell(row=r,column=4),BRL,center=True)
    # Auditoria final-4 (F4-G01): minutos, retorno e material vazios ou em texto não viram zero
    t.cell(row=r,column=5,value=f'=IF(A{r}="","",{f_custo("("+f_tempo(f"B{r}",f"C{r}",NRET,DRET)+")",f"D{r}",CH)})'); calc(t.cell(row=r,column=5),BRL)
    t.cell(row=r,column=6,value=f'=IF(A{r}="","",IF(NOT(ISNUMBER(E{r})),E{r},IF({CFGX},"margens inválidas (Config)",E{r}/(1-{IMP}-{MMIN}))))'); calc(t.cell(row=r,column=6),BRL); t.cell(row=r,column=6).font=F(bold=True,color=UVA,size=10)
    t.cell(row=r,column=7,value=f'=IF(A{r}="","",IF(NOT(ISNUMBER(E{r})),E{r},IF({CFGX},"margens inválidas (Config)",E{r}/(1-{IMP}-{MALVO}))))'); calc(t.cell(row=r,column=7),BRL)
    for j in range(NPAG): inp(t.cell(row=r,column=8+j),BRL0,center=True)
    t.cell(row=r,column=14,value=f'=IF(OR(A{r}="",H{r}=""),"",IF(NOT(ISNUMBER(H{r})),"Preço inválido",IF(NOT(ISNUMBER(E{r})),IF(ISNUMBER({CH}),"Faltam dados da linha","Falta o custo-hora"),IF(OR(NOT(ISNUMBER(F{r})),NOT(ISNUMBER(G{r}))),"Margens inválidas",IF(H{r}=0,"Sem cobrança",IF(H{r}<F{r},"Abaixo do mínimo",IF(H{r}<G{r},"Entre mínimo e alvo","No alvo ou acima")))))))'); calc(t.cell(row=r,column=14))
    t.cell(row=r,column=15,value=f'=IF(OR(A{r}="",NOT(ISNUMBER(F{r}))),"",IF(SUMPRODUCT(ISTEXT(I{r}:M{r})*1)>0,"tabela com texto",SUMPRODUCT((I{r}:M{r}<>"")*(I{r}:M{r}>0)*(I{r}:M{r}*(1-{IMP})<E{r}))))'); calc(t.cell(row=r,column=15),"0")
    inp(t.cell(row=r,column=16),"0",center=True); inp(t.cell(row=r,column=17),BRL0,center=True)
    t.cell(row=r,column=18,value=f'=IF(OR(A{r}="",AND(P{r}="",Q{r}="")),"",IF(OR(NOT(ISNUMBER(P{r})),NOT(ISNUMBER(Q{r}))),"volume ou produção faltando",IF(P{r}=0,"",Q{r}/P{r})))'); calc(t.cell(row=r,column=18),BRL)
    t.cell(row=r,column=19,value=f'=IF(OR(NOT(ISNUMBER(R{r})),NOT(ISNUMBER(F{r})),N(H{r})=0),"",IF(F{r}=0,"sem base (mínimo zero)",R{r}/F{r}-1))'); calc(t.cell(row=r,column=19),"+0%;-0%;0%")
    inp(t.cell(row=r,column=20))
dvs=lista('"Sim,Não"'); dvs.add(f"C{T0}:C{TN}"); t.add_data_validation(dvs)
t.conditional_formatting.add(f"N{T0}:N{TN}", FormulaRule(formula=[f'OR(N{T0}="Abaixo do mínimo",N{T0}="Falta o custo-hora",N{T0}="Margens inválidas",N{T0}="Faltam dados da linha",N{T0}="Preço inválido")'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
t.conditional_formatting.add(f"N{T0}:N{TN}", FormulaRule(formula=[f'N{T0}="Entre mínimo e alvo"'], fill=fill(AMARELO)))
t.conditional_formatting.add(f"N{T0}:N{TN}", FormulaRule(formula=[f'N{T0}="No alvo ou acima"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
t.conditional_formatting.add(f"I{T0}:M{TN}", FormulaRule(formula=[f'AND(ISNUMBER(I{T0}),I{T0}>0,ISNUMBER($F{T0}),I{T0}*(1-{IMP})<$E{T0})'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
t.conditional_formatting.add(f"I{T0}:M{TN}", FormulaRule(formula=[f'AND(ISNUMBER(I{T0}),ISNUMBER($F{T0}),I{T0}*(1-{IMP})>=$E{T0},I{T0}<$F{T0})'], fill=fill(AMARELO)))
t.conditional_formatting.add(f"N{T0}:N{TN}", FormulaRule(formula=[f'N{T0}="Sem cobrança"'], font=F(color="8A86A0",size=10)))
t.conditional_formatting.add(f"O{T0}:O{TN}", FormulaRule(formula=[f'AND(ISNUMBER(O{T0}),O{T0}>0)'], font=F(color="C8402E",size=10,bold=True)))
t.conditional_formatting.add(f"S{T0}:S{TN}", FormulaRule(formula=[f'AND(ISNUMBER(S{T0}),S{T0}<0)'], font=F(color="C8402E",size=10,bold=True)))
KA=f"$A${T0}:$A${TN}"; KN=f"$N${T0}:$N${TN}"; KO=f"$O${T0}:$O${TN}"; KP=f"$P${T0}:$P${TN}"; KQ=f"$Q${T0}:$Q${TN}"; KB=f"$B${T0}:$B${TN}"; KF=f"$F${T0}:$F${TN}"
# Auditoria final-4 (F4-G01/G02): contagens e médias só com todas as linhas completas; a receita
# por hora exige, em cada linha com volume ou produção, os dois números e os minutos.
NINC=(f'(SUMPRODUCT(({KA}<>"")*(ISNUMBER({KF})=FALSE))+SUMPRODUCT(ISTEXT($H${T0}:$M${TN})*1))')
INCV=(f'SUMPRODUCT((({KP}<>"")+({KQ}<>"")>0)*((ISNUMBER({KP})=FALSE)+(ISNUMBER({KQ})=FALSE)+(ISNUMBER({KB})=FALSE)>0))')
kpi(t,4,1,"Hora mínima (05)",f"={HMIN}",LAVANDA,UVA,fmt=BRL)
kpi(t,4,3,"Particular abaixo do mínimo",f'=IF(NOT(ISNUMBER({HMIN})),{HMIN},IF({NINC}>0,{NINC}&" linha(s) ou preço(s) incompletos",COUNTIF({KN},"Abaixo do mínimo")&" de "&COUNTA({KA})))',VERM,VERM_T,fmt="@")
kpi(t,4,5,"Tabelas de convênio abaixo do custo cheio + imposto",f'=IF(NOT(ISNUMBER({HMIN})),{HMIN},IF({NINC}>0,{NINC}&" linha(s) ou preço(s) incompletos",SUM({KO})))',VERM,VERM_T,fmt="0")
kpi(t,4,7,"Produção no mês (R$)",f'=IF(SUMPRODUCT(ISTEXT({KQ})*1)>0,"produção com texto",SUM({KQ}))',VERDE,VERDE_T,fmt=BRL0)
kpi(t,4,9,"Valor médio por hora no mês",f'=IF({INCV}>0,"volume incompleto em "&{INCV}&" linha(s)",IFERROR(SUM({KQ})/(SUMPRODUCT({KP},{KB})/60),""))',SOL,UVA,fmt=BRL)
kpi(t,4,11,"Contra a hora mínima",f'=IF(NOT(ISNUMBER({HMIN})),"",IF({HMIN}=0,"sem base (mínimo zero)",IFERROR(I5/{HMIN}-1,"")))',LAVANDA,UVA,fmt="+0%;-0%;0%")
NT=TN+2
notas=["Custo cheio, preço mínimo e preço alvo: a mesma conta da planilha 06 (tempo com retorno × custo-hora + material; ÷ (1 − impostos − margem)). Tabela de convênio em vermelho: depois dos impostos não cobre o custo cheio (margem negativa) — é a mesma conta e o mesmo número do KPI \"Tabelas de convênio abaixo do custo cheio + imposto\" e do Resumo da planilha 06. Amarela: cobre o custo, mas não a margem mínima.",
       "Realizados e Produção no mês: copie de \"Por procedimento\" no Painel da planilha 01 com o mês fechado escolhido em Config (no exemplo, agosto de 2026). Valor médio praticado = produção ÷ realizados: mistura particular e convênio, por isso fica abaixo do preço particular.",
       "Referência de mercado é opcional e sua: pesquise a faixa da sua região e anote; a planilha não afirma preço de mercado. A tabela de preços e a divulgação de valores seguem as regras do CFM: revise antes de publicar."]
for i,tx in enumerate(notas):
    c=t.cell(row=NT+i,column=1,value=tx); c.font=F(size=9,color=LILAS); t.merge_cells(start_row=NT+i,start_column=1,end_row=NT+i,end_column=20); c.alignment=Alignment(wrap_text=True,vertical="top"); t.row_dimensions[NT+i].height=30
# valor por hora por pagador
H0=NT+4
t.cell(row=H0,column=1,value="Valor por hora de atendimento, por pagador (tabela ÷ tempo com retorno)").font=F(bold=True,size=13,color=UVA)
hdr(t,H0+1,["Procedimento","Minutos com retorno"]+PAG_H+["Hora mínima","Hora alvo"],height=32)
for i in range(NPROC):
    r=H0+2+i; s=T0+i
    t.cell(row=r,column=1,value=f'=IF($A{s}="","",$A{s})'); calc(t.cell(row=r,column=1),center=False)
    t.cell(row=r,column=2,value=f'=IF($A{s}="","",{f_tempo(f"B{s}",f"C{s}",NRET,DRET)})'); calc(t.cell(row=r,column=2),"0.0")
    for j in range(NPAG):
        c=3+j; src=f"{L(8+j)}{s}"
        t.cell(row=r,column=c,value=f'=IF(OR($A{s}="",{src}="",NOT(ISNUMBER({src})),NOT(ISNUMBER($B{r})),N($B{r})=0),"",{src}/($B{r}/60))'); calc(t.cell(row=r,column=c),BRL0)
    t.cell(row=r,column=9,value=f'=IF($A{s}="","",{HMIN})'); calc(t.cell(row=r,column=9),BRL0)
    t.cell(row=r,column=10,value=f'=IF($A{s}="","",{HALVO})'); calc(t.cell(row=r,column=10),BRL0)
t.conditional_formatting.add(f"C{H0+2}:H{H0+1+NPROC}", FormulaRule(formula=[f'AND(ISNUMBER(C{H0+2}),C{H0+2}>0,ISNUMBER({HMIN}),C{H0+2}<{HMIN})'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
t.conditional_formatting.add(f"C{H0+2}:H{H0+1+NPROC}", FormulaRule(formula=[f'AND(ISNUMBER(C{H0+2}),ISNUMBER({HALVO}),C{H0+2}>={HALVO})'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
t.cell(row=H0+2+NPROC,column=1,value="Vermelho: abaixo da hora mínima. Verde: na hora alvo ou acima. Retorno aparece vazio (valor zero). Este quadro é o argumento na negociação de tabela: quanto cada convênio paga pela hora da clínica.").font=F(size=9,color=LILAS)
t.merge_cells(start_row=H0+2+NPROC,start_column=1,end_row=H0+2+NPROC,end_column=20)
bc=BarChart(); bc.type="bar"; bc.grouping="clustered"; bc.height=8; bc.width=16; bc.title="Valor por hora: particular × convênios (R$)"; bc.style=2
bc.add_data(Reference(t,min_col=3,max_col=6,min_row=H0+1,max_row=H0+1+len(dados.PROCEDIMENTOS)),titles_from_data=True); bc.set_categories(Reference(t,min_col=1,min_row=H0+2,max_row=H0+1+len(dados.PROCEDIMENTOS)))
for sr,cor in zip(bc.series,(UVA,"B89BE0",LILAS,SOL)): sr.graphicalProperties.solidFill=cor
bc.legend.position="b"; bc.x_axis.majorGridlines=None
t.add_chart(bc,f"L{H0}")
widths(t,(24,8,9,10,11,11,11,11,11,11,11,9,9,18,11,10,13,12,11,22)); t.freeze_panes=f"B{T0}"; t.sheet_view.showGridLines=False
# ---------- exemplo ----------
for i,(pnome,d,m,ret) in enumerate(dados.PROCEDIMENTOS):
    r=T0+i
    t.cell(row=r,column=1,value=pnome); t.cell(row=r,column=2,value=d); t.cell(row=r,column=3,value="Sim" if ret>0 else "Não"); t.cell(row=r,column=4,value=m)
    for j,pg in enumerate(dados.PAGADORES):
        v=dados.preco(pnome,pg)
        if v is not None: t.cell(row=r,column=8+j,value=v)
    L_=[x for x in dados.AGENDA_TODA if x["situacao"]=="Realizado" and x["procedimento"]==pnome and x["data"].month==8]
    t.cell(row=r,column=16,value=len(L_)); t.cell(row=r,column=17,value=sum(x["valor"] for x in L_))
como_usar(wb,"Tabela de preços e referência por procedimento",[
 ("O que esta planilha faz","É a tabela interna da clínica: para cada procedimento, o preço particular e a tabela de cada convênio ao lado do custo cheio, do preço mínimo e do preço alvo (a mesma conta da 06), mais o valor por hora que cada pagador paga e o valor médio praticado no mês fechado."),
 ("Passo 1","Em Config, copie o custo-hora da planilha 05, a alíquota, as margens e os pagadores (o primeiro é o particular)."),
 ("Passo 2","Em Tabela, uma linha por procedimento: minutos, se gera retorno, material e o valor de cada pagador. Esta é a tabela oficial da clínica: a Config da planilha 01 (agenda), a 06 e a 07 copiam daqui."),
 ("Passo 3","Copie de \"Por procedimento\" no Painel da 01 (mês fechado) os realizados e a produção do mês: a planilha mostra o valor médio praticado e quanto ele fica acima ou abaixo do mínimo."),
 ("Passo 4","Leia as cores: particular abaixo do mínimo (vermelho) é preço a revisar; tabela de convênio vermelha não cobre o custo cheio (leve para a planilha 07 antes de decidir). O quadro \"Valor por hora\" é o argumento para negociar tabela."),
 ("Ordem de atualização","Mudou um preço? Atualize sempre nesta ordem: 08 (aqui) → 06 · Precificação → 07 · Simulador → 01 · Config (a tabela que a Agenda busca). Procedimento novo: 01 e 02 (Config), 08, 06 e 07. Convênio novo: 01, 02, 06, 07, 08, 09 (categorias de entrada) e 13 (Config, com prazo e dia de envio)."),
 ("Rotina","Revise quando o custo-hora mudar (05) ou ao renegociar um convênio. Se pesquisar preços da região, anote na coluna \"Referência de mercado\"; a divulgação de preços segue as regras do CFM."),
 ("Com a IA","Copie a Tabela e use o prompt \"Preço 04 · Revisar a tabela de preços\" da biblioteca do kit, ou \"Preço 03 · Vale a pena este convênio?\" para um convênio específico."),
])
proteger(wb); salvar(wb,"08-tabela-de-precos.xlsx","Tabela de preços e referência · Kit de Gestão para Médicos")
