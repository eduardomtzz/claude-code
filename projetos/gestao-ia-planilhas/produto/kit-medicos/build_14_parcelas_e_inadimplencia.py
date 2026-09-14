#!/usr/bin/env python3
"""Planilha 14 do Kit de Gestão para Médicos: Parcelas particulares e inadimplência com régua de cobrança.
Gera 14-parcelas-e-inadimplencia.xlsx (Como usar, Painel, Config, Parcelas, Pacientes).
Só o que foi combinado a prazo (dados.PARCELAS): o à vista entra direto no caixa (09). Parcela paga = entrada no caixa na data do pagamento."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference

NP=400; R0=5; RNP=R0+NP-1          # Parcelas: linhas 5..404
NPC=400; RNC=R0+NPC-1              # Pacientes: linhas 5..404
RG0=10; RG1=13                     # Régua na Config: linhas 10..13
def off(sheet,col,r0,r1): return f"=OFFSET({sheet}!${col}${r0},0,0,MAX(1,COUNTA({sheet}!${col}${r0}:${col}${r1})),1)"
def dstr(ref): return f'TEXT(DAY({ref}),"00")&"/"&TEXT(MONTH({ref}),"00")&"/"&YEAR({ref})'
NOME=f"{dados.CLINICA} (exemplo fictício)"
HOJE="Config!$B$5"; JC="Config!$B$6"; JL="Config!$B$7"
RD=f"Config!$A${RG0}:$A${RG1}"; RA=f"Config!$B${RG0}:$B${RG1}"; RF=f"Config!$C${RG0}:$C${RG1}"

wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações e régua de cobrança","Células amarelas: você preenche. A régua diz o que fazer em cada faixa de atraso; o texto aparece no Painel, parcela por parcela.",merge_to="H")
cfg["A4"]="Clínica"; cfg["B4"]=NOME
cfg["A5"]="Data de referência (hoje)"; cfg["B5"]="=TODAY()"
cfg["A6"]="Janela curta: vence em até (dias)"; cfg["B6"]=7
cfg["A7"]="Janela longa: vence em até (dias)"; cfg["B7"]=30
for r in (4,5,6,7): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"],DATA); inp(cfg["B6"],center=True); inp(cfg["B7"],center=True)
hdr(cfg,9,["A partir de (dias de atraso)","Ação sugerida","Faixa (calculada)"])
for r in range(RG0,RG1+1):
    inp(cfg.cell(row=r,column=1),center=True); inp(cfg.cell(row=r,column=2)); cfg.cell(row=r,column=2).alignment=Alignment(wrap_text=True,vertical="top")
    if r<RG1: cfg.cell(row=r,column=3,value=f'=IF(A{r}="","",IF(A{r+1}-1=A{r},A{r}&" dia"&IF(A{r}>1,"s",""),A{r}&" a "&(A{r+1}-1)&" dias"))')
    else: cfg.cell(row=r,column=3,value=f'=IF(A{r}="","",A{r}&" dias ou mais")')
    calc(cfg.cell(row=r,column=3)); cfg.row_dimensions[r].height=44
cfg.cell(row=RG1+2,column=1,value="Os dias precisam estar em ordem crescente (ex.: 1, 7, 15, 30). Cada parcela vencida recebe a ação da maior faixa que ela já alcançou. Escreva as ações como você fala com o paciente: educado, direto e sem ameaça. Cobrança de paciente tem regras próprias (CDC e ética médica): nada de exposição ou constrangimento.").font=F(size=9,color=LILAS)
cfg.cell(row=RG1+3,column=1,value="Os modelos de mensagem do bônus \"15 modelos de mensagem\" seguem esta mesma régua.").font=F(size=9,color=LILAS)
cfg.cell(row=RG1+4,column=1,value="Regra que vale mais que a régua: quem tem parcela vencida em aberto não combina novo pagamento a prazo. O próximo atendimento é à vista (Pix, dinheiro ou cartão) ou fica para depois de acertar o que está em atraso — sem constrangimento e em particular. No exemplo do kit, nenhum paciente com parcela vencida recebeu crédito novo, e o mesmo vale para orçamento aprovado a prazo (planilha 15).").font=F(size=9,color=LILAS)
cfg.cell(row=RG1+4,column=1).alignment=Alignment(wrap_text=True,vertical="top"); cfg.merge_cells(start_row=RG1+4,start_column=1,end_row=RG1+4,end_column=3); cfg.row_dimensions[RG1+4].height=30
widths(cfg,(30,70,18,3,3,3,3,3)); cfg.sheet_view.showGridLines=False

# ---------- Pacientes ----------
pc=wb.create_sheet("Pacientes")
titulo(pc,"Pacientes (entrada)","Copie da planilha 01 (aba Pacientes; a 01 é a fonte): nome e contato. As colunas brancas somam as parcelas de cada paciente.",merge_to="I")
hdr(pc,4,["Paciente","Contato","Parcelas","Valor em parcelas (R$)","Pago (R$)","Em aberto (R$)","Vencido (R$)","Situação"])
PA=f"Parcelas!$A${R0}:$A${RNP}"; PF=f"Parcelas!$F${R0}:$F${RNP}"; PI=f"Parcelas!$I${R0}:$I${RNP}"
for r in range(R0,RNC+1):
    inp(pc.cell(row=r,column=1)); inp(pc.cell(row=r,column=2),center=True)
    pc.cell(row=r,column=3,value=f'=IF(A{r}="","",COUNTIFS({PA},A{r}))'); calc(pc.cell(row=r,column=3))
    pc.cell(row=r,column=4,value=f'=IF(A{r}="","",SUMIFS({PF},{PA},A{r}))'); calc(pc.cell(row=r,column=4),BRL0)
    pc.cell(row=r,column=5,value=f'=IF(A{r}="","",SUMIFS({PF},{PA},A{r},{PI},"Paga"))'); calc(pc.cell(row=r,column=5),BRL0)
    pc.cell(row=r,column=6,value=f'=IF(A{r}="","",D{r}-E{r})'); calc(pc.cell(row=r,column=6),BRL0)
    pc.cell(row=r,column=7,value=f'=IF(A{r}="","",SUMIFS({PF},{PA},A{r},{PI},"Vencida"))'); calc(pc.cell(row=r,column=7),BRL0)
    pc.cell(row=r,column=8,value=f'=IF(A{r}="","",IF(G{r}>0,"Em atraso",IF(F{r}>0,"Em dia","Sem parcela em aberto")))'); calc(pc.cell(row=r,column=8))
pc.conditional_formatting.add(f"H{R0}:H{RNC}", FormulaRule(formula=[f'H{R0}="Em atraso"'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
pc.cell(row=RNC+2,column=1,value="Cadastre aqui todos os pacientes que já pagaram a prazo alguma vez (a lista alimenta a escolha do paciente em Parcelas). Só nome e contato: nada clínico.").font=F(size=9,color=LILAS)
widths(pc,(28,16,9,16,14,14,14,20)); pc.freeze_panes="B5"; pc.sheet_view.showGridLines=False; pc.auto_filter.ref=f"A4:H{RNC}"

# ---------- Parcelas ----------
pr=wb.create_sheet("Parcelas")
titulo(pr,"Parcelas","Uma linha por parcela combinada a prazo. Escolha o paciente na lista; situação, dias e ação da régua são calculados.",merge_to="M")
hdr(pr,4,["Paciente","Procedimento","Data do atendimento","Nº da parcela","Vencimento","Valor (R$)","Pago?","Data do pagamento","Situação","Dias de atraso","Dias para vencer","Faixa de atraso","Ação sugerida (régua)"],height=32)
for r in range(R0,RNP+1):
    for c in (1,2,3,4,5,6,7,8): inp(pr.cell(row=r,column=c))
    for c in (3,4,5,7,8): pr.cell(row=r,column=c).alignment=Alignment(horizontal="center")
    pr.cell(row=r,column=3).number_format=DATA; pr.cell(row=r,column=5).number_format=DATA; pr.cell(row=r,column=8).number_format=DATA; pr.cell(row=r,column=6).number_format=BRL0
    pr.cell(row=r,column=9,value=f'=IF(A{r}="","",IF(G{r}="Sim","Paga",IF(E{r}="","Sem vencimento",IF(E{r}<{HOJE},"Vencida","A vencer"))))'); calc(pr.cell(row=r,column=9))
    pr.cell(row=r,column=10,value=f'=IF(I{r}="Vencida",{HOJE}-E{r},"")'); calc(pr.cell(row=r,column=10),"0")
    pr.cell(row=r,column=11,value=f'=IF(I{r}="A vencer",E{r}-{HOJE},"")'); calc(pr.cell(row=r,column=11),"0")
    nivel=f'COUNTIF({RD},"<="&J{r})'
    pr.cell(row=r,column=12,value=f'=IF(I{r}<>"Vencida","",IF({nivel}=0,"Antes da régua",INDEX({RF},{nivel})))'); calc(pr.cell(row=r,column=12))
    pr.cell(row=r,column=13,value=f'=IF(I{r}<>"Vencida","",IF({nivel}=0,"Aguardar: ainda não chegou ao primeiro passo da régua.",INDEX({RA},{nivel})))'); calc(pr.cell(row=r,column=13),center=False)
    pr.cell(row=r,column=14,value=f'=IF(I{r}="Vencida",F{r}*J{r}+ROW()/100000,0)'); pr.cell(row=r,column=14).font=F(color=CINZA,size=9)
    pr.cell(row=r,column=15,value=f'=IF(AND(I{r}="A vencer",K{r}<={JL}),100000-K{r}-ROW()/100000,0)'); pr.cell(row=r,column=15).font=F(color=CINZA,size=9)
pr.column_dimensions["N"].hidden=True; pr.column_dimensions["O"].hidden=True
dvs=[lista(off("Pacientes","A",R0,RNC),strict=False), lista('"Sim,Não"'),
     DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True),
     DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True),
     DataValidation(type="whole",operator="greaterThanOrEqual",formula1="1",allow_blank=True)]
for dv,rng in zip(dvs,[f"A{R0}:A{RNP}",f"G{R0}:G{RNP}",f"E{R0}:E{RNP}",f"F{R0}:F{RNP}",f"D{R0}:D{RNP}"]): dv.add(rng); pr.add_data_validation(dv)
pr.conditional_formatting.add(f"A{R0}:M{RNP}", FormulaRule(formula=[f'$I{R0}="Vencida"'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
pr.conditional_formatting.add(f"A{R0}:M{RNP}", FormulaRule(formula=[f'AND($I{R0}="A vencer",$K{R0}<={JC})'], fill=fill("FFF4CC")))
pr.conditional_formatting.add(f"A{R0}:M{RNP}", FormulaRule(formula=[f'$I{R0}="Paga"'], font=F(color=VERDE_T,size=10)))
pr.conditional_formatting.add(f"H{R0}:H{RNP}", FormulaRule(formula=[f'AND($G{R0}="Sim",$H{R0}="")'], fill=fill(VERM)))
pr.cell(row=RNP+2,column=1,value="Vermelho: vencida. Amarelo: vence na janela curta. Verde: paga. Data do pagamento em vermelho: marcada como paga sem data. No exemplo, as parcelas pagas têm data fixa (cada uma é uma entrada do caixa, planilha 09) e as em aberto têm vencimento relativo a hoje.").font=F(size=9,color=LILAS)
widths(pr,(28,30,12,9,12,12,8,13,13,9,9,16,60)); pr.freeze_panes="C5"; pr.sheet_view.showGridLines=False; pr.auto_filter.ref=f"A4:M{RNP}"

# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
titulo(p,f'=Config!$B$4&" · Parcelas e inadimplência · "&{dstr(HOJE)}',"Nada para digitar aqui: tudo vem de Parcelas e da régua em Config. Só o que foi combinado a prazo entra aqui.",merge_to="L")
p.merge_cells("A1:L1")
PK=f"Parcelas!$K${R0}:$K${RNP}"; PJ=f"Parcelas!$J${R0}:$J${RNP}"; PL=f"Parcelas!$L${R0}:$L${RNP}"; PN=f"Parcelas!$N${R0}:$N${RNP}"; PO=f"Parcelas!$O${R0}:$O${RNP}"
kpi(p,4,1,'="Vence em até "&'+JC+'&" dias"',f'=SUMIFS({PF},{PI},"A vencer",{PK},"<="&{JC})',SOL,UVA,fmt=BRL0)
kpi(p,4,3,'="Vence em até "&'+JL+'&" dias"',f'=SUMIFS({PF},{PI},"A vencer",{PK},"<="&{JL})',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,5,"Em aberto (total)",f'=SUMIFS({PF},{PI},"A vencer")+SUMIFS({PF},{PI},"Vencida")',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,7,"Vencido (R$)",f'=SUMIFS({PF},{PI},"Vencida")',VERM,VERM_T,fmt=BRL0)
kpi(p,4,9,"Parcelas vencidas",f'=COUNTIFS({PI},"Vencida")',VERM,VERM_T,fmt="0")
kpi(p,4,11,"Inadimplência (vencido ÷ (pago + vencido))",f'=IFERROR(G5/(G5+SUMIFS({PF},{PI},"Paga")),0)',VERM,VERM_T,fmt="0.0%")
p["A7"]="Inadimplência = vencido ÷ (pago + vencido): de tudo o que já venceu, quanto ainda não entrou. É a definição usada em todo o kit (17, 19, 20) e vale só para o que foi combinado a prazo (o à vista entra direto no caixa). Outra conta possível, vencido ÷ em aberto, mede a carteira futura e dá um número maior; não a use para comparar com estas planilhas."; nota(p["A7"]); p.merge_cells("A7:L7"); p["A7"].alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[7].height=30
p["A9"]="Vencidas por faixa de atraso"; p["A9"].font=F(bold=True,size=13,color=UVA)
hdr(p,10,["Faixa","Parcelas","Valor (R$)","% do vencido","Ação da régua"]); p.merge_cells("E10:H10")
for i in range(4):
    r=11+i; src=f"Config!$C${RG0+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",COUNTIFS({PL},{src}))'); calc(p.cell(row=r,column=2))
    p.cell(row=r,column=3,value=f'=IF({src}="","",SUMIFS({PF},{PL},{src}))'); calc(p.cell(row=r,column=3),BRL0)
    p.cell(row=r,column=4,value=f'=IF({src}="","",IFERROR(C{r}/$G$5,0))'); calc(p.cell(row=r,column=4),PCT)
    p.cell(row=r,column=5,value=f'=IF({src}="","",Config!$B${RG0+i})'); calc(p.cell(row=r,column=5),center=False); p.merge_cells(start_row=r,start_column=5,end_row=r,end_column=8)
    p.cell(row=r,column=5).alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[r].height=32
p.cell(row=15,column=1,value="Antes da régua"); calc(p.cell(row=15,column=1),center=False)
p.cell(row=15,column=2,value=f'=COUNTIFS({PL},"Antes da régua")'); calc(p.cell(row=15,column=2))
p.cell(row=15,column=3,value=f'=SUMIFS({PF},{PL},"Antes da régua")'); calc(p.cell(row=15,column=3),BRL0)
p.cell(row=15,column=4,value='=IFERROR(C15/$G$5,0)'); calc(p.cell(row=15,column=4),PCT)
p.cell(row=15,column=5,value="Vencidas há menos dias que o primeiro passo da régua."); calc(p.cell(row=15,column=5),center=False); p.merge_cells("E15:H15")
p["A17"]="Cobrar primeiro"; p["A17"].font=F(bold=True,size=13,color=UVA)
p["A18"]="Parcelas vencidas em ordem de valor × dias de atraso: as que mais pesam no caixa vêm no topo. O texto da última coluna é a ação da régua para a faixa de cada uma."; nota(p["A18"]); p.merge_cells("A18:L18")
hdr(p,19,["#","Paciente","Procedimento","Parcela","Vencimento","Valor (R$)","Dias de atraso","Faixa","Ação sugerida"]); p.merge_cells("I19:L19")
TOP=10
for k in range(1,TOP+1):
    r=19+k; m=f'MATCH(LARGE({PN},{k}),{PN},0)'; g=f'LARGE({PN},{k})>0'
    p.cell(row=r,column=1,value=k); calc(p.cell(row=r,column=1))
    for col,src in zip((2,3,4,5,6,7,8,9),("A","B","D","E","F","J","L","M")):
        p.cell(row=r,column=col,value=f'=IFERROR(IF({g},INDEX(Parcelas!${src}${R0}:${src}${RNP},{m}),""),"")'); calc(p.cell(row=r,column=col),center=(col not in (2,3,9)))
    p.cell(row=r,column=5).number_format=DATA; p.cell(row=r,column=6).number_format=BRL0; p.cell(row=r,column=7).number_format="0"
    p.merge_cells(start_row=r,start_column=9,end_row=r,end_column=12); p.cell(row=r,column=9).alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[r].height=32
p.conditional_formatting.add(f"A20:L{19+TOP}", FormulaRule(formula=['$G20>=30'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
r2=19+TOP+2
p.cell(row=r2,column=1,value="Vencem nos próximos dias").font=F(bold=True,size=13,color=UVA)
p.cell(row=r2+1,column=1,value='="Parcelas a vencer em até "&'+JL+'&" dias, a mais próxima primeiro. Bom momento para lembrar o paciente antes do vencimento."'); nota(p.cell(row=r2+1,column=1)); p.merge_cells(start_row=r2+1,start_column=1,end_row=r2+1,end_column=12)
hdr(p,r2+2,["#","Paciente","Procedimento","Parcela","Vencimento","Valor (R$)","Dias para vencer"])
for k in range(1,9):
    r=r2+2+k; m=f'MATCH(LARGE({PO},{k}),{PO},0)'; g=f'LARGE({PO},{k})>0'
    p.cell(row=r,column=1,value=k); calc(p.cell(row=r,column=1))
    for col,src in zip((2,3,4,5,6,7),("A","B","D","E","F","K")):
        p.cell(row=r,column=col,value=f'=IFERROR(IF({g},INDEX(Parcelas!${src}${R0}:${src}${RNP},{m}),""),"")'); calc(p.cell(row=r,column=col),center=(col not in (2,3)))
    p.cell(row=r,column=5).number_format=DATA; p.cell(row=r,column=6).number_format=BRL0; p.cell(row=r,column=7).number_format="0"
p.conditional_formatting.add(f"A{r2+3}:G{r2+10}", FormulaRule(formula=[f'AND(ISNUMBER($G{r2+3}),$G{r2+3}<={JC})'], fill=fill("FFF4CC")))
p.cell(row=r2+12,column=1,value="Fonte: aba Parcelas. A planilha avisa e sugere; a conversa com o paciente é sua, sempre em particular e sem constrangimento.").font=F(size=9,color=LILAS)
bc=BarChart(); bc.type="bar"; bc.height=6; bc.width=12; bc.title="Vencido por faixa de atraso (R$)"; bc.style=2
bc.add_data(Reference(p,min_col=3,min_row=10,max_row=15),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=11,max_row=15))
bc.series[0].graphicalProperties.solidFill="7A1F1F"; bc.legend=None; bc.x_axis.majorGridlines=None
p.add_chart(bc,"J9")
widths(p,(16,26,30,9,12,14,12,16,14,14,14,14)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False

# ---------- Exemplo ----------
regua=[(1,"Lembrete gentil por WhatsApp: confirmar se o Pix ou o boleto chegou e reenviar, se preciso."),
       (7,"Mensagem da recepção perguntando se houve algum imprevisto e oferecendo uma nova data para o pagamento."),
       (15,"Ligação da recepção com o demonstrativo das parcelas em aberto e proposta de nova data ou divisão do valor."),
       (30,"Conversa do médico ou da administração para combinar um plano de pagamento; registrar por escrito o que foi combinado.")]
for i,(d,t) in enumerate(regua): cfg.cell(row=RG0+i,column=1,value=d); cfg.cell(row=RG0+i,column=2,value=t)
nomes=sorted({pp["paciente"] for pp in dados.PARCELAS})
for i,n in enumerate(nomes): pc.cell(row=R0+i,column=1,value=n); pc.cell(row=R0+i,column=2,value=dados.PAC[n]["contato"])
linhas=[]
for pp in dados.PARCELAS:
    if pp["pago"]: linhas.append((pp["paciente"],pp["procedimento"],pp["atendimento"],pp["n"],pp["vencimento"],pp["valor"],"Sim",pp["pagamento"]))
    else: linhas.append((pp["paciente"],pp["procedimento"],pp["atendimento"],pp["n"],dados.prazo_formula(dados.dias(pp["vencimento"])),pp["valor"],"Não",None))
assert len(linhas)<=NP
for i,row in enumerate(linhas):
    r=R0+i
    for c,v in enumerate(row,start=1):
        if v is not None: pr.cell(row=r,column=c,value=v)
print(len(linhas),"parcelas e",len(nomes),"pacientes no exemplo")

como_usar(wb,"Parcelas particulares e inadimplência",[
 ("O que esta planilha faz","Você lança as parcelas combinadas a prazo com cada paciente e marca as pagas; ela mostra o que vence nos próximos dias, o que já venceu por faixa de atraso, a inadimplência e a lista de quem cobrar primeiro, com a ação sugerida da régua de cobrança."),
 ("Passo 1","Em Config, ajuste as janelas (7 e 30 dias) e a régua: a partir de quantos dias de atraso fazer o quê. Escreva as ações do seu jeito, educadas e diretas."),
 ("Passo 2","Em Pacientes, cole os nomes e contatos da planilha 01 (aba Pacientes; a 01 é a fonte) de quem paga a prazo. É a lista que alimenta a escolha do paciente em Parcelas."),
 ("Passo 3","Em Parcelas, uma linha por parcela: paciente, procedimento, data do atendimento, número da parcela, vencimento e valor. Quando receber, marque Pago? = Sim e a data, e lance a mesma entrada no caixa (planilha 09, categoria \"Particular a prazo\"). O à vista não entra aqui."),
 ("Passo 4","Em Painel, veja o que vence em 7 e 30 dias, o vencido por faixa, a inadimplência (vencido ÷ (pago + vencido), a mesma conta das planilhas 17, 19 e 20) e a lista \"Cobrar primeiro\" com o texto da régua para cada parcela."),
 ("Rotina","Sexta-feira, 4 minutos (rotina da semana, planilha 03): conferir o extrato, marcar as parcelas pagas, mandar as mensagens da lista \"Cobrar primeiro\" e lembrar quem vence na semana seguinte. Tudo numa passada só, na sexta."),
 ("Com a IA","Copie uma linha de \"Cobrar primeiro\" (sem o nome: procedimento, valor, dias, ação) e use o prompt \"Recebíveis 04 · Cobrança educada em três versões\" da biblioteca do kit, ou os 15 modelos de mensagem do bônus. Nunca cole dados do paciente que ele não autorizou."),
])
proteger(wb); salvar(wb,"14-parcelas-e-inadimplencia.xlsx","Parcelas particulares e inadimplência · Kit de Gestão para Médicos")
