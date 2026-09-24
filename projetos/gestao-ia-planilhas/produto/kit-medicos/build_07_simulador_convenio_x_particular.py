#!/usr/bin/env python3
"""Planilha 7 do Kit de Gestão para Médicos: Simulador convênio × particular. Gera 07-simulador-convenio-x-particular.xlsx
Para um procedimento: valor de tabela de cada pagador, glosa esperada, prazo de pagamento (custo do dinheiro a 1,5 % ao mês),
impostos, custo cheio do atendimento (custo-hora da 05 + material) e duas leituras: agenda cheia (custo cheio) e agenda vazia (contribuição)."""
from ssg import *
from openpyxl.chart import BarChart, Reference
import dados
NPAG=6; NPROC=12
P0=10; PN=P0+NPAG-1                 # pagadores na Config: linhas 10..15
Q0=19; QN=Q0+NPROC-1                # procedimentos na Config: linhas 19..30
wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Custo-hora da planilha 05; pagadores com prazo e glosa esperada; procedimentos com tempo e material (iguais à 06).",merge_to="H")
campos=[("Nome da clínica",f"{dados.CLINICA} (exemplo fictício)",None),("Data de referência",dados.HOJE,DATA),("Custo da hora de atendimento (R$)",dados.CUSTO_HORA,BRL),
 ("Impostos sobre o que entra (%)",dados.ALIQ,PCT),("Margem mínima sobre o preço (%)",dados.MARGEM,PCT),("Custo do dinheiro (% ao mês)",dados.JUROS_MES,"0.0%")]
for i,(a,v,fmt) in enumerate(campos):
    r=3+i; cfg.cell(row=r,column=1,value=a); rotulo(cfg.cell(row=r,column=1)); cfg.cell(row=r,column=2,value=v)
    if r==4: calc(cfg.cell(row=r,column=2),fmt)
    else: inp(cfg.cell(row=r,column=2),fmt,center=fmt is not None)
cfg["C5"]="Copie do Painel da planilha 05. No exemplo, R$ 200,00."; cfg["C6"]="A mesma alíquota efetiva das planilhas 05 e 06 (exemplo: 11 %). Só imposto: a taxa da maquininha não entra aqui (o convênio não passa na maquininha; o cartão é conciliado na 16)."; cfg["C7"]="A mesma margem mínima da 05 e da 06."
cfg["C8"]="O que custa esperar 30, 45 ou 60 dias pelo dinheiro: juros do cheque especial ou da antecipação, ou o rendimento que o dinheiro parado deixa de dar. 1,5 % ao mês é um exemplo; use o seu."
for r in (5,6,7,8): nota(cfg.cell(row=r,column=3))
hdr(cfg,P0-1,["Pagador","Prazo de pagamento (dias)","Glosa esperada (%)","Observação"],height=30)
for r in range(P0,PN+1):
    inp(cfg.cell(row=r,column=1)); inp(cfg.cell(row=r,column=2),"0",center=True); inp(cfg.cell(row=r,column=3),"0.0%",center=True); inp(cfg.cell(row=r,column=4))
cfg.cell(row=P0,column=1,value="Particular"); cfg.cell(row=P0,column=2,value=0); cfg.cell(row=P0,column=3,value=0); cfg.cell(row=P0,column=4,value="À vista no dia (Pix, dinheiro, cartão)")
for i,(n,prazo,g) in enumerate(dados.CONVENIOS):
    r=P0+1+i; cfg.cell(row=r,column=1,value=n); cfg.cell(row=r,column=2,value=prazo); cfg.cell(row=r,column=3,value=g); cfg.cell(row=r,column=4,value="Prazo contratual a partir do envio do lote; glosa = histórico da planilha 13")
cfg.cell(row=PN+1,column=1,value="Preencha de cima para baixo, sem pular linha. Glosa esperada: quanto o convênio costuma não pagar do que foi enviado (Painel da planilha 13, por convênio). Prazo: dias entre o envio do lote e o pagamento.").font=F(size=9,color=LILAS)
hdr(cfg,Q0-1,["Procedimento","Minutos","Gera retorno?","Material (R$)","Retornos por consulta","Minutos do retorno"],height=30)
for r in range(Q0,QN+1):
    inp(cfg.cell(row=r,column=1)); inp(cfg.cell(row=r,column=2),"0",center=True); inp(cfg.cell(row=r,column=3),center=True); inp(cfg.cell(row=r,column=4),BRL,center=True)
cfg.cell(row=Q0,column=5,value=dados.RETORNO_PROB); inp(cfg.cell(row=Q0,column=5),"0.00",center=True); cfg.cell(row=Q0,column=6,value=dados.DUR["Retorno"]); inp(cfg.cell(row=Q0,column=6),"0",center=True)
cfg.cell(row=Q0+1,column=5,value="Só na primeira linha: valem para todos os procedimentos que geram retorno (iguais à planilha 06)."); nota(cfg.cell(row=Q0+1,column=5))
for i,(pnome,d,m,ret) in enumerate(dados.PROCEDIMENTOS):
    r=Q0+i; cfg.cell(row=r,column=1,value=pnome); cfg.cell(row=r,column=2,value=d); cfg.cell(row=r,column=3,value="Sim" if ret>0 else "Não"); cfg.cell(row=r,column=4,value=m)
dvs=lista('"Sim,Não"'); dvs.add(f"C{Q0}:C{QN}"); cfg.add_data_validation(dvs)
widths(cfg,(40,16,16,50,14,14,3,3)); cfg.sheet_view.showGridLines=False
CH="Config!$B$5"; IMP="Config!$B$6"; MMIN="Config!$B$7"; JUR="Config!$B$8"; NRET=f"Config!$E${Q0}"; DRET=f"Config!$F${Q0}"
PROCS=f"Config!$A${Q0}:$A${QN}"; PAGS=f"Config!$A${P0}:$A${PN}"
# Auditoria final-2 (G01/G02): o mesmo domínio da 05 e da 06 (impostos e margem de 0 a 99 %, soma
# abaixo de 100 %, sem texto nem branco), num predicado só; custo-hora em branco não é custo zero.
OKC=f'IF(AND(ISNUMBER({IMP}),ISNUMBER({MMIN})),AND({IMP}>=0,{IMP}<1,{MMIN}>=0,{MMIN}<1,{IMP}+{MMIN}<1),FALSE)'
for _c,_o,_m in (("B6","B7","Impostos entre 0 % e 99 %, e impostos + margem abaixo de 100 %."),("B7","B6","Margem entre 0 % e 99 %, e impostos + margem abaixo de 100 %.")):
    _dv=DataValidation(type="custom",formula1=f'=IF(ISNUMBER({_c}),AND({_c}>=0,{_c}<1,{_c}+N({_o})<1),FALSE)',allow_blank=False,showErrorMessage=True,errorTitle="Percentual",error=_m); _dv.add(_c); cfg.add_data_validation(_dv)
# ---------- Simulador ----------
s=wb.create_sheet("Simulador",0)
titulo(s,'=Config!$B$3&" · Simulador convênio × particular · "&TEXT(DAY(Config!$B$4),"00")&"/"&TEXT(MONTH(Config!$B$4),"00")&"/"&YEAR(Config!$B$4)',"Escolha o procedimento e digite o valor de tabela de cada pagador (amarelo). O resto é calculado: glosa, custo do dinheiro pelo prazo, impostos, custo cheio, margem e as duas leituras (agenda cheia × agenda vazia).",merge_to="N")
s["A7"]="1. O procedimento"; s["A7"].font=F(bold=True,size=13,color=UVA)
s["A8"]="Procedimento"; s["B8"]="Consulta"; rotulo(s["A8"]); inp(s["B8"])
s["C8"]="Os valores de tabela do quadro 2 saem da planilha 08 · Tabela de preços (a fonte única da tabela no kit): copie de lá, não digite um preço diferente aqui."; nota(s["C8"]); s.merge_cells("C8:J8")
dvp=lista(f"=OFFSET(Config!$A${Q0},0,0,MAX(1,COUNTA(Config!$A${Q0}:$A${QN})),1)"); dvp.add("B8"); s.add_data_validation(dvp)
campos=[("Minutos",f'=IFERROR(INDEX(Config!$B${Q0}:$B${QN},MATCH(B8,{PROCS},0)),"")',"0"),
        ("Gera retorno?",f'=IFERROR(INDEX(Config!$C${Q0}:$C${QN},MATCH(B8,{PROCS},0)),"")',None),
        ("Material (R$)",f'=IFERROR(INDEX(Config!$D${Q0}:$D${QN},MATCH(B8,{PROCS},0)),0)',BRL),
        ("Tempo com retorno embutido (min)",f'=IF(B9="","",B9+IF(B10="Sim",{NRET}*{DRET},0))',"0.0"),
        ("Custo cheio do atendimento (R$)",f'=IF(B12="","",IF(NOT(ISNUMBER({CH})),"falta o custo-hora em Config",B12/60*{CH}+B11))',BRL),
        ("Hora mínima a cobrar (R$)",f'=IF(NOT(ISNUMBER({CH})),"falta o custo-hora em Config",IF({OKC},{CH}/(1-{IMP}-{MMIN}),"margens inválidas (Config)"))',BRL),
        ("Preço mínimo deste procedimento (R$)",f'=IF(B13="","",IF(NOT(ISNUMBER(B13)),B13,IF({OKC},B13/(1-{IMP}-{MMIN}),"margens inválidas (Config)")))',BRL)]
for i,(a,f_,fmt) in enumerate(campos):
    r=9+i; s.cell(row=r,column=1,value=a); rotulo(s.cell(row=r,column=1),bold=False); s.cell(row=r,column=1).border=borda; s.cell(row=r,column=2,value=f_); calc(s.cell(row=r,column=2),fmt)
s["B13"].font=F(bold=True,color=UVA,size=10); s["B15"].font=F(bold=True,color=UVA,size=10)
s["C13"]="Tempo (com retorno) ÷ 60 × custo-hora + material. É o custo com a estrutura inteira rateada: a leitura de agenda cheia."; nota(s["C13"]); s.merge_cells("C13:J13")
s["C15"]="Custo cheio ÷ (1 − impostos − margem mínima). O mesmo número da planilha 06."; nota(s["C15"]); s.merge_cells("C15:J15")
CUSTO="$B$13"; MAT="$B$11"; TEMPO="$B$12"; HMIN="$B$14"; PMIN="$B$15"
# 2. tabela por pagador
T0=19; TN=T0+NPAG-1
s.cell(row=T0-2,column=1,value="2. O que cada pagador paga por este procedimento").font=F(bold=True,size=13,color=UVA)
hdr(s,T0-1,["Pagador","Valor de tabela (R$)","Prazo (dias)","Glosa esperada","Valor após glosa","Custo do dinheiro (R$)","Impostos (R$)","Valor líquido (R$)","Custo cheio (R$)","Margem (R$)","Margem (%)","Líquido por hora (R$)","Contra a hora mínima","Leitura com agenda cheia"],height=40)
for i in range(NPAG):
    r=T0+i; src=f"Config!$A${P0+i}"
    s.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(s.cell(row=r,column=1),center=False)
    inp(s.cell(row=r,column=2),BRL0,center=True)
    s.cell(row=r,column=3,value=f'=IF({src}="","",Config!$B${P0+i})'); calc(s.cell(row=r,column=3),"0")
    s.cell(row=r,column=4,value=f'=IF({src}="","",Config!$C${P0+i})'); calc(s.cell(row=r,column=4),"0.0%")
    s.cell(row=r,column=5,value=f'=IF(OR({src}="",B{r}=""),"",B{r}*(1-D{r}))'); calc(s.cell(row=r,column=5),BRL)
    s.cell(row=r,column=6,value=f'=IF(E{r}="","",E{r}*{JUR}*C{r}/30)'); calc(s.cell(row=r,column=6),BRL)
    s.cell(row=r,column=7,value=f'=IF(OR(E{r}="",NOT({OKC})),"",E{r}*{IMP})'); calc(s.cell(row=r,column=7),BRL)
    s.cell(row=r,column=8,value=f'=IF(OR(E{r}="",NOT(ISNUMBER(G{r}))),"",E{r}-F{r}-G{r})'); calc(s.cell(row=r,column=8),BRL); s.cell(row=r,column=8).font=F(bold=True,color=UVA,size=10)
    s.cell(row=r,column=9,value=f'=IF(E{r}="","",{CUSTO})'); calc(s.cell(row=r,column=9),BRL)
    s.cell(row=r,column=10,value=f'=IF(OR(NOT(ISNUMBER(H{r})),NOT(ISNUMBER(I{r}))),"",H{r}-I{r})'); calc(s.cell(row=r,column=10),BRL)
    s.cell(row=r,column=11,value=f'=IF(OR(NOT(ISNUMBER(J{r})),N(B{r})=0),"",J{r}/B{r})'); calc(s.cell(row=r,column=11),PCT)
    s.cell(row=r,column=12,value=f'=IF(OR(NOT(ISNUMBER(H{r})),N({TEMPO})=0),"",H{r}/({TEMPO}/60))'); calc(s.cell(row=r,column=12),BRL)
    s.cell(row=r,column=13,value=f'=IF(OR(L{r}="",NOT(ISNUMBER({HMIN}))),"",IF({HMIN}=0,"sem base (mínimo zero)",L{r}/{HMIN}-1))'); calc(s.cell(row=r,column=13),"+0%;-0%;0%")
    s.cell(row=r,column=14,value=f'=IF(E{r}="","",IF(NOT({OKC}),"Margens inválidas em Config",IF(NOT(ISNUMBER(I{r})),"Falta o custo-hora em Config",IF(J{r}<0,"Não cobre o custo cheio",IF(K{r}<{MMIN},"Cobre o custo, não a margem mínima","Cobre custo e margem")))))'); calc(s.cell(row=r,column=14),center=False)
s.conditional_formatting.add(f"N{T0}:N{TN}", FormulaRule(formula=[f'N{T0}="Não cobre o custo cheio"'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
s.conditional_formatting.add(f"N{T0}:N{TN}", FormulaRule(formula=[f'N{T0}="Cobre o custo, não a margem mínima"'], fill=fill(AMARELO)))
s.conditional_formatting.add(f"N{T0}:N{TN}", FormulaRule(formula=[f'N{T0}="Cobre custo e margem"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
s.conditional_formatting.add(f"J{T0}:J{TN}", FormulaRule(formula=[f'AND(ISNUMBER(J{T0}),J{T0}<0)'], font=F(color="C8402E",size=10,bold=True)))
s.cell(row=TN+1,column=1,value="Valor líquido = tabela − glosa esperada − custo do dinheiro (juros × prazo ÷ 30) − impostos. Margem = líquido − custo cheio. Líquido por hora compara pagadores com a hora mínima da 05: é o número para negociar tabela ou decidir descredenciar.").font=F(size=9,color=LILAS)
s.merge_cells(start_row=TN+1,start_column=1,end_row=TN+1,end_column=14); s.cell(row=TN+1,column=1).alignment=Alignment(wrap_text=True,vertical="top"); s.row_dimensions[TN+1].height=30
# 3. agenda vazia
V0=TN+4
s.cell(row=V0,column=1,value="3. A outra leitura: e se o horário ficaria vazio?").font=F(bold=True,size=13,color=UVA)
s.cell(row=V0+1,column=1,value="Com agenda vazia o custo fixo já está pago: o que importa é a contribuição (líquido − material). Com agenda cheia, cada horário de convênio toma o lugar de um particular, e vale a leitura do quadro 2."); nota(s.cell(row=V0+1,column=1)); s.merge_cells(start_row=V0+1,start_column=1,end_row=V0+1,end_column=14)
hdr(s,V0+2,["Pagador","Valor líquido (R$)","Material (R$)","Contribuição por atendimento (R$)","Contribuição por hora (R$)","Atendimentos deste pagador = 1 particular","Leitura com agenda vazia"],height=40)
for i in range(NPAG):
    r=V0+3+i; t=T0+i
    s.cell(row=r,column=1,value=f'=A{t}'); calc(s.cell(row=r,column=1),center=False)
    s.cell(row=r,column=2,value=f'=IF(H{t}="","",H{t})'); calc(s.cell(row=r,column=2),BRL)
    s.cell(row=r,column=3,value=f'=IF(H{t}="","",{MAT})'); calc(s.cell(row=r,column=3),BRL)
    s.cell(row=r,column=4,value=f'=IF(B{r}="","",B{r}-C{r})'); calc(s.cell(row=r,column=4),BRL); s.cell(row=r,column=4).font=F(bold=True,color=UVA,size=10)
    s.cell(row=r,column=5,value=f'=IF(OR(D{r}="",{TEMPO}=0),"",D{r}/({TEMPO}/60))'); calc(s.cell(row=r,column=5),BRL)
    s.cell(row=r,column=6,value=f'=IF(OR(D{r}="",D{r}<=0,$D${V0+3}=""),"",$D${V0+3}/D{r})'); calc(s.cell(row=r,column=6),"0.0")
    s.cell(row=r,column=7,value=f'=IF(D{r}="","",IF(D{r}<=0,"Nem com agenda vazia: perde dinheiro a cada atendimento",IF(C{t}>45,"Contribui, mas o prazo pesa no caixa: só com reserva (12)","Contribui: melhor que o horário vazio")))'); calc(s.cell(row=r,column=7),center=False)
s.conditional_formatting.add(f"G{V0+3}:G{V0+2+NPAG}", FormulaRule(formula=[f'LEFT(G{V0+3},3)="Nem"'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
s.conditional_formatting.add(f"G{V0+3}:G{V0+2+NPAG}", FormulaRule(formula=[f'LEFT(G{V0+3},20)="Contribui, mas o pra"'], fill=fill(AMARELO)))
s.conditional_formatting.add(f"G{V0+3}:G{V0+2+NPAG}", FormulaRule(formula=[f'LEFT(G{V0+3},16)="Contribui: melho"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
s.cell(row=V0+3+NPAG,column=1,value="\"Atendimentos = 1 particular\": quantos atendimentos deste pagador rendem a contribuição de um particular. Compare com as horas vazias do mês (Painel da 01) antes de aceitar ou descredenciar um convênio.").font=F(size=9,color=LILAS)
s.merge_cells(start_row=V0+3+NPAG,start_column=1,end_row=V0+3+NPAG,end_column=14)
# 4. mix do mês
M0=V0+3+NPAG+3
s.cell(row=M0,column=1,value="4. O mês inteiro: mix de pagadores para este procedimento").font=F(bold=True,size=13,color=UVA)
s.cell(row=M0+1,column=1,value="Amarelo: quantos atendimentos de cada pagador no mês (no exemplo, as consultas realizadas em agosto de 2026, Painel da 01 por pagador). A planilha soma o líquido, o custo cheio e o resultado."); nota(s.cell(row=M0+1,column=1)); s.merge_cells(start_row=M0+1,start_column=1,end_row=M0+1,end_column=14)
hdr(s,M0+2,["Pagador","Atendimentos no mês","Horas","Líquido no mês (R$)","Custo cheio no mês (R$)","Resultado (R$)","% dos atendimentos","% do líquido"],height=32)
for i in range(NPAG):
    r=M0+3+i; t=T0+i
    s.cell(row=r,column=1,value=f'=A{t}'); calc(s.cell(row=r,column=1),center=False)
    inp(s.cell(row=r,column=2),"0",center=True)
    s.cell(row=r,column=3,value=f'=IF(OR(A{r}="",B{r}=""),"",B{r}*{TEMPO}/60)'); calc(s.cell(row=r,column=3),"#,##0.0")
    s.cell(row=r,column=4,value=f'=IF(OR(A{r}="",B{r}="",H{t}=""),"",B{r}*H{t})'); calc(s.cell(row=r,column=4),BRL0)
    s.cell(row=r,column=5,value=f'=IF(OR(A{r}="",B{r}="",NOT(ISNUMBER({CUSTO}))),"",B{r}*{CUSTO})'); calc(s.cell(row=r,column=5),BRL0)
    s.cell(row=r,column=6,value=f'=IF(OR(D{r}="",E{r}=""),"",D{r}-E{r})'); calc(s.cell(row=r,column=6),BRL0)
    s.cell(row=r,column=7,value=f'=IF(OR(D{r}="",$B${M0+3+NPAG}=0),"",B{r}/$B${M0+3+NPAG})'); calc(s.cell(row=r,column=7),PCT)
    s.cell(row=r,column=8,value=f'=IF(OR(D{r}="",$D${M0+3+NPAG}=0),"",D{r}/$D${M0+3+NPAG})'); calc(s.cell(row=r,column=8),PCT)
rt=M0+3+NPAG
s.cell(row=rt,column=1,value="Total"); rotulo(s.cell(row=rt,column=1)); s.cell(row=rt,column=1).border=borda
for c in (2,3,4,5,6):
    # totais: soma de linhas suspensas não é zero apurado (auditoria final-2, G01)
    _g={4:f'NOT({OKC})',5:f'NOT(ISNUMBER({CUSTO}))',6:f'OR(NOT({OKC}),NOT(ISNUMBER({CUSTO})))'}.get(c)
    _t=f"SUM({L(c)}{M0+3}:{L(c)}{rt-1})"
    _msg={4:'"margens inválidas"',5:'"falta o custo-hora"',6:f'IF(NOT(ISNUMBER({CUSTO})),"falta o custo-hora","margens inválidas")'}.get(c)
    s.cell(row=rt,column=c,value=f'=IF({_g},{_msg},{_t})' if _g else "="+_t); calc(s.cell(row=rt,column=c),"#,##0.0" if c==3 else ("0" if c==2 else BRL0)); s.cell(row=rt,column=c).font=F(bold=True,color=UVA,size=10)
s.conditional_formatting.add(f"F{M0+3}:F{rt}", FormulaRule(formula=[f'AND(ISNUMBER(F{M0+3}),F{M0+3}<0)'], font=F(color="C8402E",size=10,bold=True)))
s.cell(row=rt+1,column=1,value="Resultado negativo num pagador e positivo no total é a situação comum: o particular sustenta o convênio. Pergunta certa: o convênio traz pacientes que virariam particulares ou exames? Se não, negocie a tabela ou reduza os horários dele.").font=F(size=9,color=LILAS)
s.merge_cells(start_row=rt+1,start_column=1,end_row=rt+1,end_column=14)
# KPIs no topo
kpi(s,4,1,"Custo cheio do atendimento",f"={CUSTO}",LAVANDA,UVA,fmt=BRL)
kpi(s,4,3,"Hora mínima (05)",f"={HMIN}",LAVANDA,UVA,fmt=BRL)
kpi(s,4,5,"Particular · líquido por hora",f'=IF(L{T0}="","",L{T0})',VERDE,VERDE_T,fmt=BRL)
LR=f"$L${T0+1}:$L${TN}"; AR=f"$A${T0+1}:$A${TN}"
kpi(s,4,7,"Melhor convênio (líquido/hora)",f'=IFERROR(INDEX({AR},MATCH(MAX({LR}),{LR},0))&": R$ "&FIXED(MAX({LR}),0),"")',SOL,UVA,fmt="@",span=3)
kpi(s,4,10,"Pior convênio (líquido/hora)",f'=IFERROR(INDEX({AR},MATCH(MIN({LR}),{LR},0))&": R$ "&FIXED(MIN({LR}),0),"")',VERM,VERM_T,fmt="@",span=3)
s["G5"].font=F(size=13,bold=True,color=UVA); s["J5"].font=F(size=13,bold=True,color=VERM_T)
bc=BarChart(); bc.type="bar"; bc.height=6.5; bc.width=13; bc.title="Líquido por hora × hora mínima"; bc.style=2
bc.add_data(Reference(s,min_col=12,min_row=T0-1,max_row=T0+len(dados.PAGADORES)-1),titles_from_data=True); bc.set_categories(Reference(s,min_col=1,min_row=T0,max_row=T0+len(dados.PAGADORES)-1))
bc.series[0].graphicalProperties.solidFill="3B1F5E"; bc.legend=None; bc.x_axis.majorGridlines=None
s.add_chart(bc,"E7")
widths(s,(38,16,11,12,14,14,12,14,13,12,10,14,12,34)); s.freeze_panes="A4"; s.sheet_view.showGridLines=False
# exemplo: tabelas da consulta e mix de agosto
for i,pg in enumerate(dados.PAGADORES):
    s.cell(row=T0+i,column=2,value=dados.preco("Consulta",pg))
    n=sum(1 for r in dados.AGENDA_TODA if r["situacao"]=="Realizado" and r["procedimento"]=="Consulta" and r["pagador"]==pg and r["data"].month==8)
    s.cell(row=M0+3+i,column=2,value=n)
como_usar(wb,"Simulador convênio × particular",[
 ("O que esta planilha faz","Responde \"vale a pena este convênio?\" para um procedimento: pega o valor de tabela de cada pagador, tira a glosa esperada, o custo do dinheiro pelo prazo de pagamento e os impostos, compara com o custo cheio do atendimento (custo-hora da 05 + material) e mostra o líquido por hora contra a hora mínima. Depois faz a leitura oposta: com horários vazios, o que cada pagador ainda contribui."),
 ("Passo 1","Em Config, copie o custo-hora, a alíquota e a margem mínima da planilha 05; digite o custo do dinheiro (% ao mês); cadastre os pagadores com prazo de pagamento e glosa esperada (Painel da 13) e os procedimentos com tempo e material (iguais à 06)."),
 ("Passo 2","Em Simulador, escolha o procedimento e copie o valor de tabela de cada pagador da planilha 08 · Tabela de preços (o particular é o preço praticado). A 08 é a fonte única da tabela no kit: mudou um preço, atualize 08 → 06 → 07 → 01, nessa ordem. Leia o quadro 2 (agenda cheia): margem, líquido por hora e o veredito de cada pagador."),
 ("Passo 3","Leia o quadro 3 (agenda vazia): a contribuição por atendimento e quantos atendimentos do convênio equivalem a um particular. Compare com as horas vazias do mês (01). No quadro 4, digite os atendimentos do mês por pagador e veja o resultado do mix."),
 ("Rotina","Antes de renovar, aceitar ou descredenciar um convênio, e uma vez por semestre para os que já estão. Leva 15 minutos por procedimento. No exemplo: consulta, com as tabelas da 08 e as consultas realizadas em agosto."),
 ("Com a IA","Copie os quadros 2 e 3 e use o prompt \"Preço 03 · Vale a pena este convênio?\" da biblioteca do kit para preparar os argumentos da negociação de tabela (sem prometer descredenciamento: é decisão da clínica, com contrato)."),
])
proteger(wb); salvar(wb,"07-simulador-convenio-x-particular.xlsx","Simulador convênio × particular · Kit de Gestão para Médicos")
