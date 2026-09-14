"""Planilha 9 do Kit de Gestão para Médicos: Caixa da clínica. Gera 09-caixa-da-clinica.xlsx
Exemplo (dados.LANCAMENTOS): particular à vista pelo fechamento do dia (uma linha por dia e forma), parcelas a prazo pagas (14),
lotes de convênio pagos (13), saídas (custos fixos, pró-labore, repasse, materiais, taxas de cartão, impostos, movimentos dos sócios).
As planilhas 10, 11, 12, 17, 18 e 20 leem os mesmos totais (dados.TOTAIS)."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference

N=1500; R0=5; RN=R0+N-1; NCE=12; NCS=16; NREC=30   # ≈ 80 lançamentos/mês no exemplo → folga de mais de 12 meses
CAT_ENTRADA=dados.CAT_ENTRADA; CAT_SAIDA=dados.CAT_SAIDA
FORMAS=["Pix","Dinheiro","Cartão de débito","Cartão de crédito","Transferência","Boleto"]

def build():
    wb=Workbook()
    # ---------- Config ----------
    cfg=wb.active; cfg.title="Config"
    titulo(cfg,"Configurações","Células amarelas: você preenche. Categorias e recebedores alimentam as listas de Lançamentos e o Painel.",merge_to="L")
    cfg["A4"]="Clínica"; cfg["B4"]=f"{dados.CLINICA} (exemplo fictício)"
    cfg["A5"]="Ano do painel"; cfg["B5"]=2026
    cfg["A6"]="Mês do painel"; cfg["B6"]="Setembro"
    cfg["A7"]="Número do mês"; cfg["B7"]="=MATCH(B6,$L$5:$L$16,0)"
    cfg["A8"]="Saldo em caixa antes do 1º lançamento"; cfg["B8"]=dados.SALDO_INICIAL
    cfg["A9"]="Data de referência"; cfg["B9"]="=TODAY()"
    for r in range(4,10): rotulo(cfg.cell(row=r,column=1))
    inp(cfg["B4"]); inp(cfg["B5"],center=True); inp(cfg["B6"],center=True); calc(cfg["B7"]); inp(cfg["B8"],BRL); calc(cfg["B9"],DATA)
    cfg["D4"]="Categorias de entrada (até 12)"; cfg["F4"]="Categorias de saída (até 16)"; cfg["H4"]="Todas as categorias (automático)"; cfg["J4"]="Convênios e outros recebedores (até 30)"; cfg["L4"]="Meses"
    for c in ("D4","F4","H4","J4","L4"): rotulo(cfg[c])
    for i in range(NCE): inp(cfg.cell(row=5+i,column=4))
    for i in range(NCS): inp(cfg.cell(row=5+i,column=6))
    for i in range(NREC): inp(cfg.cell(row=5+i,column=10))
    for i,v in enumerate(CAT_ENTRADA): cfg.cell(row=5+i,column=4,value=v)
    for i,v in enumerate(CAT_SAIDA): cfg.cell(row=5+i,column=6,value=v)
    for i,(n_,_,_) in enumerate(dados.CONVENIOS): cfg.cell(row=5+i,column=10,value=n_)
    cfg.cell(row=5+len(dados.CONVENIOS),column=10,value=dados.REN)
    for i,m_ in enumerate(MESES): cfg.cell(row=5+i,column=12,value=m_).font=F(size=10,color=TINTA)
    NE=f"COUNTA($D$5:$D${4+NCE})"; NS=f"COUNTA($F$5:$F${4+NCS})"
    for i in range(NCE+NCS):
        k=i+1
        cfg.cell(row=5+i,column=8,value=f'=IF({k}<={NE},INDEX($D$5:$D${4+NCE},{k}),IF({k}<={NE}+{NS},INDEX($F$5:$F${4+NCS},{k}-{NE}),""))').font=F(size=9,color=CINZA)
    notas=["Preencha categorias e recebedores de cima para baixo, sem pular linha: as listas suspensas de Lançamentos param na última linha preenchida.",
           "Saldo antes do 1º lançamento: o que havia na conta da clínica na data do primeiro lançamento (no exemplo, 01/01/2026).",
           "Particular à vista entra pelo fechamento do dia (uma linha por dia e forma de pagamento; o detalhe por paciente fica na agenda 01). Particular a prazo entra por paciente, quando a parcela é paga (14). Convênio entra por lote pago (13).",
           "A receber (Pago? = Não): parcelas a prazo (14) e lotes de convênio enviados (13) com previsão de recebimento até o fim do mês de referência — vencidos e a vencer. O que tem previsão para depois do fim do mês fica fora e entra quando aquele mês chegar. Por isso \"A receber\" do Painel é menor que a soma do total em aberto da 13 com o da 14: aquelas duas olham a carteira inteira, esta olha o mês.",
           "Quem lança: a RECEPÇÃO lança o fechamento do dia todo dia, no fechamento (planilha 04 · Checklist do dia). Na sexta, a sócia só CONFERE os cinco fechamentos da semana contra o extrato e marca o que caiu na conta (planilha 03 · Rotina da semana).",
           "Cartão: lance o valor bruto no dia da venda e, no fim do mês, uma saída \"Taxas de cartão\" com o total das taxas (a conciliação 16 calcula). Repasse à médica parceira é a saída do dia 10 (planilha 11). Lançamentos pagos do exemplo vão até 11/09/2026 (a última sexta)."]
    for i,t in enumerate(notas): cfg.cell(row=37+i,column=1,value=t); nota(cfg.cell(row=37+i,column=1))
    dv=lista("=Config!$L$5:$L$16"); dv.add("B6"); cfg.add_data_validation(dv)
    widths(cfg,(36,22,3,28,3,34,3,34,3,30,3,12)); cfg.sheet_view.showGridLines=False
    # ---------- Lançamentos ----------
    lan=wb.create_sheet("Lançamentos")
    titulo(lan,"Lançamentos","Uma linha por entrada ou saída. Preencha o amarelo; mês e ano são calculados. Pago? = Não fica em A receber / A pagar e só entra no caixa quando virar Sim. Regra do A receber: pré-lance as parcelas (14) e os lotes de convênio (13) com previsão de recebimento ATÉ O FIM DO MÊS DE REFERÊNCIA — nem mais, nem menos. O que vence depois entra no mês seguinte.",merge_to="K")
    hdr(lan,4,["Data","Tipo","Categoria","Paciente ou convênio","Referência","Descrição","Valor","Forma","Pago?","Mês","Ano"])
    for r in range(R0,RN+1):
        for c in range(1,10): inp(lan.cell(row=r,column=c))
        lan.cell(row=r,column=1).number_format=DATA; lan.cell(row=r,column=7).number_format=BRL
        for c in (1,2,8,9): lan.cell(row=r,column=c).alignment=Alignment(horizontal="center")
        lan.cell(row=r,column=10,value=f'=IF(A{r}="","",MONTH(A{r}))'); calc(lan.cell(row=r,column=10))
        lan.cell(row=r,column=11,value=f'=IF(A{r}="","",YEAR(A{r}))'); calc(lan.cell(row=r,column=11))
    dvs=[(lista('"Entrada,Saída"'),f"B{R0}:B{RN}"),
         (lista(f"=OFFSET(Config!$H$5,0,0,MAX(1,COUNTA(Config!$D$5:$D${4+NCE})+COUNTA(Config!$F$5:$F${4+NCS})),1)",strict=False),f"C{R0}:C{RN}"),
         (lista(f"=OFFSET(Config!$J$5,0,0,MAX(1,COUNTA(Config!$J$5:$J${4+NREC})),1)",strict=False),f"D{R0}:D{RN}"),
         (lista('"'+",".join(FORMAS)+'"'),f"H{R0}:H{RN}"),
         (lista('"Sim,Não"'),f"I{R0}:I{RN}"),
         (DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True),f"A{R0}:A{RN}"),
         (DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True),f"G{R0}:G{RN}")]
    for dv,rng_ in dvs: dv.add(rng_); lan.add_data_validation(dv)
    lan.conditional_formatting.add(f"A{R0}:K{RN}", FormulaRule(formula=[f'$B{R0}="Entrada"'], font=F(color=VERDE_T,size=10)))
    lan.conditional_formatting.add(f"A{R0}:K{RN}", FormulaRule(formula=[f'AND($A{R0}<>"",$I{R0}="Não")'], fill=fill(VERM)))
    lan.cell(row=RN+2,column=1,value=f"Esta aba tem {N} linhas ({R0} a {RN}): com cerca de 80 lançamentos por mês, sobra mais de um ano. Para virar o ano, o mais simples é começar um arquivo novo e levar o saldo de 31/12 para o campo \"Saldo em caixa antes do 1º lançamento\" da Config.").font=F(size=9,color=LILAS)
    widths(lan,(12,10,30,26,18,46,14,15,8,6,7)); lan.freeze_panes="A5"; lan.sheet_view.showGridLines=False; lan.auto_filter.ref=f"A4:K{RN}"
    assert len(dados.LANCAMENTOS)<=N
    for i,row in enumerate(dados.LANCAMENTOS):
        for c,v in enumerate(row,start=1): lan.cell(row=R0+i,column=c,value=v)
    # ---------- Painel ----------
    p=wb.create_sheet("Painel",0)
    titulo(p,'=Config!B4&" · Caixa da clínica · "&Config!B6&" de "&Config!B5',"Nada para digitar aqui. Escolha o mês em Config; tudo vem de Lançamentos. Só o que está com Pago? = Sim conta como caixa.",merge_to="L")
    M="Config!$B$7"; Y="Config!$B$5"
    LA=f"Lançamentos!$A${R0}:$A${RN}"; LB=f"Lançamentos!$B${R0}:$B${RN}"; LC=f"Lançamentos!$C${R0}:$C${RN}"; LD=f"Lançamentos!$D${R0}:$D${RN}"
    LG=f"Lançamentos!$G${R0}:$G${RN}"; LH=f"Lançamentos!$H${R0}:$H${RN}"; LI=f"Lançamentos!$I${R0}:$I${RN}"; LJ=f"Lançamentos!$J${R0}:$J${RN}"; LK=f"Lançamentos!$K${R0}:$K${RN}"
    def somames(tipo,m=M,y=Y,extra=""): return f'SUMIFS({LG},{LB},"{tipo}",{LI},"Sim",{LJ},{m},{LK},{y}{extra})'
    def saldo_ate(fim): return f'Config!$B$8+SUMIFS({LG},{LB},"Entrada",{LI},"Sim",{LA},"<="&{fim})-SUMIFS({LG},{LB},"Saída",{LI},"Sim",{LA},"<="&{fim})'
    kpi(p,4,1,"Entrou no mês",f"={somames('Entrada')}",VERDE,VERDE_T,fmt=BRL0)
    kpi(p,4,3,"Saiu no mês",f"={somames('Saída')}",VERM,VERM_T,fmt=BRL0)
    kpi(p,4,5,"Sobrou no mês","=A5-C5",SOL,UVA,fmt=BRL0)
    kpi(p,4,7,"Saldo acumulado",f"={saldo_ate(f'DATE({Y},{M}+1,0)')}",LAVANDA,UVA,fmt=BRL0)
    kpi(p,4,9,"A receber (Pago? = Não)",f'=SUMIFS({LG},{LB},"Entrada",{LI},"Não")',LAVANDA,UVA,fmt=BRL0)
    kpi(p,4,11,"A pagar (Pago? = Não)",f'=SUMIFS({LG},{LB},"Saída",{LI},"Não")',LAVANDA,UVA,fmt=BRL0)
    p.conditional_formatting.add("E5", FormulaRule(formula=['E5<0'], font=F(color="C8402E",size=16,bold=True)))
    p["A7"]="Para onde foi o dinheiro"; p["A7"].font=F(bold=True,size=13,color=UVA)
    hdr(p,8,["Categoria de saída","Valor no mês","% do total","Mês anterior","Acumulado no ano","Barra"]); p.merge_cells(start_row=8,start_column=6,end_row=8,end_column=9)
    def bloco_cat(r0_,tipo,col_cfg,n,kpi_ref):
        for i in range(n):
            r=r0_+i; src=f"Config!${col_cfg}${5+i}"
            p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
            p.cell(row=r,column=2,value=f'=IF({src}="","",{somames(tipo,extra=f",{LC},{src}")})'); calc(p.cell(row=r,column=2),BRL)
            p.cell(row=r,column=3,value=f'=IF(OR({src}="",{kpi_ref}=0),"",B{r}/{kpi_ref})'); calc(p.cell(row=r,column=3),PCT)
            p.cell(row=r,column=4,value=f'=IF({src}="","",IF({M}=1,{somames(tipo,"12",f"{Y}-1",f",{LC},{src}")},{somames(tipo,f"{M}-1",Y,f",{LC},{src}")}))'); calc(p.cell(row=r,column=4),BRL)
            p.cell(row=r,column=5,value=f'=IF({src}="","",SUMIFS({LG},{LB},"{tipo}",{LI},"Sim",{LC},{src},{LK},{Y},{LJ},"<="&{M}))'); calc(p.cell(row=r,column=5),BRL)
            p.cell(row=r,column=6,value=f'=IF(C{r}="","",REPT("█",ROUND(C{r}*40,0)))'); p.cell(row=r,column=6).font=F(size=10,color=LILAS); p.cell(row=r,column=6).border=borda
            p.merge_cells(start_row=r,start_column=6,end_row=r,end_column=9)
        t=r0_+n
        p.cell(row=t,column=1,value="Total"); p.cell(row=t,column=2,value=f"=SUM(B{r0_}:B{t-1})"); p.cell(row=t,column=4,value=f"=SUM(D{r0_}:D{t-1})"); p.cell(row=t,column=5,value=f"=SUM(E{r0_}:E{t-1})")
        for c in (1,2,4,5): p.cell(row=t,column=c).font=F(bold=True,color=UVA,size=10); p.cell(row=t,column=c).border=borda
        for c in (2,4,5): p.cell(row=t,column=c).number_format=BRL
        return t
    t1=bloco_cat(9,"Saída","F",NCS,"$C$5")          # 9..24, total 25
    r_e=t1+2
    p.cell(row=r_e,column=1,value="De onde veio o dinheiro").font=F(bold=True,size=13,color=UVA)
    hdr(p,r_e+1,["Categoria de entrada","Valor no mês","% do total","Mês anterior","Acumulado no ano","Barra"]); p.merge_cells(start_row=r_e+1,start_column=6,end_row=r_e+1,end_column=9)
    t2=bloco_cat(r_e+2,"Entrada","D",NCE,"$A$5")
    # por forma de pagamento
    r_f=t2+2
    p.cell(row=r_f,column=1,value="Entradas por forma de pagamento").font=F(bold=True,size=13,color=UVA)
    hdr(p,r_f+1,["Forma","Entrou no mês","% do mês","Entrou no ano","% do ano"])
    for i,fm in enumerate(FORMAS):
        r=r_f+2+i
        p.cell(row=r,column=1,value=fm); calc(p.cell(row=r,column=1),center=False)
        p.cell(row=r,column=2,value=f'={somames("Entrada",extra=f",{LH},A{r}")}'); calc(p.cell(row=r,column=2),BRL)
        p.cell(row=r,column=3,value=f'=IF($A$5=0,"",B{r}/$A$5)'); calc(p.cell(row=r,column=3),PCT)
        p.cell(row=r,column=4,value=f'=SUMIFS({LG},{LB},"Entrada",{LI},"Sim",{LH},A{r},{LK},{Y},{LJ},"<="&{M})'); calc(p.cell(row=r,column=4),BRL)
        p.cell(row=r,column=5,value=f'=IF(SUM($D${r_f+2}:$D${r_f+1+len(FORMAS)})=0,"",D{r}/SUM($D${r_f+2}:$D${r_f+1+len(FORMAS)}))'); calc(p.cell(row=r,column=5),PCT)
    t3=r_f+2+len(FORMAS)
    p.cell(row=t3,column=1,value="Total"); p.cell(row=t3,column=2,value=f"=SUM(B{r_f+2}:B{t3-1})"); p.cell(row=t3,column=4,value=f"=SUM(D{r_f+2}:D{t3-1})")
    for c in (1,2,4): p.cell(row=t3,column=c).font=F(bold=True,color=UVA,size=10); p.cell(row=t3,column=c).border=borda
    for c in (2,4): p.cell(row=t3,column=c).number_format=BRL
    p.cell(row=t3+1,column=1,value="Cartão entra pelo valor bruto no dia da venda; a taxa é a saída \"Taxas de cartão\" do fim do mês (conciliação na planilha 16). Convênio entra por transferência, por lote (planilha 13)."); nota(p.cell(row=t3+1,column=1))
    # ano
    A0=t3+3
    p.cell(row=A0,column=1,value="O ano, mês a mês").font=F(bold=True,size=13,color=UVA)
    hdr(p,A0+1,["Mês","Entrou","Saiu","Sobrou","Saldo ao fim do mês"])
    for i in range(12):
        r=A0+2+i
        p.cell(row=r,column=1,value=MESES[i]); calc(p.cell(row=r,column=1),center=False)
        p.cell(row=r,column=2,value=f'={somames("Entrada",i+1)}'); calc(p.cell(row=r,column=2),BRL)
        p.cell(row=r,column=3,value=f'={somames("Saída",i+1)}'); calc(p.cell(row=r,column=3),BRL)
        p.cell(row=r,column=4,value=f'=B{r}-C{r}'); calc(p.cell(row=r,column=4),BRL)
        p.cell(row=r,column=5,value=f'={saldo_ate(f"DATE({Y},{i+2},0)")}'); calc(p.cell(row=r,column=5),BRL)
    p.conditional_formatting.add(f"D{A0+2}:E{A0+13}", FormulaRule(formula=[f'D{A0+2}<0'], font=F(color="C8402E",size=10,bold=True)))
    p.cell(row=A0+14,column=1,value="Meses sem lançamento pago aparecem zerados. Compare só os meses já fechados; o mês corrente ainda não tem todas as saídas (pró-labore, guia de impostos)."); nota(p.cell(row=A0+14,column=1))
    bc=BarChart(); bc.type="col"; bc.grouping="clustered"; bc.height=7.5; bc.width=18; bc.title="Entrou × saiu no ano"; bc.style=2
    bc.add_data(Reference(p,min_col=2,max_col=3,min_row=A0+1,max_row=A0+13),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=A0+2,max_row=A0+13))
    bc.series[0].graphicalProperties.solidFill=LILAS; bc.series[1].graphicalProperties.solidFill=UVA; bc.legend.position="b"; bc.y_axis.majorGridlines=None
    p.add_chart(bc,f"G{A0}")
    widths(p,(34,15,11,15,17,10,10,10,10,4,12,12)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
    como_usar(wb,"Caixa da clínica",[
     ("O que esta planilha faz","Você lança o que entra (fechamento do dia do particular, parcelas a prazo, lotes de convênio) e o que sai (custo fixo, pró-labore, repasse, materiais, taxas de cartão, impostos). Ela mostra o mês: entrou, saiu, sobrou, saldo acumulado, a receber, a pagar; para onde foi o dinheiro por categoria, de onde veio, por forma de pagamento e o ano mês a mês com gráfico."),
     ("Passo 1","Em Config, preencha o nome da clínica, o ano, o saldo que havia em caixa antes do primeiro lançamento e o mês do painel. Ajuste categorias e recebedores de cima para baixo, sem pular linha."),
     ("Passo 2","Em Lançamentos, uma linha por movimento: data, tipo (entrada ou saída), categoria, paciente ou convênio (quando houver), referência, descrição, valor, forma e Pago?. Lote de convênio enviado e ainda não pago entra com Pago? = Não, na data da previsão, e aparece em A receber; quando cair na conta, troque para Sim e ajuste o valor pela glosa. Pré-lance só o que tem previsão até o fim do mês de referência: é essa a regra de A receber."),
     ("Passo 3","Em Painel, escolha o mês em Config e leia de cima para baixo. Só o que está com Pago? = Sim conta como caixa; o resto é previsão."),
     ("Limite e como estender","A aba Lançamentos tem 1.500 linhas (5 a 1504). No exemplo são cerca de 80 lançamentos por mês (fechamento do dia por forma, parcelas, lotes, custos fixos, pró-labore, impostos), o que dá folga para mais de um ano inteiro. Perto do fim, desproteja a aba (Revisar > Desproteger planilha), copie a última linha para baixo e ajuste o número final nas fórmulas do Painel — ou, o mais simples, comece um arquivo por ano e leve o saldo de 31/12 para o campo \"Saldo em caixa antes do 1º lançamento\" da Config."),
     ("Rotina de sexta","5 minutos: conferir os cinco fechamentos do dia da semana (quem lança é a recepção, todo dia, no checklist 04) e marcar as parcelas e os lotes que caíram na conta. Marcar as parcelas recebidas é outra rotina de sexta, na planilha 14. No fechamento do mês: lançar as taxas de cartão (16), olhar Sobrou e Saldo acumulado antes de decidir retirada extra ou gasto grande."),
     ("Ligação com as outras planilhas","O total de entradas do mês alimenta a planilha 10 (provisão de impostos) e a 17 (painel); as saídas de pró-labore e repasse, a 11; o saldo acumulado e o custo fixo, a 12 (reserva); as entradas por categoria, a 18 (resultado). Cada arquivo tem a própria aba de entrada amarela: copie os números do painel."),
     ("Com a IA","Copie \"Para onde foi o dinheiro\", \"De onde veio\" e \"O ano, mês a mês\" e use o prompt \"Caixa 01 · Explicar o mês do caixa\" da biblioteca do kit. Nunca cole nome de paciente na IA: use as tabelas do Painel, que só têm totais."),
    ])
    proteger(wb); salvar(wb,"09-caixa-da-clinica.xlsx","Caixa da clínica · Kit de Gestão para Médicos")

if __name__=="__main__":
    print(len(dados.LANCAMENTOS),"lançamentos de exemplo"); build()
