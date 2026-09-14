#!/usr/bin/env python3
"""Planilha 2 do Kit de Gestão para Médicos: Faltas, remarcações e lista de retorno. Gera 02-faltas-e-retornos.xlsx
A aba Agenda é cópia da Agenda da planilha 01 (a 01 é a fonte). Painel em Setembro (em andamento)."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference
from datetime import time

N=3000; R0=5; RN=R0+N-1            # mesma capacidade da Agenda da 01 (3.000 linhas ≈ 10 meses)
NPROF=8; NPAG=6; NPROC=12; TOP=25
def off(sheet,col,r0,r1): return f"=OFFSET({sheet}!${col}${r0},0,0,MAX(1,COUNTA({sheet}!${col}${r0}:${col}${r1})),1)"
NOME=f"{dados.CLINICA} (exemplo fictício)"
wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Listas iguais às da planilha 01; o retorno esperado por procedimento monta a lista de retorno.",merge_to="L")
cfg["A4"]="Clínica"; cfg["B4"]=NOME
cfg["A5"]="Ano do painel"; cfg["B5"]=2026
cfg["A6"]="Mês do painel"; cfg["B6"]="Setembro"
cfg["A7"]="Número do mês"; cfg["B7"]="=MATCH(B6,$L$5:$L$16,0)"
cfg["A8"]="Data de referência (hoje)"; cfg["B8"]="=TODAY()"
cfg["A9"]="Lista de retorno: incluir quem passou do prazo há mais de (dias)"; cfg["B9"]=0
for r in range(4,10): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"],center=True); inp(cfg["B6"],center=True); calc(cfg["B7"]); inp(cfg["B8"],DATA); inp(cfg["B9"],"0",center=True)
cfg["C9"]="0 = entra na lista no dia em que o retorno previsto passa. 7 = só uma semana depois."; nota(cfg["C9"])
cfg["L4"]="Meses"; rotulo(cfg["L4"])
for i,m in enumerate(MESES): cfg.cell(row=5+i,column=12,value=m).font=F(size=10,color=TINTA)
dv=lista("=Config!$L$5:$L$16",allow_blank=False); dv.add("B6"); cfg.add_data_validation(dv)
cfg["D4"]="Profissionais (até 8)"; cfg["F4"]="Pagadores (até 6)"; cfg["H4"]="Procedimentos (até 12)"; cfg["I4"]="Retorno em (dias)"
for c in ("D4","F4","H4","I4"): rotulo(cfg[c])
for i in range(NPROF): inp(cfg.cell(row=5+i,column=4))
for i in range(NPAG): inp(cfg.cell(row=5+i,column=6))
for i in range(NPROC): inp(cfg.cell(row=5+i,column=8)); inp(cfg.cell(row=5+i,column=9),"0",center=True)
for i,n in enumerate(dados.MEDICOS): cfg.cell(row=5+i,column=4,value=n)
for i,pg in enumerate(dados.PAGADORES): cfg.cell(row=5+i,column=6,value=pg)
for i,(p,d,m,r) in enumerate(dados.PROCEDIMENTOS): cfg.cell(row=5+i,column=8,value=p); cfg.cell(row=5+i,column=9,value=r)
cfg["D18"]="Copie as listas da planilha 01 (Config). Retorno em (dias): em quantos dias se espera o paciente de volta depois desse procedimento (0 = não gera retorno). Preencha de cima para baixo, sem pular linha."; nota(cfg["D18"])
widths(cfg,(52,16,3,26,3,16,3,22,14,3,3,12)); cfg.sheet_view.showGridLines=False
PROF_L=off("Config","D",5,4+NPROF); PAG_L=off("Config","F",5,4+NPAG); PROC_L=off("Config","H",5,4+NPROC)
PROCS=f"Config!$H$5:$H${4+NPROC}"; RETS=f"Config!$I$5:$I${4+NPROC}"; HOJE="Config!$B$8"; TOL="Config!$B$9"
# ---------- Agenda ----------
ag=wb.create_sheet("Agenda")
titulo(ag,"Agenda (cópia da planilha 01)","Cole aqui as colunas A a K da Agenda da planilha 01 (a 01 é a fonte). As colunas brancas calculam faltas, retorno previsto e a lista de retorno.",merge_to="T")
hdr(ag,4,["Data","Hora","Profissional","Sala","Paciente","Pagador","Procedimento","Situação","Forma de pagamento","Parcelas (cartão)","Observação","Mês","Ano","Dia da semana","Período","Retorno em (dias)","Retorno previsto","Voltou ou tem horário?","Dias além do previsto","Na lista de retorno?"],height=40)
AA=f"$A${R0}:$A${RN}"; AE=f"$E${R0}:$E${RN}"; AH=f"$H${R0}:$H${RN}"
for r in range(R0,RN+1):
    for c in range(1,12): inp(ag.cell(row=r,column=c),center=(c not in (5,11)))
    ag.cell(row=r,column=1).number_format=DATA; ag.cell(row=r,column=2).number_format="hh:mm"
    ag.cell(row=r,column=12,value=f'=IF(A{r}="","",MONTH(A{r}))'); calc(ag.cell(row=r,column=12))
    ag.cell(row=r,column=13,value=f'=IF(A{r}="","",YEAR(A{r}))'); calc(ag.cell(row=r,column=13))
    ag.cell(row=r,column=14,value=f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"Segunda","Terça","Quarta","Quinta","Sexta","Sábado","Domingo"))'); calc(ag.cell(row=r,column=14))
    ag.cell(row=r,column=15,value=f'=IF(B{r}="","",IF(B{r}<TIME(12,0,0),"Manhã",IF(B{r}<TIME(18,0,0),"Tarde","Noite")))'); calc(ag.cell(row=r,column=15))
    ag.cell(row=r,column=16,value=f'=IF(G{r}="","",IFERROR(INDEX({RETS},MATCH(G{r},{PROCS},0)),0))'); calc(ag.cell(row=r,column=16),"0")
    ag.cell(row=r,column=17,value=f'=IF(OR(H{r}<>"Realizado",P{r}="",P{r}=0),"",A{r}+P{r})'); calc(ag.cell(row=r,column=17),DATA)
    ag.cell(row=r,column=18,value=f'=IF(Q{r}="","",IF(COUNTIFS({AE},E{r},{AA},">"&A{r},{AH},"<>Falta",{AH},"<>Cancelado",{AH},"<>Remarcado")>0,"Sim","Não"))'); calc(ag.cell(row=r,column=18))
    ag.cell(row=r,column=19,value=f'=IF(R{r}<>"Não","",IF(Q{r}+{TOL}<={HOJE},{HOJE}-Q{r},""))'); calc(ag.cell(row=r,column=19),"0")
    ag.cell(row=r,column=20,value=f'=IF(S{r}="","","Sim")'); calc(ag.cell(row=r,column=20))
    ag.cell(row=r,column=21,value=f'=IF(T{r}="Sim",S{r}+ROW()/100000,0)'); ag.cell(row=r,column=21).font=F(color=CINZA,size=9)
    # faltas do paciente no mês do painel (só na primeira falta dele no mês, para a lista de reincidentes não repetir o nome)
    ag.cell(row=r,column=22,value=f'=IF(AND(H{r}="Falta",L{r}=Config!$B$7,M{r}=Config!$B$5,COUNTIFS($E${R0}:E{r},E{r},$H${R0}:H{r},"Falta",$L${R0}:L{r},Config!$B$7,$M${R0}:M{r},Config!$B$5)=1),COUNTIFS({AE},E{r},{AH},"Falta",$L${R0}:$L${RN},Config!$B$7,$M${R0}:$M${RN},Config!$B$5)*1000-ROW()/100000,0)'); ag.cell(row=r,column=22).font=F(color=CINZA,size=9)
ag.column_dimensions["U"].hidden=True; ag.column_dimensions["V"].hidden=True
dvs=[(lista(PROF_L,strict=False),f"C{R0}:C{RN}"),(lista(PAG_L,strict=False),f"F{R0}:F{RN}"),(lista(PROC_L,strict=False),f"G{R0}:G{RN}"),
     (lista('"'+",".join(dados.SITUACOES)+'"'),f"H{R0}:H{RN}"),(DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True),f"A{R0}:A{RN}")]
for dv,rng in dvs: dv.add(rng); ag.add_data_validation(dv)
ag.conditional_formatting.add(f"A{R0}:T{RN}", FormulaRule(formula=[f'$H{R0}="Falta"'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
ag.conditional_formatting.add(f"A{R0}:T{RN}", FormulaRule(formula=[f'$T{R0}="Sim"'], fill=fill("FFF4CC")))
ag.conditional_formatting.add(f"A{R0}:T{RN}", FormulaRule(formula=[f'OR($H{R0}="Cancelado",$H{R0}="Remarcado")'], font=F(color="8A86A0",size=10)))
ag.cell(row=RN+3,column=1,value=f"Esta aba tem {N} linhas ({R0} a {RN}), a mesma capacidade da Agenda da planilha 01: com cerca de 300 atendimentos por mês, dá 10 meses. Para estender, desproteja a aba, copie a última linha para baixo e ajuste o número final nas fórmulas do Painel — ou comece um arquivo por ano, junto com a 01.").font=F(size=9,color=LILAS)
ag.cell(row=RN+2,column=1,value="Vermelho: falta. Amarelo: paciente na lista de retorno (passou do retorno previsto e não tem nenhum horário depois). \"Voltou ou tem horário?\" olha qualquer atendimento posterior do paciente que não seja falta, cancelamento ou remarcação.").font=F(size=9,color=LILAS)
widths(ag,(11,7,22,8,26,12,18,11,16,9,24,6,6,11,8,9,12,12,11,11)); ag.freeze_panes="F5"; ag.sheet_view.showGridLines=False; ag.auto_filter.ref=f"A4:T{RN}"
for i,r_ in enumerate(dados.AGENDA):
    r=R0+i
    ag.cell(row=r,column=1,value=r_["data"] if r_["data"]<dados.HOJE else dados.prazo_formula(dados.dias(r_["data"])))
    ag.cell(row=r,column=2,value=time(r_["hora"]//60,r_["hora"]%60))
    for c,k in ((3,"profissional"),(4,"sala"),(5,"paciente"),(6,"pagador"),(7,"procedimento"),(8,"situacao")): ag.cell(row=r,column=c,value=r_[k])
    if r_["forma"]: ag.cell(row=r,column=9,value=r_["forma"])
    if r_["parcelas"]: ag.cell(row=r,column=10,value=r_["parcelas"])
    if r_["checkup"]: ag.cell(row=r,column=11,value="Check-up cardiológico (consulta + ECG + teste ergométrico)")
    elif r_["situacao"]=="Remarcado": ag.cell(row=r,column=11,value="Paciente pediu outra data")
# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
titulo(p,'=Config!$B$4&" · Faltas, remarcações e lista de retorno · "&Config!$B$6&" de "&Config!$B$5',"Nada para digitar aqui. Escolha o mês em Config; tudo vem de Agenda. Taxa de falta = faltas ÷ (faltas + realizados).",merge_to="L")
M="Config!$B$7"; Y="Config!$B$5"
A=lambda col: f"Agenda!${col}${R0}:${col}${RN}"
MES=f"{A('L')},{M},{A('M')},{Y}"
kpi(p,4,1,"Taxa de falta no mês",'=IF(C5+E5=0,"",C5/(C5+E5))',VERM,VERM_T,fmt="0.0%")
kpi(p,4,3,"Faltas",f'=COUNTIFS({A("H")},"Falta",{MES})',VERM,VERM_T,fmt="0")
kpi(p,4,5,"Realizados",f'=COUNTIFS({A("H")},"Realizado",{MES})',VERDE,VERDE_T,fmt="0")
kpi(p,4,7,"Cancelamentos",f'=COUNTIFS({A("H")},"Cancelado",{MES})',LAVANDA,UVA,fmt="0")
kpi(p,4,9,"Remarcações",f'=COUNTIFS({A("H")},"Remarcado",{MES})',LAVANDA,UVA,fmt="0")
kpi(p,4,11,"Na lista de retorno",f'=COUNTIF({A("T")},"Sim")',SOL,UVA,fmt="0")
p["A7"]="Cancelamento = o paciente avisou; remarcação = pediu outra data (o novo horário é outra linha). Falta = não veio e não avisou: é a que custa, porque o horário já não pode ser oferecido a ninguém. A lista de retorno vale para hoje, não só para o mês."; nota(p["A7"]); p.merge_cells("A7:L7"); p["A7"].alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[7].height=30
def tabela(r0,nome,itens,col_ag,cfg_col=None,n=None):
    p.cell(row=r0,column=1,value=nome).font=F(bold=True,size=13,color=UVA)
    hdr(p,r0+1,[nome.replace("Por ","").capitalize(),"Realizados","Faltas","Taxa de falta","Cancelamentos","Remarcações","Barra"]); p.merge_cells(start_row=r0+1,start_column=7,end_row=r0+1,end_column=9)
    n=n or len(itens)
    for i in range(n):
        r=r0+2+i
        if cfg_col: src=f"Config!${cfg_col}${5+i}"; p.cell(row=r,column=1,value=f'=IF({src}="","",{src})')
        else: src=f'"{itens[i]}"'; p.cell(row=r,column=1,value=itens[i])
        calc(p.cell(row=r,column=1),center=False)
        crit=f'{A(col_ag)},A{r}'
        p.cell(row=r,column=2,value=f'=IF(A{r}="","",COUNTIFS({crit},{A("H")},"Realizado",{MES}))'); calc(p.cell(row=r,column=2))
        p.cell(row=r,column=3,value=f'=IF(A{r}="","",COUNTIFS({crit},{A("H")},"Falta",{MES}))'); calc(p.cell(row=r,column=3))
        p.cell(row=r,column=4,value=f'=IF(A{r}="","",IF(B{r}+C{r}=0,"",C{r}/(B{r}+C{r})))'); calc(p.cell(row=r,column=4),"0.0%")
        p.cell(row=r,column=5,value=f'=IF(A{r}="","",COUNTIFS({crit},{A("H")},"Cancelado",{MES}))'); calc(p.cell(row=r,column=5))
        p.cell(row=r,column=6,value=f'=IF(A{r}="","",COUNTIFS({crit},{A("H")},"Remarcado",{MES}))'); calc(p.cell(row=r,column=6))
        p.cell(row=r,column=7,value=f'=IF(D{r}="","",REPT("█",ROUND(MIN(1,D{r}/0.25)*24,0)))'); p.cell(row=r,column=7).font=F(size=10,color="C0392B"); p.cell(row=r,column=7).border=borda; p.merge_cells(start_row=r,start_column=7,end_row=r,end_column=9)
    p.conditional_formatting.add(f"D{r0+2}:D{r0+1+n}", FormulaRule(formula=[f'AND(ISNUMBER(D{r0+2}),D{r0+2}>0.1)'], font=F(color="C8402E",size=10,bold=True)))
    return r0+2+n
r=tabela(9,"Por dia da semana",["Segunda","Terça","Quarta","Quinta","Sexta","Sábado"],"N")
r=tabela(r+1,"Por período",["Manhã","Tarde","Noite"],"O")
r=tabela(r+1,"Por pagador",None,"F",cfg_col="F",n=NPAG)
r=tabela(r+1,"Por profissional",None,"C",cfg_col="D",n=NPROF)
p.cell(row=r,column=1,value="Taxa acima de 10 % em vermelho. A barra vai até 25 %. Compare pagadores e dias antes de decidir a regra de confirmação (mensagem de véspera, lista de espera, política de remarcação).").font=F(size=9,color=LILAS)
# últimas 8 semanas
W0=r+2
p.cell(row=W0,column=1,value="Últimas 8 semanas (segunda a domingo)").font=F(bold=True,size=13,color=UVA)
hdr(p,W0+1,["Semana","Segunda-feira","Realizados","Faltas","Taxa de falta","Cancelamentos","Remarcações","Barra"])
for k in range(8):
    rr=W0+2+k; j=7-k
    p.cell(row=rr,column=2,value=f'={HOJE}-WEEKDAY({HOJE},2)+1-{7*j}'); calc(p.cell(row=rr,column=2),DATA)
    p.cell(row=rr,column=1,value=f'="S"&TEXT(B{rr},"dd/mm")'); calc(p.cell(row=rr,column=1))
    for c,st in ((3,"Realizado"),(4,"Falta"),(6,"Cancelado"),(7,"Remarcado")):
        p.cell(row=rr,column=c,value=f'=COUNTIFS({A("A")},">="&B{rr},{A("A")},"<="&B{rr}+6,{A("H")},"{st}")'); calc(p.cell(row=rr,column=c))
    p.cell(row=rr,column=5,value=f'=IF(C{rr}+D{rr}=0,"",D{rr}/(C{rr}+D{rr}))'); calc(p.cell(row=rr,column=5),"0.0%")
    p.cell(row=rr,column=8,value=f'=IF(E{rr}="","",REPT("█",ROUND(MIN(1,E{rr}/0.25)*24,0)))'); p.cell(row=rr,column=8).font=F(size=10,color="C0392B"); p.cell(row=rr,column=8).border=borda; p.merge_cells(start_row=rr,start_column=8,end_row=rr,end_column=10)
p.conditional_formatting.add(f"A{W0+2}:H{W0+9}", FormulaRule(formula=[f'$B{W0+2}>{HOJE}-7'], fill=fill(SOL)))
p.cell(row=W0+10,column=1,value="A semana atual (destacada) ainda está em andamento. A mensagem de confirmação de véspera costuma derrubar a taxa de falta; meça aqui se derrubou.").font=F(size=9,color=LILAS)
# lista de retorno
L0=W0+13
p.cell(row=L0,column=1,value="Lista de retorno: quem passou do retorno previsto e não tem horário marcado").font=F(bold=True,size=13,color=UVA)
p.cell(row=L0+1,column=1,value=f"Quem está há mais tempo sem voltar primeiro. Mostra os {TOP} primeiros; os demais ficam na aba Agenda (filtre \"Na lista de retorno?\" = Sim). Ligar ou mandar mensagem é decisão da clínica; a planilha só avisa."); nota(p.cell(row=L0+1,column=1)); p.merge_cells(start_row=L0+1,start_column=1,end_row=L0+1,end_column=12)
hdr(p,L0+2,["#","Paciente","Profissional","Pagador","Último atendimento","Procedimento","Retorno previsto","Dias além do previsto"])
KEY=A("U")
for k in range(1,TOP+1):
    rr=L0+2+k; m=f'MATCH(LARGE({KEY},{k}),{KEY},0)'; g=f'LARGE({KEY},{k})>0'
    p.cell(row=rr,column=1,value=k); calc(p.cell(row=rr,column=1))
    for col,src in zip((2,3,4,5,6,7,8),("E","C","F","A","G","Q","S")):
        p.cell(row=rr,column=col,value=f'=IFERROR(IF({g},INDEX({A(src)},{m}),""),"")'); calc(p.cell(row=rr,column=col),center=(col not in (2,3,6)))
    p.cell(row=rr,column=5).number_format=DATA; p.cell(row=rr,column=7).number_format=DATA; p.cell(row=rr,column=8).number_format="0"
p.conditional_formatting.add(f"A{L0+3}:H{L0+2+TOP}", FormulaRule(formula=[f'AND(ISNUMBER($H{L0+3}),$H{L0+3}>=30)'], fill=fill("FFF4CC")))
# reincidentes
Z0=L0+2+TOP+3
p.cell(row=Z0,column=1,value="Faltaram mais de uma vez no mês").font=F(bold=True,size=13,color=UVA)
hdr(p,Z0+1,["#","Paciente","Faltas no mês","Pagador","Profissional"])
KF=A("V")
for k in range(1,11):
    rr=Z0+1+k; m=f'MATCH(LARGE({KF},{k}),{KF},0)'; g=f'LARGE({KF},{k})>=2000'
    p.cell(row=rr,column=1,value=k); calc(p.cell(row=rr,column=1))
    p.cell(row=rr,column=2,value=f'=IFERROR(IF({g},INDEX({A("E")},{m}),""),"")'); calc(p.cell(row=rr,column=2),center=False)
    p.cell(row=rr,column=3,value=f'=IFERROR(IF({g},INT(LARGE({KF},{k})/1000),""),"")'); calc(p.cell(row=rr,column=3))
    p.cell(row=rr,column=4,value=f'=IFERROR(IF({g},INDEX({A("F")},{m}),""),"")'); calc(p.cell(row=rr,column=4))
    p.cell(row=rr,column=5,value=f'=IFERROR(IF({g},INDEX({A("C")},{m}),""),"")'); calc(p.cell(row=rr,column=5),center=False)
p.cell(row=Z0+12,column=1,value="Duas faltas no mesmo mês pedem uma conversa antes do próximo agendamento (confirmação obrigatória, horário de encaixe). Sem cobrança de multa aqui: é decisão da clínica e tem regras próprias.").font=F(size=9,color=LILAS)
p.merge_cells(start_row=Z0+12,start_column=1,end_row=Z0+12,end_column=12)
bc=BarChart(); bc.type="col"; bc.height=7; bc.width=14; bc.title="Taxa de falta por dia da semana"; bc.style=2
bc.add_data(Reference(p,min_col=4,min_row=10,max_row=15),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=11,max_row=15))
bc.series[0].graphicalProperties.solidFill="C0392B"; bc.legend=None; bc.y_axis.majorGridlines=None; bc.y_axis.number_format="0%"
p.add_chart(bc,"J9")
widths(p,(24,14,12,12,14,14,14,14,12,12,12,12)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
como_usar(wb,"Faltas, remarcações e lista de retorno",[
 ("O que esta planilha faz","Mede por que a agenda esvazia: taxa de falta por dia da semana, período, pagador e profissional, a série das últimas 8 semanas, quem faltou mais de uma vez no mês e a lista de retorno (pacientes que passaram do retorno previsto e não têm horário marcado)."),
 ("Passo 1","Em Config, confira as listas (iguais às da planilha 01) e, para cada procedimento, em quantos dias se espera o retorno. A data de referência fica em =HOJE()."),
 ("Passo 2","Em Agenda, cole as colunas A a K da Agenda da planilha 01 (a 01 é a fonte; não digite aqui uma agenda diferente). Toda sexta, cole de novo a agenda atualizada."),
 ("Passo 3","Em Painel, escolha o mês em Config. Leia de cima para baixo: taxa de falta, onde ela é maior, a tendência semanal, a lista de retorno e os reincidentes."),
 ("Rotina de segunda","4 minutos: a recepção liga ou manda mensagem para a lista de retorno (os 25 primeiros), oferece as vagas dos próximos 7 dias (Painel da 01) e registra as faltas da semana passada. Marque o retorno na 01: o paciente sai da lista sozinho. A rotina completa da semana (12 min na segunda, 18 na sexta) está na planilha 03."),
 ("Limite e como estender","A aba Agenda tem 3.000 linhas (5 a 3004), a mesma da planilha 01: cerca de 10 meses com 300 atendimentos por mês. Perto do fim, copie a última linha para baixo (desproteja a aba antes) e ajuste o número final nas fórmulas do Painel, ou comece um arquivo por ano junto com a 01."),
 ("Ligação com as outras planilhas","O Painel da clínica (17) copia a taxa de falta do mês e o tamanho da lista de retorno daqui. As metas do trimestre (19) acompanham os dois."),
 ("Com a IA","Copie \"Por dia da semana\" e \"Por pagador\" e use o prompt \"Agenda 02 · Reduzir faltas sem brigar com o paciente\" da biblioteca do kit; para a lista de retorno, \"Agenda 03 · Mensagem de retorno educada\" (sem nomes: a IA não precisa deles)."),
])
proteger(wb); salvar(wb,"02-faltas-e-retornos.xlsx","Faltas, remarcações e lista de retorno · Kit de Gestão para Médicos")
