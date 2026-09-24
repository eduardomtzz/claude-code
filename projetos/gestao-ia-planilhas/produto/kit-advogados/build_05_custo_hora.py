#!/usr/bin/env python3
"""Planilha 5 do Kit de Gestão para Advogados: Custo-hora do escritório. Gera 05-custo-hora.xlsx"""
from ssg import *
INC='"cadastro incompleto"'; MINV='"margens inválidas (Config)"'   # textos das guardas (auditoria final G01/G02)
from openpyxl.chart import BarChart, Reference
import dados
NCF=30; NP=10                       # linhas de custos fixos e de pessoas
CF0=5; CFN=CF0+NCF-1; CFT=CFN+2     # custos fixos: linhas 5..34, total na 36
P0=5; PN=P0+NP-1                    # pessoas: linhas 5..14
wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Margem e impostos definem a hora mínima a cobrar.",merge_to="F")
cfg["A4"]="Nome do escritório"; cfg["B4"]=f"{dados.ESCRITORIO} (exemplo fictício)"
cfg["A5"]="Mês de referência"; cfg["B5"]="Setembro de 2026"
cfg["A6"]="Data de referência"; cfg["B6"]=dados.HOJE
cfg["A7"]="Margem desejada sobre o preço (%)"; cfg["B7"]=dados.MARGEM
cfg["A8"]="Impostos e taxas sobre o que entra (%)"; cfg["B8"]=dados.ALIQ
cfg["A9"]="Arredondar a hora mínima para múltiplos de (R$)"; cfg["B9"]=5
for r in range(4,10): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"]); calc(cfg["B6"],DATA); inp(cfg["B7"],PCT,center=True); inp(cfg["B8"],PCT,center=True); inp(cfg["B9"],BRL0,center=True)
# Auditoria final (G02): imposto e margem são percentuais de 0 a 99 % e a soma tem de ficar
# abaixo de 100 %; senão o denominador (1 − margem − impostos) zera ou inverte e o preço
# mínimo sai negativo. A validação barra a digitação; as fórmulas guardam a colagem.
def _val_pct(ws,cel,outro,msg):
    dv=DataValidation(type="custom",formula1=f'=IF(ISNUMBER({cel}),AND({cel}>=0,{cel}<1,{cel}+N({outro})<1),FALSE)',allow_blank=False,showErrorMessage=True,errorTitle="Percentual",error=msg)
    dv.add(cel); ws.add_data_validation(dv)
_val_pct(cfg,"B7","B8","Margem entre 0 % e 99 %, e margem + impostos abaixo de 100 %: senão nenhum preço cobre o custo.")
_val_pct(cfg,"B8","B7","Impostos entre 0 % e 99 %, e margem + impostos abaixo de 100 %: senão nenhum preço cobre o custo.")
# Auditoria final-2 (G02): um predicado só, visível em Config!B10, decide se margem e impostos
# valem; toda hora mínima e todo preço desta planilha consultam essa célula (a validação acima
# barra a digitação; esta célula cobre colagem, texto, branco e negativo).
cfg["A10"]="Margem e impostos conferem? (calculado)"; rotulo(cfg["A10"])
cfg["B10"]='=IF(AND(ISNUMBER(B7),ISNUMBER(B8)),IF(AND(B7>=0,B7<1,B8>=0,B8<1,B7+B8<1),"Sim","Não"),"Não")'; calc(cfg["B10"])
cfg["C10"]="\"Não\" suspende a hora mínima em todas as abas: cada percentual precisa ficar entre 0 % e 99 %, e margem + impostos abaixo de 100 %."; nota(cfg["C10"])
cfg["C7"]="Quanto do preço deve sobrar depois de pagar o custo do escritório. 30% é um ponto de partida; ajuste ao seu mercado."
cfg["C8"]="Percentual que sai de cada real recebido (imposto do escritório e taxa de recebimento). Exemplo; confira com o seu contador."
cfg["C9"]="Só para o número ficar redondo na proposta."
for r in (7,8,9): nota(cfg.cell(row=r,column=3))
cfg["A11"]="A hora mínima a cobrar é: custo-hora ÷ (1 − margem − impostos). Assim a margem é calculada sobre o preço, e não sobre o custo."; nota(cfg["A11"])
cfg["A12"]="Esta é a hora mínima do kit (no exemplo: R$ 106,57, arredondada para R$ 110). A planilha 06 usa a mesma conta; a planilha 08 acrescenta uma folga de 20 % para horas não previstas e por isso mostra um valor maior (\"hora mínima com folga\")."; nota(cfg["A12"])
widths(cfg,(46,26,90)); cfg.sheet_view.showGridLines=False
# ---------- Custos fixos ----------
cfx=wb.create_sheet("Custos fixos")
titulo(cfx,"Custos fixos do mês","Tudo que o escritório paga todo mês mesmo sem nenhum caso novo. Preencha de cima para baixo, sem pular linha.",merge_to="D")
hdr(cfx,4,["Custo fixo","Valor mensal (R$)","Observação"])
for r in range(CF0,CFN+1):
    inp(cfx.cell(row=r,column=1)); inp(cfx.cell(row=r,column=2),BRL,center=True); inp(cfx.cell(row=r,column=3))
cfx.cell(row=CFT,column=1,value="Total de custos fixos"); rotulo(cfx.cell(row=CFT,column=1)); cfx.cell(row=CFT,column=1).border=borda
# Auditoria final-4 (F4-G01): custo fixo com nome e sem valor, ou com texto, não é zero
cfx.cell(row=CFT,column=2,value=f'=IF(SUMPRODUCT(((A{CF0}:A{CFN}<>"")+(B{CF0}:B{CFN}<>"")>0)*(ISNUMBER(B{CF0}:B{CFN})=FALSE))>0,"custo fixo sem valor numérico",SUM(B{CF0}:B{CFN}))'); calc(cfx.cell(row=CFT,column=2),BRL); cfx.cell(row=CFT,column=2).font=F(bold=True,color=UVA,size=10)
cfx.cell(row=CFT+2,column=1,value="Pró-labore dos sócios vai na aba Pessoas. Se alguém da equipe já aparece aqui (bolsa de estágio, salário), marque \"Sim\" na aba Pessoas para não contar duas vezes.").font=F(size=9,color=LILAS)
widths(cfx,(38,18,50)); cfx.freeze_panes="A5"; cfx.sheet_view.showGridLines=False
# ---------- Pessoas ----------
pes=wb.create_sheet("Pessoas")
titulo(pes,"Pessoas, pró-labore e horas faturáveis","Uma linha por pessoa que trabalha nos casos. Horas faturáveis: as que podem ser cobradas de algum cliente (reunião interna e administração não contam).",merge_to="J")
hdr(pes,4,["Pessoa","Papel","Pró-labore ou salário (R$/mês)","Já está nos custos fixos?","Horas de trabalho no mês","Horas faturáveis no mês","Ocupação","Custo direto por hora","Custo-hora completo","Hora mínima a cobrar"],height=40)
IND_H="Painel!$B$19"   # custo indireto por hora faturável (calculado no Painel)
DIV="(1-Config!$B$7-Config!$B$8)"; CFGX='Config!$B$10<>"Sim"'
for r in range(P0,PN+1):
    inp(pes.cell(row=r,column=1)); inp(pes.cell(row=r,column=2)); inp(pes.cell(row=r,column=3),BRL,center=True); inp(pes.cell(row=r,column=4),center=True)
    inp(pes.cell(row=r,column=5),"0",center=True); inp(pes.cell(row=r,column=6),"0",center=True)
    pes.cell(row=r,column=7,value=f'=IF(OR(A{r}="",NOT(ISNUMBER(E{r})),N(E{r})=0,NOT(ISNUMBER(F{r}))),"",F{r}/E{r})'); calc(pes.cell(row=r,column=7),PCT)
    # Auditoria final (G01): pró-labore ou horas em branco não viram custo zero — a pessoa
    # escreve o que falta, e a coluna K (oculta) marca a linha para o Painel suspender o
    # custo-hora, a hora mínima e a sensibilidade com "cadastro incompleto".
    pes.cell(row=r,column=8,value=f'=IF(A{r}="","",IF(NOT(ISNUMBER(C{r})),"falta o pró-labore",IF(AND(D{r}<>"Sim",D{r}<>"Não"),"falta: já está nos custos fixos?",IF(OR(NOT(ISNUMBER(F{r})),F{r}=0),"faltam as horas faturáveis",C{r}/F{r}))))'); calc(pes.cell(row=r,column=8),BRL)
    pes.cell(row=r,column=9,value=f'=IF(H{r}="","",IF(NOT(ISNUMBER(H{r})),H{r},IF(NOT(ISNUMBER({IND_H})),{INC},H{r}+{IND_H})))'); calc(pes.cell(row=r,column=9),BRL)
    pes.cell(row=r,column=10,value=f'=IF(I{r}="","",IF(NOT(ISNUMBER(I{r})),I{r},IF({CFGX},{MINV},I{r}/{DIV})))'); calc(pes.cell(row=r,column=10),BRL)
    pes.cell(row=r,column=11,value=f'=IF(AND(A{r}<>"",OR(NOT(ISNUMBER(C{r})),NOT(ISNUMBER(F{r})),AND(D{r}<>"Sim",D{r}<>"Não"))),1,0)'); pes.cell(row=r,column=11).font=F(color=CINZA,size=9)
dv=lista('"Sim,Não"'); dv.add(f"D{P0}:D{PN}"); pes.add_data_validation(dv)
pes.cell(row=PN+2,column=1,value="Custo direto por hora = pró-labore ou salário ÷ horas faturáveis. Custo-hora completo = custo direto + rateio dos custos indiretos do escritório por hora faturável (o rateio é calculado no Painel). Hora mínima = custo-hora completo ÷ (1 − margem − impostos).").font=F(size=9,color=LILAS)
pes.cell(row=PN+3,column=1,value="Horas de trabalho no mês: 160 para tempo integral; 120 para meio período. Horas faturáveis: uma média realista; 60% a 70% das horas de trabalho já é bom para quem também administra o escritório.").font=F(size=9,color=LILAS)
widths(pes,(22,30,18,14,14,14,11,14,14,14)); pes.column_dimensions["K"].hidden=True; pes.freeze_panes="A5"; pes.sheet_view.showGridLines=False
# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
titulo(p,'=Config!$B$4&" · Custo-hora · "&Config!$B$5',"Nada para digitar aqui, exceto os percentuais da sensibilidade. Custos vêm da aba Custos fixos, pessoas da aba Pessoas, margem e impostos de Config.",merge_to="J")
PC=f"Pessoas!$C${P0}:$C${PN}"; PD=f"Pessoas!$D${P0}:$D${PN}"; PE=f"Pessoas!$E${P0}:$E${PN}"; PF=f"Pessoas!$F${P0}:$F${PN}"
FLAGP=f"SUM(Pessoas!$K${P0}:$K${PN})"   # pessoas com pró-labore, horas ou "já nos fixos?" em branco
FLAGC=f"IF(ISNUMBER('Custos fixos'!$B${CFT}),0,1)"   # custo fixo sem valor numérico (auditoria final-4)
FLAG=f"({FLAGP}+{FLAGC})"
p["A6"]=(f'=IF({FLAG}=0,"",IF({FLAGC}>0,"Atenção: há custo fixo sem valor numérico na aba Custos fixos. ","")&IF({FLAGP}>0,"Atenção: "&{FLAGP}&" pessoa(s) na aba Pessoas com pró-labore, horas faturáveis ou \'já está nos custos fixos?\' em branco. ","")'
         f'&"Custo do mês, custo-hora e hora mínima ficam como \'cadastro incompleto\' até você completar: uma soma parcial daria um custo-hora MENOR do que o real.")')
p["A6"].font=F(size=10,bold=True,color=VERM_T); p.merge_cells("A6:J6"); p["A6"].alignment=Alignment(wrap_text=True,vertical="top")
kpi(p,4,1,"Custo total do mês","=B11",LAVANDA,UVA,fmt=BRL0)
kpi(p,4,3,"Horas faturáveis no mês","=B12",LAVANDA,UVA,fmt="#,##0")
kpi(p,4,5,"Custo-hora do escritório","=B13",SOL,UVA,fmt=BRL)
kpi(p,4,7,"Hora mínima a cobrar","=B17",VERDE,VERDE_T,fmt=BRL)
kpi(p,4,9,"Tempo faturável (planejado)","=B21",LAVANDA,UVA,fmt=PCT)
p["A7"]="Como chegamos ao número"; p["A7"].font=F(bold=True,size=13,color=UVA)
hdr(p,8,["Passo","Valor","De onde vem"])
linhas=[
 ("Custos fixos do mês",f"='Custos fixos'!B{CFT}",BRL,"Aba Custos fixos"),
 ("Pró-labore e salários fora dos custos fixos",f'=IF({FLAG}>0,{INC},SUMIFS({PC},{PD},"Não"))',BRL,"Aba Pessoas, quem está marcado \"Não\""),
 ("Custo total do mês",f'=IF({FLAG}>0,{INC},B9+B10)',BRL,"Soma dos dois acima"),
 ("Horas faturáveis no mês (todas as pessoas)",f'=IF({FLAG}>0,{INC},SUM({PF}))',"#,##0","Aba Pessoas"),
 ("Custo-hora do escritório",f'=IF({FLAG}>0,{INC},IF(B12=0,"",B11/B12))',BRL,"Custo total ÷ horas faturáveis"),
 ("Margem desejada sobre o preço","=Config!$B$7",PCT,"Config"),
 ("Impostos e taxas sobre o que entra","=Config!$B$8",PCT,"Config"),
 ("Hora mínima a cobrar (exata)",f'=IF(B13="","",IF(NOT(ISNUMBER(B13)),B13,IF({CFGX},{MINV},B13/{DIV})))',BRL,"Custo-hora ÷ (1 − margem − impostos)"),
 ("Hora mínima a cobrar (arredondada)",'=IF(B16="","",IF(NOT(ISNUMBER(B16)),B16,CEILING(B16,Config!$B$9)))',BRL,"Arredondada para cima, no múltiplo de Config"),
 ("Custos indiretos (custos fixos menos pessoas já contadas neles)",f'=IF({FLAG}>0,{INC},B9-SUMIFS({PC},{PD},"Sim"))',BRL,"Para ratear entre as pessoas"),
 ("Custo indireto por hora faturável",f'=IF({FLAG}>0,{INC},IF(B12=0,0,B18/B12))',BRL,"Custos indiretos ÷ horas faturáveis"),
 ("Horas de trabalho no mês (todas as pessoas)",f'=IF(SUMPRODUCT((Pessoas!$A${P0}:$A${PN}<>"")*(ISNUMBER({PE})=FALSE))>0,"faltam horas de trabalho",SUM({PE}))',"#,##0","Aba Pessoas"),
 ("Tempo faturável planejado (horas faturáveis ÷ horas de trabalho)",'=IF(OR(NOT(ISNUMBER(B20)),N(B20)=0,NOT(ISNUMBER(B12))),"",B12/B20)',PCT,"Quanto do tempo pago vira hora cobrável (a planilha 16 mede o realizado contra a meta)"),
]
for i,(a,f_,fmt,c) in enumerate(linhas):
    r=9+i
    p.cell(row=r,column=1,value=a); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f_); calc(p.cell(row=r,column=2),fmt)
    p.cell(row=r,column=3,value=c); calc(p.cell(row=r,column=3),center=False); nota(p.cell(row=r,column=3))
for r in (11,13,17): p.cell(row=r,column=2).font=F(bold=True,color=UVA,size=10)
for r in range(9,22): p.merge_cells(start_row=r,start_column=3,end_row=r,end_column=5)
# por pessoa
r0=24
p.cell(row=r0,column=1,value="Por pessoa").font=F(bold=True,size=13,color=UVA)
hdr(p,r0+1,["Pessoa","Papel","Horas faturáveis","Ocupação","Custo direto por hora","Custo-hora completo","Hora mínima a cobrar"],height=30)
for i in range(NP):
    r=r0+2+i; src=f"Pessoas!$A${P0+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",Pessoas!$B${P0+i})'); calc(p.cell(row=r,column=2),center=False)
    p.cell(row=r,column=3,value=f'=IF({src}="","",Pessoas!$F${P0+i})'); calc(p.cell(row=r,column=3),"#,##0")
    p.cell(row=r,column=4,value=f'=IF({src}="","",Pessoas!$G${P0+i})'); calc(p.cell(row=r,column=4),PCT)
    p.cell(row=r,column=5,value=f'=IF({src}="","",Pessoas!$H${P0+i})'); calc(p.cell(row=r,column=5),BRL)
    p.cell(row=r,column=6,value=f'=IF({src}="","",Pessoas!$I${P0+i})'); calc(p.cell(row=r,column=6),BRL)
    p.cell(row=r,column=7,value=f'=IF({src}="","",Pessoas!$J${P0+i})'); calc(p.cell(row=r,column=7),BRL)
p.conditional_formatting.add(f"D{r0+2}:D{r0+1+NP}", FormulaRule(formula=[f'AND(ISNUMBER(D{r0+2}),D{r0+2}<0.5)'], fill=fill(AMARELO)))
p.cell(row=r0+2+NP,column=1,value="Ocupação abaixo de 50% fica destacada: ou há horas trabalhadas que não estão sendo cobradas, ou falta caso para a pessoa.").font=F(size=9,color=LILAS)
# sensibilidade
s0=r0+NP+4
p.cell(row=s0,column=1,value="Sensibilidade: e se as horas faturáveis caírem?").font=F(bold=True,size=13,color=UVA)
p.cell(row=s0+1,column=1,value="Mês fraco, férias, um sócio doente: o custo fixo continua e cada hora fica mais cara. Os percentuais em amarelo podem ser trocados.").font=F(size=9,color=LILAS)
hdr(p,s0+2,["Cenário","Queda nas horas","Horas faturáveis","Custo-hora","Hora mínima a cobrar","Diferença por hora"],height=30)
for i,q in enumerate((0,0.10,0.20,0.30)):
    r=s0+3+i
    p.cell(row=r,column=1,value="Hoje" if i==0 else f'="Queda de "&ROUND(B{r}*100,0)&"%"'); calc(p.cell(row=r,column=1),center=False)
    if i==0:
        p.cell(row=r,column=2,value=0); calc(p.cell(row=r,column=2),PCT)
    else:
        p.cell(row=r,column=2,value=q); inp(p.cell(row=r,column=2),PCT,center=True)
    p.cell(row=r,column=3,value=f'=IF(ISNUMBER($B$12),$B$12*(1-B{r}),"")'); calc(p.cell(row=r,column=3),"#,##0")
    p.cell(row=r,column=4,value=f'=IF(OR(C{r}="",C{r}=0,NOT(ISNUMBER($B$11))),"",$B$11/C{r})'); calc(p.cell(row=r,column=4),BRL)
    p.cell(row=r,column=5,value=f'=IF(D{r}="","",IF({CFGX},{MINV},D{r}/{DIV}))'); calc(p.cell(row=r,column=5),BRL)
    p.cell(row=r,column=6,value=f'=IF(D{r}="","",D{r}-$D${s0+3})'); calc(p.cell(row=r,column=6),BRL)
bc=BarChart(); bc.type="col"; bc.height=7; bc.width=14; bc.title="Custo-hora por cenário"; bc.style=2
bc.add_data(Reference(p,min_col=4,min_row=s0+2,max_row=s0+6),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=s0+3,max_row=s0+6))
bc.series[0].graphicalProperties.solidFill="3B1F5E"; bc.legend=None; bc.y_axis.majorGridlines=None
p.add_chart(bc,f"H{s0}")
widths(p,(44,16,16,16,16,16,16,14,14,14)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
# ---------- exemplos (dados.py) ----------
for i,(nome,valor) in enumerate(dados.CUSTOS_FIXOS):
    cfx.cell(row=CF0+i,column=1,value=nome); cfx.cell(row=CF0+i,column=2,value=valor)
cfx.cell(row=CF0+6,column=3,value="Mesma pessoa da aba Pessoas (marcada \"Sim\" lá)")
horas_trab=dados.HORAS_TRABALHO
nomes_fixos=[n.lower() for n,_ in dados.CUSTOS_FIXOS]
for i,(nome,papel,custo,hf) in enumerate(dados.PESSOAS):
    r=P0+i; ja="Sim" if any(papel.split(" ")[0].lower() in n for n in nomes_fixos) else "Não"
    for c,v in zip((1,2,3,4,5,6),(nome,papel,custo,ja,horas_trab[nome],hf)): pes.cell(row=r,column=c,value=v)
como_usar(wb,"Custo-hora do escritório",[
 ("O que esta planilha faz","Responde \"quanto custa uma hora do meu escritório?\". Soma custos fixos e pró-labore, divide pelas horas que podem ser cobradas de clientes e mostra o custo-hora, a hora mínima a cobrar (com margem e impostos) e o que acontece se as horas faturáveis caírem."),
 ("Passo 1","Em Custos fixos, liste o que o escritório paga todo mês (aluguel, contador, sistemas, bolsa de estágio, marketing). Uma linha por item, de cima para baixo."),
 ("Passo 2","Em Pessoas, cada pessoa com pró-labore ou salário, horas de trabalho e horas faturáveis por mês. Se o custo dela já está em Custos fixos, marque \"Sim\" para não contar duas vezes."),
 ("Passo 3","Em Config, a margem que você quer sobre o preço e o percentual de impostos e taxas que sai de cada recebimento (confira com o contador)."),
 ("Passo 4","Em Painel: custo-hora do escritório, hora mínima a cobrar, custo-hora de cada pessoa e a sensibilidade. Leve o CUSTO-HORA do escritório (Painel B13) como entrada das planilhas 06 (Simulador) e 08 (Tabela de referência): é sobre ele que elas aplicam imposto e margem. A hora mínima serve para comparar preço, não como entrada — usá-la no lugar do custo-hora aplica margem e imposto duas vezes."),
 ("Rotina","Revise uma vez por trimestre ou quando mudar aluguel, equipe ou pró-labore. Leva 10 minutos."),
 ("Com a IA","Copie \"Como chegamos ao número\" e use o prompt \"Honorários 01 · Entender o meu custo-hora\" da biblioteca do kit para explicar o custo-hora ao sócio em linguagem simples, ou para sugerir o que fazer se a ocupação estiver abaixo de 50%."),
])
proteger(wb); salvar(wb,"05-custo-hora.xlsx","Custo-hora do escritório · Kit de Gestão para Advogados")
