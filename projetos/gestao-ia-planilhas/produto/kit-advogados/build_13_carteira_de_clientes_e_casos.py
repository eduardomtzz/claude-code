#!/usr/bin/env python3
"""Planilha 13 do Kit de Gestão para Advogados: Carteira de Clientes e Casos.
Gera 13-carteira-de-clientes-e-casos.xlsx (Como usar, Painel, Config, Casos, Clientes)."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference

NC=200; R0=5; RNC=R0+NC-1          # Casos: linhas 5..204
NCL=60; RNL=R0+NCL-1               # Clientes: linhas 5..64
NL=10; L0=9; L1=L0+NL-1            # listas da Config: linhas 9..18
def off(sheet,col,r0,r1): return f"=OFFSET({sheet}!${col}${r0},0,0,MAX(1,COUNTA({sheet}!${col}${r0}:${col}${r1})),1)"
def dstr(ref): return f'TEXT(DAY({ref}),"00")&"/"&TEXT(MONTH({ref}),"00")&"/"&YEAR({ref})'
NOME_ESC=f"{dados.ESCRITORIO} (exemplo fictício)"

wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. As listas alimentam os campos de escolha em Casos e Clientes.",merge_to="H")
cfg["A4"]="Escritório"; cfg["B4"]=NOME_ESC
cfg["A5"]="Data de referência (hoje)"; cfg["B5"]="=TODAY()"
cfg["A6"]="Fase que encerra o caso"; cfg["B6"]="Encerrado"
for r in (4,5,6): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"],DATA); inp(cfg["B6"],center=True)
hdr(cfg,8,["Áreas de atuação"],start=1); hdr(cfg,8,["Fases do caso"],start=3); hdr(cfg,8,["Responsáveis"],start=5); hdr(cfg,8,["Modalidades de honorário"],start=7)
for r in range(L0,L1+1):
    for c in (1,3,5,7): inp(cfg.cell(row=r,column=c))
cfg.cell(row=L1+2,column=1,value="Preencha cada lista de cima para baixo, sem pular linha: as listas suspensas das outras abas param na primeira célula vazia. Até 10 itens por lista.").font=F(size=9,color=LILAS)
cfg.cell(row=L1+3,column=1,value="Fase que encerra o caso: os casos nessa fase contam como Encerrados no Painel; todos os outros contam como Ativos.").font=F(size=9,color=LILAS)
dvf=lista(off("Config","C",L0,L1)); dvf.add("B6"); cfg.add_data_validation(dvf)
widths(cfg,(28,32,20,3,20,3,26,3)); cfg.sheet_view.showGridLines=False
AREAS_L=off("Config","A",L0,L1); FASES_L=off("Config","C",L0,L1); RESP_L=off("Config","E",L0,L1); MOD_L=off("Config","G",L0,L1)
CLI_L=off("Clientes","A",R0,RNL)

# ---------- Clientes ----------
cl=wb.create_sheet("Clientes")
titulo(cl,"Clientes","Uma linha por cliente (pessoa física ou jurídica). Casos, valores e último caso aberto vêm de Casos.",merge_to="J")
hdr(cl,4,["Cliente","Tipo","Área principal","Casos","Casos ativos","Contratado (R$)","Recebido (R$)","A receber (R$)","% recebido","Último caso aberto"])
CB=f"Casos!$B${R0}:$B${RNC}"; CH=f"Casos!$H${R0}:$H${RNC}"; CI=f"Casos!$I${R0}:$I${RNC}"; CJ=f"Casos!$J${R0}:$J${RNC}"; CK=f"Casos!$K${R0}:$K${RNC}"; CL_=f"Casos!$L${R0}:$L${RNC}"
for r in range(R0,RNL+1):
    inp(cl.cell(row=r,column=1)); inp(cl.cell(row=r,column=2),center=True); inp(cl.cell(row=r,column=3),center=True)
    cl.cell(row=r,column=4,value=f'=IF(A{r}="","",COUNTIFS({CB},A{r}))'); calc(cl.cell(row=r,column=4))
    cl.cell(row=r,column=5,value=f'=IF(A{r}="","",COUNTIFS({CB},A{r},{CL_},"Ativo"))'); calc(cl.cell(row=r,column=5))
    cl.cell(row=r,column=6,value=f'=IF(A{r}="","",SUMIFS({CH},{CB},A{r}))'); calc(cl.cell(row=r,column=6),BRL0)
    cl.cell(row=r,column=7,value=f'=IF(A{r}="","",SUMIFS({CI},{CB},A{r}))'); calc(cl.cell(row=r,column=7),BRL0)
    cl.cell(row=r,column=8,value=f'=IF(A{r}="","",F{r}-G{r})'); calc(cl.cell(row=r,column=8),BRL0)
    cl.cell(row=r,column=9,value=f'=IF(OR(A{r}="",F{r}=0),"",G{r}/F{r})'); calc(cl.cell(row=r,column=9),PCT)
    cl.cell(row=r,column=10,value=f'=IF(A{r}="","",IF(_xlfn.MAXIFS({CK},{CB},A{r})=0,"",_xlfn.MAXIFS({CK},{CB},A{r})))'); calc(cl.cell(row=r,column=10),DATA)
    cl.cell(row=r,column=11,value=f'=IF(OR(A{r}="",F{r}=0),0,F{r}-ROW()/100000)'); cl.cell(row=r,column=11).font=F(color=CINZA,size=9)
cl.column_dimensions["K"].hidden=True
dvt=lista('"PF,PJ"'); dvt.add(f"B{R0}:B{RNL}"); cl.add_data_validation(dvt)
dva=lista(AREAS_L,strict=False); dva.add(f"C{R0}:C{RNL}"); cl.add_data_validation(dva)
cl.conditional_formatting.add(f"H{R0}:H{RNL}", FormulaRule(formula=[f'AND(ISNUMBER(H{R0}),H{R0}<0)'], font=F(color="C8402E",size=10,bold=True)))
widths(cl,(30,8,16,8,11,16,16,16,11,16)); cl.freeze_panes="B5"; cl.sheet_view.showGridLines=False; cl.auto_filter.ref=f"A4:J{RNL}"

# ---------- Casos ----------
cs=wb.create_sheet("Casos")
titulo(cs,"Casos","Uma linha por caso. Cliente, área, fase, responsável e modalidade vêm de listas; tipo, a receber e situação são calculados.",merge_to="M")
hdr(cs,4,["Número do processo ou referência","Cliente","Tipo","Área","Fase","Responsável","Modalidade","Valor contratado (R$)","Recebido (R$)","A receber (R$)","Abertura","Situação","% recebido"])
for r in range(R0,RNC+1):
    for c in (1,2,4,5,6,7,8,9,11): inp(cs.cell(row=r,column=c))
    for c in (4,5,6,7,11): cs.cell(row=r,column=c).alignment=Alignment(horizontal="center")
    cs.cell(row=r,column=8).number_format=BRL0; cs.cell(row=r,column=9).number_format=BRL0; cs.cell(row=r,column=11).number_format=DATA
    cs.cell(row=r,column=3,value=f'=IF(B{r}="","",IFERROR(INDEX(Clientes!$B${R0}:$B${RNL},MATCH(B{r},Clientes!$A${R0}:$A${RNL},0)),""))'); calc(cs.cell(row=r,column=3))
    cs.cell(row=r,column=10,value=f'=IF(B{r}="","",H{r}-I{r})'); calc(cs.cell(row=r,column=10),BRL0)
    cs.cell(row=r,column=12,value=f'=IF(B{r}="","",IF(E{r}=Config!$B$6,"Encerrado","Ativo"))'); calc(cs.cell(row=r,column=12))
    cs.cell(row=r,column=13,value=f'=IF(OR(B{r}="",H{r}=0,H{r}=""),"",I{r}/H{r})'); calc(cs.cell(row=r,column=13),PCT)
dvs=[lista(CLI_L,strict=False), lista(AREAS_L), lista(FASES_L), lista(RESP_L), lista(MOD_L),
     DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True),
     DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True)]
for dv,rng in zip(dvs,[f"B{R0}:B{RNC}",f"D{R0}:D{RNC}",f"E{R0}:E{RNC}",f"F{R0}:F{RNC}",f"G{R0}:G{RNC}",f"K{R0}:K{RNC}",f"H{R0}:I{RNC}"]): dv.add(rng); cs.add_data_validation(dv)
cs.conditional_formatting.add(f"I{R0}:I{RNC}", FormulaRule(formula=[f'AND(ISNUMBER(I{R0}),I{R0}>H{R0})'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
cs.conditional_formatting.add(f"A{R0}:M{RNC}", FormulaRule(formula=[f'$L{R0}="Encerrado"'], font=F(color="8A86A0",size=10)))
cs.cell(row=RNC+2,column=1,value="Recebido em vermelho: está maior que o valor contratado; confira o contrato ou o valor lançado. Linhas cinza: casos encerrados.").font=F(size=9,color=LILAS)
widths(cs,(28,28,7,15,13,16,12,16,14,14,12,11,10)); cs.freeze_panes="C5"; cs.sheet_view.showGridLines=False; cs.auto_filter.ref=f"A4:M{RNC}"

# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
titulo(p,f'=Config!$B$4&" · Carteira de clientes e casos · "&{dstr("Config!$B$5")}',"Nada para digitar aqui: tudo vem de Casos, Clientes e Config.",merge_to="L")
p.merge_cells("A1:L1")
CD=f"Casos!$D${R0}:$D${RNC}"; CF=f"Casos!$F${R0}:$F${RNC}"; CG=f"Casos!$G${R0}:$G${RNC}"
kpi(p,4,1,"Contratado (R$)",f'=SUM({CH})',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,3,"Recebido (R$)",f'=SUM({CI})',VERDE,VERDE_T,fmt=BRL0)
kpi(p,4,5,"A receber (R$)",f'=SUM({CH})-SUM({CI})',SOL,UVA,fmt=BRL0)
kpi(p,4,7,"% recebido",f'=IFERROR(C5/A5,0)',LAVANDA,UVA,fmt=PCT)
kpi(p,4,9,"Casos ativos",f'=COUNTIFS({CL_},"Ativo")',LAVANDA,UVA,fmt="0")
kpi(p,4,11,"Casos encerrados",f'=COUNTIFS({CL_},"Encerrado")',LAVANDA,UVA,fmt="0")
def tabela(r_titulo,nome,col_cfg,col_casos):
    p.cell(row=r_titulo,column=1,value=nome).font=F(bold=True,size=13,color=UVA)
    hdr(p,r_titulo+1,[nome.replace("Por ","").capitalize(),"Casos","Ativos","Contratado (R$)","Recebido (R$)","A receber (R$)","% recebido","Barra"])
    for i in range(NL):
        r=r_titulo+2+i; src=f"Config!${col_cfg}${L0+i}"
        p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
        p.cell(row=r,column=2,value=f'=IF({src}="","",COUNTIFS({col_casos},{src}))'); calc(p.cell(row=r,column=2))
        p.cell(row=r,column=3,value=f'=IF({src}="","",COUNTIFS({col_casos},{src},{CL_},"Ativo"))'); calc(p.cell(row=r,column=3))
        p.cell(row=r,column=4,value=f'=IF({src}="","",SUMIFS({CH},{col_casos},{src}))'); calc(p.cell(row=r,column=4),BRL0)
        p.cell(row=r,column=5,value=f'=IF({src}="","",SUMIFS({CI},{col_casos},{src}))'); calc(p.cell(row=r,column=5),BRL0)
        p.cell(row=r,column=6,value=f'=IF({src}="","",D{r}-E{r})'); calc(p.cell(row=r,column=6),BRL0)
        p.cell(row=r,column=7,value=f'=IF(OR({src}="",D{r}=0),"",E{r}/D{r})'); calc(p.cell(row=r,column=7),PCT)
        a=r_titulo+2; b=r_titulo+1+NL
        p.cell(row=r,column=8,value=f'=IF({src}="","",REPT("█",ROUND(IFERROR(D{r}/MAX($D${a}:$D${b}),0)*20,0)))'); p.cell(row=r,column=8).font=F(size=10,color=LILAS); p.cell(row=r,column=8).border=borda
    return r_titulo+2, r_titulo+1+NL
a1,b1=tabela(7,"Por área","A",CD)
a2,b2=tabela(b1+2,"Por responsável","E",CF)
a3,b3=tabela(b2+2,"Por modalidade de honorário","G",CG)
# Top 5 clientes
rt=b3+2
p.cell(row=rt,column=1,value="Top 5 clientes por valor contratado").font=F(bold=True,size=13,color=UVA)
hdr(p,rt+1,["#","Cliente","Tipo","Casos","Contratado (R$)","Recebido (R$)","A receber (R$)","% recebido"])
PK=f"Clientes!$K${R0}:$K${RNL}"
for k in range(1,6):
    r=rt+1+k; m=f'MATCH(LARGE({PK},{k}),{PK},0)'; g=f'LARGE({PK},{k})>0'
    p.cell(row=r,column=1,value=k); calc(p.cell(row=r,column=1))
    for col,src in zip((2,3,4,5,6,7,8),("A","B","D","F","G","H","I")):
        p.cell(row=r,column=col,value=f'=IFERROR(IF({g},INDEX(Clientes!${src}${R0}:${src}${RNL},{m}),""),"")'); calc(p.cell(row=r,column=col),center=(col!=2))
    for col in (5,6,7): p.cell(row=r,column=col).number_format=BRL0
    p.cell(row=r,column=8).number_format=PCT
# Ativos × encerrados
rs=rt+8
p.cell(row=rs,column=1,value="Casos ativos × encerrados").font=F(bold=True,size=13,color=UVA)
hdr(p,rs+1,["Situação","Casos","Contratado (R$)","Recebido (R$)","A receber (R$)","% recebido"])
for i,s in enumerate(("Ativo","Encerrado")):
    r=rs+2+i
    p.cell(row=r,column=1,value=s); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=COUNTIFS({CL_},A{r})'); calc(p.cell(row=r,column=2))
    p.cell(row=r,column=3,value=f'=SUMIFS({CH},{CL_},A{r})'); calc(p.cell(row=r,column=3),BRL0)
    p.cell(row=r,column=4,value=f'=SUMIFS({CI},{CL_},A{r})'); calc(p.cell(row=r,column=4),BRL0)
    p.cell(row=r,column=5,value=f'=C{r}-D{r}'); calc(p.cell(row=r,column=5),BRL0)
    p.cell(row=r,column=6,value=f'=IFERROR(D{r}/C{r},0)'); calc(p.cell(row=r,column=6),PCT)
p.cell(row=rs+4,column=1,value="A receber dos casos encerrados é honorário devido e ainda não recebido: vale conferir antes de arquivar. Fonte: abas Casos e Clientes.").font=F(size=9,color=LILAS)
bc=BarChart(); bc.type="bar"; bc.height=7; bc.width=15; bc.title="Contratado × recebido por área"; bc.style=2
bc.add_data(Reference(p,min_col=4,max_col=5,min_row=a1-1,max_row=b1),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=a1,max_row=b1))
bc.series[0].graphicalProperties.solidFill="B89BE0"; bc.series[1].graphicalProperties.solidFill="3B1F5E"; bc.legend.position="b"; bc.x_axis.majorGridlines=None
p.add_chart(bc,"J7")
widths(p,(26,26,12,15,15,15,12,24,3,12,12,12)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False

# ---------- Exemplo (Ferraz & Lima) ----------
for i,v in enumerate(dados.AREAS): cfg.cell(row=L0+i,column=1,value=v)
for i,v in enumerate(dados.FASES): cfg.cell(row=L0+i,column=3,value=v)
for i,(nome,_,_,_) in enumerate(dados.PESSOAS): cfg.cell(row=L0+i,column=5,value=nome)
for i,v in enumerate(dados.TIPOS_HON): cfg.cell(row=L0+i,column=7,value=v)
for i,(nome,tipo,area) in enumerate(dados.CLIENTES):
    cl.cell(row=R0+i,column=1,value=nome); cl.cell(row=R0+i,column=2,value=tipo); cl.cell(row=R0+i,column=3,value=area)
for i,c in enumerate(dados.CASOS):
    r=R0+i
    for col,v in zip((1,2,4,5,6,7,8,9,11),(c["numero"],c["cliente"],c["area"],c["fase"],c["responsavel"],c["tipo_hon"],c["valor_contratado"],c["recebido"],c["abertura"])):
        cs.cell(row=r,column=col,value=v)

como_usar(wb,"Carteira de Clientes e Casos",[
 ("O que esta planilha faz","Você cadastra os clientes e os casos com valor contratado e recebido; ela calcula o que falta receber, agrupa por área, responsável e modalidade, mostra os cinco maiores clientes e separa casos ativos de encerrados. Responde \"quanto ainda vai entrar?\"."),
 ("Passo 1","Em Config, confira o nome do escritório e as listas: áreas, fases, responsáveis e modalidades de honorário. Preencha de cima para baixo, sem pular linha."),
 ("Passo 2","Em Clientes, uma linha por cliente, com tipo (PF ou PJ) e área principal. O nome daqui alimenta a lista de clientes em Casos."),
 ("Passo 3","Em Casos, uma linha por caso: número do processo (ou uma referência sua, para consultivo), cliente, área, fase, responsável, modalidade, valor contratado, recebido até hoje e data de abertura. Atualize o recebido a cada pagamento."),
 ("Passo 4","Em Painel, leia contratado, recebido e a receber, os totais por área e por responsável, o top 5 de clientes e ativos × encerrados. Nada para digitar lá."),
 ("Rotina","Sexta-feira, 10 minutos: atualizar recebidos e fases. Dia 1 do mês: copiar o Painel para o Resumo do mês (planilha 20)."),
 ("Com a IA","Copie a tabela \"Por área\" e o \"Top 5 clientes\" e use o prompt \"Resumir a carteira\" da biblioteca para preparar a reunião de sócios ou a conversa com o contador."),
])
proteger(wb); salvar(wb,"13-carteira-de-clientes-e-casos.xlsx","Carteira de Clientes e Casos · Kit de Gestão para Advogados")
