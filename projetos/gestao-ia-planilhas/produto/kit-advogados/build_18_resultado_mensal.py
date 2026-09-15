#!/usr/bin/env python3
"""Planilha 18 do Kit de Gestão para Advogados: Resultado Mensal (DRE simplificada do escritório). Gera 18-resultado-mensal.xlsx
Receita por categoria do caixa (as seis categorias de entrada da 09), custos fixos (linhas de dados.CUSTOS_FIXOS), despesas de casos
e viagens, pró-labore fixo, impostos provisionados como % da receita (8 %), resultado e margem, mês a mês. Compara com o mês anterior
e com o previsto. Exemplo: janeiro a agosto de 2026 (agosto = último mês fechado), valores iguais ao Painel mensal da 09 por categoria."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference

MODAL=dados.CAT_ENTRADA                                              # Honorários fixos, por hora, de êxito, consultoria, reembolso, outras
VARIAVEIS=["Custas e despesas de processo","Deslocamento e viagens"]  # despesas de casos e viagens (categorias de saída da 09)
SOCIOS=dados.PRO_LABORE                                              # Marina 6.000, Rafael 6.000
DRE={m:dados.dre(m) for m in range(1,9)}                             # jan..ago (meses fechados)
NMESES=8
PREV_RECEITA=[22000,22000,24000,24000,26000,26000,27000,27000,27000,27000,28000,28000]
PREV_VARIAVEIS=[400]*12

wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche.",merge_to="E")
cfg["A4"]="Escritório"; cfg["B4"]=f"{dados.ESCRITORIO} (exemplo fictício)"
cfg["A5"]="Ano"; cfg["B5"]=2026
cfg["A6"]="Mês de análise"; cfg["B6"]="Agosto"
cfg["A7"]="Impostos provisionados (% da receita)"; cfg["B7"]=dados.ALIQ
cfg["A8"]="Número do mês"; cfg["B8"]="=MATCH(B6,$D$5:$D$16,0)"
cfg["A9"]="Mês anterior"; cfg["B9"]='=IF(B8>1,INDEX($D$5:$D$16,B8-1),"")'
for r in range(4,10): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"]); inp(cfg["B6"]); inp(cfg["B7"],"0.0%"); calc(cfg["B8"]); calc(cfg["B9"])
cfg["C6"]="Analise o último mês fechado (no exemplo, agosto: setembro ainda está em andamento no caixa)."; nota(cfg["C6"])
cfg["C7"]="Use a alíquota efetiva que o contador informar (Simples, ISS + IR ou lucro presumido). A provisão sai todo mês, mesmo que o imposto seja pago depois. No exemplo, 8 %, a mesma das planilhas 05, 06, 08, 09 e 10."; nota(cfg["C7"])
cfg["D4"]="Meses"; rotulo(cfg["D4"])
for i,m in enumerate(MESES): cfg.cell(row=5+i,column=4,value=m).font=F(size=10,color=TINTA)
dv=lista("=Config!$D$5:$D$16",allow_blank=False); dv.add("B6"); cfg.add_data_validation(dv)
widths(cfg,(36,18,4,14)); cfg.sheet_view.showGridLines=False

# ---------- Resultado (DRE) ----------
rs=wb.create_sheet("Resultado")
titulo(rs,"Resultado mensal do escritório","Preencha as linhas amarelas de cada mês com os totais por categoria do Painel do caixa (planilha 09). Totais, impostos provisionados, resultado e margem são calculados. Meses futuros ficam em branco.",merge_to="O")
hdr(rs,4,["Linha"]+[m[:3] for m in MESES]+["Total do ano","Média dos meses preenchidos"],height=30)
NC=len(dados.CUSTOS_FIXOS); NS=len(SOCIOS); NV=len(VARIAVEIS)
R_REC0=6; R_RECT=R_REC0+len(MODAL)             # 6..11 categorias, 12 receita total
R_CF0=R_RECT+2; R_CFT=R_CF0+NC                  # 14..21 custos, 22 total custos
R_VA0=R_CFT+2; R_VAT=R_VA0+NV                   # 24..25 despesas de casos e viagens, 26 total
R_PL0=R_VAT+2; R_PLT=R_PL0+NS                   # 28..29 sócios, 30 total pró-labore
R_IMP=R_PLT+1; R_SAI=R_IMP+1; R_RES=R_SAI+1; R_MAR=R_RES+1   # 31 impostos, 32 saídas, 33 resultado, 34 margem
def secao(r,txt): c=rs.cell(row=r,column=1,value=txt); c.font=F(bold=True,size=11,color=UVA); c.fill=fill(LAVANDA); [setattr(rs.cell(row=r,column=k),"fill",fill(LAVANDA)) for k in range(2,16)]
def linha_total(r,txt,fn,fmt=BRL0,bold=True):
    rs.cell(row=r,column=1,value=txt); rotulo(rs.cell(row=r,column=1),bold=bold); rs.cell(row=r,column=1).border=borda
    for m in range(12):
        col=L(2+m); cell=rs.cell(row=r,column=2+m,value=fn(col)); calc(cell,fmt)
        if bold: cell.font=F(bold=True,color=UVA,size=10)
def linhas_entrada(r0,nomes,valor):
    for j,nome in enumerate(nomes):
        r=r0+j; rs.cell(row=r,column=1,value=nome); rotulo(rs.cell(row=r,column=1),bold=False); rs.cell(row=r,column=1).border=borda
        for m in range(12):
            cell=rs.cell(row=r,column=2+m); inp(cell,BRL0,center=True)
            if m<NMESES: cell.value=valor(m+1,nome)
PREENCH=lambda col: f'COUNT({col}${R_REC0}:{col}${R_RECT-1},{col}${R_CF0}:{col}${R_CFT-1},{col}${R_VA0}:{col}${R_VAT-1},{col}${R_PL0}:{col}${R_PLT-1})=0'
secao(R_REC0-1,"Receita (o que entrou no mês, por categoria do caixa)")
linhas_entrada(R_REC0,MODAL,lambda m,n: DRE[m]["receita"][n])
linha_total(R_RECT,"Receita total",lambda col: f'=IF({PREENCH(col)},"",SUM({col}{R_REC0}:{col}{R_RECT-1}))')
secao(R_CF0-1,"Custos fixos")
linhas_entrada(R_CF0,[n for n,_ in dados.CUSTOS_FIXOS],lambda m,n: DRE[m]["fixos"][n])
linha_total(R_CFT,"Total de custos fixos",lambda col: f'=IF({PREENCH(col)},"",SUM({col}{R_CF0}:{col}{R_CFT-1}))')
secao(R_VA0-1,"Despesas de casos e viagens (pagas pelo escritório)")
linhas_entrada(R_VA0,VARIAVEIS,lambda m,n: DRE[m]["variaveis"][n])
linha_total(R_VAT,"Total de despesas de casos e viagens",lambda col: f'=IF({PREENCH(col)},"",SUM({col}{R_VA0}:{col}{R_VAT-1}))')
secao(R_PL0-1,"Pró-labore fixo dos sócios")
linhas_entrada(R_PL0,[n for n,_ in SOCIOS],lambda m,n: DRE[m]["pro_labore"][n])
linha_total(R_PLT,"Total de pró-labore",lambda col: f'=IF({PREENCH(col)},"",SUM({col}{R_PL0}:{col}{R_PLT-1}))')
linha_total(R_IMP,"Impostos provisionados (% da receita, Config)",lambda col: f'=IF({col}{R_RECT}="","",ROUND({col}{R_RECT}*Config!$B$7,0))',bold=False)
linha_total(R_SAI,"Total de saídas (custos + despesas + pró-labore + impostos)",lambda col: f'=IF({col}{R_RECT}="","",{col}{R_CFT}+{col}{R_VAT}+{col}{R_PLT}+{col}{R_IMP})')
linha_total(R_RES,"Resultado do mês (receita − saídas)",lambda col: f'=IF({col}{R_RECT}="","",{col}{R_RECT}-{col}{R_SAI})')
linha_total(R_MAR,"Margem (resultado ÷ receita)",lambda col: f'=IF(OR({col}{R_RECT}="",{col}{R_RECT}=0),"",{col}{R_RES}/{col}{R_RECT})',fmt="0.0%")
for r in list(range(R_REC0,R_RECT+1))+list(range(R_CF0,R_CFT+1))+list(range(R_VA0,R_VAT+1))+list(range(R_PL0,R_PLT+1))+[R_IMP,R_SAI,R_RES]:
    rs.cell(row=r,column=14,value=f'=IF(COUNT(B{r}:M{r})=0,"",SUM(B{r}:M{r}))'); calc(rs.cell(row=r,column=14),BRL0)
    rs.cell(row=r,column=15,value=f'=IF(COUNT(B{r}:M{r})=0,"",AVERAGE(B{r}:M{r}))'); calc(rs.cell(row=r,column=15),BRL0)
    if r in (R_RECT,R_CFT,R_VAT,R_PLT,R_SAI,R_RES):
        for k in (14,15): rs.cell(row=r,column=k).font=F(bold=True,color=UVA,size=10)
rs.cell(row=R_MAR,column=14,value=f'=IF(OR(N{R_RECT}="",N{R_RECT}=0),"",N{R_RES}/N{R_RECT})'); calc(rs.cell(row=R_MAR,column=14),"0.0%")
rs.cell(row=R_MAR,column=15,value=f'=IF(COUNT(B{R_MAR}:M{R_MAR})=0,"",AVERAGE(B{R_MAR}:M{R_MAR}))'); calc(rs.cell(row=R_MAR,column=15),"0.0%")
rs.conditional_formatting.add(f"B{R_RES}:O{R_RES}", FormulaRule(formula=[f'AND(ISNUMBER(B{R_RES}),B{R_RES}<0)'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
rs.conditional_formatting.add(f"B{R_RES}:O{R_RES}", FormulaRule(formula=[f'AND(ISNUMBER(B{R_RES}),B{R_RES}>=0)'], fill=fill(VERDE), font=F(color=VERDE_T,size=10,bold=True)))
rs.conditional_formatting.add("B4:M4", FormulaRule(formula=['COLUMN()-1=Config!$B$8'], fill=fill(SOL), font=F(color=UVA,size=10,bold=True)))
notas=["Receita por categoria: copie de \"De onde veio o dinheiro\" (Painel da 09). Honorário misto: a parte fixa entra em \"Honorários fixos\" e a de êxito em \"Honorários de êxito\" quando cada uma entrar. Custos e despesas: de \"Para onde foi o dinheiro\".",
       "Impostos são provisão (8 % do que entrou no mês), não a guia paga; no caixa aparece a guia do mês anterior. Retiradas extras, distribuição de lucro e despesas pessoais dos sócios não entram aqui: são movimentos sócio × escritório (planilha 11). Resultado negativo fica vermelho.",
       "Os valores de exemplo (janeiro a agosto de 2026, meses fechados) são os totais por categoria do Painel do Caixa (09) com Pago? = Sim; setembro fica em branco porque está em andamento. O Painel do escritório (17) e o Resumo do mês (20) usam estes mesmos números."]
for i,t in enumerate(notas):
    c=rs.cell(row=R_MAR+2+i,column=1,value=t); c.font=F(size=9,color=LILAS); rs.merge_cells(start_row=R_MAR+2+i,start_column=1,end_row=R_MAR+2+i,end_column=15)
    c.alignment=Alignment(wrap_text=True,vertical="top"); rs.row_dimensions[R_MAR+2+i].height=28
widths(rs,[48]+[11]*12+[14,16]); rs.freeze_panes="B5"; rs.sheet_view.showGridLines=False

# ---------- Previsto ----------
pv=wb.create_sheet("Previsto")
titulo(pv,"Previsto para o ano","Preencha o que espera de receita, custos fixos, despesas de casos e pró-labore em cada mês. O Painel compara o realizado com este previsto.",merge_to="N")
hdr(pv,4,["Linha"]+[m[:3] for m in MESES]+["Total do ano"],height=30)
PV=["Receita prevista","Custos fixos previstos","Despesas de casos e viagens previstas","Pró-labore previsto","Impostos previstos (% da receita, Config)","Saídas previstas","Resultado previsto","Margem prevista"]
PV_EX=[PREV_RECEITA,[dados.CUSTOS_FIXOS_TOTAL]*12,PREV_VARIAVEIS,[dados.PRO_LABORE_TOTAL]*12]
P_REC,P_CF,P_VA,P_PL,P_IMP,P_SAI,P_RES,P_MAR=range(5,13)
for j,nome in enumerate(PV):
    r=5+j; pv.cell(row=r,column=1,value=nome); rotulo(pv.cell(row=r,column=1),bold=(j>=5)); pv.cell(row=r,column=1).border=borda
    for m in range(12):
        col=L(2+m); cell=pv.cell(row=r,column=2+m)
        if j<4: inp(cell,BRL0,center=True); cell.value=PV_EX[j][m]
        elif j==4: cell.value=f'=IF({col}{P_REC}="","",ROUND({col}{P_REC}*Config!$B$7,0))'; calc(cell,BRL0)
        elif j==5: cell.value=f'=IF({col}{P_REC}="","",N({col}{P_CF})+N({col}{P_VA})+N({col}{P_PL})+{col}{P_IMP})'; calc(cell,BRL0)
        elif j==6: cell.value=f'=IF({col}{P_REC}="","",{col}{P_REC}-{col}{P_SAI})'; calc(cell,BRL0)
        else: cell.value=f'=IF(OR({col}{P_REC}="",{col}{P_REC}=0),"",{col}{P_RES}/{col}{P_REC})'; calc(cell,"0.0%")
    if j<7: pv.cell(row=r,column=14,value=f'=IF(COUNT(B{r}:M{r})=0,"",SUM(B{r}:M{r}))'); calc(pv.cell(row=r,column=14),BRL0)
    else: pv.cell(row=r,column=14,value=f'=IF(OR(N{P_REC}="",N{P_REC}=0),"",N{P_RES}/N{P_REC})'); calc(pv.cell(row=r,column=14),"0.0%")
pv.cell(row=14,column=1,value="Uma boa previsão de custos fixos é a média dos últimos três meses; de receita, o que já está contratado para o mês (parcelas a vencer na planilha 14) mais o que costuma fechar.").font=F(size=9,color=LILAS)
pv.merge_cells("A14:N14")
widths(pv,[40]+[11]*12+[14]); pv.freeze_panes="B5"; pv.sheet_view.showGridLines=False

# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
p["A1"]='=Config!B4&" · Resultado mensal · "&Config!B6&" de "&Config!B5'; p["A1"].font=F(bold=True,size=16,color=UVA); p.merge_cells("A1:H1")
p["A2"]="Nada para digitar aqui. Escolha o mês em Config; os números vêm de Resultado e de Previsto."; nota(p["A2"]); p.merge_cells("A2:H2")
M="Config!$B$8"
def mes(r,off=0):  # valor da linha r de Resultado no mês (ou mês anterior)
    idx=f"{M}-{off}" if off else M
    return f'INDEX(Resultado!$B${r}:$M${r},{idx})'
def val(r,fmt_blank='"—"'): return f'=IF({mes(r)}="",{fmt_blank},{mes(r)})'
kpi(p,4,1,"Receita do mês",val(R_RECT),LAVANDA,UVA,fmt=BRL0)
kpi(p,4,3,"Saídas do mês",val(R_SAI),LAVANDA,UVA,fmt=BRL0)
kpi(p,4,5,"Resultado",val(R_RES),VERDE,VERDE_T,fmt=BRL0)
kpi(p,4,7,"Margem",val(R_MAR),SOL,UVA,fmt="0.0%")
p.conditional_formatting.add("E5", FormulaRule(formula=['AND(ISNUMBER(E5),E5<0)'], fill=fill(VERM), font=F(size=16,bold=True,color=VERM_T)))
p["A7"]="Comparação com o mês anterior e com o previsto"; p["A7"].font=F(bold=True,size=13,color=UVA)
hdr(p,8,["Linha","","Mês","Mês anterior","Variação","Previsto","Vs. previsto","Situação"]); p.merge_cells("A8:B8")
COMP=[("Receita total",R_RECT,P_REC,"Maior é melhor"),("Custos fixos",R_CFT,P_CF,"Menor é melhor"),("Despesas de casos e viagens",R_VAT,P_VA,"Menor é melhor"),
      ("Pró-labore",R_PLT,P_PL,""),("Impostos provisionados",R_IMP,P_IMP,""),
      ("Total de saídas",R_SAI,P_SAI,"Menor é melhor"),("Resultado do mês",R_RES,P_RES,"Maior é melhor"),("Margem",R_MAR,P_MAR,"Maior é melhor")]
C0=9; CN=C0+len(COMP)-1
for i,(nome,r,pr,sent) in enumerate(COMP):
    rr=C0+i; pct=(r==R_MAR); fmt="0.0%" if pct else BRL0
    p.cell(row=rr,column=1,value=nome); rotulo(p.cell(row=rr,column=1),bold=(r in (R_RECT,R_SAI,R_RES))); p.cell(row=rr,column=1).border=borda; p.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=2)
    p.cell(row=rr,column=3,value=val(r)); calc(p.cell(row=rr,column=3),fmt)
    p.cell(row=rr,column=4,value=f'=IF({M}<=1,"—",IF({mes(r,1)}="","—",{mes(r,1)}))'); calc(p.cell(row=rr,column=4),fmt)
    if pct:
        p.cell(row=rr,column=5,value=f'=IF(OR(C{rr}="—",D{rr}="—"),"—",(C{rr}-D{rr})*100)'); calc(p.cell(row=rr,column=5),'+0.0" p.p.";-0.0" p.p.";0.0" p.p."')
    else:
        p.cell(row=rr,column=5,value=f'=IF(OR(C{rr}="—",D{rr}="—",D{rr}=0),"—",(C{rr}-D{rr})/ABS(D{rr}))'); calc(p.cell(row=rr,column=5),"+0.0%;-0.0%;0.0%")
    p.cell(row=rr,column=6,value=f'=IF(INDEX(Previsto!$B${pr}:$M${pr},{M})="","—",INDEX(Previsto!$B${pr}:$M${pr},{M}))'); calc(p.cell(row=rr,column=6),fmt)
    if pct:
        p.cell(row=rr,column=7,value=f'=IF(OR(C{rr}="—",F{rr}="—"),"—",(C{rr}-F{rr})*100)'); calc(p.cell(row=rr,column=7),'+0.0" p.p.";-0.0" p.p.";0.0" p.p."')
    else:
        p.cell(row=rr,column=7,value=f'=IF(OR(C{rr}="—",F{rr}="—",F{rr}=0),"—",(C{rr}-F{rr})/ABS(F{rr}))'); calc(p.cell(row=rr,column=7),"+0.0%;-0.0%;0.0%")
    if sent:
        cond=f'C{rr}<=F{rr}' if sent=="Menor é melhor" else f'C{rr}>=F{rr}'
        p.cell(row=rr,column=8,value=f'=IF(OR(C{rr}="—",F{rr}="—"),"—",IF({cond},"Dentro do previsto",IF({"C"+str(rr)+"<=F"+str(rr)+"*1.05" if sent=="Menor é melhor" else "C"+str(rr)+">=F"+str(rr)+"*0.95"},"Perto","Fora do previsto")))')
    else:
        p.cell(row=rr,column=8,value="Informativo")
    calc(p.cell(row=rr,column=8))
for cor,txt,fnt in ((VERDE,"Dentro do previsto",VERDE_T),(AMARELO,"Perto","7A5200"),(VERM,"Fora do previsto",VERM_T)):
    p.conditional_formatting.add(f"H{C0}:H{CN}", FormulaRule(formula=[f'H{C0}="{txt}"'], fill=fill(cor), font=F(color=fnt,size=10,bold=(txt=="Fora do previsto"))))
M0=CN+2
p.cell(row=M0,column=1,value="Receita do mês por categoria").font=F(bold=True,size=13,color=UVA)
hdr(p,M0+1,["Categoria","","Mês","Mês anterior","Variação","% da receita"]); p.merge_cells(start_row=M0+1,start_column=1,end_row=M0+1,end_column=2)
for j,nome in enumerate(MODAL):
    rr=M0+2+j; r=R_REC0+j
    p.cell(row=rr,column=1,value=nome); rotulo(p.cell(row=rr,column=1),bold=False); p.cell(row=rr,column=1).border=borda; p.merge_cells(start_row=rr,start_column=1,end_row=rr,end_column=2)
    p.cell(row=rr,column=3,value=val(r)); calc(p.cell(row=rr,column=3),BRL0)
    p.cell(row=rr,column=4,value=f'=IF({M}<=1,"—",IF({mes(r,1)}="","—",{mes(r,1)}))'); calc(p.cell(row=rr,column=4),BRL0)
    p.cell(row=rr,column=5,value=f'=IF(OR(C{rr}="—",D{rr}="—",D{rr}=0),"—",(C{rr}-D{rr})/ABS(D{rr}))'); calc(p.cell(row=rr,column=5),"+0.0%;-0.0%;0.0%")
    p.cell(row=rr,column=6,value=f'=IF(OR(C{rr}="—",C{C0}="—",C{C0}=0),"—",C{rr}/C{C0})'); calc(p.cell(row=rr,column=6),PCT)
G0=M0+2+len(MODAL)+1
p.cell(row=G0,column=1,value="Receita, saídas e resultado no ano (meses em branco não aparecem)"); nota(p.cell(row=G0,column=1))
bc=BarChart(); bc.type="col"; bc.grouping="clustered"; bc.title=None; bc.height=8.5; bc.width=24; bc.style=2
for r,cor in ((R_RECT,UVA),(R_SAI,LILAS),(R_RES,SOL)):
    bc.add_data(Reference(rs,min_col=1,max_col=13,min_row=r,max_row=r),from_rows=True,titles_from_data=True)
for sr,cor in zip(bc.series,(UVA,LILAS,SOL)): sr.graphicalProperties.solidFill=cor; sr.graphicalProperties.line.solidFill=cor
bc.set_categories(Reference(rs,min_col=2,max_col=13,min_row=4,max_row=4))
bc.y_axis.majorGridlines=None; bc.legend.position="b"; bc.y_axis.number_format='"R$" #,##0'
p.add_chart(bc,f"A{G0+1}")
p.cell(row=G0+19,column=1,value='Variação da margem em pontos percentuais (p.p.). Para explicar o resultado ao sócio ou ao contador, use a planilha 20 · Resumo do mês e o prompt "Painel 01 · Explicar o mês ao sócio".').font=F(size=9,color=LILAS)
p.merge_cells(start_row=G0+19,start_column=1,end_row=G0+19,end_column=8)
widths(p,(26,16,14,14,13,14,13,18)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False

como_usar(wb,"Resultado Mensal",[
 ("O que esta planilha faz","Uma DRE simplificada do escritório: receita por categoria (as mesmas do caixa), custos fixos, despesas de casos e viagens, pró-labore fixo, impostos provisionados como % da receita, resultado e margem, mês a mês. Compara o mês com o anterior e com o que você previu."),
 ("Passo 1","Em Config, informe o ano, o mês de análise (o último mês fechado) e a % de impostos a provisionar (peça ao contador a alíquota efetiva)."),
 ("Passo 2","Em Resultado, preencha por mês o que entrou em cada categoria e o que saiu em cada linha de custo fixo, despesa de casos e pró-labore. Os totais por categoria estão no Painel da planilha 09 · Caixa do escritório (\"De onde veio\" e \"Para onde foi o dinheiro\")."),
 ("Passo 3","Em Previsto, escreva o que espera de receita, custos fixos, despesas e pró-labore para cada mês. Pode ser o mesmo valor o ano inteiro."),
 ("Passo 4","Abra Painel: receita, saídas, resultado e margem do mês, comparação com o mês anterior e com o previsto (variação da margem em p.p.), receita por categoria e o gráfico do ano."),
 ("Fim do mês","Depois de fechar o caixa, lance o mês aqui e troque o mês em Config. Leva 10 minutos. No exemplo, o mês analisado é agosto de 2026 (setembro ainda está em andamento)."),
 ("Com a IA","Copie a tabela do Painel para o prompt \"Painel 01 · Explicar o mês ao sócio\", \"Caixa 02 · Comparar dois meses do resultado\" ou \"Caixa 04 · Preparar a reunião mensal com o contador\" da biblioteca do kit. A planilha 20 · Resumo do mês já escreve o texto-base."),
])
proteger(wb); salvar(wb,"18-resultado-mensal.xlsx","Resultado Mensal · Kit de Gestão para Advogados")
