#!/usr/bin/env python3
"""Planilha 17 do Kit de Gestão para Advogados: Painel do Escritório. Gera 17-painel-do-escritorio.xlsx
Uma tela com prazos, horas, caixa, recebíveis e propostas. Os números vêm da aba Dados (copiados toda sexta das
planilhas 01, 09, 13, 14, 15 e 16); o Histórico guarda uma linha por mês.
Exemplo: sexta 11/09/2026 (setembro em andamento), números copiados exatamente dos painéis das planilhas 01, 09, 13, 14, 15 e 16;
Histórico jan–ago do Painel mensal da 09 e da 16 (dados.HISTORICO). Fonte única: dados.py."""
from ssg import *
import dados
from openpyxl.chart import BarChart, LineChart, Reference

# ---------- exemplo (fonte única: dados.py) ----------
# Dados = a sexta 11/09/2026 (setembro em andamento): prazos da 01, horas da 16 (Config = Setembro), caixa da 09 (Painel de setembro),
# carteira da 13, parcelas da 14, propostas da 15. Histórico jan–ago = Painel mensal da 09 (entrou/saiu reais do caixa) e 16 (horas,
# só de julho em diante: o controle de horas começou em julho); carteira, parcelas e propostas no fim de cada mês (dados.estado).
PZ=dados.resumo_prazos(); E=dados.estado(dados.HOJE); T9=dados.TOTAIS[9]; H9=dados.horas_mes(9); HIST=dados.HISTORICO
wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche.",merge_to="E")
cfg["A4"]="Escritório"; cfg["B4"]=f"{dados.ESCRITORIO} (exemplo fictício)"
cfg["A5"]="Mês do painel"; cfg["B5"]="Setembro"
cfg["A6"]="Ano"; cfg["B6"]=2026
cfg["A7"]="Data de referência (hoje)"; cfg["B7"]="=TODAY()"
cfg["A8"]="Número do mês"; cfg["B8"]="=MATCH(B5,$D$5:$D$16,0)"
cfg["A9"]="Mês anterior"; cfg["B9"]='=IF(B8>1,INDEX($D$5:$D$16,B8-1),"")'
for r in range(4,10): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"]); inp(cfg["B6"]); inp(cfg["B7"],DATA); calc(cfg["B8"]); calc(cfg["B9"])
cfg["D4"]="Meses"; rotulo(cfg["D4"])
for i,m in enumerate(MESES): cfg.cell(row=5+i,column=4,value=m).font=F(size=10,color=TINTA)
dv=lista("=Config!$D$5:$D$16",allow_blank=False); dv.add("B5"); cfg.add_data_validation(dv)
cfg["A11"]="Use a mesma planilha o ano inteiro: troque o mês aqui e, no fim de cada mês, copie a linha de Dados para o Histórico."; nota(cfg["A11"])
widths(cfg,(28,40,4,14)); cfg.sheet_view.showGridLines=False

# ---------- Dados (entrada semanal) ----------
d=wb.create_sheet("Dados")
titulo(d,"Dados da semana","Toda sexta, copie os totais das outras planilhas do kit para as células amarelas. Limites e metas também são seus. O resto é calculado.",merge_to="G")
hdr(d,4,["Indicador","Valor desta sexta","Limite ou meta","Sentido","Copie de (planilha do kit)","Situação"])
R0=5
# (chave, rótulo, valor do exemplo ou fórmula, limite, sentido, fonte, formato)
LINHAS=[
 ("prazos","Prazos de hoje e dos próximos 7 dias",PZ["hoje"]+PZ["semana"],None,"","01 · Agenda de prazos (Painel: Hoje + Próximos 7 dias)","#,##0"),
 ("atras","Prazos atrasados",PZ["atrasados"],0,"Menor é melhor","01 · Agenda de prazos (Painel)","#,##0"),
 ("horas","Horas registradas no mês",H9[0],None,"","16 · Horas por caso e por pessoa (Painel, Config = mês do painel)","#,##0.0"),
 ("fatur","Horas faturáveis no mês",H9[1],None,"","16 · Horas por caso e por pessoa (Painel)","#,##0.0"),
 ("pfat","% de horas faturáveis","=IF(OR(B7=\"\",B7=0,B8=\"\"),\"\",B8/B7)",0.75,"Maior é melhor","calculado aqui",PCT),
 ("entrou","Entrou no mês (recebimentos)",T9["ent"],None,"","09 · Caixa do escritório (Painel: Entrou no mês)",BRL0),
 ("saiu","Saiu no mês (tudo o que saiu do caixa)",T9["sai_total"],None,"","09 · Caixa do escritório (Painel: Saiu no mês)",BRL0),
 ("sobrou","Sobrou no mês","=IF(OR(B10=\"\",B11=\"\"),\"\",B10-B11)",None,"","calculado aqui",BRL0),
 ("areceb","A receber (carteira; inclui êxito de casos ativos)",E["a_receber"],None,"","13 · Carteira de clientes e casos (Painel: A receber)",BRL0),
 ("venc","Vencido (parcelas em atraso)",E["vencido"],10000,"Menor é melhor","14 · Parcelas e inadimplência (Painel: Vencido)",BRL0),
 ("inad","Inadimplência (vencido ÷ (pago + vencido))",E["inadimplencia"],0.04,"Menor é melhor","14 · Parcelas e inadimplência (Painel: Inadimplência)","0.0%"),
 ("propn","Propostas abertas (quantidade)",E["propostas_n"],None,"","15 · Funil de propostas (Painel: soma do funil por etapa)","#,##0"),
 ("propv","Propostas abertas (valor)",E["propostas_valor"],40000,"Maior é melhor","15 · Funil de propostas (Painel: Em aberto)",BRL0),
 ("casos","Casos ativos",E["casos_ativos"],30,"Maior é melhor","13 · Carteira de clientes e casos (Painel: Casos ativos)","#,##0"),
]
ROW={k:R0+i for i,(k,*_) in enumerate(LINHAS)}
RN=R0+len(LINHAS)-1
for i,(k,rot,val,lim,sent,fonte,fmt) in enumerate(LINHAS):
    r=R0+i
    d.cell(row=r,column=1,value=rot); rotulo(d.cell(row=r,column=1),bold=False); d.cell(row=r,column=1).border=borda
    b=d.cell(row=r,column=2,value=val)
    if isinstance(val,str) and val.startswith("="): calc(b,fmt)
    else: inp(b,fmt,center=True)
    inp(d.cell(row=r,column=3,value=lim),fmt,center=True); inp(d.cell(row=r,column=4,value=sent),center=True)
    d.cell(row=r,column=5,value=fonte); nota(d.cell(row=r,column=5)); d.cell(row=r,column=5).border=borda
    d.cell(row=r,column=6,value=(f'=IF(OR(B{r}="",C{r}=""),"Informativo",IF(D{r}="Menor é melhor",IF(B{r}<=C{r},"No alvo",IF(B{r}<=C{r}*1.1,"Perto","Fora")),'
                                 f'IF(B{r}>=C{r},"No alvo",IF(B{r}>=C{r}*0.9,"Perto","Fora"))))')); calc(d.cell(row=r,column=6))
dvs=lista('"Maior é melhor,Menor é melhor"'); dvs.add(f"D{R0}:D{RN}"); d.add_data_validation(dvs)
for cor,txt,fnt in ((VERDE,"No alvo",VERDE_T),(AMARELO,"Perto","7A5200"),(VERM,"Fora",VERM_T)):
    d.conditional_formatting.add(f"F{R0}:F{RN}", FormulaRule(formula=[f'F{R0}="{txt}"'], fill=fill(cor), font=F(color=fnt,size=10,bold=(txt=="Fora"))))
d.cell(row=RN+2,column=1,value="Linha sem limite fica \"Informativo\". \"Perto\" é até 10% do limite. Inadimplência = vencido ÷ (pago + vencido), a mesma conta da planilha 14. A receber é a carteira contratada (inclui êxito esperado de casos ativos): informativo, não é atraso. Os valores desta sexta valem para o mês escolhido em Config; no fim do mês, copie a coluna B para a linha do mês em Histórico.").font=F(size=9,color=LILAS)
d.merge_cells(start_row=RN+2,start_column=1,end_row=RN+2,end_column=6)
d.cell(row=RN+3,column=1,value="As planilhas indicadas em \"Copie de\" são as do próprio kit. Se ainda não usa alguma, deixe a linha em branco: o painel mostra \"—\". No exemplo (sexta 11/09/2026) horas e caixa são do mês em andamento: metas mensais para eles ficam em branco até o fechamento.").font=F(size=9,color=LILAS)
d.merge_cells(start_row=RN+3,start_column=1,end_row=RN+3,end_column=6)
for r in (RN+2,RN+3): d.cell(row=r,column=1).alignment=Alignment(wrap_text=True,vertical="top"); d.row_dimensions[r].height=30
widths(d,(50,16,14,16,52,12)); d.freeze_panes="B5"; d.sheet_view.showGridLines=False

# ---------- Histórico ----------
h=wb.create_sheet("Histórico")
titulo(h,"Histórico mensal","No fim de cada mês, copie os valores da aba Dados para a linha do mês. Colunas brancas são calculadas; deixe em branco os meses que ainda não aconteceram.",merge_to="O")
HCOLS=[("prazos","Prazos hoje + 7 dias","#,##0"),("atras","Atrasados","#,##0"),("horas","Horas do mês","#,##0.0"),("fatur","Faturáveis","#,##0.0"),("pfat","% faturável",PCT),
       ("entrou","Entrou",BRL0),("saiu","Saiu",BRL0),("sobrou","Sobrou",BRL0),("areceb","A receber",BRL0),("venc","Vencido",BRL0),("inad","Inadimplência","0.0%"),
       ("propn","Propostas abertas","#,##0"),("propv","Valor das propostas",BRL0),("casos","Casos ativos","#,##0")]
HCOL={k:2+i for i,(k,*_) in enumerate(HCOLS)}
hdr(h,4,["Mês"]+[t for _,t,_ in HCOLS],height=32)
EX={"prazos":[None]*8+[PZ["hoje"]+PZ["semana"]],"atras":[None]*8+[PZ["atrasados"]],
    "horas":[HIST[m]["horas"] for m in range(1,10)],"fatur":[HIST[m]["faturaveis"] for m in range(1,10)],
    "entrou":[HIST[m]["entrou"] for m in range(1,10)],"saiu":[HIST[m]["saiu"] for m in range(1,10)],"areceb":[HIST[m]["a_receber"] for m in range(1,10)],
    "venc":[HIST[m]["vencido"] for m in range(1,10)],"inad":[HIST[m]["inadimplencia"] for m in range(1,10)],"propn":[HIST[m]["propostas_n"] for m in range(1,10)],
    "propv":[HIST[m]["propostas_valor"] for m in range(1,10)],"casos":[HIST[m]["casos_ativos"] for m in range(1,10)]}
for m in range(12):
    r=5+m
    h.cell(row=r,column=1,value=MESES[m]); rotulo(h.cell(row=r,column=1),bold=False); h.cell(row=r,column=1).border=borda
    for k,t,fmt in HCOLS:
        col=HCOL[k]; cell=h.cell(row=r,column=col)
        if k=="pfat": cell.value=f'=IF(OR({L(HCOL["horas"])}{r}="",{L(HCOL["horas"])}{r}=0,{L(HCOL["fatur"])}{r}=""),"",{L(HCOL["fatur"])}{r}/{L(HCOL["horas"])}{r})'; calc(cell,fmt)
        elif k=="sobrou": cell.value=f'=IF(OR({L(HCOL["entrou"])}{r}="",{L(HCOL["saiu"])}{r}=""),"",{L(HCOL["entrou"])}{r}-{L(HCOL["saiu"])}{r})'; calc(cell,fmt)
        else:
            inp(cell,fmt,center=True)
            if m<9 and EX[k][m] is not None: cell.value=EX[k][m]
h.conditional_formatting.add("A5:A16", FormulaRule(formula=['ROW()-4=Config!$B$8'], fill=fill(SOL), font=F(color=UVA,size=10,bold=True)))
h.cell(row=18,column=1,value="A linha do mês escolhido em Config fica destacada. Setembro repete os valores da aba Dados (mês em andamento, até 11/09). Entrou e Saiu = Painel mensal da planilha 09 (caixa real). Horas: o controle de horas (16) começou em julho; antes fica em branco. Prazos: a agenda (01) não guarda histórico; o escritório começou a anotar em setembro. Carteira, vencido, inadimplência, propostas e casos: o que o escritório copiou no fim de cada mês.").font=F(size=9,color=LILAS)
h.merge_cells("A18:O18"); h["A18"].alignment=Alignment(wrap_text=True,vertical="top"); h.row_dimensions[18].height=44
# gráficos apontam direto para as colunas de entrada: mês ainda vazio fica sem barra/ponto (sem NA(), sem linha auxiliar)
def serie_cols(k1,k2): return Reference(h,min_col=HCOL[k1],max_col=HCOL[k2],min_row=4,max_row=16)
cats=Reference(h,min_col=1,max_col=1,min_row=5,max_row=16)
bc=BarChart(); bc.type="col"; bc.grouping="clustered"; bc.title="Caixa do mês: entrou, saiu, sobrou"; bc.height=8; bc.width=22; bc.style=2
bc.add_data(serie_cols("entrou","sobrou"),titles_from_data=True); bc.set_categories(cats)  # mês vazio = barra de altura zero (invisível)
for sr,cor in zip(bc.series,(UVA,LILAS,SOL)): sr.graphicalProperties.solidFill=cor; sr.graphicalProperties.line.solidFill=cor
bc.y_axis.majorGridlines=None; bc.legend.position="b"; bc.y_axis.number_format='"R$" #,##0'; bc.dispBlanksAs="gap"
h.add_chart(bc,"A20")
lc=LineChart(); lc.title="Recebíveis: a receber e vencido"; lc.height=8; lc.width=22; lc.style=2
lc.add_data(serie_cols("areceb","venc"),titles_from_data=True); lc.set_categories(cats)
lc.series[0].graphicalProperties.line.solidFill=UVA; lc.series[0].graphicalProperties.line.width=28000
lc.series[1].graphicalProperties.line.solidFill="C0392B"; lc.series[1].graphicalProperties.line.width=28000
for sr in lc.series: sr.smooth=False
lc.y_axis.majorGridlines=None; lc.legend.position="b"; lc.dispBlanksAs="gap"; lc.y_axis.number_format='"R$" #,##0'
h.add_chart(lc,"A38")
widths(h,[12]+[12]*14); h.freeze_panes="B5"; h.sheet_view.showGridLines=False

# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
p["A1"]='=Config!B4&" · Painel do escritório · "&Config!B5&" de "&Config!B6'; p["A1"].font=F(bold=True,size=16,color=UVA); p.merge_cells("A1:J1")
p["A2"]='="Nada para digitar aqui. Os números vêm da aba Dados (atualizada toda sexta) e o mês anterior vem do Histórico. Referência: "&TEXT(Config!B7,"dd/mm/yyyy")&"."'; nota(p["A2"]); p.merge_cells("A2:J2")
def dv_(k): return f"Dados!$B${ROW[k]}"
def val(k): return f'=IF({dv_(k)}="","—",{dv_(k)})'
tiles=[(4,[("prazos","Prazos hoje + 7 dias","#,##0"),("atras","Prazos atrasados","#,##0"),("horas","Horas do mês","#,##0.0"),("fatur","Horas faturáveis","#,##0.0"),("casos","Casos ativos","#,##0")]),
       (7,[("entrou","Entrou no mês",BRL0),("saiu","Saiu no mês",BRL0),("sobrou","Sobrou no mês",BRL0),("areceb","A receber",BRL0),("venc","Vencido",BRL0)]),
       (10,[("propn","Propostas abertas","#,##0"),("propv","Valor das propostas",BRL0),("pfat","% de horas faturáveis",PCT),("inad","Inadimplência","0.0%")])]
for row,items in tiles:
    for j,(k,rot,fmt) in enumerate(items):
        col=1+2*j; kpi(p,row,col,rot,val(k),LAVANDA,UVA,fmt=fmt)
        ref=f"Dados!$F${ROW[k]}"; cell=f"{L(col)}{row+1}"
        p.conditional_formatting.add(cell, FormulaRule(formula=[f'{ref}="Fora"'], fill=fill(VERM), font=F(size=16,bold=True,color=VERM_T)))
        p.conditional_formatting.add(cell, FormulaRule(formula=[f'{ref}="Perto"'], fill=fill(AMARELO), font=F(size=16,bold=True,color="7A5200")))
        p.conditional_formatting.add(cell, FormulaRule(formula=[f'{ref}="No alvo"'], fill=fill(VERDE), font=F(size=16,bold=True,color=VERDE_T)))
p["I10"]="Cores dos cartões"; p["I10"].font=F(size=9,bold=True,color=UVA); p["I11"]="Verde: no alvo · Amarelo: perto do limite · Vermelho: fora · Lilás: informativo"; nota(p["I11"]); p["I11"].alignment=Alignment(wrap_text=True,vertical="top"); p.merge_cells("I11:J11")
p["A13"]="Semáforos da semana e comparação com o mês anterior"; p["A13"].font=F(bold=True,size=13,color=UVA); p.merge_cells("A13:G13")
hdr(p,14,["Indicador","","Esta sexta","Limite ou meta","Situação","Mês anterior (Histórico)","Variação"]); p.merge_cells("A14:B14")
T0=15
for i,(k,rot,val_,lim,sent,fonte,fmt) in enumerate(LINHAS):
    r=T0+i; s=ROW[k]
    p.cell(row=r,column=1,value=f"=Dados!A{s}"); calc(p.cell(row=r,column=1),center=False); p.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
    p.cell(row=r,column=3,value=f'=IF(Dados!B{s}="","—",Dados!B{s})'); calc(p.cell(row=r,column=3),fmt)
    p.cell(row=r,column=4,value=f'=IF(Dados!C{s}="","—",Dados!C{s})'); calc(p.cell(row=r,column=4),fmt)
    p.cell(row=r,column=5,value=f"=Dados!F{s}"); calc(p.cell(row=r,column=5))
    hc=L(HCOL[k])
    p.cell(row=r,column=6,value=f'=IF(Config!$B$8<=1,"—",IF(INDEX(Histórico!${hc}$5:${hc}$16,Config!$B$8-1)="","—",INDEX(Histórico!${hc}$5:${hc}$16,Config!$B$8-1)))'); calc(p.cell(row=r,column=6),fmt)
    if fmt in (PCT,"0.0%"):   # indicador em %: variação em pontos percentuais, como na planilha 18
        p.cell(row=r,column=7,value=f'=IF(OR(C{r}="—",F{r}="—"),"—",(C{r}-F{r})*100)'); calc(p.cell(row=r,column=7),'+0.0" p.p.";-0.0" p.p.";0.0" p.p."')
    else:
        p.cell(row=r,column=7,value=f'=IF(OR(C{r}="—",F{r}="—",F{r}=0),"—",(C{r}-F{r})/ABS(F{r}))'); calc(p.cell(row=r,column=7),"+0.0%;-0.0%;0.0%")
TN=T0+len(LINHAS)-1
for cor,txt,fnt in ((VERDE,"No alvo",VERDE_T),(AMARELO,"Perto","7A5200"),(VERM,"Fora",VERM_T)):
    p.conditional_formatting.add(f"E{T0}:E{TN}", FormulaRule(formula=[f'E{T0}="{txt}"'], fill=fill(cor), font=F(color=fnt,size=10,bold=(txt=="Fora"))))
p.cell(row=TN+2,column=1,value="Fora do alvo primeiro: prazos atrasados e vencido pedem ação na segunda-feira (planilhas 01 e 14). Variação de indicadores em % é em pontos percentuais (p.p.). No meio do mês, horas e caixa ainda estão parciais: compare com o mês anterior só no fechamento. A planilha avisa; o controle é do escritório.").font=F(size=9,color=LILAS)
p.cell(row=TN+2,column=1).alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[TN+2].height=30
p.merge_cells(start_row=TN+2,start_column=1,end_row=TN+2,end_column=10)
p.cell(row=TN+3,column=1,value='Para transformar esta tela em texto para o sócio ou o contador, use a planilha 20 · Resumo do mês e o prompt "Painel 01 · Explicar o mês ao sócio".').font=F(size=9,color=LILAS)
p.merge_cells(start_row=TN+3,start_column=1,end_row=TN+3,end_column=10)
widths(p,(22,22,14,14,14,18,12,12,12,12)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False

como_usar(wb,"Painel do Escritório",[
 ("O que esta planilha faz","Reúne em uma tela os números que já estão nas outras planilhas do kit: prazos da semana e atrasados, horas do mês e faturáveis, o que entrou, saiu e sobrou, a receber e vencido, propostas abertas e casos ativos. Cada número ganha um semáforo contra o limite que você definir e a comparação com o mês anterior."),
 ("Toda sexta (10 minutos)","Abra as planilhas 01 (prazos), 16 (horas), 09 (caixa), 13 (carteira), 14 (inadimplência) e 15 (propostas), copie os totais do Painel de cada uma para a coluna B da aba Dados. Pronto: o Painel se refaz. No exemplo, a sexta é 11/09/2026 e cada número é exatamente o que está no Painel da planilha de origem."),
 ("Limites e metas","Na aba Dados, coluna C, escreva o limite de cada indicador (ex.: 0 prazos atrasados, 75% de horas faturáveis, inadimplência de 4%) e diga se maior ou menor é melhor. Linha sem limite fica informativa (A receber é a carteira contratada, inclui êxito esperado: não é problema)."),
 ("Fim do mês","Copie a coluna B da aba Dados para a linha do mês em Histórico. Os gráficos de caixa e de recebíveis usam essa aba; o Painel busca ali o mês anterior."),
 ("Config","Escolha o mês do painel e mantenha a data de referência em =HOJE(). O título do Painel se ajusta."),
 ("Com a IA","O Painel é a tela da rotina de sexta. Para o texto do mês (sócio, contador), use a planilha 20 · Resumo do mês e o prompt \"Painel 01 · Explicar o mês ao sócio\" da biblioteca."),
])
proteger(wb); salvar(wb,"17-painel-do-escritorio.xlsx","Painel do Escritório · Kit de Gestão para Advogados")
