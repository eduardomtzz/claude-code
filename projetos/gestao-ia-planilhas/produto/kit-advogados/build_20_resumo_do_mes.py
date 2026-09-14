#!/usr/bin/env python3
"""Planilha 20 do Kit de Gestão para Advogados: Resumo do Mês para a IA e para o contador. Gera 20-resumo-do-mes.xlsx
12 indicadores do mês (valor, mês anterior, meta), painel com variação e situação, e a aba Resumo que escreve sozinha
as frases do mês, os destaques e um bloco único para colar no prompt "Explicar o mês".
Exemplo: setembro × agosto de 2026, mesmos números das planilhas 17 e 18."""
from ssg import *
import dados
from openpyxl.formatting.rule import Rule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.styles.numbers import NumberFormat
NI=12; R0=5; RN=R0+NI-1   # indicadores nas linhas 5..16

wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche.",merge_to="E")
cfg["A4"]="Escritório"; cfg["B4"]=f"{dados.ESCRITORIO} (exemplo fictício)"
cfg["A5"]="Ano"; cfg["B5"]=2026
cfg["A6"]="Mês do resumo"; cfg["B6"]="Setembro"
cfg["A7"]="Número do mês"; cfg["B7"]="=MATCH(B6,$D$5:$D$16,0)"
cfg["A8"]="Mês anterior"; cfg["B8"]='=IF(B7>1,INDEX($D$5:$D$16,B7-1),INDEX($D$5:$D$16,12))'
for r in range(4,9): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"]); inp(cfg["B6"]); calc(cfg["B7"]); calc(cfg["B8"])
cfg["D4"]="Meses"; rotulo(cfg["D4"])
for i,m in enumerate(MESES): cfg.cell(row=5+i,column=4,value=m).font=F(size=10,color=TINTA)
dv=lista("=Config!$D$5:$D$16",allow_blank=False); dv.add("B6"); cfg.add_data_validation(dv)
cfg["A10"]="Use a mesma planilha todo mês: troque o mês aqui e, em Indicadores, mova o \"Mês atual\" para \"Mês anterior\" antes de digitar os novos valores."; nota(cfg["A10"])
widths(cfg,(28,40,4,14)); cfg.sheet_view.showGridLines=False

# ---------- Indicadores ----------
ind=wb.create_sheet("Indicadores")
titulo(ind,"Indicadores do mês","Até 12 indicadores (de cima para baixo, sem pular linha): nome, unidade, casas decimais, meta, se maior é melhor, valor do mês anterior e do mês atual. Copie os números do Painel do escritório (17) e do Resultado mensal (18).",merge_to="H")
hdr(ind,4,["Indicador","Unidade","Casas decimais","Meta do mês","Maior é melhor?","Mês anterior","Mês atual","Copie de (planilha do kit)"],height=30)
# exemplo: agosto → setembro de 2026 (mesmos números de 17 e 18)
ex=[("Entrou no mês (recebimentos)","R$",0,30000,"Sim",26660,29400,"09 · Caixa ou 17 · Painel"),
    ("Saiu no mês (custos, pró-labore e impostos provisionados)","R$",0,20500,"Não",20100,20264,"18 · Resultado mensal"),
    ("Sobrou no mês (resultado)","R$",0,9500,"Sim",6560,9136,"18 · Resultado mensal"),
    ("Margem do mês","%",1,30,"Sim",24.6,31.1,"18 · Resultado mensal"),
    ("Horas registradas no mês","h",0,300,"Sim",298,318,"16 · Horas por caso e por pessoa"),
    ("Horas faturáveis","%",1,70,"Sim",77.9,71.1,"16 · Horas por caso e por pessoa"),
    ("Prazos atrasados na última sexta","un",0,0,"Não",5,7,"01 · Agenda de prazos"),
    ("A receber (carteira)","R$",0,120000,"Não",124500,128280,"13 · Carteira de clientes e casos"),
    ("Vencido (parcelas em atraso)","R$",0,15000,"Não",24900,26590,"14 · Parcelas e inadimplência"),
    ("Inadimplência (vencido ÷ parcelas em aberto)","%",1,10,"Não",17.1,17.4,"14 · Parcelas e inadimplência"),
    ("Propostas abertas (valor)","R$",0,40000,"Sim",51300,68400,"15 · Propostas enviadas × fechadas"),
    ("Casos ativos","un",0,32,"Sim",30,30,"13 · Carteira de clientes e casos")]
assert round(9136/29400*100,1)==31.1 and round(6560/26660*100,1)==24.6 and round(226/318*100,1)==71.1 and round(232/298*100,1)==77.9
for i in range(NI):
    r=R0+i
    for c in range(1,8): inp(ind.cell(row=r,column=c),center=(c>1))
    for c in (4,6,7): ind.cell(row=r,column=c).number_format="#,##0.00"
    ind.cell(row=r,column=1).alignment=Alignment(horizontal="left")
    ind.cell(row=r,column=8); inp(ind.cell(row=r,column=8)); ind.cell(row=r,column=8).font=F(size=9,color=LILAS)
    if i<len(ex):
        for c,v in enumerate(ex[i],start=1): ind.cell(row=r,column=c,value=v)
dv_sn=lista('"Sim,Não"'); dv_sn.add(f"E{R0}:E{RN}"); ind.add_data_validation(dv_sn)
dv_cd=lista('"0,1,2"'); dv_cd.add(f"C{R0}:C{RN}"); ind.add_data_validation(dv_cd)
dv_un=lista('"R$,%,un,h,pts"'); dv_un.add(f"B{R0}:B{RN}"); ind.add_data_validation(dv_un)
ind.cell(row=RN+2,column=1,value='Para indicadores em %, digite 16,8 (não 0,168). "Maior é melhor?" = Não para saídas, vencido, inadimplência e prazos atrasados. Meta 0 (ex.: prazos atrasados) funciona: a situação aparece, a distância em % não.').font=F(size=9,color=LILAS)
ind.merge_cells(start_row=RN+2,start_column=1,end_row=RN+2,end_column=8)
widths(ind,(46,9,10,13,12,13,13,34)); ind.freeze_panes="B5"; ind.sheet_view.showGridLines=False

# ---------- Painel ----------
p=wb.create_sheet("Painel")
p["A1"]='=Config!B4&" · Resumo do mês · "&Config!B6&" de "&Config!B5'; p["A1"].font=F(bold=True,size=16,color=UVA); p.merge_cells("A1:H1")
p["A2"]="Nada para digitar aqui. Escolha o mês em Config; os números vêm de Indicadores."; nota(p["A2"]); p.merge_cells("A2:H2")
hdr(p,4,["Indicador","Mês","Mês anterior","Variação","Meta","Vs. meta","Situação","Unidade"])
for i in range(NI):
    r=R0+i; s=R0+i
    p.cell(row=r,column=1,value=f'=IF(Indicadores!A{s}="","",Indicadores!A{s})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF(OR(Indicadores!A{s}="",Indicadores!G{s}=""),"",Indicadores!G{s})'); calc(p.cell(row=r,column=2),"#,##0.00")
    p.cell(row=r,column=3,value=f'=IF(OR(Indicadores!A{s}="",Indicadores!F{s}=""),"",Indicadores!F{s})'); calc(p.cell(row=r,column=3),"#,##0.00")
    p.cell(row=r,column=4,value=f'=IF(OR(B{r}="",C{r}="",C{r}=0),"",(B{r}-C{r})/ABS(C{r}))'); calc(p.cell(row=r,column=4),"+0.0%;-0.0%;0.0%")
    p.cell(row=r,column=5,value=f'=IF(OR(Indicadores!A{s}="",Indicadores!D{s}=""),"",Indicadores!D{s})'); calc(p.cell(row=r,column=5),"#,##0.00")
    p.cell(row=r,column=6,value=f'=IF(OR(B{r}="",E{r}="",E{r}=0),"",(B{r}-E{r})/ABS(E{r}))'); calc(p.cell(row=r,column=6),"+0.0%;-0.0%;0.0%")
    p.cell(row=r,column=7,value=f'=IF(OR(B{r}="",E{r}=""),"",IF(Indicadores!E{s}="Não",IF(B{r}<=E{r},"No alvo","Acima da meta"),IF(B{r}>=E{r},"No alvo","Abaixo da meta")))'); calc(p.cell(row=r,column=7))
    p.cell(row=r,column=8,value=f'=IF(Indicadores!A{s}="","",Indicadores!B{s})'); calc(p.cell(row=r,column=8))
    # N: casas decimais (formato condicional). O: distância da meta só se fora do alvo, sempre positiva
    # (sinal invertido quando "Maior é melhor?" = Não), para "mais longe da meta" respeitar o sentido.
    p.cell(row=r,column=14,value=f'=IF(Indicadores!C{s}="",2,Indicadores!C{s})'); p.cell(row=r,column=14).font=F(size=9,color=CINZA)
    p.cell(row=r,column=15,value=f'=IF(OR(F{r}="",G{r}="",G{r}="No alvo"),"",IF(Indicadores!E{s}="Não",F{r},-F{r}))'); p.cell(row=r,column=15).font=F(size=9,color=CINZA)
    # P: variação com sinal de melhora (positivo = melhorou), para "maior melhora" e "maior piora" respeitarem o sentido
    p.cell(row=r,column=16,value=f'=IF(D{r}="","",IF(Indicadores!E{s}="Não",-D{r},D{r}))'); p.cell(row=r,column=16).font=F(size=9,color=CINZA)
rng=f"A{R0}:H{RN}"
p.conditional_formatting.add(rng, FormulaRule(formula=[f'OR($G{R0}="Abaixo da meta",$G{R0}="Acima da meta")'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
p.conditional_formatting.add(rng, FormulaRule(formula=[f'$G{R0}="No alvo"'], fill=fill(VERDE)))
for cols_,fmt,fid,casas in ((f"B{R0}:C{RN} E{R0}:E{RN}","#,##0",3,0),(f"B{R0}:C{RN} E{R0}:E{RN}","#,##0.0",201,1)):
    p.conditional_formatting.add(cols_, Rule(type="expression",formula=[f'$N{R0}={casas}'],dxf=DifferentialStyle(numFmt=NumberFormat(numFmtId=fid,formatCode=fmt))))
p.cell(row=RN+2,column=1,value="Verde: no alvo. Vermelho: fora da meta. As colunas N, O e P (ocultas) guardam casas decimais, a distância da meta e a variação com sinal de melhora, usadas nos destaques.").font=F(size=9,color=LILAS)
p.merge_cells(start_row=RN+2,start_column=1,end_row=RN+2,end_column=8)
p.column_dimensions["N"].hidden=True; p.column_dimensions["O"].hidden=True; p.column_dimensions["P"].hidden=True
widths(p,(46,12,12,10,12,10,15,9)); p.freeze_panes="A5"; p.sheet_view.showGridLines=False

# ---------- Resumo ----------
rs=wb.create_sheet("Resumo",0)
rs["A1"]="Resumo do mês, pronto para a IA e para o contador"; rs["A1"].font=F(bold=True,size=16,color=UVA)
rs["A2"]="Só a célula amarela é para digitar (observações do escritório). Copie o bloco único e cole no prompt \"Explicar o mês\" da biblioteca do kit, ou envie ao sócio e ao contador."; nota(rs["A2"]); rs.merge_cells("A2:D2")
rs["A4"]='="Resumo de "&Config!B6&" de "&Config!B5&" · "&Config!B4'; rs["A4"].font=F(bold=True,color=UVA); rs.merge_cells("A4:D4")
rs["A5"]="Observações do escritório (opcional; entram no fim do bloco):"; rotulo(rs["A5"],bold=False)
rs["A6"]="Dois casos de êxito encerraram com acordo em setembro; o consultivo mensal ganhou um contrato; dois clientes PJ pediram renegociação das parcelas vencidas."
inp(rs["A6"]); [inp(rs[c+"6"]) for c in "BCD"]; rs.merge_cells("A6:D6"); rs["A6"].alignment=Alignment(wrap_text=True,vertical="top"); rs.row_dimensions[6].height=44
F0=8
def frase(i):
    r=R0+i; s=R0+i
    v=f"Painel!B{r}"; a=f"Painel!C{r}"; var=f"Painel!D{r}"; meta=f"Painel!E{r}"; st=f"Painel!G{r}"; un=f"Indicadores!B{s}"; cd=f"Indicadores!C{s}"; nome=f"Indicadores!A{s}"
    num=lambda x: f'IF({un}="R$","R$ ","")&FIXED({x},IF({cd}="",0,{cd}))&IF({un}="%","%",IF(OR({un}="R$",{un}="",{un}="un"),""," "&{un}))'
    return (f'=IF(OR({nome}="",{v}=""),"",'
            f'"• "&{nome}&": "&{num(v)}'
            f'&IF({var}<>"",IF({var}=0," (igual a "&Config!$B$8&")",IF({var}>0," (+"," (")&FIXED({var}*100,0)&"% em relação a "&Config!$B$8&")"),IF({a}<>""," (igual a "&Config!$B$8&")",""))'
            f'&IF({meta}<>"","; meta "&{num(meta)}&", "&LOWER({st}),"")&".")')
for i in range(NI):
    rs.cell(row=F0+i,column=1,value=frase(i)).font=F(size=10,color=TINTA); rs.merge_cells(start_row=F0+i,start_column=1,end_row=F0+i,end_column=4)
    rs.cell(row=F0+i,column=1).alignment=Alignment(wrap_text=True,vertical="top")
D0=F0+NI+1
rs.cell(row=D0,column=1,value="Destaques automáticos").font=F(bold=True,color=UVA)
VR=f"Painel!$D${R0}:$D${RN}"; NR=f"Painel!$A${R0}:$A${RN}"; MR=f"Painel!$F${R0}:$F${RN}"; AR=f"Painel!$O${R0}:$O${RN}"; SR=f"Painel!$G${R0}:$G${RN}"
PR=f"Painel!$P${R0}:$P${RN}"
def pick(agg): return f'INDEX({NR},MATCH({agg}({PR}),{PR},0))'
def varde(agg): return f'INDEX({VR},MATCH({agg}({PR}),{PR},0))'
rs.cell(row=D0+1,column=1,value=f'=IF(COUNT({PR})=0,"• Melhora e piora: preencha o mês anterior para comparar.",IF(MAX({PR})<=0,"• Nenhum indicador melhorou contra "&Config!$B$8&".","• Maior melhora contra "&Config!$B$8&": "&{pick("MAX")}&" ("&IF({varde("MAX")}>=0,"+","")&FIXED({varde("MAX")}*100,0)&"%)."))')
rs.cell(row=D0+2,column=1,value=f'=IF(COUNT({PR})=0,"",IF(MIN({PR})>=0,"• Nenhum indicador piorou contra "&Config!$B$8&".","• Maior piora contra "&Config!$B$8&": "&{pick("MIN")}&" ("&IF({varde("MIN")}>=0,"+","")&FIXED({varde("MIN")}*100,0)&"%)."))')
rs.cell(row=D0+3,column=1,value=f'=IF(COUNT({AR})=0,"• Nenhum indicador fora da meta.","• Mais longe da meta: "&INDEX({NR},MATCH(MAX({AR}),{AR},0))&" ("&IF(INDEX({MR},MATCH(MAX({AR}),{AR},0))>=0,"+","")&FIXED(INDEX({MR},MATCH(MAX({AR}),{AR},0))*100,0)&"% da meta, "&LOWER(INDEX({SR},MATCH(MAX({AR}),{AR},0)))&").")')
rs.cell(row=D0+4,column=1,value=f'="• Indicadores no alvo: "&COUNTIF({SR},"No alvo")&" de "&COUNTIF({SR},"<>")&"."')
for k in range(1,5):
    rs.cell(row=D0+k,column=1).font=F(size=10,color=TINTA); rs.merge_cells(start_row=D0+k,start_column=1,end_row=D0+k,end_column=4); rs.cell(row=D0+k,column=1).alignment=Alignment(wrap_text=True,vertical="top")
T0=D0+6
rs.cell(row=T0,column=1,value="Bloco único para copiar").font=F(bold=True,color=UVA)
partes=" & ".join([f'IF(A{F0+i}="","",A{F0+i}&CHAR(10))' for i in range(NI)])
destaques=" & ".join([f'IF(A{D0+k}="","",A{D0+k}&CHAR(10))' for k in range(1,5)])
rs.cell(row=T0+1,column=1,value=f'=A4&CHAR(10)&CHAR(10)&"Indicadores do mês:"&CHAR(10)&{partes}&CHAR(10)&"Destaques:"&CHAR(10)&{destaques}&IF(A6="","",CHAR(10)&"Observações do escritório: "&A6)')
rs.merge_cells(start_row=T0+1,start_column=1,end_row=T0+1,end_column=4); rs.cell(row=T0+1,column=1).alignment=Alignment(wrap_text=True,vertical="top"); rs.cell(row=T0+1,column=1).font=F(size=10,color=TINTA); rs.cell(row=T0+1,column=1).fill=fill(LAVANDA)
rs.row_dimensions[T0+1].height=330
rs.cell(row=T0+3,column=1,value="Os textos são gerados por fórmula. Revise antes de enviar: a IA e o contador precisam do porquê dos números, e isso só o escritório sabe. Nenhum dado de cliente entra aqui: só totais.").font=F(size=9,color=LILAS)
rs.merge_cells(start_row=T0+3,start_column=1,end_row=T0+3,end_column=4)
widths(rs,(44,30,30,30)); rs.sheet_view.showGridLines=False
wb.move_sheet("Painel",offset=-(len(wb.sheetnames)-2))  # ordem: Resumo, Painel, Config, Indicadores

como_usar(wb,"Resumo do Mês",[
 ("O que esta planilha faz","Você digita até 12 indicadores do mês (valor, mês anterior, meta). Ela calcula variação e situação, monta o Painel e escreve sozinha as frases do mês, os destaques (maior melhora, maior piora, mais longe da meta, respeitando se maior ou menor é melhor) e um bloco único para colar na IA ou mandar ao sócio e ao contador."),
 ("Passo 1","Em Config, preencha o escritório, o ano e escolha o mês do resumo."),
 ("Passo 2","Em Indicadores, troque os exemplos pelos seus (de cima para baixo, sem pular linha): nome, unidade, casas decimais, meta e se maior é melhor. Depois, o valor do mês anterior e do mês atual, copiados do Painel do escritório (17) e do Resultado mensal (18). Em %, digite 16,8 e não 0,168."),
 ("Passo 3","Abra Painel: cada indicador com mês, mês anterior, variação, meta e situação (verde no alvo, vermelho fora)."),
 ("Passo 4","Abra Resumo: as frases já estão escritas. Escreva as observações do mês na célula amarela, copie o bloco único e cole no prompt \"Explicar o mês\" (para o sócio) ou \"Preparar a reunião com o contador\" da biblioteca do kit."),
 ("Todo mês","Em Indicadores, mova o \"Mês atual\" para \"Mês anterior\", digite os novos valores e troque o mês em Config. O resto se refaz sozinho."),
 ("Cuidado com dados","O bloco só tem totais do escritório. Não cole nome de cliente, número de processo ou detalhe de caso na IA: veja o guia LGPD do kit."),
])
proteger(wb); salvar(wb,"20-resumo-do-mes.xlsx","Resumo do Mês · Kit de Gestão para Advogados")
