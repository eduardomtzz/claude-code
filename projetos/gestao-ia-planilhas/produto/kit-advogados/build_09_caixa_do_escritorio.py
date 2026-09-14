#!/usr/bin/env python3
"""Planilha 9 do Kit de Gestão para Advogados: Caixa do escritório. Gera 09-caixa-do-escritorio.xlsx
As funções de exemplo (lancamentos_exemplo, totais_mensais) são reaproveitadas por build_10, 11 e 12
para que as quatro planilhas do Núcleo 3 contem a mesma história."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference
from datetime import date, timedelta
import random

N=600; R0=5; RN=R0+N-1; NCE=12; NCS=16; NCLI=30
CAT_ENTRADA=["Honorários fixos","Honorários por hora","Honorários de êxito","Consultoria e pareceres","Reembolso de custas","Outras entradas"]
CAT_SAIDA=["Pró-labore dos sócios"]+[c for c,_ in dados.CUSTOS_FIXOS]+["Impostos e taxas","Custas e despesas de processo","Deslocamento e viagens","Outras saídas"]
CAT_HON={"Fixo":"Honorários fixos","Hora":"Honorários por hora","Êxito":"Honorários de êxito","Misto":"Honorários fixos"}
PROLABORE=[(n,v) for n,p,v,_ in dados.PESSOAS if p.startswith("Sóci")]
DISTRIB_EXEMPLO=0.5  # parte do lucro do trimestre distribuída aos sócios no exemplo (regra fictícia dos sócios)
ALIQ_EXEMPLO=0.08  # alíquota efetiva usada só no exemplo (escolha fictícia do escritório com o contador)
DIA_FIXO={"Aluguel e condomínio":5,"Contador":10,"Sistemas e assinaturas":8,"Telefone e internet":12,"Anuidades OAB e cursos":15,
          "Marketing e site":20,"Estagiária (bolsa)":5,"Material, correio e outros":18}

def lancamentos_exemplo():
    """Lista de tuplas (data, tipo, categoria, cliente, caso, descrição, valor, forma, pago) de jan a set/2026."""
    rng=random.Random(9); ex=[]
    def add(d,t,c,cli,caso,desc,v,f="Pix",p="Sim"): ex.append((d,t,c,cli,caso,desc,v,f,p))
    ini=date(2026,1,1); lim=dados.HOJE-timedelta(days=5)
    for k in dados.CASOS:
        rec=k["recebido"]
        if rec<=0: continue
        n=3 if rec>=8000 else (2 if rec>=3500 else 1)
        primeiro=k["abertura"]+timedelta(days=rng.randint(7,20))
        passo=rng.randint(75,130) if k["abertura"].year<2026 else rng.randint(30,70)
        partes=[round(rec/n/100)*100]*n; partes[-1]=rec-sum(partes[:-1])
        for i in range(n):
            d=primeiro+timedelta(days=i*passo)
            if d>lim: d=lim-timedelta(days=rng.randint(2,50))
            if d<ini: continue  # recebido em 2025 já está no saldo inicial
            add(d,"Entrada",CAT_HON[k["tipo_hon"]],k["cliente"],k["numero"],f"Honorários · parcela {i+1}/{n}",partes[i],rng.choice(["Pix","Pix","Transferência","Boleto"]))
    # honorários de casos encerrados antes de 2026 ainda parcelados e pareceres avulsos (fora da carteira de 38 casos)
    for m in range(1,7): add(date(2026,m,10),"Entrada","Honorários fixos","Loja Verde Comércio","Encerrado em 2025 · parcelamento",f"Honorários · parcela {m+6}/12",2000,"Boleto")
    for m in range(1,10): add(date(2026,m,15),"Entrada","Honorários fixos","Oficina Mecânica Central","Encerrado em 2025 · parcelamento",f"Honorários · parcela {m+3}/12",1500,"Pix")
    pareceres=[(1,"Padaria do Sol Ltda",1800),(2,"Clínica Bem-Estar",3200),(3,"Agência Prisma",2400),(4,"Escola Aurora",1600),(5,"Transportadora Rota Sul",2900),
               (6,"Loja Verde Comércio",2200),(7,"Construtora Horizonte",3600),(8,"Bistrô 42",1900),(9,"Padaria do Sol Ltda",2600)]
    for m,cli,v in pareceres:
        if m<9 or v<=2600: add(date(2026,m,min(22,4+m*2)),"Entrada","Consultoria e pareceres",cli,"Consultivo avulso","Parecer e reunião de orientação",v,"Pix")
    add(date(2026,3,9),"Entrada","Reembolso de custas","Construtora Horizonte",dados.CASOS[2]["numero"],"Reembolso de custas periciais",780,"Transferência")
    add(date(2026,8,12),"Entrada","Reembolso de custas","Roberto Almeida",dados.CASOS[25]["numero"],"Reembolso de custas de distribuição",310,"Pix")
    add(date(2026,9,25),"Entrada","Honorários fixos","Construtora Horizonte",dados.CASOS[2]["numero"],"Honorários · parcela final (a receber)",4500,"Boleto","Não")
    add(date(2026,9,30),"Entrada","Consultoria e pareceres","Clínica Bem-Estar",dados.CASOS[8]["numero"],"Parecer societário (a receber)",3000,"Pix","Não")
    custas=[(1,"Ana Beatriz Moreira",1,"Custas iniciais",260),(2,"Transportadora Rota Sul",10,"Custas de distribuição",420),(3,"Construtora Horizonte",2,"Honorários periciais adiantados",780),
            (4,"Agência Prisma",16,"Custas de recurso",540),(5,"Ana Beatriz Moreira",1,"Diligência de oficial de justiça",130),(6,"Roberto Almeida",25,"Custas de distribuição",310),
            (7,"Luciana Farias",15,"Certidões e cópias",95),(8,"Bistrô 42",24,"Custas de execução",380),(9,"Escola Aurora",12,"Certidões",70)]
    for m,cli,idx,desc,v in custas: add(date(2026,m,rng.randint(3,26)),"Saída","Custas e despesas de processo",cli,dados.CASOS[idx]["numero"],desc,v,"Boleto")
    for m in range(1,10):
        for c,v in dados.CUSTOS_FIXOS:
            d=date(2026,m,DIA_FIXO[c]); pago="Não" if (m==9 and c=="Contador") else "Sim"
            add(d,"Saída",c,"","",{"Contador":"Honorários do contador","Estagiária (bolsa)":"Bolsa da estagiária · Júlia Prado","Aluguel e condomínio":"Aluguel e condomínio da sala"}.get(c,c),v,"Boleto" if c in("Aluguel e condomínio","Contador","Telefone e internet") else "Cartão",pago)
        for n_,v in PROLABORE: add(date(2026,m,28),"Saída","Pró-labore dos sócios","","",f"Pró-labore · {n_}",v,"Transferência","Não" if m==9 else "Sim")
        if m in(2,5,8): add(date(2026,m,rng.randint(6,22)),"Saída","Deslocamento e viagens","","","Combustível e estacionamento (audiências)",rng.choice([180,240,310]),"Cartão")
    # imposto do exemplo: pago dia 20 sobre as entradas recebidas no mês anterior (alíquota efetiva fictícia de 8 %)
    ent_mes={m:sum(v for d,t,c,cli,caso,desc,v,f,p in ex if t=="Entrada" and p=="Sim" and d.month==m) for m in range(1,10)}
    for m in range(1,10):
        base=ent_mes.get(m-1,20600)  # janeiro: sobre dezembro/2025 (fictício)
        add(date(2026,m,20),"Saída","Impostos e taxas","","","Guia de impostos do mês anterior (alíquota efetiva combinada com o contador)",round(base*ALIQ_EXEMPLO),"Boleto","Não" if m==9 else "Sim")
    # movimentos sócio × escritório (detalhados na planilha 11): despesa pessoal paga pela conta do escritório, devolução, retirada extra
    for m in (2,5,8): add(date(2026,m,6),"Saída","Outras saídas","","","Despesa pessoal · Marina Ferraz · plano de saúde particular (a acertar)",480,"Boleto")
    for m in (3,7): add(date(2026,m,11),"Saída","Outras saídas","","","Despesa pessoal · Rafael Lima · combustível particular (a acertar)",300,"Cartão")
    add(date(2026,4,14),"Entrada","Outras entradas","","","Devolução de despesa pessoal · Marina Ferraz",480,"Pix")
    add(date(2026,3,16),"Saída","Pró-labore dos sócios","","","Retirada extra · Rafael Lima · adiantamento para viagem",1500,"Transferência")
    add(date(2026,6,19),"Saída","Pró-labore dos sócios","","","Retirada extra · Marina Ferraz · reforma em casa",2000,"Transferência")
    # distribuição de lucro: 50 % do resultado do trimestre fechado (após pró-labore fixo), dividido meio a meio, paga no dia 10 do mês seguinte
    for q,(m_ini,m_pag) in enumerate(((1,4),(4,7)),start=1):
        lucro=sum(v if t=="Entrada" else -v for d,t,c,cli,caso,desc,v,f,p in ex if p=="Sim" and m_ini<=d.month<=m_ini+2 and not desc.startswith(("Retirada extra","Despesa pessoal","Devolução","Distribuição")))
        cota=round(lucro*DISTRIB_EXEMPLO/len(PROLABORE))
        for n_,_ in PROLABORE: add(date(2026,m_pag,10),"Saída","Pró-labore dos sócios","","",f"Distribuição de lucro · {q}º trimestre · {n_}",cota,"Transferência")
    ex.sort(key=lambda x:(x[0],x[1]))
    return ex

def totais_mensais(ex=None):
    """Dict mês -> dict(ent, devol, sai, pessoal, pro, extra, distrib, fixo, imp) só com o que está Pago? = Sim, jan a dez/2026.
    ent: entradas (inclui devol); sai: saídas sem pró-labore/retiradas (inclui pessoal); pro: pró-labore fixo; extra: retiradas extras;
    distrib: distribuição de lucro; fixo: custos fixos + pró-labore fixo; imp: impostos."""
    ex=ex or lancamentos_exemplo(); out={}
    fixos={c for c,_ in dados.CUSTOS_FIXOS}
    for m in range(1,13):
        L_=[x for x in ex if x[0].month==m and x[8]=="Sim"]
        pro=sum(x[6] for x in L_ if x[5].startswith("Pró-labore"))
        out[m]=dict(ent=sum(x[6] for x in L_ if x[1]=="Entrada"),devol=sum(x[6] for x in L_ if x[5].startswith("Devolução")),
                    sai=sum(x[6] for x in L_ if x[1]=="Saída" and x[2]!="Pró-labore dos sócios"),pessoal=sum(x[6] for x in L_ if x[5].startswith("Despesa pessoal")),
                    pro=pro,extra=sum(x[6] for x in L_ if x[5].startswith("Retirada extra")),distrib=sum(x[6] for x in L_ if x[5].startswith("Distribuição")),
                    fixo=sum(x[6] for x in L_ if x[2] in fixos)+pro,imp=sum(x[6] for x in L_ if x[2]=="Impostos e taxas"))
    return out

def build():
    wb=Workbook()
    # ---------- Config ----------
    cfg=wb.active; cfg.title="Config"
    titulo(cfg,"Configurações","Células amarelas: você preenche. Categorias e clientes alimentam as listas de Lançamentos e o Painel.",merge_to="L")
    cfg["A4"]="Escritório"; cfg["B4"]=f"{dados.ESCRITORIO} (exemplo fictício)"
    cfg["A5"]="Ano do painel"; cfg["B5"]=2026
    cfg["A6"]="Mês do painel"; cfg["B6"]="Setembro"
    cfg["A7"]="Número do mês"; cfg["B7"]="=MATCH(B6,$L$5:$L$16,0)"
    cfg["A8"]="Saldo em caixa antes do 1º lançamento"; cfg["B8"]=22000
    cfg["A9"]="Data de referência"; cfg["B9"]="=TODAY()"
    for r in range(4,10): rotulo(cfg.cell(row=r,column=1))
    inp(cfg["B4"]); inp(cfg["B5"],center=True); inp(cfg["B6"],center=True); calc(cfg["B7"]); inp(cfg["B8"],BRL); calc(cfg["B9"],DATA)
    cfg["D4"]="Categorias de entrada (até 12)"; cfg["F4"]="Categorias de saída (até 16)"; cfg["H4"]="Todas as categorias (automático)"; cfg["J4"]="Clientes (até 30)"; cfg["L4"]="Meses"
    for c in ("D4","F4","H4","J4","L4"): rotulo(cfg[c])
    for i in range(NCE): inp(cfg.cell(row=5+i,column=4))
    for i in range(NCS): inp(cfg.cell(row=5+i,column=6))
    for i in range(NCLI): inp(cfg.cell(row=5+i,column=10))
    for i,v in enumerate(CAT_ENTRADA): cfg.cell(row=5+i,column=4,value=v)
    for i,v in enumerate(CAT_SAIDA): cfg.cell(row=5+i,column=6,value=v)
    for i,(n_,_,_) in enumerate(dados.CLIENTES): cfg.cell(row=5+i,column=10,value=n_)
    for i,m_ in enumerate(MESES): cfg.cell(row=5+i,column=12,value=m_).font=F(size=10,color=TINTA)
    NE=f"COUNTA($D$5:$D${4+NCE})"; NS=f"COUNTA($F$5:$F${4+NCS})"
    for i in range(NCE+NCS):
        k=i+1
        cfg.cell(row=5+i,column=8,value=f'=IF({k}<={NE},INDEX($D$5:$D${4+NCE},{k}),IF({k}<={NE}+{NS},INDEX($F$5:$F${4+NCS},{k}-{NE}),""))').font=F(size=9,color=CINZA)
    notas=["Preencha categorias e clientes de cima para baixo, sem pular linha: as listas suspensas de Lançamentos param na última linha preenchida.",
           "Saldo antes do 1º lançamento: o que havia na conta do escritório na data do primeiro lançamento (no exemplo, 01/01/2026, já contando honorários recebidos em 2025).",
           "Pró-labore dos sócios é uma categoria de saída como qualquer outra: o que sai do caixa do escritório para a pessoa física. A separação detalhada fica na planilha 11."]
    for i,t in enumerate(notas): cfg.cell(row=37+i,column=1,value=t); nota(cfg.cell(row=37+i,column=1))
    dv=lista("=Config!$L$5:$L$16"); dv.add("B6"); cfg.add_data_validation(dv)
    widths(cfg,(36,22,3,26,3,30,3,30,3,28,3,12)); cfg.sheet_view.showGridLines=False
    # ---------- Lançamentos ----------
    lan=wb.create_sheet("Lançamentos")
    titulo(lan,"Lançamentos","Uma linha por entrada ou saída. Preencha o amarelo; mês e ano são calculados. Pago? = Não fica em A receber / A pagar e só entra no caixa quando virar Sim.",merge_to="K")
    hdr(lan,4,["Data","Tipo","Categoria","Cliente","Caso (nº ou referência)","Descrição","Valor","Forma","Pago?","Mês","Ano"])
    for r in range(R0,RN+1):
        for c in range(1,10): inp(lan.cell(row=r,column=c))
        lan.cell(row=r,column=1).number_format=DATA; lan.cell(row=r,column=7).number_format=BRL
        for c in (1,2,8,9): lan.cell(row=r,column=c).alignment=Alignment(horizontal="center")
        lan.cell(row=r,column=10,value=f'=IF(A{r}="","",MONTH(A{r}))'); calc(lan.cell(row=r,column=10))
        lan.cell(row=r,column=11,value=f'=IF(A{r}="","",YEAR(A{r}))'); calc(lan.cell(row=r,column=11))
    dvs=[(lista('"Entrada,Saída"'),f"B{R0}:B{RN}"),
         (lista(f"=OFFSET(Config!$H$5,0,0,MAX(1,COUNTA(Config!$D$5:$D${4+NCE})+COUNTA(Config!$F$5:$F${4+NCS})),1)",strict=False),f"C{R0}:C{RN}"),
         (lista(f"=OFFSET(Config!$J$5,0,0,MAX(1,COUNTA(Config!$J$5:$J${4+NCLI})),1)",strict=False),f"D{R0}:D{RN}"),
         (lista('"Pix,Transferência,Boleto,Cartão,Dinheiro"'),f"H{R0}:H{RN}"),
         (lista('"Sim,Não"'),f"I{R0}:I{RN}"),
         (DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True),f"A{R0}:A{RN}"),
         (DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True),f"G{R0}:G{RN}")]
    for dv,rng_ in dvs: dv.add(rng_); lan.add_data_validation(dv)
    lan.conditional_formatting.add(f"A{R0}:K{RN}", FormulaRule(formula=[f'$B{R0}="Entrada"'], font=F(color=VERDE_T,size=10)))
    lan.conditional_formatting.add(f"A{R0}:K{RN}", FormulaRule(formula=[f'AND($A{R0}<>"",$I{R0}="Não")'], fill=fill(VERM)))
    widths(lan,(12,10,28,26,26,40,14,13,8,6,7)); lan.freeze_panes="A5"; lan.sheet_view.showGridLines=False; lan.auto_filter.ref=f"A4:K{RN}"
    for i,row in enumerate(lancamentos_exemplo()):
        for c,v in enumerate(row,start=1): lan.cell(row=R0+i,column=c,value=v)
    # ---------- Painel ----------
    p=wb.create_sheet("Painel",0)
    titulo(p,'=Config!B4&" · Caixa do escritório · "&Config!B6&" de "&Config!B5',"Nada para digitar aqui. Escolha o mês em Config; tudo vem de Lançamentos. Só o que está com Pago? = Sim conta como caixa.",merge_to="L")
    M="Config!$B$7"; Y="Config!$B$5"
    LA=f"Lançamentos!$A${R0}:$A${RN}"; LB=f"Lançamentos!$B${R0}:$B${RN}"; LC=f"Lançamentos!$C${R0}:$C${RN}"; LD=f"Lançamentos!$D${R0}:$D${RN}"
    LG=f"Lançamentos!$G${R0}:$G${RN}"; LI=f"Lançamentos!$I${R0}:$I${RN}"; LJ=f"Lançamentos!$J${R0}:$J${RN}"; LK=f"Lançamentos!$K${R0}:$K${RN}"
    def somames(tipo,m=M,y=Y,extra=""): return f'SUMIFS({LG},{LB},"{tipo}",{LI},"Sim",{LJ},{m},{LK},{y}{extra})'
    def saldo_ate(fim): return f'Config!$B$8+SUMIFS({LG},{LB},"Entrada",{LI},"Sim",{LA},"<="&{fim})-SUMIFS({LG},{LB},"Saída",{LI},"Sim",{LA},"<="&{fim})'
    kpi(p,4,1,"Entrou no mês",f"={somames('Entrada')}",VERDE,VERDE_T,fmt=BRL0)
    kpi(p,4,3,"Saiu no mês",f"={somames('Saída')}",VERM,VERM_T,fmt=BRL0)
    kpi(p,4,5,"Sobrou no mês","=A5-C5",SOL,UVA,fmt=BRL0)
    kpi(p,4,7,"Saldo acumulado",f"={saldo_ate(f'DATE({Y},{M}+1,0)')}",LAVANDA,UVA,fmt=BRL0)
    kpi(p,4,9,"A receber (Pago? = Não)",f'=SUMIFS({LG},{LB},"Entrada",{LI},"Não")',LAVANDA,UVA,fmt=BRL0)
    kpi(p,4,11,"A pagar (Pago? = Não)",f'=SUMIFS({LG},{LB},"Saída",{LI},"Não")',LAVANDA,UVA,fmt=BRL0)
    p.conditional_formatting.add("E5", FormulaRule(formula=['E5<0'], font=F(color="C8402E",size=16,bold=True)))
    # por categoria de saída
    p["A7"]="Para onde foi o dinheiro"; p["A7"].font=F(bold=True,size=13,color=UVA)
    hdr(p,8,["Categoria de saída","Valor no mês","% do total","Mês anterior","Acumulado no ano","Barra"]); p.merge_cells(start_row=8,start_column=6,end_row=8,end_column=9)
    def bloco_cat(r0_,tipo,col_cfg,n,kpi_ref):
        for i in range(n):
            r=r0_+i; src=f"Config!${col_cfg}${5+i}"
            p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
            p.cell(row=r,column=2,value=f'=IF({src}="","",{somames(tipo,extra=f",{LC},{src}")})'); calc(p.cell(row=r,column=2),BRL)
            p.cell(row=r,column=3,value=f'=IF(OR({src}="",{kpi_ref}=0),"",B{r}/{kpi_ref})'); calc(p.cell(row=r,column=3),PCT)
            p.cell(row=r,column=4,value=f'=IF({src}="","",IF({M}=1,{somames(tipo,"12",f"{Y}-1",f",{LC},{src}")},{somames(tipo,f"{M}-1",Y,f",{LC},{src}")}))'); calc(p.cell(row=r,column=4),BRL)
            p.cell(row=r,column=5,value=f'=IF({src}="","",SUMIFS({LG},{LB},"{tipo}",{LI},"Sim",{LC},{src},{LK},{Y},{LJ},"<="&{M}))'); calc(p.cell(row=r,column=5),BRL)
            p.cell(row=r,column=6,value=f'=IF(C{r}="","",REPT("█",ROUND(C{r}*40,0)))'); p.cell(row=r,column=6).font=F(size=10,color=LILAS); p.cell(row=r,column=6).border=borda
            p.merge_cells(start_row=r,start_column=6,end_row=r,end_column=9)
        t=r0_+n
        p.cell(row=t,column=1,value="Total"); p.cell(row=t,column=2,value=f"=SUM(B{r0_}:B{t-1})"); p.cell(row=t,column=4,value=f"=SUM(D{r0_}:D{t-1})"); p.cell(row=t,column=5,value=f"=SUM(E{r0_}:E{t-1})")
        for c in (1,2,4,5): p.cell(row=t,column=c).font=F(bold=True,color=UVA,size=10); p.cell(row=t,column=c).border=borda
        for c in (2,4,5): p.cell(row=t,column=c).number_format=BRL
        return t
    t1=bloco_cat(9,"Saída","F",NCS,"$C$5")          # 9..24, total 25
    r_e=t1+2
    p.cell(row=r_e,column=1,value="De onde veio o dinheiro").font=F(bold=True,size=13,color=UVA)
    hdr(p,r_e+1,["Categoria de entrada","Valor no mês","% do total","Mês anterior","Acumulado no ano","Barra"]); p.merge_cells(start_row=r_e+1,start_column=6,end_row=r_e+1,end_column=9)
    t2=bloco_cat(r_e+2,"Entrada","D",NCE,"$A$5")
    # por cliente
    r_c=t2+2
    p.cell(row=r_c,column=1,value="Por cliente").font=F(bold=True,size=13,color=UVA)
    hdr(p,r_c+1,["Cliente","Entrou no mês","Entrou no ano","% do ano","A receber (Pago? = Não)","Custas pagas no ano"])
    for i in range(NCLI):
        r=r_c+2+i; src=f"Config!$J${5+i}"
        p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
        p.cell(row=r,column=2,value=f'=IF({src}="","",{somames("Entrada",extra=f",{LD},{src}")})'); calc(p.cell(row=r,column=2),BRL)
        p.cell(row=r,column=3,value=f'=IF({src}="","",SUMIFS({LG},{LB},"Entrada",{LI},"Sim",{LD},{src},{LK},{Y}))'); calc(p.cell(row=r,column=3),BRL)
        p.cell(row=r,column=4,value=f'=IF(OR({src}="",$C${r_c+2+NCLI}=0),"",C{r}/$C${r_c+2+NCLI})'); calc(p.cell(row=r,column=4),PCT)
        p.cell(row=r,column=5,value=f'=IF({src}="","",SUMIFS({LG},{LB},"Entrada",{LI},"Não",{LD},{src}))'); calc(p.cell(row=r,column=5),BRL)
        p.cell(row=r,column=6,value=f'=IF({src}="","",SUMIFS({LG},{LB},"Saída",{LI},"Sim",{LD},{src},{LK},{Y}))'); calc(p.cell(row=r,column=6),BRL)
    t3=r_c+2+NCLI
    p.cell(row=t3,column=1,value="Total"); p.cell(row=t3,column=2,value=f"=SUM(B{r_c+2}:B{t3-1})"); p.cell(row=t3,column=3,value=f"=SUM(C{r_c+2}:C{t3-1})"); p.cell(row=t3,column=5,value=f"=SUM(E{r_c+2}:E{t3-1})"); p.cell(row=t3,column=6,value=f"=SUM(F{r_c+2}:F{t3-1})")
    for c in (1,2,3,5,6): p.cell(row=t3,column=c).font=F(bold=True,color=UVA,size=10); p.cell(row=t3,column=c).border=borda
    for c in (2,3,5,6): p.cell(row=t3,column=c).number_format=BRL
    p.cell(row=t3+1,column=1,value="Lançamentos sem cliente (aluguel, pró-labore, impostos) não aparecem aqui; estão em \"Para onde foi o dinheiro\"."); nota(p.cell(row=t3+1,column=1))
    # ano
    A0=t3+3
    p.cell(row=A0,column=1,value="O ano, mês a mês").font=F(bold=True,size=13,color=UVA)
    hdr(p,A0+1,["Mês","Entrou","Saiu","Sobrou","Saldo ao fim do mês"])
    for i in range(12):
        r=A0+2+i
        p.cell(row=r,column=1,value=MESES[i]); calc(p.cell(row=r,column=1),center=False)
        p.cell(row=r,column=2,value=f'={somames("Entrada",i+1)}'); calc(p.cell(row=r,column=2),BRL)
        p.cell(row=r,column=3,value=f'={somames("Saída",i+1)}'); calc(p.cell(row=r,column=3),BRL)
        p.cell(row=r,column=4,value=f'=B{r}-C{r}'); calc(p.cell(row=r,column=4),BRL)
        p.cell(row=r,column=5,value=f'={saldo_ate(f"DATE({Y},{i+2},0)")}'); calc(p.cell(row=r,column=5),BRL)
    p.conditional_formatting.add(f"D{A0+2}:E{A0+13}", FormulaRule(formula=[f'D{A0+2}<0'], font=F(color="C8402E",size=10,bold=True)))
    p.cell(row=A0+14,column=1,value="Meses sem lançamento pago aparecem zerados. Compare só os meses já fechados; o mês corrente ainda não tem todas as saídas."); nota(p.cell(row=A0+14,column=1))
    bc=BarChart(); bc.type="col"; bc.grouping="clustered"; bc.height=7.5; bc.width=18; bc.title="Entrou × saiu no ano"; bc.style=2
    bc.add_data(Reference(p,min_col=2,max_col=3,min_row=A0+1,max_row=A0+13),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=A0+2,max_row=A0+13))
    bc.series[0].graphicalProperties.solidFill=LILAS; bc.series[1].graphicalProperties.solidFill=UVA; bc.legend.position="b"; bc.y_axis.majorGridlines=None
    p.add_chart(bc,f"G{A0}")
    widths(p,(30,15,11,15,17,10,10,10,10,4,12,12)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
    como_usar(wb,"Caixa do escritório",[
     ("O que esta planilha faz","Você lança o que entra (honorários por cliente e caso) e o que sai (custo fixo, pró-labore, impostos, custas). Ela mostra o mês: entrou, saiu, sobrou, saldo acumulado, a receber, a pagar; para onde foi o dinheiro por categoria, quanto veio de cada cliente e o ano mês a mês com gráfico."),
     ("Passo 1","Em Config, preencha o nome do escritório, o ano, o saldo que havia em caixa antes do primeiro lançamento e o mês do painel. Ajuste categorias e clientes de cima para baixo, sem pular linha."),
     ("Passo 2","Em Lançamentos, uma linha por movimento: data, tipo (entrada ou saída), categoria, cliente, caso, descrição, valor, forma e Pago?. Honorário faturado e ainda não recebido entra com Pago? = Não e aparece em A receber; quando cair na conta, troque para Sim."),
     ("Passo 3","Em Painel, escolha o mês em Config e leia de cima para baixo. Só o que está com Pago? = Sim conta como caixa; o resto é previsão."),
     ("Rotina de sexta","15 minutos: lançar o que entrou e saiu na semana, conferir A receber contra o extrato e marcar Sim no que caiu. No fechamento do mês, olhar Sobrou e Saldo acumulado antes de decidir retirada extra ou gasto grande."),
     ("Ligação com as outras planilhas","O total de entradas do mês alimenta a planilha 10 (provisão de impostos); as saídas de pró-labore, a 11; o saldo acumulado e o custo fixo, a 12 (reserva). Cada arquivo tem a própria aba de entrada amarela: copie os números do painel."),
     ("Com a IA","Copie \"Para onde foi o dinheiro\" e \"O ano, mês a mês\" e use o prompt \"Explicar o mês ao sócio\" da biblioteca do kit. Nunca cole nome de cliente ou número de processo na IA sem necessidade: use a coluna Categoria."),
    ])
    proteger(wb); salvar(wb,"09-caixa-do-escritorio.xlsx","Caixa do escritório · Kit de Gestão para Advogados")

if __name__=="__main__":
    ex=lancamentos_exemplo(); print(len(ex),"lançamentos de exemplo")
    for m,t in totais_mensais(ex).items():
        if t["ent"] or t["sai"]: print(MESES[m-1][:3],t)
    build()
