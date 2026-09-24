#!/usr/bin/env python3
"""Planilha 6 do Kit de Gestão para Advogados: Simulador de honorários. Gera 06-simulador-de-honorarios.xlsx"""
from ssg import *
from openpyxl.chart import BarChart, Reference
import dados
NE=10; ND=6                      # etapas e despesas
wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Impostos, margem e as regras de risco valem para todos os casos simulados.",merge_to="H")
cfg["A4"]="Nome do escritório"; cfg["B4"]=f"{dados.ESCRITORIO} (exemplo fictício)"
cfg["A5"]="Mês de referência"; cfg["B5"]="Setembro de 2026"
cfg["A6"]="Data de referência"; cfg["B6"]=dados.HOJE
cfg["A7"]="Impostos e taxas sobre o que entra (%)"; cfg["B7"]=dados.ALIQ
cfg["A8"]="Margem desejada sobre o preço (%)"; cfg["B8"]=dados.MARGEM
cfg["A9"]="Chance de êxito mínima para aceitar êxito puro (%)"; cfg["B9"]=0.60
cfg["A10"]="Folga mínima de chance acima do ponto de equilíbrio (pontos)"; cfg["B10"]=0.15
cfg["A11"]="Folga mínima de horas no valor fixo (%)"; cfg["B11"]=0.20
cfg["A12"]="Aceitar risco médio na recomendação?"; cfg["B12"]="Sim"
for r in range(4,13): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"]); calc(cfg["B6"],DATA)
for r in (7,8,9,10,11): inp(cfg.cell(row=r,column=2),PCT,center=True)
inp(cfg["B12"],center=True)
# Auditoria final (G02): imposto e margem são percentuais de 0 a 99 % e a soma tem de ficar
# abaixo de 100 %; senão o denominador (1 − margem − impostos) zera ou inverte e o preço
# mínimo sai negativo. A validação barra a digitação; as fórmulas guardam a colagem.
def _val_pct(ws,cel,outro,msg):
    dv=DataValidation(type="custom",formula1=f'=IF(ISNUMBER({cel}),AND({cel}>=0,{cel}<1,{cel}+N({outro})<1),FALSE)',allow_blank=False,showErrorMessage=True,errorTitle="Percentual",error=msg)
    dv.add(cel); ws.add_data_validation(dv)
_val_pct(cfg,"B7","B8","Impostos entre 0 % e 99 %, e impostos + margem abaixo de 100 %: senão nenhum preço cobre o custo.")
_val_pct(cfg,"B8","B7","Margem entre 0 % e 99 %, e impostos + margem abaixo de 100 %: senão nenhum preço cobre o custo.")
for _c in ("B9","B10","B11"):
    _dv=DataValidation(type="decimal",operator="between",formula1="0",formula2="1",allow_blank=False,showErrorMessage=True,error="Percentual entre 0 % e 100 %."); _dv.add(_c); cfg.add_data_validation(_dv)
dvsn=lista('"Sim,Não"'); dvsn.add("B12"); cfg.add_data_validation(dvsn)
notas={7:"Exemplo; confira com o contador. O mesmo percentual da planilha 05.",8:"Quanto do preço deve sobrar depois do custo. Igual à planilha 05.",
 9:"Abaixo disso, êxito puro é marcado como risco alto: o escritório trabalha meses e pode não receber.",
 10:"Se a chance estimada está a menos de 15 pontos da chance mínima para não ter prejuízo, o risco é alto.",
 11:"No fixo, o risco é o caso consumir mais horas que o estimado. Com 20% de folga ou mais, risco baixo.",
 12:"\"Não\" faz a recomendação escolher só entre modalidades de risco baixo."}
for r,t in notas.items(): cfg.cell(row=r,column=3,value=t); nota(cfg.cell(row=r,column=3))
# Auditoria final-2 (G02): um predicado só, visível em Config!B13, decide se impostos e margem
# valem (0 a 99 %, soma abaixo de 100 %, sem texto nem branco); hora mínima, impostos, margem,
# ponto de equilíbrio, risco e recomendação consultam essa célula.
cfg["A13"]="Impostos e margem conferem? (calculado)"; rotulo(cfg["A13"])
cfg["B13"]='=IF(AND(ISNUMBER(B7),ISNUMBER(B8)),IF(AND(B7>=0,B7<1,B8>=0,B8<1,B7+B8<1),"Sim","Não"),"Não")'; calc(cfg["B13"])
cfg["C13"]="\"Não\" suspende hora mínima, margens, risco e recomendação no Simulador até a Config ser corrigida."; nota(cfg["C13"])
# Auditoria final-3 (F3-G01): as regras de risco (B9:B12) têm conferência própria; "Não" suspende a
# recomendação e as modalidades que usam o parâmetro inválido.
cfg["A14"]="Regras de risco conferem? (calculado)"; rotulo(cfg["A14"])
cfg["B14"]='=IF(AND(ISNUMBER(B9),ISNUMBER(B10),ISNUMBER(B11)),IF(AND(B9>=0,B9<=1,B10>=0,B10<=1,B11>=0,B11<=1,OR(B12="Sim",B12="Não")),"Sim","Não"),"Não")'; calc(cfg["B14"])
cfg["C14"]="\"Não\" suspende a recomendação: chance mínima e folgas de 0 % a 100 %, e \"Aceitar risco médio\" = Sim ou Não."; nota(cfg["C14"])
cfg["E4"]="Áreas"; rotulo(cfg["E4"]); cfg["E3"]="Preencha de cima para baixo, sem pular linha."; nota(cfg["E3"])
for i in range(26): inp(cfg.cell(row=5+i,column=5))
for i,a in enumerate(dados.AREAS): cfg.cell(row=5+i,column=5,value=a)
cfg["G4"]="Modalidades"; rotulo(cfg["G4"])
for i,m in enumerate(dados.TIPOS_HON): cfg.cell(row=5+i,column=7,value=m).font=F(size=10,color=TINTA)
widths(cfg,(52,14,80,3,22,3,14)); cfg.sheet_view.showGridLines=False
# ---------- Simulador ----------
s=wb.create_sheet("Simulador",0)
titulo(s,'=Config!$B$4&" · Simulador de honorários · "&Config!$B$5',"Preencha o amarelo: dados do caso, horas por etapa, despesas e o que pretende cobrar em cada modalidade. A comparação e a recomendação são calculadas.",merge_to="K")
IMP="Config!$B$7"; MARG="Config!$B$8"; CHMIN="Config!$B$9"; FOLGA="Config!$B$10"; FOLGAH="Config!$B$11"; ACMED="Config!$B$12"
# 1. o caso
s["A7"]="1. O caso"; s["A7"].font=F(bold=True,size=13,color=UVA)
campos=[("Cliente","Cliente",None),("Área","Cível",None),("O que será feito (resumo)","Discussão de contrato de prestação de serviços",None),
 ("Custo-hora do escritório (R$)",dados.CUSTO_HORA,BRL),("Valor em discussão (R$)",60000,BRL0),("Chance de êxito estimada (%)",0.55,PCT)]
for i,(a,v,fmt) in enumerate(campos):
    r=8+i; s.cell(row=r,column=1,value=a); rotulo(s.cell(row=r,column=1)); s.cell(row=r,column=2,value=v); inp(s.cell(row=r,column=2),fmt,center=fmt is not None)
s["B8"]=dados.CLIENTES[7][0]
s["C11"]="Copie da planilha 05 · Custo-hora (Painel, \"Custo-hora do escritório\"). No exemplo, R$ 66,0714 (18.500 ÷ 280 h). A hora mínima do caso (abaixo) é MAIOR do que a hora mínima do escritório na 05 sempre que o caso tiver despesa que o escritório absorve: a da 05 cobre só a estrutura, a daqui cobre também essas despesas rateadas pelas horas do caso."; nota(s["C11"])
s["C12"]="Quanto o cliente recebe, deixa de pagar ou economiza se o caso der certo. Base do cálculo de êxito."; nota(s["C12"])
s["C13"]="Sua estimativa honesta. Ela define o valor esperado e o risco das modalidades com êxito."; nota(s["C13"])
# Auditoria final-4 (F4-G01): valor em discussão, chance e valores das modalidades são dados
# obrigatórios; vazio ou texto não vira zero nem classificação.
def _num(c,mx=None): return f"IF(ISNUMBER({c}),AND({c}>=0{','+c+'<='+mx if mx else ''}),FALSE)"
OK_VC=_num("$B$12"); OK_PCH=_num("$B$13","1")
s["A14"]="Valor esperado da causa (R$)"; rotulo(s["A14"]); s["B14"]=f'=IF(AND({OK_VC},{OK_PCH}),B12*B13,"falta valor ou chance")'; calc(s["B14"],BRL0); s["C14"]="Valor em discussão × chance. É o que se espera receber, em média."; nota(s["C14"])
dva=lista("=OFFSET(Config!$E$5,0,0,MAX(1,COUNTA(Config!$E$5:$E$30)),1)"); dva.add("B9"); s.add_data_validation(dva)
CH="$B$11"; VC="$B$12"; PCH="$B$13"; VE="$B$14"
# 2. horas por etapa
E0=19; EN=E0+NE-1; ET=EN+1
s.cell(row=E0-2,column=1,value="2. Horas estimadas por etapa").font=F(bold=True,size=13,color=UVA)
hdr(s,E0-1,["Etapa","Horas estimadas","Custo (R$)"])
for r in range(E0,EN+1):
    inp(s.cell(row=r,column=1)); inp(s.cell(row=r,column=2),"0.0",center=True)
    # Auditoria final (G01): sem o custo-hora (B11) o custo das horas não é zero, é pendência
    s.cell(row=r,column=3,value=f'=IF(AND(A{r}="",B{r}=""),"",IF(NOT(IF(ISNUMBER(B{r}),B{r}>=0,FALSE)),"faltam as horas",IF(IF(ISNUMBER({CH}),{CH}>=0,FALSE),B{r}*{CH},"falta o custo-hora")))'); calc(s.cell(row=r,column=3),BRL)
s.cell(row=ET,column=1,value="Total de horas e custo das horas"); rotulo(s.cell(row=ET,column=1)); s.cell(row=ET,column=1).border=borda
s.cell(row=ET,column=2,value=f'=IF(SUMPRODUCT(((A{E0}:A{EN}<>"")+(B{E0}:B{EN}<>"")>0)*((ISNUMBER(B{E0}:B{EN})=FALSE)+ISNUMBER(B{E0}:B{EN})*(B{E0}:B{EN}<0)>0))=0,SUM(B{E0}:B{EN}),"horas faltando ou inválidas")'); calc(s.cell(row=ET,column=2),"0.0"); s.cell(row=ET,column=3,value=f'=IF(NOT(IF(ISNUMBER({CH}),{CH}>=0,FALSE)),"falta o custo-hora",IF(SUMPRODUCT(((A{E0}:A{EN}<>"")+(B{E0}:B{EN}<>"")>0)*((ISNUMBER(B{E0}:B{EN})=FALSE)+ISNUMBER(B{E0}:B{EN})*(B{E0}:B{EN}<0)>0))>0,"horas faltando ou inválidas",SUM(C{E0}:C{EN})))'); calc(s.cell(row=ET,column=3),BRL)
for c in (2,3): s.cell(row=ET,column=c).font=F(bold=True,color=UVA,size=10)
HT=f"$B${ET}"; CHT=f"$C${ET}"
HOK=f'SUMPRODUCT((($A${E0}:$A${EN}<>"")+($B${E0}:$B${EN}<>"")>0)*((ISNUMBER($B${E0}:$B${EN})=FALSE)+ISNUMBER($B${E0}:$B${EN})*($B${E0}:$B${EN}<0)>0))=0'   # todas as etapas com horas ≥ 0 (auditoria final-5)
# 3. despesas
D0=ET+4; DN=D0+ND-1; DT=DN+1
s.cell(row=D0-2,column=1,value="3. Despesas do caso").font=F(bold=True,size=13,color=UVA)
hdr(s,D0-1,["Despesa","Valor (R$)","Cliente reembolsa?","Custo para o escritório (R$)"])
for r in range(D0,DN+1):
    inp(s.cell(row=r,column=1)); inp(s.cell(row=r,column=2),BRL,center=True); inp(s.cell(row=r,column=3),center=True)
    s.cell(row=r,column=4,value=f'=IF(AND(A{r}="",B{r}="",C{r}=""),"",IF(NOT(IF(ISNUMBER(B{r}),B{r}>=0,FALSE)),"falta o valor",IF(AND(C{r}<>"Sim",C{r}<>"Não"),"falta: reembolsa?",IF(C{r}="Sim",0,B{r}))))'); calc(s.cell(row=r,column=4),BRL)
dvr=lista('"Sim,Não"'); dvr.add(f"C{D0}:C{DN}"); s.add_data_validation(dvr)
# Auditoria final-5 (F5-G01): horas e valores de despesa são números ≥ 0 (zero digitado vale)
for _rg in (f"B{E0}:B{EN}",f"B{D0}:B{DN}"):
    _dv=DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True,showErrorMessage=True,errorTitle="Valor",error="Número maior ou igual a zero."); _dv.add(_rg); s.add_data_validation(_dv)
s.cell(row=DT,column=1,value="Total de despesas · custo para o escritório"); rotulo(s.cell(row=DT,column=1)); s.cell(row=DT,column=1).border=borda
# Auditoria final-6 (G03): o total bruto não soma por cima de despesa sem valor, com texto ou negativa
s.cell(row=DT,column=2,value=f'=IF(SUMPRODUCT(((A{D0}:A{DN}<>"")+(B{D0}:B{DN}<>"")+(C{D0}:C{DN}<>"")>0)*((ISNUMBER(B{D0}:B{DN})=FALSE)+ISNUMBER(B{D0}:B{DN})*(B{D0}:B{DN}<0)>0))>0,"despesa incompleta",SUM(B{D0}:B{DN}))'); calc(s.cell(row=DT,column=2),BRL); s.cell(row=DT,column=4,value=f'=IF(SUMPRODUCT(((A{D0}:A{DN}<>"")+(B{D0}:B{DN}<>"")+(C{D0}:C{DN}<>"")>0)*(ISNUMBER(D{D0}:D{DN})=FALSE))>0,"despesa incompleta",SUM(D{D0}:D{DN}))'); calc(s.cell(row=DT,column=4),BRL)
s.cell(row=DT,column=4).font=F(bold=True,color=UVA,size=10)
DESP=f"$D${DT}"
s.cell(row=DT+1,column=1,value="Custo total do caso (horas + despesas que o escritório absorve)"); rotulo(s.cell(row=DT+1,column=1)); s.cell(row=DT+1,column=1).border=borda
s.cell(row=DT+1,column=4,value=f'=IF(NOT(IF(ISNUMBER({CH}),{CH}>=0,FALSE)),"Falta o custo-hora (B11)",IF(NOT(ISNUMBER({CHT})),"Horas faltando ou inválidas (bloco 2)",IF(NOT(ISNUMBER({DESP})),"Despesa incompleta (bloco 3)",{CHT}+{DESP})))'); calc(s.cell(row=DT+1,column=4),BRL); s.cell(row=DT+1,column=4).font=F(bold=True,color=UVA,size=10)
CUSTO=f"$D${DT+1}"
# Hora mínima do caso: depende de CUSTO e HT, por isso fica aqui e não junto do rótulo.
s["A15"]="Hora mínima para este caso (R$)"; rotulo(s["A15"])
# Inclui as despesas que o escritório absorve, rateadas pelas horas estimadas: sem
# elas o simulador aprovava preço que não entrega a margem informada.
MINV_='(Config!$B$13<>"Sim")'   # predicado único de Config (auditoria final-2, G02)
s["B15"]=f'=IF(NOT(IF(ISNUMBER({CH}),{CH}>=0,FALSE)),"Falta o custo-hora (B11)",IF(NOT(ISNUMBER({CUSTO})),{CUSTO},IF({MINV_},"Margens inválidas (Config)",IFERROR(IF({HT}=0,B11/(1-{IMP}-{MARG}),({CUSTO}/{HT})/(1-{IMP}-{MARG})),""))))'; calc(s["B15"],BRL)
s["C15"]="(custo das horas + despesas absorvidas) ÷ horas estimadas ÷ (1 − impostos − margem). Sem caso montado, cai no custo-hora do escritório."; nota(s["C15"])
# 4. o que cobrar
M0=DT+5
s.cell(row=M0-2,column=1,value="4. O que você pretende cobrar em cada modalidade").font=F(bold=True,size=13,color=UVA)
hdr(s,M0-1,["Modalidade","Valor","Como funciona"])
mod_in=[("Fixo · valor fechado (R$)",7000,BRL0,"Um valor pelo caso inteiro, independente das horas. Risco: o caso consumir mais horas que o estimado."),
 ("Hora · valor da hora cobrada (R$)",dados.VALOR_HORA_COBRADA,BRL,"Cobra as horas trabalhadas: o escritório não absorve hora a mais, mas o cliente pode questionar as horas. Se a hora cobre a hora mínima do caso ou não, quem diz é a coluna Risco ao lado."),
 ("Êxito · % sobre o resultado",0.25,PCT,"Só recebe se o caso der certo, no fim. Risco: trabalhar e não receber."),
 ("Misto · entrada fixa (R$)",4500,BRL0,"Uma parte fixa no início (cobre o custo) mais um percentual sobre o resultado."),
 ("Misto · % sobre o resultado",0.15,PCT,"")]
for i,(a,v,fmt,t) in enumerate(mod_in):
    r=M0+i; s.cell(row=r,column=1,value=a); calc(s.cell(row=r,column=1),center=False); s.cell(row=r,column=2,value=v); inp(s.cell(row=r,column=2),fmt,center=True)
    s.cell(row=r,column=3,value=t); calc(s.cell(row=r,column=3),center=False); nota(s.cell(row=r,column=3)); s.merge_cells(start_row=r,start_column=3,end_row=r,end_column=8)
FX=f"$B${M0}"; HR=f"$B${M0+1}"; EX=f"$B${M0+2}"; ME=f"$B${M0+3}"; MP=f"$B${M0+4}"
OK_FX=_num(FX); OK_HR=_num(HR); OK_EX=_num(EX,"1"); OK_ME=_num(ME); OK_MP=_num(MP,"1")
s["A16"]="Dados do caso conferem?"; rotulo(s["A16"])
s["B16"]=(f'=IF(NOT({OK_VC}),"falta o valor em discussão (B12)",IF(NOT({OK_PCH}),"falta a chance de êxito, de 0 a 100 % (B13)",IF(NOT({OK_FX}),"falta o valor do fixo ({FX[1]}{FX[3:]})",'
          f'IF(NOT({OK_HR}),"falta a hora cobrada ({HR[1]}{HR[3:]})",IF(NOT({OK_EX}),"falta o % de êxito, de 0 a 100 % ({EX[1]}{EX[3:]})",IF(NOT({OK_ME}),"falta a entrada do misto ({ME[1]}{ME[3:]})",IF(NOT({OK_MP}),"falta o % do misto, de 0 a 100 % ({MP[1]}{MP[3:]})","Sim")))))))')
calc(s["B16"]); s["C16"]="\"Sim\" quando valor em discussão, chance e os valores das quatro modalidades são números no domínio certo. Zero digitado vale; vazio ou texto suspende a modalidade que usa o dado e a recomendação."; nota(s["C16"])
for _c,_mx in ((FX,None),(HR,None),(EX,"1"),(ME,None),(MP,"1"),("$B$12",None),("$B$13","1")):
    _dv=DataValidation(type="decimal",operator="between" if _mx else "greaterThanOrEqual",formula1="0",formula2=_mx,allow_blank=False,showErrorMessage=True,errorTitle="Valor",error="Número de 0 a 100 %." if _mx else "Número maior ou igual a zero.")
    _dv.add(_c.replace("$","")); s.add_data_validation(_dv)
# 5. comparação
C0=M0+8
s.cell(row=C0-2,column=1,value="5. Comparação das modalidades").font=F(bold=True,size=13,color=UVA)
hdr(s,C0-1,["Modalidade","Receita esperada (R$)","Impostos e taxas (R$)","Custo do caso (R$)","Margem esperada (R$)","Margem (%)","Se o caso for perdido (R$)","Ponto de equilíbrio","Risco","Pontuação","Por que este risco"],height=40)
LIQ=f"(1-{IMP})"
# Auditoria final-3 (F3-G01): cada regra de risco vale só com o parâmetro dela numérico e de 0 a
# 100 % (chance mínima B9, folga de chance B10, folga de horas B11); Hora não usa nenhum.
# (F3-G02): nenhuma classificação divide por custo-hora, percentual ou valor que podem ser zero
# legítimos; as desigualdades foram multiplicadas pelos denominadores (positivos quando não zero).
def _okp(c): return f"IF(ISNUMBER({c}),AND({c}>=0,{c}<=1),FALSE)"
OK_FOLGAH=_okp(FOLGAH); OK_CHMIN=_okp(CHMIN); OK_FOLGA=_okp(FOLGA)
SOBRA_FX=f"({FX}*{LIQ}-{DESP})"          # o que o fixo deixa depois de impostos e despesas absorvidas
D_EX=f"({EX}*{VC}*{LIQ})"; D_MP=f"({MP}*{VC}*{LIQ})"; N_MX=f"({CUSTO}-{ME}*{LIQ})"
rows=[
 ("Fixo",f"={FX}",f"=E{{r}}",
  f'=IF({CH}=0,IF({SOBRA_FX}>=0,"Custo-hora zero: o fixo cobre as despesas absorvidas","Custo-hora zero, mas o fixo não cobre as despesas absorvidas"),"Até "&ROUND({SOBRA_FX}/{CH},0)&" horas sem prejuízo (estimadas: "&ROUND({HT},0)&")")',
  f'=IF({SOBRA_FX}>={CH}*{HT}*(1+{FOLGAH}),"Baixo",IF({SOBRA_FX}>={CH}*{HT},"Médio","Alto"))',
  f'=IF({CH}=0,IF({SOBRA_FX}>=0,"Sem custo de hora: só as despesas absorvidas contam, e o fixo as cobre","O fixo não cobre nem as despesas absorvidas"),"Folga de horas: "&ROUND(({SOBRA_FX}/{CH}/{HT}-1)*100,0)&"% (mínimo para risco baixo: "&ROUND({FOLGAH}*100,0)&"%)")',
  OK_FOLGAH, OK_FX, "Falta o valor do fixo"),
 ("Hora",f'=IF({HOK},{HR}*{HT},"")',f"=E{{r}}",
  f'="Hora mínima sem prejuízo: R$ "&FIXED({CUSTO}/({HT}*{LIQ}),2)&" (você cobra R$ "&FIXED({HR},2)&")"',   # FIXED usa o separador do idioma do Excel (pt-BR: 1.234,56)
  f'=IF({HR}>=$B$15,"Baixo",IF({HR}*{HT}*{LIQ}>={CUSTO},"Médio","Alto"))',
  f'=IF({HR}>=$B$15,"A hora cobrada cobre custo, impostos e margem desejada","A hora cobrada fica abaixo da hora mínima com margem (R$ "&FIXED($B$15,2)&")")',
  "TRUE", OK_HR, "Falta a hora cobrada"),
 ("Êxito",f"={EX}*{VE}",f"=-{CUSTO}",
  f'=IF({D_EX}<=0,"Sem percentual de êxito ou valor em discussão: o êxito não paga o custo","Chance mínima de êxito para não ter prejuízo: "&ROUND({CUSTO}/{D_EX}*100,0)&"% (estimada: "&ROUND({PCH}*100,0)&"%)")',
  f'=IF(OR({PCH}<{CHMIN},({PCH}-{FOLGA})*{D_EX}<{CUSTO}),"Alto","Médio")',
  f'=IF({PCH}<{CHMIN},"Chance estimada abaixo do mínimo aceitável para êxito puro ("&ROUND({CHMIN}*100,0)&"%)",IF(({PCH}-{FOLGA})*{D_EX}<{CUSTO},"Chance estimada perto demais do ponto de equilíbrio (ou sem êxito que pague o custo)","Só recebe no fim; o caixa precisa aguentar até lá"))',
  f"AND({OK_CHMIN},{OK_FOLGA})", f"AND({OK_EX},{OK_VC},{OK_PCH})", "Falta % de êxito, valor ou chance"),
 ("Misto",f"={ME}+{MP}*{VE}",f"={ME}*{LIQ}-{CUSTO}",
  f'=IF({N_MX}<=0,"A entrada já cobre o custo do caso; o êxito é margem",IF({D_MP}<=0,"A entrada não cobre o custo e não há êxito que complete","Chance mínima de êxito para não ter prejuízo: "&ROUND({N_MX}/{D_MP}*100,0)&"% (estimada: "&ROUND({PCH}*100,0)&"%)"))',
  f'=IF(IF({N_MX}<=0,{PCH}<{FOLGA},({PCH}-{FOLGA})*{D_MP}<{N_MX}),"Alto",IF({N_MX}>0,"Médio","Baixo"))',
  f'=IF({N_MX}<=0,"Entrada cobre o custo: mesmo perdendo, não há prejuízo","Entrada não cobre todo o custo; parte depende do êxito")',
  OK_FOLGA, f"AND({OK_ME},{OK_MP},{OK_VC},{OK_PCH})", "Falta entrada, %, valor ou chance"),
]
for i,(nome,rec,perde,pe,risco,pq,okr,okin,msgin) in enumerate(rows):
    r=C0+i
    s.cell(row=r,column=1,value=nome); calc(s.cell(row=r,column=1),center=False); s.cell(row=r,column=1).font=F(bold=True,color=UVA,size=10)
    s.cell(row=r,column=2,value=f'=IF({okin},{rec[1:]},"")'); calc(s.cell(row=r,column=2),BRL)
    s.cell(row=r,column=3,value=f'=IF(OR({MINV_},NOT(ISNUMBER(B{r}))),"",B{r}*{IMP})'); calc(s.cell(row=r,column=3),BRL)
    s.cell(row=r,column=4,value=f"={CUSTO}"); calc(s.cell(row=r,column=4),BRL)
    # Auditoria final (G01/G02): custo-hora ausente ou margens impossíveis suspendem margem,
    # ponto de equilíbrio, risco e recomendação — nada de número plausível com dado faltando.
    s.cell(row=r,column=5,value=f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),B{r}-C{r}-D{r},"")'); calc(s.cell(row=r,column=5),BRL)
    s.cell(row=r,column=6,value=f'=IF(OR(B{r}=0,NOT(ISNUMBER(E{r}))),"",E{r}/B{r})'); calc(s.cell(row=r,column=6),PCT)
    s.cell(row=r,column=7,value=f'=IF(AND(ISNUMBER({CUSTO}),NOT({MINV_}),{okin}),{perde.format(r=r)[1:]},"")'); calc(s.cell(row=r,column=7),BRL)
    s.cell(row=r,column=8,value=f'=IF(NOT(ISNUMBER({CUSTO})),{CUSTO},IF({MINV_},"Margens inválidas (Config)",IF(NOT({okin}),"{msgin}",IFERROR({pe[1:]},"Preencha horas e valores"))))'); calc(s.cell(row=r,column=8),center=False); s.cell(row=r,column=8).alignment=Alignment(wrap_text=True,vertical="center")
    # sem horas estimadas não se classifica risco: não há custo por hora nem margem
    s.cell(row=r,column=9,value=f'=IF(NOT(ISNUMBER({CUSTO})),IF(IF(ISNUMBER({CH}),{CH}>=0,FALSE),"Horas ou despesas incompletas","Falta o custo-hora"),IF({MINV_},"Margens inválidas",IF({HT}<=0,"Faltam as horas estimadas",IF(NOT({okr}),"Regra de risco inválida (Config)",IF(NOT({okin}),"{msgin}",IFERROR({risco[1:]},"Sem base para classificar"))))))'); calc(s.cell(row=r,column=9))
    s.cell(row=r,column=10,value=f'=IF(OR(I{r}="Baixo",AND(I{r}="Médio",{ACMED}="Sim")),E{r},-1E+9)'); calc(s.cell(row=r,column=10),BRL0); s.cell(row=r,column=10).font=F(size=9,color=CINZA)
    s.cell(row=r,column=11,value=f'=IF(NOT(ISNUMBER({CUSTO})),IF(IF(ISNUMBER({CH}),{CH}>=0,FALSE),"Corrija o bloco 2 ou 3: etapa sem horas (ou negativa), despesa sem valor ou sem a resposta de reembolso.","Preencha o custo-hora do escritório (B11): sem ele não há custo do caso, margem nem risco."),IF({MINV_},"Impostos e margem em Config fora do permitido (cada um de 0 % a 99 %, soma abaixo de 100 %). Corrija a Config.",IF({HT}<=0,"Preencha as horas estimadas no bloco 2: sem elas não há custo do caso por hora, nem margem, nem ponto de equilíbrio.",IF(NOT({okr}),"A regra de risco desta modalidade em Config está vazia ou fora de 0 % a 100 %. Corrija a Config.",IF(NOT({okin}),"Preencha no bloco 1 ou 4 o dado que falta: "&$B$16&".",IFERROR({pq[1:]},""))))))'); calc(s.cell(row=r,column=11),center=False); s.cell(row=r,column=11).alignment=Alignment(wrap_text=True,vertical="center"); nota(s.cell(row=r,column=11))
    s.row_dimensions[r].height=42
CN=C0+3
s.conditional_formatting.add(f"I{C0}:I{CN}", FormulaRule(formula=[f'AND(I{C0}<>"",I{C0}<>"Baixo",I{C0}<>"Médio",I{C0}<>"Faltam as horas estimadas")'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
s.conditional_formatting.add(f"I{C0}:I{CN}", FormulaRule(formula=[f'I{C0}="Médio"'], fill=fill(AMARELO)))
s.conditional_formatting.add(f"I{C0}:I{CN}", FormulaRule(formula=[f'I{C0}="Baixo"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
s.conditional_formatting.add(f"E{C0}:G{CN}", FormulaRule(formula=[f'AND(ISNUMBER(E{C0}),E{C0}<0)'], font=F(color="C8402E",size=10,bold=True)))
s.conditional_formatting.add(f"A{C0}:A{CN}", FormulaRule(formula=[f'$A{C0}=$A$5'], fill=fill(SOL)))
s.cell(row=CN+1,column=1,value="Pontuação: coluna auxiliar da recomendação (margem esperada; modalidade com risco recusado vale −1 bilhão). Receita esperada de êxito e misto = percentual × valor esperado da causa.").font=F(size=9,color=LILAS)
# recomendação no topo
PONT=f"$J${C0}:$J${CN}"; MODS=f"$A${C0}:$A${CN}"; RISC=f"$I${C0}:$I${CN}"
kpi(s,4,1,"Modalidade recomendada",f'=IF(NOT(IF(ISNUMBER({CH}),{CH}>=0,FALSE)),"Preencha o custo-hora primeiro",IF(NOT(ISNUMBER({CUSTO})),"Corrija horas ou despesas",IF($B$16<>"Sim","Complete os dados do caso",IF({MINV_},"Corrija as margens em Config",IF({HT}<=0,"Estime as horas primeiro",IF(Config!$B$14<>"Sim","Corrija as regras de risco em Config",IF(MAX({PONT})<=-1E+9,"Nenhuma com risco aceitável",INDEX({MODS},MATCH(MAX({PONT}),{PONT},0)))))))))',SOL,UVA,fmt="@")
kpi(s,4,3,"Margem esperada",f'=IF(OR(Config!$B$14<>"Sim",$B$16<>"Sim",MAX({PONT})<=-1E+9),"",INDEX($E${C0}:$E${CN},MATCH(MAX({PONT}),{PONT},0)))',VERDE,VERDE_T,fmt=BRL0)
kpi(s,4,5,"Risco da recomendada",f'=IF(OR(Config!$B$14<>"Sim",$B$16<>"Sim",MAX({PONT})<=-1E+9),"",INDEX({RISC},MATCH(MAX({PONT}),{PONT},0)))',LAVANDA,UVA,fmt="@")
kpi(s,4,7,"Custo total do caso",f"={CUSTO}",LAVANDA,UVA,fmt=BRL0)
kpi(s,4,9,"Hora mínima para o caso","=$B$15",LAVANDA,UVA,fmt=BRL)
fora=f'IF($I${C0}="Alto","Fixo, ","")&IF($I${C0+1}="Alto","Hora, ","")&IF($I${C0+2}="Alto","Êxito, ","")&IF($I${C0+3}="Alto","Misto, ","")'
s["A6"]=(f'=IF(OR(COUNT($E${C0}:$E${CN})=0,{HT}<=0,Config!$B$14<>"Sim",$B$16<>"Sim"),"Sem recomendação ainda: complete o custo-hora, as horas, os dados do caso (B16) e a Config (margens e regras de risco).",'
         f'IF(MAX({PONT})<=-1E+9,"Nenhuma modalidade ficou com risco aceitável. Reveja horas, valores ou a chance de êxito; a maior margem em números seria "&INDEX({MODS},MATCH(MAX($E${C0}:$E${CN}),$E${C0}:$E${CN},0))&".",'
         f'"Regra: maior margem esperada entre as modalidades com risco aceitável. Fora por risco alto: "&IF(COUNTIF({RISC},"Alto")=0,"nenhuma",LEFT({fora},LEN({fora})-2))&". Compare também a coluna \'Se o caso for perdido\' antes de decidir."))')
nota(s["A6"]); s.merge_cells("A6:K6"); s["A6"].alignment=Alignment(wrap_text=True,vertical="top"); s.row_dimensions[6].height=30
bc=BarChart(); bc.type="col"; bc.height=7; bc.width=13; bc.title="Margem esperada × se perder"; bc.style=2
bc.add_data(Reference(s,min_col=5,min_row=C0-1,max_row=CN),titles_from_data=True); bc.add_data(Reference(s,min_col=7,min_row=C0-1,max_row=CN),titles_from_data=True)
bc.set_categories(Reference(s,min_col=1,min_row=C0,max_row=CN)); bc.series[0].graphicalProperties.solidFill="3B1F5E"; bc.series[1].graphicalProperties.solidFill="B89BE0"; bc.legend.position="b"; bc.y_axis.majorGridlines=None
s.add_chart(bc,f"F{E0-2}")
widths(s,(40,16,16,16,16,10,16,40,10,12,44)); s.freeze_panes="A7"; s.sheet_view.showGridLines=False
# ---------- exemplo ----------
etapas=[("Reunião inicial e análise dos documentos",4),("Estudo do caso e estratégia",6),("Redação e protocolo inicial",12),
        ("Acompanhamento e manifestações",20),("Audiências e reuniões",8),("Encerramento e prestação de contas",3)]
for i,(a,h) in enumerate(etapas): s.cell(row=E0+i,column=1,value=a); s.cell(row=E0+i,column=2,value=h)
desp=[("Taxas e custas",800,"Sim"),("Deslocamentos",300,"Não"),("Cópias e certidões",150,"Não"),("Terceiros (perito, tradutor)",0,"Sim")]
for i,(a,v,rb) in enumerate(desp): s.cell(row=D0+i,column=1,value=a); s.cell(row=D0+i,column=2,value=v); s.cell(row=D0+i,column=3,value=rb)
como_usar(wb,"Simulador de honorários",[
 ("O que esta planilha faz","Para um caso, estima as horas por etapa e as despesas, calcula o custo do caso e compara quatro formas de cobrar: fixo, por hora, êxito e misto. Mostra a margem esperada, o que acontece se o caso for perdido, o ponto de equilíbrio e o risco de cada uma, e recomenda a de maior margem com risco aceitável."),
 ("Passo 1","Em Config, confira impostos, margem e as regras de risco (chance mínima para êxito puro, folgas). Uma vez só."),
 ("Passo 2","Em Simulador, bloco 1: cliente, área, custo-hora (copie da planilha 05), valor em discussão e a chance de êxito que você estima."),
 ("Passo 3","Blocos 2 e 3: horas por etapa e despesas do caso, marcando o que o cliente reembolsa."),
 ("Passo 4","Bloco 4: o valor que você pretende cobrar em cada modalidade. Mude os números e veja a comparação e a recomendação no topo mudarem na hora."),
 ("Rotina","Antes de cada proposta: 10 minutos. Depois, leve a modalidade escolhida para a planilha 07 (Proposta de honorários)."),
 ("Com a IA","Copie o bloco 5 e use o prompt \"Honorários 07 · Fixo, hora, êxito ou misto: perguntas antes de escolher\" da biblioteca do kit para listar o que falta perguntar antes de fechar o valor. A decisão continua sua."),
 ("Números em texto","As frases da comparação (\"Hora mínima sem prejuízo: R$ ...\") usam a função FIXED: o separador de milhar e de decimal segue o idioma do Excel (no Excel em português: 1.234,56)."),
])
# Auditoria final-6 (G01): custo-hora aceita só número >= 0 ao digitar; a fórmula também confere.
_dvch=DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True,showErrorMessage=True,errorTitle="Custo-hora",error="Número maior ou igual a zero."); _dvch.add("B11"); s.add_data_validation(_dvch)
proteger(wb); salvar(wb,"06-simulador-de-honorarios.xlsx","Simulador de honorários · Kit de Gestão para Advogados")
