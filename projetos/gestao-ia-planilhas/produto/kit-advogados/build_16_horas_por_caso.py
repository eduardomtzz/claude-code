#!/usr/bin/env python3
"""Planilha 16 do Kit de Gestão para Advogados: Horas por Caso e por Pessoa.
Adaptação da planilha de horas do Kit Completo (09). Gera 16-horas-por-caso.xlsx (Como usar, Painel, Config, Casos, Lançamentos)."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference

N=600; R0=5; RN=R0+N-1             # Lançamentos: linhas 5..604
NC=200; RNC=R0+NC-1                # Casos (entrada): linhas 5..204
NPES=10; P0=12; P1=P0+NPES-1       # Pessoas na Config: linhas 12..21
NATV=12; A0=12; A1=A0+NATV-1       # Atividades na Config: linhas 12..23; Fases (col H) e Modalidades (col J): linhas 12..21
NPAINEL=60                         # casos mostrados no Painel
def off(sheet,col,r0,r1): return f"=OFFSET({sheet}!${col}${r0},0,0,MAX(1,COUNTA({sheet}!${col}${r0}:${col}${r1})),1)"
NOME_ESC=f"{dados.ESCRITORIO} (exemplo fictício)"
M="Config!$B$7"; Y="Config!$B$5"; CHM="Config!$B$23"

wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Custo-hora de cada pessoa = custo mensal ÷ horas faturáveis + rateio dos custos fixos por hora.",merge_to="H")
cfg["A4"]="Escritório"; cfg["B4"]=NOME_ESC
cfg["A5"]="Ano"; cfg["B5"]=2026
cfg["A6"]="Mês do painel"; cfg["B6"]="Setembro"
cfg["A7"]="Número do mês"; cfg["B7"]="=MATCH(B6,$L$5:$L$16,0)"
cfg["A8"]="Custos fixos mensais, sem a equipe (R$)"; cfg["B8"]=dados.CUSTOS_SEM_EQUIPE
for r in range(4,9): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"],center=True); inp(cfg["B6"],center=True); calc(cfg["B7"]); inp(cfg["B8"],BRL0)
cfg["C8"]="Copie da planilha 05 (Custo-hora): aluguel, contador, sistemas, telefone, anuidades, marketing e outros. Sem pró-labore e sem salários."; nota(cfg["C8"])
cfg["L4"]="Meses"; rotulo(cfg["L4"])
for i,m_ in enumerate(MESES): cfg.cell(row=5+i,column=12,value=m_).font=F(size=10,color=TINTA)
dv=lista("=Config!$L$5:$L$16"); dv.add("B6"); cfg.add_data_validation(dv)
cfg["A10"]="Pessoas (até 10)"; rotulo(cfg["A10"])
hdr(cfg,P0-1,["Pessoa","Custo mensal (R$)","Horas faturáveis por mês (meta)","Custo-hora (R$)"])
for r in range(P0,P1+1):
    inp(cfg.cell(row=r,column=1)); inp(cfg.cell(row=r,column=2),BRL0,center=True); inp(cfg.cell(row=r,column=3),center=True)
    cfg.cell(row=r,column=4,value=f'=IF(OR(A{r}="",C{r}=0,C{r}=""),"",B{r}/C{r}+IFERROR($B$8/SUM($C${P0}:$C${P1}),0))'); calc(cfg.cell(row=r,column=4),BRL)
cfg["A23"]="Custo-hora médio do escritório (R$)"; rotulo(cfg["A23"])
cfg["B23"]=f'=IFERROR(($B$8+SUMIFS($B${P0}:$B${P1},$A${P0}:$A${P1},"<>"))/SUM($C${P0}:$C${P1}),0)'; calc(cfg["B23"],BRL)
cfg["A26"]="Custo mensal: pró-labore do sócio, ou salário mais encargos, ou a bolsa. Horas faturáveis por mês: quantas horas a pessoa consegue dedicar a casos (110 para um sócio que também administra; 60 para estagiário). O custo-hora médio vale para as horas lançadas antes de começar a usar esta planilha."; nota(cfg["A26"]); cfg.merge_cells("A26:H26"); cfg["A26"].alignment=Alignment(wrap_text=True,vertical="top"); cfg.row_dimensions[26].height=42
hdr(cfg,A0-1,["Atividades (até 12)"],start=6); hdr(cfg,A0-1,["Fases (até 10)"],start=8); hdr(cfg,A0-1,["Modalidades (até 10)"],start=10)
for r in range(A0,A1+1): inp(cfg.cell(row=r,column=6))
for r in range(A0,A0+10): inp(cfg.cell(row=r,column=8)); inp(cfg.cell(row=r,column=10))
cfg.cell(row=A1+1,column=6,value="Preencha de cima para baixo, sem pular linha. Fases e modalidades: as mesmas da planilha 13 (Carteira); \"Interno\" é a modalidade do trabalho sem cliente.").font=F(size=9,color=LILAS)
for i,v in enumerate(dados.FASES): cfg.cell(row=A0+i,column=8,value=v)
for i,v in enumerate(dados.TIPOS_HON+["Interno"]): cfg.cell(row=A0+i,column=10,value=v)
widths(cfg,(36,18,24,16,3,30,3,14,3,14)); cfg.sheet_view.showGridLines=False
PES_L=off("Config","A",P0,P1); ATV_L=off("Config","F",A0,A1); FAS_L=off("Config","H",A0,A0+9); MOD_L=off("Config","J",A0,A0+9)

# ---------- Casos (entrada) ----------
cs=wb.create_sheet("Casos")
titulo(cs,"Casos (entrada)","Copie da planilha 13 · Carteira (aba Casos; a 13 é a fonte): número, cliente, área, responsável, modalidade, fase e valor; horas estimadas são suas. Trabalho interno entra como modalidade Interno.",merge_to="I")
hdr(cs,4,["Número do processo ou referência","Cliente","Área","Responsável","Modalidade","Valor contratado (R$)","Horas estimadas","Horas gastas antes de usar esta planilha","Fase"])
for r in range(R0,RNC+1):
    for c in range(1,10): inp(cs.cell(row=r,column=c))
    for c in (3,4,5,7,8,9): cs.cell(row=r,column=c).alignment=Alignment(horizontal="center")
    cs.cell(row=r,column=6).number_format=BRL0; cs.cell(row=r,column=7).number_format="0"; cs.cell(row=r,column=8).number_format="0.0"
dvs=[lista(PES_L,strict=True), lista(MOD_L), lista(FAS_L,strict=True),
     DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True,showErrorMessage=True)]
for dv,rng in zip(dvs,[f"D{R0}:D{RNC}",f"E{R0}:E{RNC}",f"I{R0}:I{RNC}",f"F{R0}:H{RNC}"]): dv.add(rng); cs.add_data_validation(dv)
cs.cell(row=RNC+2,column=1,value="Horas gastas antes: se o caso já vinha de antes, anote quantas horas ele consumiu até o dia em que você começou a lançar aqui (uma estimativa honesta serve). Elas entram no total do caso pelo custo-hora médio.").font=F(size=9,color=LILAS)
widths(cs,(28,28,15,16,12,16,11,18,12)); cs.freeze_panes="B5"; cs.sheet_view.showGridLines=False; cs.auto_filter.ref=f"A4:I{RNC}"

# ---------- Lançamentos ----------
h=wb.create_sheet("Lançamentos")
titulo(h,"Lançamentos de horas","Uma linha por pessoa, por dia, por caso. Pessoa, caso e atividade vêm de listas; cliente, custo e mês são calculados.",merge_to="J")
hdr(h,4,["Data","Pessoa","Caso","Atividade","Horas","Faturável?","Cliente","Custo (R$)","Mês","Ano"])
for r in range(R0,RN+1):
    for c in (1,2,3,4,5,6): inp(h.cell(row=r,column=c))
    for c in (1,2,5,6): h.cell(row=r,column=c).alignment=Alignment(horizontal="center")
    h.cell(row=r,column=1).number_format=DATA; h.cell(row=r,column=5).number_format="0.0"
    h.cell(row=r,column=7,value=f'=IF(C{r}="","",IFERROR(INDEX(Casos!$B${R0}:$B${RNC},MATCH(C{r},Casos!$A${R0}:$A${RNC},0)),""))'); calc(h.cell(row=r,column=7),center=False)
    h.cell(row=r,column=8,value=f'=IF(OR(B{r}="",E{r}=""),"",E{r}*IFERROR(INDEX(Config!$D${P0}:$D${P1},MATCH(B{r},Config!$A${P0}:$A${P1},0)),{CHM}))'); calc(h.cell(row=r,column=8),BRL)
    h.cell(row=r,column=9,value=f'=IF(A{r}="","",MONTH(A{r}))'); calc(h.cell(row=r,column=9))
    h.cell(row=r,column=10,value=f'=IF(A{r}="","",YEAR(A{r}))'); calc(h.cell(row=r,column=10))
dvs=[lista(PES_L,strict=True), lista(off("Casos","A",R0,RNC),strict=True), lista(ATV_L,strict=True), lista('"Sim,Não"'),
     DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True,showErrorMessage=True), DataValidation(type="decimal",operator="between",formula1="0",formula2="24",allow_blank=True,showErrorMessage=True)]
for dv,rng in zip(dvs,[f"B{R0}:B{RN}",f"C{R0}:C{RN}",f"D{R0}:D{RN}",f"F{R0}:F{RN}",f"A{R0}:A{RN}",f"E{R0}:E{RN}"]): dv.add(rng); h.add_data_validation(dv)
h.conditional_formatting.add(f"F{R0}:F{RN}", FormulaRule(formula=[f'AND($C{R0}<>"",$F{R0}="")'], fill=fill(VERM)))
h.cell(row=RN+2,column=1,value="Faturável em vermelho: lançamento sem dizer se a hora é faturável. Faturável = hora que o cliente paga (direta ou dentro do valor fixo); reunião interna, deslocamento não cobrado e administração não são.").font=F(size=9,color=LILAS)
widths(h,(12,16,28,30,8,11,26,13,6,7)); h.freeze_panes="A5"; h.sheet_view.showGridLines=False; h.auto_filter.ref=f"A4:J{RN}"

# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
titulo(p,'=Config!$B$4&" · Horas por caso · "&Config!$B$6&" de "&Config!$B$5',"Nada para digitar aqui. Escolha o mês em Config; tudo vem de Lançamentos e Casos.",merge_to="L")
p.merge_cells("A1:L1")
LB=f"Lançamentos!$B${R0}:$B${RN}"; LC=f"Lançamentos!$C${R0}:$C${RN}"; LE=f"Lançamentos!$E${R0}:$E${RN}"; LF=f"Lançamentos!$F${R0}:$F${RN}"; LH=f"Lançamentos!$H${R0}:$H${RN}"; LI=f"Lançamentos!$I${R0}:$I${RN}"; LJ=f"Lançamentos!$J${R0}:$J${RN}"
# Colunas ocultas de classificação: existem para o indicador do Painel cobrir TODOS os
# 200 casos cadastrados, e não só os 60 que a tabela mostra.
cs.cell(row=4,column=11,value="(horas gastas)").font=F(size=9,color=CINZA)
cs.cell(row=4,column=12,value="(custo)").font=F(size=9,color=CINZA)
cs.cell(row=4,column=13,value="(situação)").font=F(size=9,color=CINZA)
for r in range(R0,RNC+1):
    cs.cell(row=r,column=11,value=f'=IF(A{r}="","",N(H{r})+SUMIFS({LE},{LC},$A{r}))').font=F(size=9,color=CINZA)
    cs.cell(row=r,column=12,value=f'=IF(A{r}="","",N(H{r})*{CHM}+SUMIFS({LH},{LC},$A{r}))').font=F(size=9,color=CINZA)
    cs.cell(row=r,column=13,value=(f'=IF(A{r}="","",IF(E{r}="Interno","Interno",IF(K{r}=0,"Sem horas",'
                                   f'IF(L{r}>N(F{r}),"Consome mais do que paga",'
                                   f'IF(AND(N(G{r})>0,K{r}/G{r}>1),"Estourou as horas",'
                                   f'IF(AND(N(F{r})>0,(F{r}-L{r})/F{r}<0.2),"Margem baixa",'
                                   f'IF(AND(N(G{r})>0,K{r}/G{r}>0.85),"Perto do limite","Saudável")))))))')).font=F(size=9,color=CINZA)
for col in ("K","L","M"): cs.column_dimensions[col].hidden=True
MES=f'{LI},{M},{LJ},{Y}'
RT=30                                  # início da tabela "Todos os casos"
T0=RT+2; T1=T0+NPAINEL-1
kpi(p,4,1,"Horas no mês",f'=SUMIFS({LE},{MES})',LAVANDA,UVA,fmt="#,##0.0")
kpi(p,4,3,"Horas faturáveis no mês",f'=SUMIFS({LE},{MES},{LF},"Sim")',VERDE,VERDE_T,fmt="#,##0.0")
kpi(p,4,5,"% de horas faturáveis",f'=IFERROR(C5/A5,0)',VERDE,VERDE_T,fmt=PCT)
kpi(p,4,7,"Custo das horas no mês",f'=SUMIFS({LH},{MES})',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,9,"Meta de faturáveis atingida",f'=IFERROR(C5/SUMIFS(Config!$C${P0}:$C${P1},Config!$A${P0}:$A${P1},"<>"),0)',SOL,UVA,fmt=PCT)
kpi(p,4,11,"Consomem mais do que pagam",f'=COUNTIF(Casos!$M${R0}:$M${RNC},"Consome mais do que paga")',VERM,VERM_T,fmt="0")
p["A7"]="Meta atingida = horas faturáveis lançadas no mês ÷ soma das metas de horas faturáveis (Config); é o realizado, diferente do \"tempo faturável planejado\" da planilha 05. Custo das horas = horas × custo-hora de quem lançou. As horas do mês e as faturáveis são as que o Painel do escritório (17) usa."; nota(p["A7"]); p.merge_cells("A7:L7"); p["A7"].alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[7].height=30
# Casos para olhar primeiro
p["A9"]="Casos para olhar primeiro (ativos, todo o período)"; p["A9"].font=F(bold=True,size=13,color=UVA)
p["A10"]="Casos ativos em que o custo das horas passou do valor contratado, a margem ficou baixa ou as horas estimadas estouraram. Ordem: custo ÷ contratado, do maior para o menor. Encerrados ficam só na tabela completa, com o semáforo."; nota(p["A10"]); p.merge_cells("A10:L10")
hdr(p,11,["#","Caso","Cliente","Responsável","Horas estimadas","Horas gastas","Custo das horas","Valor contratado","Margem (R$)","Situação"])
SC=f"$M${T0}:$M${T1}"
for k in range(1,9):
    r=11+k; m=f'MATCH(LARGE({SC},{k}),{SC},0)'; g=f'LARGE({SC},{k})>0'
    p.cell(row=r,column=1,value=k); calc(p.cell(row=r,column=1))
    for col,src in zip((2,3,4,5,6,7,8,9,10),("A","B","C","E","F","H","I","J","L")):
        p.cell(row=r,column=col,value=f'=IFERROR(IF({g},INDEX(${src}${T0}:${src}${T1},{m}),""),"")'); calc(p.cell(row=r,column=col),center=(col not in (2,3)))
    p.cell(row=r,column=5).number_format="0"; p.cell(row=r,column=6).number_format="0.0"
    for col in (7,8,9): p.cell(row=r,column=col).number_format=BRL0
p.conditional_formatting.add("A12:J19", FormulaRule(formula=['$J12="Consome mais do que paga"'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
p.conditional_formatting.add("A12:J19", FormulaRule(formula=['OR($J12="Margem baixa",$J12="Estourou as horas")'], fill=fill("FFF4CC")))
# Por pessoa no mês
p["A21"]='="Por pessoa em "&Config!$B$6'; p["A21"].font=F(bold=True,size=13,color=UVA)
hdr(p,22,["Pessoa","Horas lançadas","Faturáveis","% faturável","Meta faturável/mês","% da meta","Custo no mês (R$)","Custo-hora (R$)"])
for i in range(NPES):
    r=23+i; src=f"Config!$A${P0+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",SUMIFS({LE},{LB},{src},{MES}))'); calc(p.cell(row=r,column=2),"0.0")
    p.cell(row=r,column=3,value=f'=IF({src}="","",SUMIFS({LE},{LB},{src},{MES},{LF},"Sim"))'); calc(p.cell(row=r,column=3),"0.0")
    p.cell(row=r,column=4,value=f'=IF(OR({src}="",B{r}=0),"",C{r}/B{r})'); calc(p.cell(row=r,column=4),PCT)
    p.cell(row=r,column=5,value=f'=IF({src}="","",Config!$C${P0+i})'); calc(p.cell(row=r,column=5),"0")
    p.cell(row=r,column=6,value=f'=IF(OR({src}="",E{r}=0,E{r}=""),"",C{r}/E{r})'); calc(p.cell(row=r,column=6),PCT)
    p.cell(row=r,column=7,value=f'=IF({src}="","",SUMIFS({LH},{LB},{src},{MES}))'); calc(p.cell(row=r,column=7),BRL0)
    p.cell(row=r,column=8,value=f'=IF({src}="","",Config!$D${P0+i})'); calc(p.cell(row=r,column=8),BRL)
p.conditional_formatting.add(f"F23:F{22+NPES}", FormulaRule(formula=['AND(ISNUMBER(F23),F23>1.1)'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
p.conditional_formatting.add(f"F23:F{22+NPES}", FormulaRule(formula=['AND(ISNUMBER(F23),F23<0.6)'], fill=fill("FFF4CC")))
# Todos os casos
p.cell(row=RT,column=1,value="Todos os casos (todo o período)").font=F(bold=True,size=13,color=UVA)
p.cell(row=RT,column=5,value=f"Mostra os {NPAINEL} primeiros casos da aba Casos; horas gastas = horas de antes + lançadas aqui.").font=F(size=9,color=LILAS)
hdr(p,RT+1,["Caso","Cliente","Responsável","Modalidade","Horas estimadas","Horas gastas","% das estimadas","Custo das horas (R$)","Valor contratado (R$)","Margem (R$)","Margem (%)","Situação"])
for i in range(NPAINEL):
    r=T0+i; src=f"Casos!$A${R0+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",Casos!$B${R0+i})'); calc(p.cell(row=r,column=2),center=False)
    p.cell(row=r,column=3,value=f'=IF({src}="","",Casos!$D${R0+i})'); calc(p.cell(row=r,column=3))
    p.cell(row=r,column=4,value=f'=IF({src}="","",Casos!$E${R0+i})'); calc(p.cell(row=r,column=4))
    p.cell(row=r,column=5,value=f'=IF({src}="","",N(Casos!$G${R0+i}))'); calc(p.cell(row=r,column=5),"0")
    p.cell(row=r,column=6,value=f'=IF({src}="","",N(Casos!$H${R0+i})+SUMIFS({LE},{LC},{src}))'); calc(p.cell(row=r,column=6),"0.0")
    p.cell(row=r,column=7,value=f'=IF(OR({src}="",E{r}=0),"",F{r}/E{r})'); calc(p.cell(row=r,column=7),PCT)
    p.cell(row=r,column=8,value=f'=IF({src}="","",N(Casos!$H${R0+i})*{CHM}+SUMIFS({LH},{LC},{src}))'); calc(p.cell(row=r,column=8),BRL0)
    p.cell(row=r,column=9,value=f'=IF({src}="","",N(Casos!$F${R0+i}))'); calc(p.cell(row=r,column=9),BRL0)
    p.cell(row=r,column=10,value=f'=IF(OR({src}="",D{r}="Interno"),"",I{r}-H{r})'); calc(p.cell(row=r,column=10),BRL0)
    p.cell(row=r,column=11,value=f'=IF(OR({src}="",D{r}="Interno",I{r}=0),"",J{r}/I{r})'); calc(p.cell(row=r,column=11),PCT)
    p.cell(row=r,column=12,value=f'=IF({src}="","",IF(D{r}="Interno","Interno",IF(F{r}=0,"Sem horas",IF(H{r}>I{r},"Consome mais do que paga",IF(AND(G{r}<>"",G{r}>1),"Estourou as horas",IF(AND(K{r}<>"",K{r}<0.2),"Margem baixa",IF(AND(G{r}<>"",G{r}>0.85),"Perto do limite","Saudável")))))))'); calc(p.cell(row=r,column=12))
    p.cell(row=r,column=13,value=f'=IF(AND(Casos!$I${R0+i}<>"Encerrado",OR(L{r}="Consome mais do que paga",L{r}="Margem baixa",L{r}="Estourou as horas")),IFERROR(H{r}/I{r},9)+ROW()/1000000,0)'); p.cell(row=r,column=13).font=F(color=CINZA,size=9)
p.column_dimensions["M"].hidden=True
p.conditional_formatting.add(f"L{T0}:L{T1}", FormulaRule(formula=[f'L{T0}="Consome mais do que paga"'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
p.conditional_formatting.add(f"L{T0}:L{T1}", FormulaRule(formula=[f'OR(L{T0}="Margem baixa",L{T0}="Estourou as horas")'], fill=fill("FFF4CC")))
p.conditional_formatting.add(f"L{T0}:L{T1}", FormulaRule(formula=[f'L{T0}="Perto do limite"'], fill=fill("FFF4CC")))
p.conditional_formatting.add(f"L{T0}:L{T1}", FormulaRule(formula=[f'L{T0}="Saudável"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
p.conditional_formatting.add(f"J{T0}:J{T1}", FormulaRule(formula=[f'AND(ISNUMBER(J{T0}),J{T0}<0)'], font=F(color="C8402E",size=10,bold=True)))
p.cell(row=T1+2,column=1,value="Semáforo: Consome mais do que paga = custo das horas maior que o contratado. Estourou as horas = gastas acima das estimadas. Margem baixa = menos de 20%. Perto do limite = mais de 85% das horas estimadas. Em caso por êxito, o contratado é o honorário esperado no fim.").font=F(size=9,color=LILAS)
bc=BarChart(); bc.type="bar"; bc.height=6; bc.width=12; bc.title="Horas faturáveis × meta, por pessoa"; bc.style=2
bc.add_data(Reference(p,min_col=3,min_row=22,max_row=22+NPES),titles_from_data=True); bc.add_data(Reference(p,min_col=5,min_row=22,max_row=22+NPES),titles_from_data=True)
bc.set_categories(Reference(p,min_col=1,min_row=23,max_row=22+NPES))
bc.series[0].graphicalProperties.solidFill="3B1F5E"; bc.series[1].graphicalProperties.solidFill="B89BE0"; bc.legend.position="b"; bc.x_axis.majorGridlines=None
p.add_chart(bc,"L9")
widths(p,(28,26,16,12,12,12,14,16,16,14,11,24)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False

# ---------- Exemplo (Ferraz & Lima) ----------
for i,(nome,papel,custo,horas) in enumerate(dados.PESSOAS):
    cfg.cell(row=P0+i,column=1,value=nome); cfg.cell(row=P0+i,column=2,value=custo); cfg.cell(row=P0+i,column=3,value=horas)
ATV=["Reunião com cliente","Audiência","Elaboração de documento","Pesquisa e estudo do caso","Análise de documentos","Atendimento por telefone ou WhatsApp","Diligência externa","Deslocamento","Administrativo do caso","Gestão do escritório","Captação e propostas"]
for i,v in enumerate(ATV): cfg.cell(row=A0+i,column=6,value=v)
INTERNOS=dados.INTERNOS
rows=dados.HORAS; antes=dados.HORAS_ANTES   # lançamentos de 01/07 a 11/09/2026 (fonte única do exemplo)
assert len(rows)<=N, len(rows)
for i,c in enumerate(dados.CASOS):
    for col,v in zip((1,2,3,4,5,6,7,8,9),(c["numero"],c["cliente"],c["area"],c["responsavel"],c["tipo_hon"],c["valor_contratado"],c["horas_estimadas"],antes[c["numero"]],c["fase"])): cs.cell(row=R0+i,column=col,value=v)
for j,(num,cli) in enumerate(INTERNOS):
    r=R0+len(dados.CASOS)+j
    for col,v in zip((1,2,3,4,5,6,7,8),(num,cli,"—","Marina Ferraz","Interno",0,0,0)): cs.cell(row=r,column=col,value=v)
for i,row in enumerate(rows):
    for c,v in enumerate(row,start=1): h.cell(row=R0+i,column=c,value=v)
print(len(rows),"lançamentos;",sum(r[4] for r in rows),"horas; jul/ago/set:",dados.horas_mes(7),dados.horas_mes(8),dados.horas_mes(9))

como_usar(wb,"Horas por Caso e por Pessoa",[
 ("O que esta planilha faz","Cada pessoa lança as horas por caso; ela transforma em custo pelo custo-hora, compara com as horas estimadas e o valor contratado, mostra a margem de cada caso com um semáforo (\"consome mais do que paga\") e a ocupação de cada pessoa no mês."),
 ("Passo 1","Em Config, cadastre as pessoas com custo mensal e meta de horas faturáveis, os custos fixos do escritório (da planilha 05), as atividades e as listas de fases e modalidades (as mesmas da planilha 13). O custo-hora é calculado."),
 ("Passo 2","Em Casos, cole a lista da planilha 13 · Carteira (a 13 é a fonte do cadastro) com valor contratado e horas estimadas. Se o caso já vinha de antes, anote as horas gastas até aqui. Trabalho interno (gestão, captação) entra como modalidade Interno."),
 ("Passo 3","Em Lançamentos, uma linha por pessoa, por dia, por caso: data, pessoa, caso, atividade, horas e se é faturável. Dois minutos no fim do dia bastam."),
 ("Passo 4","Em Painel, escolha o mês em Config. \"Casos para olhar primeiro\" mostra os que consomem mais do que pagam; \"Por pessoa\" mostra horas, % faturável e ocupação contra a meta; \"Todos os casos\" traz o semáforo completo."),
 ("Rotina","Cada pessoa lança no fim do dia ou na sexta (10 minutos). Dia 1 do mês: olhar a margem por caso antes de precificar a próxima proposta (planilha 06)."),
 ("Exemplo","No exemplo o mês do painel é setembro de 2026 (o mês da referência, lançado até sexta 11/09, como o caixa). O controle de horas começou em julho: os lançamentos vão de 01/07 a 11/09; o que veio antes está em \"horas gastas antes\". Para ver um mês completo, troque Config para Agosto."),
 ("Com a IA","Copie \"Casos para olhar primeiro\" e use o prompt \"Honorários 03 · Casos que consomem mais do que pagam\" da biblioteca; antes de precificar o próximo caso, \"Honorários 02 · Revisar a proposta pela margem\". Para justificar horas a um cliente, copie os lançamentos do caso e peça um resumo por atividade."),
])
proteger(wb); salvar(wb,"16-horas-por-caso.xlsx","Horas por Caso e por Pessoa · Kit de Gestão para Advogados")
