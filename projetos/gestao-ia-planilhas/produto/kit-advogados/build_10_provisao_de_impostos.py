#!/usr/bin/env python3
"""Planilha 10 do Kit de Gestão para Advogados: Provisão de impostos, 13º e férias. Gera 10-provisao-de-impostos.xlsx
Exemplo alimentado pelos totais mensais da planilha 09 (build_09_caixa_do_escritorio.totais_mensais)."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference

NP=8  # pessoas na equipe (linhas amarelas)
PCT1="0.0%"
wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações","Células amarelas: você preenche. Alíquotas e regras são as que o escritório combinou com o contador.",merge_to="H")
cfg["A4"]="Escritório"; cfg["B4"]=f"{dados.ESCRITORIO} (exemplo fictício)"
cfg["A5"]="Ano"; cfg["B5"]=2026
cfg["A6"]="Mês do painel"; cfg["B6"]="Setembro"
cfg["A7"]="Número do mês"; cfg["B7"]="=MATCH(B6,$H$5:$H$16,0)"
cfg["A8"]="Alíquota efetiva de impostos sobre as entradas"; cfg["B8"]=dados.ALIQ
cfg["A9"]="Adicional sobre férias / recesso"; cfg["B9"]=1/3
cfg["A10"]="Encargos sobre 13º e férias"; cfg["B10"]=0
cfg["A11"]="Mês em que o 13º é pago"; cfg["B11"]="Dezembro"
cfg["A12"]="Mês previsto de férias / recesso"; cfg["B12"]="Janeiro"
for r in range(4,13): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"],center=True); inp(cfg["B6"],center=True); calc(cfg["B7"]); inp(cfg["B8"],PCT1,center=True); inp(cfg["B9"],PCT1,center=True); inp(cfg["B10"],PCT1,center=True); inp(cfg["B11"],center=True); inp(cfg["B12"],center=True)
cfg["D8"]="Percentual que o escritório separa de cada real recebido, combinado com o contador. Escritórios pequenos costumam usar algo entre 6 % e 15,5 %, conforme o enquadramento; esta planilha não afirma qual é o seu. Confirme a alíquota com o contador."
cfg["D9"]="Padrão usual: um terço. Confirme com o contador."
cfg["D10"]="Se houver encargos sobre a folha (INSS, FGTS), o contador informa o percentual. Deixe 0 se não houver."
cfg["D11"]="Se o 13º é pago em duas parcelas, escolha Novembro para se antecipar."
for r in (8,9,10,11): nota(cfg.cell(row=r,column=4)); cfg.cell(row=r,column=4).alignment=Alignment(wrap_text=True,vertical="top"); cfg.merge_cells(start_row=r,start_column=4,end_row=r,end_column=6)
cfg.row_dimensions[8].height=58
cfg["H4"]="Meses"; rotulo(cfg["H4"])
for i,m_ in enumerate(MESES): cfg.cell(row=5+i,column=8,value=m_).font=F(size=10,color=TINTA)
for c in ("B6","B11","B12"):
    dv=lista("=Config!$H$5:$H$16"); dv.add(c); cfg.add_data_validation(dv)
cfg["A14"]="Equipe: 13º e férias / recesso a provisionar (até 8 pessoas)"; rotulo(cfg["A14"])
hdr(cfg,15,["Pessoa","Custo mensal","Provisionar 13º?","Provisionar férias / recesso?","13º por mês","Férias por mês"])
E0=16
for i in range(NP):
    r=E0+i
    inp(cfg.cell(row=r,column=1)); inp(cfg.cell(row=r,column=2),BRL); inp(cfg.cell(row=r,column=3),center=True); inp(cfg.cell(row=r,column=4),center=True)
    cfg.cell(row=r,column=5,value=f'=IF(OR(A{r}="",C{r}<>"Sim"),0,B{r}/12*(1+$B$10))'); calc(cfg.cell(row=r,column=5),BRL)
    cfg.cell(row=r,column=6,value=f'=IF(OR(A{r}="",D{r}<>"Sim"),0,B{r}*(1+$B$9)/12*(1+$B$10))'); calc(cfg.cell(row=r,column=6),BRL)
ET=E0+NP
cfg.cell(row=ET,column=1,value="Total por mês"); cfg.cell(row=ET,column=5,value=f"=SUM(E{E0}:E{ET-1})"); cfg.cell(row=ET,column=6,value=f"=SUM(F{E0}:F{ET-1})")
cfg.cell(row=ET+1,column=1,value="Total no ano (12 meses)"); cfg.cell(row=ET+1,column=5,value=f"=E{ET}*12"); cfg.cell(row=ET+1,column=6,value=f"=F{ET}*12")
for r in (ET,ET+1):
    for c in (1,5,6): cfg.cell(row=r,column=c).font=F(bold=True,color=UVA,size=10); cfg.cell(row=r,column=c).border=borda
    for c in (5,6): cfg.cell(row=r,column=c).number_format=BRL
dvsn=lista('"Sim,Não"'); dvsn.add(f"C{E0}:D{ET-1}"); cfg.add_data_validation(dvsn)
notas=["Sócios: marque Sim no 13º se os sócios quiserem guardar uma retirada extra de dezembro (decisão do escritório, não obrigação). Estagiária: recesso conforme o contrato.",
       "O que é obrigatório provisionar para cada vínculo depende do contrato e da legislação: confirme com o contador. A planilha só separa o dinheiro que você mandar separar."]
for i,t in enumerate(notas): cfg.cell(row=ET+3+i,column=1,value=t); nota(cfg.cell(row=ET+3+i,column=1))
for i,(n_,papel,custo,_) in enumerate(dados.PESSOAS):
    soc=papel.startswith("Sóci")
    for c,v in zip((1,2,3,4),(n_,custo,"Sim" if soc else "Não","Não" if soc else "Sim")): cfg.cell(row=E0+i,column=c,value=v)
widths(cfg,(40,16,18,26,16,16,3,12)); cfg.sheet_view.showGridLines=False
# ---------- Entradas e pagamentos ----------
en=wb.create_sheet("Entradas e pagamentos")
titulo(en,"Entradas e pagamentos do ano","Amarelo: o que entrou no caixa em cada mês (Entrou no mês do Painel da planilha 09, sem a categoria Outras entradas) e o que foi efetivamente pago de impostos, 13º e férias.",merge_to="H")
hdr(en,4,["Mês","Entradas recebidas","Impostos pagos","13º pago","Férias / recesso pagos","Total usado no mês"])
for i in range(12):
    r=5+i
    en.cell(row=r,column=1,value=MESES[i]); calc(en.cell(row=r,column=1),center=False)
    for c in (2,3,4,5): inp(en.cell(row=r,column=c),BRL)
    en.cell(row=r,column=6,value=f"=SUM(C{r}:E{r})"); calc(en.cell(row=r,column=6),BRL)
en.cell(row=17,column=1,value="Total"); en.cell(row=17,column=1).font=F(bold=True,color=UVA,size=10); en.cell(row=17,column=1).border=borda
for c in (2,3,4,5,6):
    en.cell(row=17,column=c,value=f"=SUM({L(c)}5:{L(c)}16)"); en.cell(row=17,column=c).font=F(bold=True,color=UVA,size=10); en.cell(row=17,column=c).border=borda; en.cell(row=17,column=c).number_format=BRL
en["A19"]="Deixe em branco os meses que ainda não fecharam. Impostos pagos: o valor das guias quitadas no mês (no exemplo, a guia de setembro ainda não foi paga: ver planilha 09). Setembro mostra o que entrou até 11/09."; nota(en["A19"])
en["A20"]="Entradas recebidas = Entrou no mês (Painel da 09) menos \"Outras entradas\" (no exemplo, a devolução de despesa pessoal de abril, R$ 480, que não é receita e não paga imposto)."; nota(en["A20"])
widths(en,(14,18,16,14,20,18)); en.sheet_view.showGridLines=False
tot=dados.TOTAIS
for m in range(1,10):
    en.cell(row=4+m,column=2,value=tot[m]["ent_sem_devol"])
    if tot[m]["imp"]: en.cell(row=4+m,column=3,value=tot[m]["imp"])
# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
titulo(p,'=Config!B4&" · Provisão de impostos, 13º e férias · "&Config!B6&" de "&Config!B5',"Nada para digitar aqui. Alíquota e equipe vêm de Config; entradas e pagamentos, da aba ao lado.",merge_to="M")
M="Config!$B$7"; AL="Config!$B$8"; M13="MATCH(Config!$B$11,Config!$H$5:$H$16,0)"; MFE="MATCH(Config!$B$12,Config!$H$5:$H$16,0)"
T13=f"Config!$E${ET}"; TFE=f"Config!$F${ET}"; A13=f"Config!$E${ET+1}"; AFE=f"Config!$F${ET+1}"
T0=9  # primeira linha da tabela mensal
hdr(p,T0-1,["Mês","Entradas","Provisão de impostos","Provisão de 13º","Provisão de férias","Total a separar no mês","Usado no mês","Separado (acumulado)","Usado (acumulado)","Saldo provisionado","Compromisso do mês seguinte","Alerta","Quanto falta"],height=40)
for i in range(12):
    r=T0+i; er=5+i
    p.cell(row=r,column=1,value=MESES[i]); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f"='Entradas e pagamentos'!B{er}"); calc(p.cell(row=r,column=2),BRL0)
    p.cell(row=r,column=3,value=f"=B{r}*{AL}"); calc(p.cell(row=r,column=3),BRL0)
    p.cell(row=r,column=4,value=f"={T13}"); calc(p.cell(row=r,column=4),BRL0)
    p.cell(row=r,column=5,value=f"={TFE}"); calc(p.cell(row=r,column=5),BRL0)
    p.cell(row=r,column=6,value=f"=SUM(C{r}:E{r})"); calc(p.cell(row=r,column=6),BRL0)
    p.cell(row=r,column=7,value=f"='Entradas e pagamentos'!F{er}"); calc(p.cell(row=r,column=7),BRL0)
    p.cell(row=r,column=8,value=f"=SUM($F${T0}:F{r})"); calc(p.cell(row=r,column=8),BRL0)
    p.cell(row=r,column=9,value=f"=SUM($G${T0}:G{r})"); calc(p.cell(row=r,column=9),BRL0)
    p.cell(row=r,column=10,value=f"=H{r}-I{r}"); calc(p.cell(row=r,column=10),BRL0)
    p.cell(row=r,column=11,value=f"=C{r}+IF(MOD({i+1},12)+1={M13},{A13},0)+IF(MOD({i+1},12)+1={MFE},{AFE},0)"); calc(p.cell(row=r,column=11),BRL0)
    p.cell(row=r,column=12,value=f'=IF(J{r}<K{r},"Falta separar","Coberto")'); calc(p.cell(row=r,column=12))
    p.cell(row=r,column=13,value=f"=MAX(0,K{r}-J{r})"); calc(p.cell(row=r,column=13),BRL0)
    for c in (6,10): p.cell(row=r,column=c).font=F(bold=True,color=UVA,size=10)
TR=T0+12
p.conditional_formatting.add(f"L{T0}:L{TR-1}", FormulaRule(formula=[f'L{T0}="Falta separar"'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
p.conditional_formatting.add(f"L{T0}:L{TR-1}", FormulaRule(formula=[f'L{T0}="Coberto"'], fill=fill(VERDE), font=F(color=VERDE_T,size=10)))
p.conditional_formatting.add(f"A{T0}:M{TR-1}", FormulaRule(formula=[f'ROW()-{T0}+1={M}'], fill=fill(LAVANDA)))
p.conditional_formatting.add(f"J{T0}:J{TR-1}", FormulaRule(formula=[f'J{T0}<0'], font=F(color="C8402E",size=10,bold=True)))
# KPIs do mês
kpi(p,4,1,"Entradas do mês",f"=INDEX(B{T0}:B{TR-1},{M})",VERDE,VERDE_T,fmt=BRL0)
kpi(p,4,3,"A separar no mês",f"=INDEX(F{T0}:F{TR-1},{M})",SOL,UVA,fmt=BRL0)
kpi(p,4,5,"Saldo provisionado ao fim do mês",f"=INDEX(J{T0}:J{TR-1},{M})",LAVANDA,UVA,fmt=BRL0)
kpi(p,4,7,"Compromisso do mês seguinte",f"=INDEX(K{T0}:K{TR-1},{M})",VERM,VERM_T,fmt=BRL0)
kpi(p,4,9,"Quanto falta separar",f"=INDEX(M{T0}:M{TR-1},{M})",LAVANDA,UVA,fmt=BRL0)
kpi(p,4,11,"Situação",f'=IF(E5<G5,"Atenção: falta separar","Coberto")',LAVANDA,UVA,fmt="General",span=3)
p["K5"].font=F(size=13,bold=True,color=UVA)
p.conditional_formatting.add("K5", FormulaRule(formula=['LEFT(K5,7)="Atenção"'], fill=fill(VERM), font=F(color=VERM_T,size=13,bold=True)))
p.conditional_formatting.add("K5", FormulaRule(formula=['K5="Coberto"'], fill=fill(VERDE), font=F(color=VERDE_T,size=13,bold=True)))
p.cell(row=T0-2,column=1,value="O ano, mês a mês").font=F(bold=True,size=13,color=UVA)
notas=["Como ler: Provisão de impostos = entradas do mês × alíquota efetiva de Config. Compromisso do mês seguinte = imposto sobre as entradas deste mês (pago no mês que vem) + 13º e férias inteiros nos meses escolhidos em Config.",
       "Saldo provisionado é o que deveria estar numa conta separada da conta corrente do escritório. Compare com o extrato dessa conta todo fechamento de mês. Meses futuros mostram só 13º e férias até você lançar as entradas.",
       "Confirme alíquotas, encargos e regras de 13º e férias com o contador. Esta planilha separa dinheiro; ela não calcula nem substitui a apuração de imposto."]
for i,t in enumerate(notas): p.cell(row=TR+1+i,column=1,value=t); nota(p.cell(row=TR+1+i,column=1))
bc=BarChart(); bc.type="col"; bc.grouping="clustered"; bc.height=7.5; bc.width=22; bc.title="Saldo provisionado × compromisso do mês seguinte"; bc.style=2
bc.add_data(Reference(p,min_col=10,max_col=11,min_row=T0-1,max_row=TR-1),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=T0,max_row=TR-1))
bc.series[0].graphicalProperties.solidFill=LILAS; bc.series[1].graphicalProperties.solidFill=UVA; bc.legend.position="b"; bc.y_axis.majorGridlines=None
p.add_chart(bc,f"A{TR+5}")
widths(p,(13,13,13,13,13,14,13,14,13,14,15,14,13)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False
como_usar(wb,"Provisão de impostos, 13º e férias",[
 ("O que esta planilha faz","Separa, todo mês, o dinheiro do imposto sobre o que entrou, do 13º e das férias / recesso da equipe. Mostra quanto deveria estar guardado, quanto já foi usado e avisa quando o que está separado não cobre o compromisso do mês seguinte."),
 ("Passo 1","Em Config, digite a alíquota efetiva que o escritório combinou com o contador (a planilha não afirma qual é a sua), o adicional de férias, os encargos (se houver), os meses de pagamento do 13º e das férias e a equipe com Sim / Não em cada provisão."),
 ("Passo 2","Em Entradas e pagamentos, uma linha por mês: o total que entrou (copie de Entrou no mês do Painel da planilha 09, sem Outras entradas) e o que foi pago de impostos, 13º e férias. Meses não fechados ficam em branco."),
 ("Passo 3","Em Painel, escolha o mês em Config. A tabela mostra o ano inteiro; a linha do mês fica destacada. Alerta em vermelho = separar mais dinheiro antes do próximo compromisso."),
 ("Rotina","No fechamento de cada mês (dia 5): lançar as entradas do mês anterior, transferir o valor de A separar no mês para a conta de provisão e conferir se o extrato dessa conta bate com Saldo provisionado."),
 ("Com o contador","Leve esta aba para a reunião mensal: alíquota, provisões e o que foi pago. Ajuste Config sempre que o contador mudar a orientação. Confirme alíquotas com o contador."),
 ("Com a IA","Copie a tabela do Painel e use o prompt \"Caixa 04 · Preparar a reunião mensal com o contador\" ou \"Caixa 05 · Perguntas sobre a provisão de impostos, 13º e férias\" da biblioteca do kit para montar a pauta e as perguntas."),
])
proteger(wb); salvar(wb,"10-provisao-de-impostos.xlsx","Provisão de impostos, 13º e férias · Kit de Gestão para Advogados")
