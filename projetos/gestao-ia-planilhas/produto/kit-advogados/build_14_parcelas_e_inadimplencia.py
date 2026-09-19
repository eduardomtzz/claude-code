#!/usr/bin/env python3
"""Planilha 14 do Kit de Gestão para Advogados: Parcelas e Inadimplência com régua de cobrança.
Gera 14-parcelas-e-inadimplencia.xlsx (Como usar, Painel, Config, Parcelas, Casos)."""
from ssg import *
import dados
from openpyxl.chart import BarChart, Reference

NP=400; R0=5; RNP=R0+NP-1          # Parcelas: linhas 5..404
NC=200; RNC=R0+NC-1                # Casos (entrada): linhas 5..204
RG0=10; RG1=13                     # Régua na Config: linhas 10..13
def off(sheet,col,r0,r1): return f"=OFFSET({sheet}!${col}${r0},0,0,MAX(1,COUNTA({sheet}!${col}${r0}:${col}${r1})),1)"
def dstr(ref): return f'TEXT(DAY({ref}),"00")&"/"&TEXT(MONTH({ref}),"00")&"/"&YEAR({ref})'
NOME_ESC=f"{dados.ESCRITORIO} (exemplo fictício)"
HOJE="Config!$B$5"; JC="Config!$B$6"; JL="Config!$B$7"
RD=f"Config!$A${RG0}:$A${RG1}"; RA=f"Config!$B${RG0}:$B${RG1}"; RF=f"Config!$C${RG0}:$C${RG1}"

wb=Workbook()
# ---------- Config ----------
cfg=wb.active; cfg.title="Config"
titulo(cfg,"Configurações e régua de cobrança","Células amarelas: você preenche. A régua diz o que fazer em cada faixa de atraso; o texto aparece no Painel, parcela por parcela.",merge_to="H")
cfg["A4"]="Escritório"; cfg["B4"]=NOME_ESC
cfg["A5"]="Data de referência (hoje)"; cfg["B5"]=dados.HOJE
cfg["A6"]="Janela curta: vence em até (dias)"; cfg["B6"]=7
cfg["A7"]="Janela longa: vence em até (dias)"; cfg["B7"]=30
for r in (4,5,6,7): rotulo(cfg.cell(row=r,column=1))
inp(cfg["B4"]); inp(cfg["B5"],DATA); inp(cfg["B6"],center=True); inp(cfg["B7"],center=True)
hdr(cfg,9,["A partir de (dias de atraso)","Ação sugerida","Faixa (calculada)"])
for r in range(RG0,RG1+1):
    inp(cfg.cell(row=r,column=1),center=True); inp(cfg.cell(row=r,column=2)); cfg.cell(row=r,column=2).alignment=Alignment(wrap_text=True,vertical="top")
    if r<RG1: cfg.cell(row=r,column=3,value=f'=IF(A{r}="","",IF(A{r+1}-1=A{r},A{r}&" dia"&IF(A{r}>1,"s",""),A{r}&" a "&(A{r+1}-1)&" dias"))')
    else: cfg.cell(row=r,column=3,value=f'=IF(A{r}="","",A{r}&" dias ou mais")')
    calc(cfg.cell(row=r,column=3)); cfg.row_dimensions[r].height=44
cfg.cell(row=RG1+2,column=1,value="Os dias precisam estar em ordem crescente (ex.: 1, 7, 15, 30). Cada parcela vencida recebe a ação da maior faixa que ela já alcançou. Escreva as ações como você fala com o cliente: educado, direto e sem ameaça.").font=F(size=9,color=LILAS)
cfg.cell(row=RG1+3,column=1,value="Os modelos de mensagem do bônus \"15 modelos de cobrança e confirmação\" seguem esta mesma régua.").font=F(size=9,color=LILAS)
widths(cfg,(30,70,18,3,3,3,3,3)); cfg.sheet_view.showGridLines=False

# ---------- Casos (entrada) ----------
cs=wb.create_sheet("Casos")
titulo(cs,"Casos (entrada)","Copie da planilha 13 · Carteira (aba Casos; a 13 é a fonte): número, cliente, modalidade e valor. As colunas brancas somam as parcelas de cada caso.",merge_to="J")
hdr(cs,4,["Número do processo ou referência","Cliente","Modalidade","Valor contratado (R$)","Parcelas","Valor em parcelas (R$)","Pago (R$)","Em aberto (R$)","Vencido (R$)","Diferença contratado − parcelas"])
PA=f"Parcelas!$A${R0}:$A${RNP}"; PE=f"Parcelas!$E${R0}:$E${RNP}"; PH=f"Parcelas!$H${R0}:$H${RNP}"
for r in range(R0,RNC+1):
    inp(cs.cell(row=r,column=1)); inp(cs.cell(row=r,column=2)); inp(cs.cell(row=r,column=3),center=True); inp(cs.cell(row=r,column=4),BRL0)
    cs.cell(row=r,column=5,value=f'=IF(A{r}="","",COUNTIFS({PA},A{r}))'); calc(cs.cell(row=r,column=5))
    cs.cell(row=r,column=6,value=f'=IF(A{r}="","",SUMIFS({PE},{PA},A{r}))'); calc(cs.cell(row=r,column=6),BRL0)
    cs.cell(row=r,column=7,value=f'=IF(A{r}="","",SUMIFS({PE},{PA},A{r},{PH},"Paga"))'); calc(cs.cell(row=r,column=7),BRL0)
    cs.cell(row=r,column=8,value=f'=IF(A{r}="","",F{r}-G{r})'); calc(cs.cell(row=r,column=8),BRL0)
    cs.cell(row=r,column=9,value=f'=IF(A{r}="","",SUMIFS({PE},{PA},A{r},{PH},"Vencida"))'); calc(cs.cell(row=r,column=9),BRL0)
    cs.cell(row=r,column=10,value=f'=IF(OR(A{r}="",E{r}=0),"",D{r}-F{r})'); calc(cs.cell(row=r,column=10),BRL0)
dvm=lista('"Fixo,Hora,Êxito,Misto"'); dvm.add(f"C{R0}:C{RNC}"); cs.add_data_validation(dvm)
cs.conditional_formatting.add(f"J{R0}:J{RNC}", FormulaRule(formula=[f'AND(ISNUMBER(J{R0}),J{R0}<>0)'], fill=fill("FFF4CC")))
cs.conditional_formatting.add(f"I{R0}:I{RNC}", FormulaRule(formula=[f'AND(ISNUMBER(I{R0}),I{R0}>0)'], fill=fill(VERM), font=F(color=VERM_T,size=10,bold=True)))
cs.cell(row=RNC+2,column=1,value="Casos por hora: lance uma parcela por fatura mensal (a diferença em amarelo é o que ainda não foi faturado). Êxito e a parte de êxito do misto só viram parcela no fim do caso: até lá a diferença é o êxito esperado. Cadastre todos os casos, mesmo sem parcela, para a lista ficar completa.").font=F(size=9,color=LILAS)
widths(cs,(28,28,12,16,9,16,14,14,14,18)); cs.freeze_panes="B5"; cs.sheet_view.showGridLines=False; cs.auto_filter.ref=f"A4:J{RNC}"

# ---------- Parcelas ----------
pr=wb.create_sheet("Parcelas")
titulo(pr,"Parcelas","Uma linha por parcela. Escolha o caso na lista; cliente, situação, dias e ação da régua são calculados.",merge_to="L")
hdr(pr,4,["Caso","Cliente","Nº da parcela","Vencimento","Valor (R$)","Pago?","Data do pagamento","Situação","Dias de atraso","Dias para vencer","Faixa de atraso","Ação sugerida (régua)"])
for r in range(R0,RNP+1):
    for c in (1,3,4,5,6,7): inp(pr.cell(row=r,column=c))
    for c in (3,4,6,7): pr.cell(row=r,column=c).alignment=Alignment(horizontal="center")
    pr.cell(row=r,column=4).number_format=DATA; pr.cell(row=r,column=7).number_format=DATA; pr.cell(row=r,column=5).number_format=BRL0
    pr.cell(row=r,column=2,value=f'=IF(A{r}="","",IFERROR(INDEX(Casos!$B${R0}:$B${RNC},MATCH(A{r},Casos!$A${R0}:$A${RNC},0)),""))'); calc(pr.cell(row=r,column=2),center=False)
    pr.cell(row=r,column=8,value=f'=IF(A{r}="","",IF(F{r}="Sim","Paga",IF(D{r}="","Sem vencimento",IF(D{r}<{HOJE},"Vencida","A vencer"))))'); calc(pr.cell(row=r,column=8))
    pr.cell(row=r,column=9,value=f'=IF(H{r}="Vencida",{HOJE}-D{r},"")'); calc(pr.cell(row=r,column=9),"0")
    pr.cell(row=r,column=10,value=f'=IF(H{r}="A vencer",D{r}-{HOJE},"")'); calc(pr.cell(row=r,column=10),"0")
    nivel=f'COUNTIF({RD},"<="&I{r})'
    pr.cell(row=r,column=11,value=f'=IF(H{r}<>"Vencida","",IF({nivel}=0,"Antes da régua",INDEX({RF},{nivel})))'); calc(pr.cell(row=r,column=11))
    pr.cell(row=r,column=12,value=f'=IF(H{r}<>"Vencida","",IF({nivel}=0,"Aguardar: ainda não chegou ao primeiro passo da régua.",INDEX({RA},{nivel})))'); calc(pr.cell(row=r,column=12),center=False)
    pr.cell(row=r,column=13,value=f'=IF(H{r}="Vencida",E{r}*I{r}+ROW()/100000,0)'); pr.cell(row=r,column=13).font=F(color=CINZA,size=9)
    pr.cell(row=r,column=14,value=f'=IF(AND(H{r}="A vencer",J{r}<={JL}),100000-J{r}-ROW()/100000,0)'); pr.cell(row=r,column=14).font=F(color=CINZA,size=9)
pr.column_dimensions["M"].hidden=True; pr.column_dimensions["N"].hidden=True   # colunas auxiliares das listas do Painel
dvs=[lista(off("Casos","A",R0,RNC),strict=True), lista('"Sim,Não"'),
     DataValidation(type="date",operator="greaterThan",formula1="1",allow_blank=True,showErrorMessage=True),
     DataValidation(type="decimal",operator="greaterThanOrEqual",formula1="0",allow_blank=True,showErrorMessage=True),
     DataValidation(type="whole",operator="greaterThanOrEqual",formula1="1",allow_blank=True,showErrorMessage=True)]
for dv,rng in zip(dvs,[f"A{R0}:A{RNP}",f"F{R0}:F{RNP}",f"D{R0}:D{RNP}",f"E{R0}:E{RNP}",f"C{R0}:C{RNP}"]): dv.add(rng); pr.add_data_validation(dv)
pr.conditional_formatting.add(f"A{R0}:L{RNP}", FormulaRule(formula=[f'$H{R0}="Vencida"'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
pr.conditional_formatting.add(f"A{R0}:L{RNP}", FormulaRule(formula=[f'AND($H{R0}="A vencer",$J{R0}<={JC})'], fill=fill("FFF4CC")))
pr.conditional_formatting.add(f"A{R0}:L{RNP}", FormulaRule(formula=[f'$H{R0}="Paga"'], font=F(color=VERDE_T,size=10)))
pr.conditional_formatting.add(f"G{R0}:G{RNP}", FormulaRule(formula=[f'AND($F{R0}="Sim",$G{R0}="")'], fill=fill(VERM)))
pr.cell(row=RNP+2,column=1,value="Vermelho: vencida. Amarelo: vence na janela curta. Verde: paga. Data do pagamento em vermelho: marcada como paga sem data. No exemplo, todas as datas são fixas na data-base de 14/09/2026, para os números fecharem com o caixa (planilha 09) e com a carteira (13).").font=F(size=9,color=LILAS)
widths(pr,(28,26,9,12,13,8,13,13,9,9,16,60)); pr.freeze_panes="C5"; pr.sheet_view.showGridLines=False; pr.auto_filter.ref=f"A4:L{RNP}"

# ---------- Painel ----------
p=wb.create_sheet("Painel",0)
titulo(p,f'=Config!$B$4&" · Parcelas e inadimplência · "&{dstr(HOJE)}',"Nada para digitar aqui: tudo vem de Parcelas e da régua em Config.",merge_to="L")
p.merge_cells("A1:L1")
PJ=f"Parcelas!$J${R0}:$J${RNP}"; PI=f"Parcelas!$I${R0}:$I${RNP}"; PK=f"Parcelas!$K${R0}:$K${RNP}"; PM=f"Parcelas!$M${R0}:$M${RNP}"; PN=f"Parcelas!$N${R0}:$N${RNP}"
kpi(p,4,1,'="Vence em até "&'+JC+'&" dias"',f'=SUMIFS({PE},{PH},"A vencer",{PJ},"<="&{JC})',SOL,UVA,fmt=BRL0)
kpi(p,4,3,'="Vence em até "&'+JL+'&" dias"',f'=SUMIFS({PE},{PH},"A vencer",{PJ},"<="&{JL})',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,5,"Em aberto (total)",f'=SUMIFS({PE},{PH},"A vencer")+SUMIFS({PE},{PH},"Vencida")',LAVANDA,UVA,fmt=BRL0)
kpi(p,4,7,"Vencido (R$)",f'=SUMIFS({PE},{PH},"Vencida")',VERM,VERM_T,fmt=BRL0)
kpi(p,4,9,"Parcelas vencidas",f'=COUNTIFS({PH},"Vencida")',VERM,VERM_T,fmt="0")
kpi(p,4,11,"Inadimplência (vencido ÷ (pago + vencido))",f'=IFERROR(G5/(G5+SUMIFS({PE},{PH},"Paga")),0)',VERM,VERM_T,fmt="0.0%")
p["A7"]="Inadimplência = vencido ÷ (pago + vencido): de tudo o que já venceu, quanto ainda não entrou. É a definição usada em todo o kit (17, 19, 20). Outra conta possível, vencido ÷ em aberto (vencido + a vencer), mede a carteira futura e dá um número maior; não a use para comparar com estas planilhas."; nota(p["A7"]); p.merge_cells("A7:L7"); p["A7"].alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[7].height=30
p["A9"]="Vencidas por faixa de atraso"; p["A9"].font=F(bold=True,size=13,color=UVA)
hdr(p,10,["Faixa","Parcelas","Valor (R$)","% do vencido","Ação da régua"]); p.merge_cells("E10:H10")
for i in range(4):
    r=11+i; src=f"Config!$C${RG0+i}"
    p.cell(row=r,column=1,value=f'=IF({src}="","",{src})'); calc(p.cell(row=r,column=1),center=False)
    p.cell(row=r,column=2,value=f'=IF({src}="","",COUNTIFS({PK},{src}))'); calc(p.cell(row=r,column=2))
    p.cell(row=r,column=3,value=f'=IF({src}="","",SUMIFS({PE},{PK},{src}))'); calc(p.cell(row=r,column=3),BRL0)
    p.cell(row=r,column=4,value=f'=IF({src}="","",IFERROR(C{r}/$G$5,0))'); calc(p.cell(row=r,column=4),PCT)
    p.cell(row=r,column=5,value=f'=IF({src}="","",Config!$B${RG0+i})'); calc(p.cell(row=r,column=5),center=False); p.merge_cells(start_row=r,start_column=5,end_row=r,end_column=8)
    p.cell(row=r,column=5).alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[r].height=32
p.cell(row=15,column=1,value="Antes da régua"); calc(p.cell(row=15,column=1),center=False)
p.cell(row=15,column=2,value=f'=COUNTIFS({PK},"Antes da régua")'); calc(p.cell(row=15,column=2))
p.cell(row=15,column=3,value=f'=SUMIFS({PE},{PK},"Antes da régua")'); calc(p.cell(row=15,column=3),BRL0)
p.cell(row=15,column=4,value='=IFERROR(C15/$G$5,0)'); calc(p.cell(row=15,column=4),PCT)
p.cell(row=15,column=5,value="Vencidas há menos dias que o primeiro passo da régua."); calc(p.cell(row=15,column=5),center=False); p.merge_cells("E15:H15")
p["A17"]="Cobrar primeiro"; p["A17"].font=F(bold=True,size=13,color=UVA)
p["A18"]="Parcelas vencidas em ordem de valor × dias de atraso: as que mais pesam no caixa vêm no topo. O texto da última coluna é a ação da régua para a faixa de cada uma."; nota(p["A18"]); p.merge_cells("A18:L18")
hdr(p,19,["#","Cliente","Caso","Parcela","Vencimento","Valor (R$)","Dias de atraso","Faixa","Ação sugerida"]); p.merge_cells("I19:L19")
TOP=10
for k in range(1,TOP+1):
    r=19+k; m=f'MATCH(LARGE({PM},{k}),{PM},0)'; g=f'LARGE({PM},{k})>0'
    p.cell(row=r,column=1,value=k); calc(p.cell(row=r,column=1))
    for col,src in zip((2,3,4,5,6,7,8,9),("B","A","C","D","E","I","K","L")):
        p.cell(row=r,column=col,value=f'=IFERROR(IF({g},INDEX(Parcelas!${src}${R0}:${src}${RNP},{m}),""),"")'); calc(p.cell(row=r,column=col),center=(col not in (2,3,9)))
    p.cell(row=r,column=5).number_format=DATA; p.cell(row=r,column=6).number_format=BRL0; p.cell(row=r,column=7).number_format="0"
    p.merge_cells(start_row=r,start_column=9,end_row=r,end_column=12); p.cell(row=r,column=9).alignment=Alignment(wrap_text=True,vertical="top"); p.row_dimensions[r].height=32
p.conditional_formatting.add(f"A20:L{19+TOP}", FormulaRule(formula=['$G20>=30'], fill=fill(VERM), font=F(color=VERM_T,size=10)))
r2=19+TOP+2
p.cell(row=r2,column=1,value="Vencem nos próximos dias").font=F(bold=True,size=13,color=UVA)
p.cell(row=r2+1,column=1,value='="Parcelas a vencer em até "&'+JL+'&" dias, a mais próxima primeiro. Bom momento para confirmar o pagamento antes do vencimento."'); nota(p.cell(row=r2+1,column=1)); p.merge_cells(start_row=r2+1,start_column=1,end_row=r2+1,end_column=12)
hdr(p,r2+2,["#","Cliente","Caso","Parcela","Vencimento","Valor (R$)","Dias para vencer"])
for k in range(1,9):
    r=r2+2+k; m=f'MATCH(LARGE({PN},{k}),{PN},0)'; g=f'LARGE({PN},{k})>0'
    p.cell(row=r,column=1,value=k); calc(p.cell(row=r,column=1))
    for col,src in zip((2,3,4,5,6,7),("B","A","C","D","E","J")):
        p.cell(row=r,column=col,value=f'=IFERROR(IF({g},INDEX(Parcelas!${src}${R0}:${src}${RNP},{m}),""),"")'); calc(p.cell(row=r,column=col),center=(col not in (2,3)))
    p.cell(row=r,column=5).number_format=DATA; p.cell(row=r,column=6).number_format=BRL0; p.cell(row=r,column=7).number_format="0"
p.conditional_formatting.add(f"A{r2+3}:G{r2+10}", FormulaRule(formula=[f'AND(ISNUMBER($G{r2+3}),$G{r2+3}<={JC})'], fill=fill("FFF4CC")))
p.cell(row=r2+12,column=1,value="Fonte: aba Parcelas. A planilha avisa e sugere; a conversa com o cliente é sua.").font=F(size=9,color=LILAS)
bc=BarChart(); bc.type="bar"; bc.height=6; bc.width=12; bc.title="Vencido por faixa de atraso (R$)"; bc.style=2
bc.add_data(Reference(p,min_col=3,min_row=10,max_row=15),titles_from_data=True); bc.set_categories(Reference(p,min_col=1,min_row=11,max_row=15))
bc.series[0].graphicalProperties.solidFill="7A1F1F"; bc.legend=None; bc.x_axis.majorGridlines=None
p.add_chart(bc,"J9")
widths(p,(16,26,28,9,12,14,12,16,14,14,14,14)); p.freeze_panes="A4"; p.sheet_view.showGridLines=False

# ---------- Exemplo (Ferraz & Lima) ----------
regua=[(1,"Lembrete gentil por WhatsApp ou e-mail: confirmar se o boleto ou o link de pagamento chegou e reenviar, se preciso."),
       (7,"Mensagem do responsável pelo caso perguntando se houve algum imprevisto e oferecendo uma nova data para o pagamento."),
       (15,"E-mail com o demonstrativo das parcelas em aberto e proposta de renegociação (nova data ou divisão do valor)."),
       (30,"Ligação ou reunião para combinar um plano de pagamento; registrar por escrito o que foi combinado e a nova data.")]
for i,(d,t) in enumerate(regua): cfg.cell(row=RG0+i,column=1,value=d); cfg.cell(row=RG0+i,column=2,value=t)
for i,c in enumerate(dados.CASOS):
    for col,v in zip((1,2,3,4),(c["numero"],c["cliente"],c["tipo_hon"],c["valor_contratado"])): cs.cell(row=R0+i,column=col,value=v)
# parcelas: o cronograma contratado de cada caso (dados.PARCELAS): pagas com data fixa; em aberto com vencimento relativo a hoje
linhas=[]
for p in dados.PARCELAS():
    c=p["caso"]
    if p["pago"]: linhas.append((c["numero"],p["n"],p["vencimento"],p["valor"],"Sim",p["pagamento"]))
    else: linhas.append((c["numero"],p["n"],dados.prazo_formula(dados.dias(p["vencimento"])),p["valor"],"Não",None))
assert len(linhas)<=NP
for i,(num,k,venc,val,pago,pg) in enumerate(linhas):
    r=R0+i; pr.cell(row=r,column=1,value=num); pr.cell(row=r,column=3,value=k); pr.cell(row=r,column=4,value=venc); pr.cell(row=r,column=5,value=val); pr.cell(row=r,column=6,value=pago)
    if pg: pr.cell(row=r,column=7,value=pg)
print(len(linhas),"parcelas no exemplo")

como_usar(wb,"Parcelas e Inadimplência",[
 ("O que esta planilha faz","Você lança as parcelas combinadas com cada cliente e marca as pagas; ela mostra o que vence nos próximos dias, o que já venceu por faixa de atraso, a inadimplência e a lista de quem cobrar primeiro, com a ação sugerida da régua de cobrança."),
 ("Passo 1","Em Config, ajuste as janelas (7 e 30 dias) e a régua: a partir de quantos dias de atraso fazer o quê. Escreva as ações do seu jeito, educadas e diretas."),
 ("Passo 2","Em Casos, cole a lista de casos da planilha 13 · Carteira (número, cliente, modalidade, valor contratado); a 13 é a fonte do cadastro. É a lista que alimenta a escolha do caso em Parcelas."),
 ("Passo 3","Em Parcelas, uma linha por parcela: caso, número da parcela, vencimento e valor. Quando receber, marque Pago? = Sim e a data do pagamento, e lance a mesma entrada no caixa (planilha 09). Honorário por hora entra como fatura mensal; êxito, só no fim do caso."),
 ("Passo 4","Em Painel, veja o que vence em 7 e 30 dias, o vencido por faixa, a inadimplência (vencido ÷ (pago + vencido), a mesma conta das planilhas 17, 19 e 20) e a lista \"Cobrar primeiro\" com o texto da régua para cada parcela."),
 ("Rotina","Segunda-feira, 10 minutos: conferir o extrato, marcar as pagas, mandar as mensagens da lista \"Cobrar primeiro\". Sexta: confirmar as que vencem na semana seguinte."),
 ("Com a IA","Copie uma linha de \"Cobrar primeiro\" (cliente, parcela, valor, dias, ação) e use o prompt \"Clientes 03 · Cobrança educada em três versões\" da biblioteca (ou \"Clientes 02 · Quem cobrar primeiro\" para a lista inteira) ou os 15 modelos de mensagem do bônus. Nunca cole dados que o cliente não autorizou compartilhar."),
])
proteger(wb); salvar(wb,"14-parcelas-e-inadimplencia.xlsx","Parcelas e Inadimplência · Kit de Gestão para Advogados")
