#!/usr/bin/env python3
"""Regera as capturas docs/tela-*.png dos dois kits a partir dos .xlsx de entrega/,
com LibreOffice em locale pt-BR (separadores brasileiros) e recálculo forçado.
Uso: python3 telas.py [essencial|completo|todos] [filtro-de-nome]"""
import subprocess, pathlib, sys, os, json, shutil
ROOT=pathlib.Path(__file__).resolve().parent
S=pathlib.Path('/tmp/claude-0/-home-user-claude-code/a6ac5a85-3495-54eb-8ee2-ee26ae091f81/scratchpad')
PROFILE=S/'lo-ptbr'; WORK=S/'telas-work'; WORK.mkdir(exist_ok=True)
# nome do png -> (arquivo, aba, largura css, altura css)   (sizes = png/1.5)
ESS={
 'tela-semana-tarefas':('01-semana-organizada.xlsx','Tarefas',1300,681),
 'tela-semana-hoje':('01-semana-organizada.xlsx','Hoje',1300,1118),
 'tela-relatorio-painel':('02-relatorio-mensal-pronto.xlsx','Painel',1100,1458),
 'tela-relatorio-resumo':('02-relatorio-mensal-pronto.xlsx','Resumo',1300,681),
 'tela-ganhos-painel':('03-ganhos-e-gastos.xlsx','Painel',1400,1458),
 'tela-ganhos-lancamentos':('03-ganhos-e-gastos.xlsx','Lançamentos',1300,681),
}
COMP={
 'tela-projetos-painel':('04-projetos-e-prazos.xlsx','Painel',1100,1080),
 'tela-projetos-linha':('04-projetos-e-prazos.xlsx','Linha do tempo',1200,800),
 'tela-projetos-etapas':('04-projetos-e-prazos.xlsx','Etapas',1300,700),
 'tela-ata-aberto':('05-ata-e-pendencias.xlsx','Em aberto',1100,1100),
 'tela-ata-resumo':('05-ata-e-pendencias.xlsx','Resumo',1100,1000),
 'tela-ata-pendencias':('05-ata-e-pendencias.xlsx','Pendências',1300,600),
 'tela-metas-painel':('06-metas-do-trimestre.xlsx','Painel',1100,840),
 'tela-metas-metas':('06-metas-do-trimestre.xlsx','Metas',1400,700),
 'tela-orcamento-painel':('07-orcamento-previsto-x-realizado.xlsx','Painel',1200,1170),
 'tela-orcamento-previsto':('07-orcamento-previsto-x-realizado.xlsx','Previsto',1400,700),
 'tela-funil-painel':('08-funil-de-propostas.xlsx','Painel',1200,1040),
 'tela-funil-propostas':('08-funil-de-propostas.xlsx','Propostas',1400,650),
 'tela-horas-painel':('09-horas-e-custo-por-projeto.xlsx','Painel',1200,930),
 'tela-horas-lancamento':('09-horas-e-custo-por-projeto.xlsx','Horas',1200,650),
 'tela-base-base':('10-base-limpa.xlsx','Base',1400,700),
 'tela-base-checklist':('10-base-limpa.xlsx','Checklist',1200,600),
 'tela-base-resumo':('10-base-limpa.xlsx','Resumo',1300,688),
}
JS=r"""const {chromium}=require('playwright');
(async()=>{const [,, html, spec]=process.argv; const itens=JSON.parse(spec);
const b=await chromium.launch({executablePath:'/opt/pw-browsers/chromium',headless:true,args:['--no-sandbox','--disable-gpu']});
const p=await b.newPage({viewport:{width:1400,height:900},deviceScaleFactor:1.5});
await p.goto('file://'+html); await p.waitForTimeout(600);
const hs=await p.$$('h1');
for(const it of itens){ let ok=false;
  for(const el of hs){const t=(await el.innerText()).replace(/\s+/g,' ').trim();
    if(/^Sheet \d+: /.test(t) && t.slice(t.indexOf(': ')+2)===it.aba){const bb=await el.boundingBox();
      let h=it.h; const k=hs.indexOf(el); if(k+1<hs.length){const nb=await hs[k+1].boundingBox(); h=Math.min(h,Math.max(200,Math.floor(nb.y-bb.y-40-24)));}
      await p.screenshot({path:it.out,fullPage:true,clip:{x:0,y:bb.y+40,width:it.w,height:h}}); ok=true; break;}}
  console.log(ok?'ok':'FALTOU', it.aba, it.out);}
await b.close();})();"""
(WORK/'shot.js').write_text(JS)
def gera(kit,tabela,filtro=None):
    ent=ROOT/f'kit-{kit}'/'entrega'; docs=ROOT/f'kit-{kit}'/'docs'
    por_arquivo={}
    for nome,(arq,aba,w,h) in tabela.items():
        if filtro and filtro not in nome: continue
        por_arquivo.setdefault(arq,[]).append(dict(aba=aba,w=w,h=h,out=str(docs/f'{nome}.png')))
    for arq,itens in por_arquivo.items():
        src=ent/arq; dst=WORK/arq; shutil.copy2(src,dst); html=WORK/(dst.stem+'.html'); html.unlink(missing_ok=True)
        import openpyxl
        na_cells=[(ws.title,c.coordinate) for ws in openpyxl.load_workbook(src).worksheets for row in ws.iter_rows() for c in row if isinstance(c.value,str) and 'NA()' in c.value]
        if na_cells:
            # Cópia só para captura: células que dão #N/A (lacuna do gráfico no Excel) ficam VAZIAS, porque o
            # LibreOffice plota erro/"" como zero na exportação. 1) recalcula, 2) acha os erros, 3) esvazia.
            rc=WORK/'recalc'; rc.mkdir(exist_ok=True)
            subprocess.run(['soffice','--headless',f'-env:UserInstallation=file://{PROFILE}','--convert-to','xlsx','--outdir',str(rc),str(dst)],env={**os.environ,'SAL_USE_VCLPLUGIN':'svp'},check=True,capture_output=True,timeout=300)
            vals=openpyxl.load_workbook(rc/dst.name,data_only=True)
            wb=openpyxl.load_workbook(dst)
            for t,coord in na_cells:
                v=vals[t][coord].value
                if v is None or v=='' or (isinstance(v,str) and v.startswith('#')): wb[t][coord].value=None
            wb.save(dst)
        subprocess.run(['soffice','--headless',f'-env:UserInstallation=file://{PROFILE}','--convert-to','html','--outdir',str(WORK),str(dst)],
                       env={**os.environ,'SAL_USE_VCLPLUGIN':'svp'},check=True,capture_output=True,timeout=300)
        r=subprocess.run(['node',str(WORK/'shot.js'),str(html),json.dumps(itens)],env={**os.environ,'NODE_PATH':str(S/'pw'/'node_modules')},capture_output=True,text=True)
        print(arq,r.stdout.strip(),r.stderr.strip()[:300])
def gera_auto(kit,largura=1300,altura=900,filtro=None):
    """Captura automática: a aba principal (primeira que não é 'Como usar') de cada .xlsx de entrega/, como tela-NN.png,
    e também a aba Painel quando existir e não for a principal (tela-NN-painel.png)."""
    import openpyxl
    ent=ROOT/f'kit-{kit}'/'entrega'; tabela={}
    for arq in sorted(ent.glob('[0-9][0-9]-*.xlsx')):
        nn=arq.name[:2]
        if filtro and filtro not in arq.name: continue
        abas=[a for a in openpyxl.load_workbook(arq,read_only=True).sheetnames if a!='Como usar']
        if not abas: continue
        tabela[f'tela-{nn}']=(arq.name,abas[0],largura,altura)
        if 'Painel' in abas and abas[0]!='Painel': tabela[f'tela-{nn}-painel']=(arq.name,'Painel',largura,altura)
    gera(kit,tabela)
kit=sys.argv[1] if len(sys.argv)>1 else 'todos'; filtro=sys.argv[2] if len(sys.argv)>2 else None
if kit=='advogados': gera_auto('advogados',filtro=filtro); sys.exit()
if kit in('essencial','todos'): gera('essencial',ESS,filtro)
if kit in('completo','todos'):
    gera('completo',COMP,filtro)
    # as capturas das planilhas 01-03 são as mesmas do Essencial: copia para o Completo (manual e aulas usam)
    for png in (ROOT/'kit-essencial'/'docs').glob('tela-*.png'):
        if (ROOT/'kit-completo'/'docs'/png.name).exists(): shutil.copy2(png, ROOT/'kit-completo'/'docs'/png.name)
